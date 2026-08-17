#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""琪哥发票入金蝶：改样发票簿 → 明细 + 金蝶凭证引入表。金额只由本脚本取表列。"""
from __future__ import annotations

import argparse
import json
import shutil
import sys
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook, load_workbook

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

HERE = Path(__file__).resolve().parent
SKILL_DIR = HERE.parent
CONFIG_DIR = SKILL_DIR / "config"
TEMPLATE_PATH = CONFIG_DIR / "凭证引入空模.xlsx"
KINGDEE_RESULT_NAME = "凭证引入_结果.xlsx"
TWOPLACES = Decimal("0.01")
KINGDEE_SHEET = "sheet1（名称勿改）"


def log(msg: str) -> None:
    print(msg, file=sys.stderr)


def money(v) -> Decimal | None:
    if v is None or v == "":
        return None
    if isinstance(v, Decimal):
        return v.quantize(TWOPLACES, rounding=ROUND_HALF_UP)
    if isinstance(v, bool):
        return None
    if isinstance(v, int):
        return Decimal(v).quantize(TWOPLACES, rounding=ROUND_HALF_UP)
    if isinstance(v, float):
        return Decimal(str(round(v, 2))).quantize(TWOPLACES, rounding=ROUND_HALF_UP)
    s = str(v).strip().replace(",", "")
    if not s or s.startswith("="):
        return None
    try:
        return Decimal(s).quantize(TWOPLACES, rounding=ROUND_HALF_UP)
    except Exception:
        return None


def code_str(v) -> str:
    if v is None or v == "":
        return ""
    if isinstance(v, bool):
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, int):
        return str(v)
    s = str(v).strip()
    if s.endswith(".0") and s[:-2].isdigit():
        return s[:-2]
    return s


def norm_customer(name: str) -> str:
    s = (name or "").strip().replace("(", "（").replace(")", "）")
    return "".join(s.split())


def load_json(path: Path, default):
    if not path.is_file():
        return default
    with path.open(encoding="utf-8") as f:
        return json.load(f)


def load_rules() -> dict:
    data = load_json(CONFIG_DIR / "rules.json", {})
    return {
        "tax_account": str(data.get("tax_account") or "21710105"),
        "pack_size": int(data.get("pack_size") or 5),
        "voucher_word": str(data.get("voucher_word") or "记"),
        "currency": str(data.get("currency") or "RMB"),
        "exchange_rate": Decimal(str(data.get("exchange_rate") or 1)),
        "booking_date": str(data.get("booking_date") or "run_day"),
        "tax_rate_divisor": Decimal(str(data.get("tax_rate_divisor") or "1.06")),
        "currency_name": str(data.get("currency_name") or "人民币"),
        "account_names": {
            str(k): str(v) for k, v in (data.get("account_names") or {}).items() if k and v
        },
    }


def load_aliases() -> dict:
    data = load_json(CONFIG_DIR / "列名别名.json", {})
    return {
        "发票": data.get("发票_列别名") or {},
        "组织架构": data.get("组织架构_列别名") or {},
    }


def header_index(headers: list, aliases: dict) -> dict:
    idx = {}
    cleaned = [(i, str(h).strip()) for i, h in enumerate(headers) if h is not None and str(h).strip()]
    for field, names in aliases.items():
        for i, h in cleaned:
            if h in names:
                idx[field] = i
                break
    return idx


def cell_at(row: tuple, idx: dict, field: str):
    i = idx.get(field)
    if i is None or i >= len(row):
        return None
    return row[i]


@dataclass
class Line:
    source_row: int
    invoice_no: str
    invoice_type: str
    unit_name: str
    applicant: str
    total: Decimal | None = None
    amount: Decimal | None = None
    tax: Decimal | None = None
    ar: str = ""
    rev: str = ""
    status: str = ""
    reason: str = ""
    customer_code: str = ""
    customer_name: str = ""
    dept_code: str = ""
    dept_name: str = ""
    emp_code: str = ""
    emp_name: str = ""
    voucher_no: int | None = None


@dataclass
class ConvertResult:
    source_count: int
    bookable_count: int
    hold_count: int
    holds: list = field(default_factory=list)
    bookable: list = field(default_factory=list)
    kingdee_path: Path | None = None
    detail_path: Path | None = None


class Master:
    def __init__(self, data: dict):
        self.emp: dict[str, list[tuple[str, str]]] = {}
        self.dept: dict[str, str] = {}
        self.cus: dict[str, list[tuple[str, str]]] = {}
        for e in data.get("employee") or []:
            name = str(e.get("name") or "").strip()
            code = code_str(e.get("code"))
            if name and code:
                self.emp.setdefault(name, []).append((code, name))
        for d in data.get("department") or []:
            code = code_str(d.get("code"))
            name = str(d.get("name") or "").strip()
            if code:
                self.dept[code] = name
        for c in data.get("customer") or []:
            name = str(c.get("name") or "")
            code = code_str(c.get("code"))
            key = norm_customer(name)
            if key and code:
                self.cus.setdefault(key, []).append((code, str(c.get("name") or "").strip()))

    def employee(self, name: str):
        hits = self.emp.get((name or "").strip()) or []
        if len(hits) == 1:
            return hits[0], None
        if not hits:
            return None, "职员档案没有此人"
        return None, "职员档案一对多"

    def customer(self, name: str):
        hits = self.cus.get(norm_customer(name)) or []
        if len(hits) == 1:
            return hits[0], None
        if not hits:
            return None, "客户档案没有此抬头"
        return None, "客户档案一对多"

    def department(self, code: str):
        code = code_str(code)
        name = self.dept.get(code)
        if not code:
            return None, "缺部门编码"
        if not name:
            return None, "部门编码档案没有"
        return (code, name), None


def load_org(ws, aliases: dict) -> tuple[dict[str, list[str]], str | None]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {}, "组织架构是空的"
    idx = header_index(list(rows[0]), aliases)
    if "姓名" not in idx or "部门编码" not in idx:
        return {}, "组织架构缺姓名或部门编码列"
    mapping: dict[str, list[str]] = {}
    for row in rows[1:]:
        if not row:
            continue
        name = str(cell_at(row, idx, "姓名") or "").strip()
        dept = code_str(cell_at(row, idx, "部门编码"))
        if not name:
            continue
        mapping.setdefault(name, []).append(dept)
    return mapping, None


def pick_amount(formula_cell, value_cell, fallback: Decimal | None) -> Decimal | None:
    got = money(value_cell)
    if got is not None:
        return got
    got = money(formula_cell)
    if got is not None:
        return got
    return fallback


def compute_amount_tax(total: Decimal | None, amount, tax, divisor: Decimal):
    amt = money(amount)
    tx = money(tax)
    if total is None:
        return amt, tx
    if amt is None:
        amt = (total / divisor).quantize(TWOPLACES, rounding=ROUND_HALF_UP)
    if tx is None and amt is not None:
        tx = (total - amt).quantize(TWOPLACES, rounding=ROUND_HALF_UP)
    return amt, tx


def read_invoices(path: Path, aliases: dict, divisor: Decimal) -> tuple[list[dict], str | None]:
    wb_f = load_workbook(path, data_only=False)
    wb_v = load_workbook(path, data_only=True)
    if "发票" not in wb_f.sheetnames:
        return [], "找不到 sheet「发票」"
    if "组织架构" not in wb_f.sheetnames:
        return [], "找不到 sheet「组织架构」"
    ws_f = wb_f["发票"]
    ws_v = wb_v["发票"]
    headers = [c.value for c in next(ws_f.iter_rows(min_row=1, max_row=1))]
    idx = header_index(headers, aliases["发票"])
    needed = ["单位名称", "价税合计", "应收账款编码", "主营业务收入编码", "申请人"]
    missing = [k for k in needed if k not in idx]
    if missing:
        return [], f"发票表缺列：{'、'.join(missing)}"
    org, org_err = load_org(wb_f["组织架构"], aliases["组织架构"])
    if org_err:
        return [], org_err
    items = []
    max_row = ws_f.max_row or 1
    for r in range(2, max_row + 1):
        row_f = [ws_f.cell(r, c + 1).value for c in range(len(headers))]
        row_v = [ws_v.cell(r, c + 1).value for c in range(len(headers))]
        unit = str(cell_at(row_f, idx, "单位名称") or "").strip()
        if not unit:
            continue
        total = pick_amount(cell_at(row_f, idx, "价税合计"), cell_at(row_v, idx, "价税合计"), None)
        amt_raw = pick_amount(cell_at(row_f, idx, "金额"), cell_at(row_v, idx, "金额"), None)
        tax_raw = pick_amount(cell_at(row_f, idx, "税额"), cell_at(row_v, idx, "税额"), None)
        amount, tax = compute_amount_tax(total, amt_raw, tax_raw, divisor)
        items.append(
            {
                "source_row": r,
                "invoice_no": str(cell_at(row_f, idx, "发票号") or "").strip(),
                "invoice_type": str(cell_at(row_f, idx, "发票类型") or "").strip(),
                "unit_name": unit,
                "applicant": str(cell_at(row_f, idx, "申请人") or "").strip(),
                "total": total,
                "amount": amount,
                "tax": tax,
                "ar": code_str(cell_at(row_f, idx, "应收账款编码")),
                "rev": code_str(cell_at(row_f, idx, "主营业务收入编码")),
                "org": org,
            }
        )
    wb_f.close()
    wb_v.close()
    return items, None


def classify(item: dict, master: Master | None, tax_account: str) -> Line:
    line = Line(
        source_row=item["source_row"],
        invoice_no=item["invoice_no"],
        invoice_type=item["invoice_type"],
        unit_name=item["unit_name"],
        applicant=item["applicant"],
        total=item["total"],
        amount=item["amount"],
        tax=item["tax"],
        ar=item["ar"],
        rev=item["rev"],
    )
    if not line.ar or not line.rev:
        line.status = "待确认"
        line.reason = "缺科目编码"
        return line
    if line.total is None or line.amount is None or line.tax is None:
        line.status = "待确认"
        line.reason = "缺金额"
        return line
    if line.total != (line.amount + line.tax):
        line.status = "待确认"
        line.reason = "价税合计不等于金额加税额"
        return line
    if not line.invoice_type:
        line.status = "待确认"
        line.reason = "缺发票类型"
        return line
    if not line.applicant:
        line.status = "待确认"
        line.reason = "缺申请人"
        return line
    depts = item["org"].get(line.applicant) or []
    if not depts:
        line.status = "待确认"
        line.reason = "组织架构无此申请人"
        return line
    uniq = list(dict.fromkeys(depts))
    if len(uniq) != 1 or not uniq[0]:
        line.status = "待确认"
        line.reason = "组织架构重名或部门编码空"
        return line
    line.dept_code = uniq[0]
    line.dept_name = ""
    line.emp_code = ""
    line.emp_name = line.applicant
    line.customer_code = ""
    line.customer_name = line.unit_name
    if master:
        dept, err = master.department(line.dept_code)
        if dept:
            line.dept_code, line.dept_name = dept
        emp, err = master.employee(line.applicant)
        if err and "一对多" in err:
            line.status = "待确认"
            line.reason = err
            return line
        if emp:
            line.emp_code, line.emp_name = emp
        cus, err = master.customer(line.unit_name)
        if err and "一对多" in err:
            line.status = "待确认"
            line.reason = err
            return line
        if cus:
            line.customer_code, line.customer_name = cus
    line.status = "可入账"
    line.reason = ""
    _ = tax_account
    return line


def summary_of(line: Line) -> str:
    typ = line.invoice_type
    if line.total is not None and line.total < 0 and "红字" not in typ:
        typ = "红字" + typ
    return f"{typ}：{line.unit_name}"


def col_by_label(ws, needle: str) -> int:
    for cell in ws[3]:
        if needle in str(cell.value or ""):
            return cell.column
    raise KeyError(needle)


def write_kingdee(path: Path, bookable: list[Line], rules: dict, booking: str, template: Path) -> Path:
    shutil.copy2(template, path)
    wb = load_workbook(path)
    if KINGDEE_SHEET not in wb.sheetnames:
        wb.close()
        raise SystemExit("金蝶空模缺 sheet1（名称勿改）")
    ws = wb[KINGDEE_SHEET]
    cols = {
        "date": col_by_label(ws, "*记账日期"),
        "word": col_by_label(ws, "*凭证字号"),
        "number": col_by_label(ws, "凭证号 #"),
        "expl": col_by_label(ws, "摘要 #"),
        "account": col_by_label(ws, "*科目.编码"),
        "account_name": col_by_label(ws, "科目.名称"),
        "currency": col_by_label(ws, "*币别.编码"),
        "currency_name": col_by_label(ws, "币别.名称"),
        "rate": col_by_label(ws, "*汇率"),
        "amountfor": col_by_label(ws, "原币金额"),
        "debit": col_by_label(ws, "借方 #"),
        "credit": col_by_label(ws, "贷方 #"),
        "cus_code": col_by_label(ws, "辅助核算.客户.编码"),
        "cus_name": col_by_label(ws, "辅助核算.客户.名称"),
        "dep_code": col_by_label(ws, "辅助核算.部门.编码"),
        "dep_name": col_by_label(ws, "辅助核算.部门.名称"),
        "emp_code": col_by_label(ws, "辅助核算.职员.编码"),
        "emp_name": col_by_label(ws, "辅助核算.职员.名称"),
    }
    pack = rules["pack_size"]
    row_i = 4
    for batch, start in enumerate(range(0, len(bookable), pack), start=1):
        chunk = bookable[start : start + pack]
        for line in chunk:
            line.voucher_no = batch
            expl = summary_of(line)
            entries = [
                (line.ar, line.total, True, True),
                (line.rev, line.amount, False, True),
                (rules["tax_account"], line.tax, False, False),
            ]
            for account, amt, is_debit, with_aux in entries:
                ws.cell(row_i, cols["date"], booking)
                ws.cell(row_i, cols["word"], rules["voucher_word"])
                ws.cell(row_i, cols["number"], batch)
                ws.cell(row_i, cols["expl"], expl)
                ws.cell(row_i, cols["account"], account)
                acc_name = rules["account_names"].get(str(account), "")
                if acc_name:
                    ws.cell(row_i, cols["account_name"], acc_name)
                ws.cell(row_i, cols["currency"], rules["currency"])
                if rules.get("currency_name"):
                    ws.cell(row_i, cols["currency_name"], rules["currency_name"])
                ws.cell(row_i, cols["rate"], float(rules["exchange_rate"]))
                val = float(amt)
                ws.cell(row_i, cols["amountfor"], val)
                if is_debit:
                    ws.cell(row_i, cols["debit"], val)
                else:
                    ws.cell(row_i, cols["credit"], val)
                if with_aux:
                    ws.cell(row_i, cols["cus_code"], line.customer_code)
                    ws.cell(row_i, cols["cus_name"], line.customer_name)
                    ws.cell(row_i, cols["dep_code"], line.dept_code)
                    ws.cell(row_i, cols["dep_name"], line.dept_name)
                    ws.cell(row_i, cols["emp_code"], line.emp_code)
                    ws.cell(row_i, cols["emp_name"], line.emp_name)
                row_i += 1
    wb.save(path)
    wb.close()
    return path


def write_detail(path: Path, lines: list[Line]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "明细"
    ws.append(
        [
            "状态",
            "原因",
            "源行号",
            "发票号",
            "单位名称",
            "申请人",
            "价税合计",
            "金额",
            "税额",
            "应收账款编码",
            "主营业务收入编码",
            "客户编码",
            "客户档案名",
            "部门编码",
            "职员编码",
            "凭证批次",
        ]
    )
    for line in lines:
        ws.append(
            [
                line.status,
                line.reason,
                line.source_row,
                line.invoice_no,
                line.unit_name,
                line.applicant,
                float(line.total) if line.total is not None else None,
                float(line.amount) if line.amount is not None else None,
                float(line.tax) if line.tax is not None else None,
                line.ar,
                line.rev,
                line.customer_code,
                line.customer_name,
                line.dept_code,
                line.emp_code,
                line.voucher_no,
            ]
        )
    wb.save(path)
    wb.close()
    return path


def convert(
    invoice_path: Path,
    out_dir: Path,
    booking_date: str | None = None,
    template_path: Path | None = None,
    master_path: Path | None = None,
) -> ConvertResult:
    invoice_path = Path(invoice_path)
    out_dir = Path(out_dir)
    if not invoice_path.is_file():
        raise SystemExit("找不到发票 Excel")
    template = Path(template_path) if template_path else TEMPLATE_PATH
    if not template.is_file():
        raise SystemExit("缺金蝶引入空模")
    rules = load_rules()
    aliases = load_aliases()
    booking = booking_date or date.today().isoformat()
    if rules["booking_date"] != "run_day" and not booking_date:
        booking = str(rules["booking_date"])
    items, err = read_invoices(invoice_path, aliases, rules["tax_rate_divisor"])
    if err:
        raise SystemExit(err)
    master = None
    if master_path:
        master_path = Path(master_path)
        if not master_path.is_file():
            raise SystemExit("金蝶档案 json 不存在")
        try:
            master = Master(json.loads(master_path.read_text(encoding="utf-8")))
        except Exception:
            raise SystemExit("金蝶档案 json 读失败")
    lines = [classify(item, master, rules["tax_account"]) for item in items]
    bookable = [x for x in lines if x.status == "可入账"]
    holds = [x for x in lines if x.status != "可入账"]
    out_dir.mkdir(parents=True, exist_ok=True)
    detail_path = out_dir / f"{invoice_path.stem}_明细结果.xlsx"
    kingdee_path = out_dir / KINGDEE_RESULT_NAME
    if kingdee_path.resolve() == template.resolve():
        kingdee_path = out_dir / "凭证引入_填写结果.xlsx"
    write_detail(detail_path, lines)
    write_kingdee(kingdee_path, bookable, rules, booking, template)
    return ConvertResult(
        source_count=len(lines),
        bookable_count=len(bookable),
        hold_count=len(holds),
        holds=[
            {
                "发票号": h.invoice_no,
                "单位名称": h.unit_name,
                "申请人": h.applicant,
                "原因": h.reason,
            }
            for h in holds
        ],
        bookable=[{"客户编码": b.customer_code, "发票号": b.invoice_no} for b in bookable],
        kingdee_path=kingdee_path,
        detail_path=detail_path,
    )


def is_result_file(path: Path) -> bool:
    return "结果" in path.stem


def sniff_invoice(path: Path) -> bool:
    if is_result_file(path):
        return False
    try:
        wb = load_workbook(path, read_only=True, data_only=False)
        ok = "发票" in wb.sheetnames and "组织架构" in wb.sheetnames
        wb.close()
        return ok
    except Exception:
        return False


def sniff_master(path: Path) -> bool:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        return isinstance(data, dict) and "employee" in data and "customer" in data
    except Exception:
        return False


def inspect_dir(input_dir: Path) -> dict:
    found = {"invoice": None, "master": None}
    for p in sorted(Path(input_dir).iterdir()):
        if p.suffix.lower() in {".xlsx", ".xlsm"} and sniff_invoice(p) and not found["invoice"]:
            found["invoice"] = str(p)
        if p.suffix.lower() == ".json" and sniff_master(p) and not found["master"]:
            found["master"] = str(p)
    return found


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="琪哥发票入金蝶")
    parser.add_argument("--inspect", action="store_true")
    parser.add_argument("--input-dir")
    parser.add_argument("--invoice")
    parser.add_argument("--master")
    parser.add_argument("--out-dir", "--out", dest="out_dir")
    parser.add_argument("--template")
    parser.add_argument("--date")
    args = parser.parse_args(argv)
    if args.inspect:
        target = Path(args.input_dir or SKILL_DIR / "工作区" / "input")
        found = inspect_dir(target)
        print(json.dumps(found, ensure_ascii=False, indent=2))
        return 0 if found["invoice"] else 2
    invoice = Path(args.invoice) if args.invoice else None
    template = Path(args.template) if args.template else None
    master = Path(args.master) if args.master else None
    input_dir = Path(args.input_dir) if args.input_dir else None
    if input_dir:
        found = inspect_dir(input_dir)
        invoice = invoice or (Path(found["invoice"]) if found["invoice"] else None)
        master = master or (Path(found["master"]) if found["master"] else None)
    if not invoice:
        log("缺发票 Excel。把改样发票簿放进文件夹即可，金蝶空模技能自带。")
        return 2
    if not template:
        template = TEMPLATE_PATH
    if not template.is_file():
        log("技能缺金蝶引入空模 config/凭证引入空模.xlsx。")
        return 2
    out_dir = Path(args.out_dir) if args.out_dir else (input_dir or invoice.parent)
    result = convert(
        invoice_path=invoice,
        out_dir=out_dir,
        booking_date=args.date,
        template_path=template,
        master_path=master,
    )
    log(
        f"源有效行 {result.source_count}：可入账 {result.bookable_count}，待确认 {result.hold_count}。"
        f"金蝶表={result.kingdee_path}；明细={result.detail_path}。未点引入。"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
