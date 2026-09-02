#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""琪哥发票入金蝶 · 合成回归 + 本机 30 张金标（有文件才跑）。"""
from __future__ import annotations

import json
import shutil
import sys
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
SCRIPTS = SKILL / "scripts"
sys.path.insert(0, str(SCRIPTS))

import convert  # noqa: E402

GOLD_INVOICE = Path(
    "/Users/evanlee/Documents/甲骨易实习/项目/长期项目/财务部skills"
    "/技能/金蝶/琪哥发票入金蝶/原始素材/20260814_改样发票/发票明昊.xlsx"
)
GOLD_MASTER = Path(
    "/Users/evanlee/Documents/甲骨易实习/项目/长期项目/自动化记账（金蝶）"
    "/原始素材/实操批次/20260803_金蝶补辅助核算_七月前半/调研/金蝶基础资料_master.json"
)


def _master():
    return {
        "employee": [
            {"code": "103", "name": "于占国"},
            {"code": "011", "name": "陈霞"},
            {"code": "036", "name": "梁玲玲"},
        ],
        "department": [
            {"code": "15", "name": "本地化事业部"},
            {"code": "0302", "name": "营销二部"},
        ],
        "customer": [
            {"code": "1001", "name": "甲科技有限公司"},
            {"code": "1002", "name": "乙事务所"},
            {"code": "2460", "name": "乙中心"},
            {"code": "3843", "name": "丙科技（深圳）有限公司"},
        ],
    }


def _write_invoice(path: Path, rows, org=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "发票"
    ws.append(
        [
            "日期",
            "发票类型",
            "发票号",
            "单位名称",
            "价税合计",
            "金额",
            "税额",
            "申请人",
            "部门编码",
            "应收账款编码",
            "主营业务收入编码",
        ]
    )
    for r in rows:
        ws.append(r)
    org_ws = wb.create_sheet("组织架构")
    org_ws.append(["姓名", "部门编码"])
    for pair in org or [("于占国", "15"), ("陈霞", "15"), ("梁玲玲", "0302"), ("不存在职员甲", "15")]:
        org_ws.append(list(pair))
    wb.save(path)
    wb.close()


def _write_master(path: Path, data=None):
    path.write_text(json.dumps(data or _master(), ensure_ascii=False), encoding="utf-8")


def _copy_template(path: Path):
    shutil.copy2(convert.TEMPLATE_PATH, path)


def _run(tmp_path: Path, rows, org=None, master=None, booking="2026-08-14", with_master=True, start_voucher_no=1):
    inv = tmp_path / "发票.xlsx"
    out = tmp_path / "out"
    out.mkdir()
    _write_invoice(inv, rows, org=org)
    master_path = None
    if with_master:
        master_path = tmp_path / "master.json"
        _write_master(master_path, master)
    return convert.convert(
        invoice_path=inv,
        out_dir=out,
        booking_date=booking,
        master_path=master_path,
        start_voucher_no=start_voucher_no,
    )


def _ok_row(
    name="甲科技有限公司",
    app="于占国",
    tot=1060,
    amt=1000,
    tax=60,
    ar="113103",
    rev="510103",
    typ="专票",
    inv="1",
    day="2026-08-01",
):
    return [day, typ, inv, name, tot, amt, tax, app, None, ar, rev]


def test_skip_empty_unit_name(tmp_path):
    rows = [_ok_row(), [None, None, None, None, None, None, None, None, None, None, None]]
    result = _run(tmp_path, rows)
    assert result.source_count == 1
    assert result.bookable_count == 1
    assert result.hold_count == 0


def test_hold_missing_account(tmp_path):
    rows = [_ok_row(ar=None, rev="510103")]
    result = _run(tmp_path, rows)
    assert result.bookable_count == 0
    assert result.hold_count == 1
    assert "科目" in result.holds[0]["原因"]


def test_unknown_customer_still_bookable(tmp_path):
    rows = [_ok_row(name="不存在客户甲有限公司")]
    result = _run(tmp_path, rows)
    assert result.bookable_count == 1
    assert result.hold_count == 0
    kd = load_workbook(result.kingdee_path)
    ws = kd[convert.KINGDEE_SHEET]
    assert ws.cell(4, 19).value == "不存在客户甲有限公司"
    assert ws.cell(4, 18).value in (None, "")
    kd.close()


def test_customer_suffix_mismatch_does_not_guess_code(tmp_path):
    rows = [_ok_row(name="乙中心有限公司")]
    result = _run(tmp_path, rows)
    assert result.bookable_count == 1
    kd = load_workbook(result.kingdee_path)
    ws = kd[convert.KINGDEE_SHEET]
    assert ws.cell(4, 18).value in (None, "")
    assert ws.cell(4, 19).value == "乙中心有限公司"
    kd.close()


def test_unknown_employee_still_bookable(tmp_path):
    rows = [_ok_row(app="不存在职员甲")]
    result = _run(tmp_path, rows)
    assert result.bookable_count == 1
    kd = load_workbook(result.kingdee_path)
    ws = kd[convert.KINGDEE_SHEET]
    assert ws.cell(4, 25).value == "不存在职员甲"
    assert ws.cell(4, 24).value in (None, "")
    kd.close()


def test_customer_one_to_many_holds(tmp_path):
    master = _master()
    master["customer"].append({"code": "9999", "name": "甲科技有限公司"})
    result = _run(tmp_path, [_ok_row()], master=master)
    assert result.bookable_count == 0
    assert "一对多" in result.holds[0]["原因"]


def test_normalize_brackets_and_spaces_match(tmp_path):
    rows = [
        _ok_row(name="丙科技(深圳)有限公司", inv="a"),
        _ok_row(name="甲科技有限公司 ", inv="b"),
    ]
    result = _run(tmp_path, rows)
    assert result.bookable_count == 2
    codes = {r["客户编码"] for r in result.bookable}
    assert codes == {"3843", "1001"}


def test_three_lines_balance_and_tax_no_aux(tmp_path):
    result = _run(tmp_path, [_ok_row()])
    kd = load_workbook(result.kingdee_path)
    ws = kd["sheet1（名称勿改）"]
    lines = []
    for row in ws.iter_rows(min_row=4, max_row=6, max_col=25, values_only=True):
        lines.append(row)
    assert len(lines) == 3
    debit = sum(Decimal(str(r[15] or 0)) for r in lines)
    credit = sum(Decimal(str(r[16] or 0)) for r in lines)
    assert debit == credit == Decimal("1060.00")
    assert lines[0][6] == "113103"
    assert lines[0][7] == "应收账款_多语本地化服务"
    assert lines[1][6] == "510103"
    assert lines[1][7] == "主营业务收入_多语本地化服务"
    assert lines[2][6] == "21710105"
    assert lines[2][7] == "应交税费_应交增值税_销项税额"
    assert lines[0][12] == "人民币"
    assert lines[2][17] in (None, "")
    assert lines[2][21] in (None, "")
    assert lines[2][23] in (None, "")
    assert lines[0][17] == "1001"
    assert str(lines[0][21]) == "15"
    assert str(lines[0][23]) == "103"
    assert lines[0][5] == "专票：甲科技有限公司"
    assert str(lines[0][0]) == "2026-08-14"
    assert lines[0][1] == "记"
    kd.close()


def _voucher_nums(path, invoices):
    kd = load_workbook(path)
    ws = kd["sheet1（名称勿改）"]
    nums = []
    for row in ws.iter_rows(min_row=4, max_row=3 + invoices * 3, max_col=3, values_only=True):
        nums.append(row[2])
    kd.close()
    return nums


def test_pack_five_per_voucher(tmp_path):
    rows = [
        _ok_row(name=f"客户{i}有限公司", inv=str(i), tot=1060, amt=1000, tax=60)
        for i in range(12)
    ]
    result = _run(tmp_path, rows)
    assert result.bookable_count == 12
    nums = _voucher_nums(result.kingdee_path, 12)
    assert nums.count(1) == 15
    assert nums.count(2) == 15
    assert nums.count(3) == 6


def test_start_voucher_no_shifts_batches(tmp_path):
    rows = [_ok_row(name=f"客户{i}有限公司", inv=str(i)) for i in range(6)]
    result = _run(tmp_path, rows, start_voucher_no=11)
    nums = _voucher_nums(result.kingdee_path, 6)
    assert 1 not in nums
    assert nums.count(11) == 15
    assert nums.count(12) == 3


def test_pack_keeps_consecutive_company_across_five(tmp_path):
    rows = (
        [_ok_row(name=f"散户{i}", inv=str(i)) for i in range(5)]
        + [_ok_row(name="连号甲", inv=f"a{i}") for i in range(3)]
        + [_ok_row(name="连号乙", inv=f"b{i}") for i in range(4)]
    )
    result = _run(tmp_path, rows)
    nums = _voucher_nums(result.kingdee_path, 12)
    assert nums.count(1) == 15
    assert nums.count(2) == 21
    assert 3 not in nums


def test_formula_account_without_cache_holds(tmp_path):
    result = _run(tmp_path, [_ok_row(ar="=VLOOKUP(1,A1:B2,1,0)")])
    assert result.bookable_count == 0
    assert result.hold_count == 1
    assert result.holds[0]["原因"] == "缺科目编码"


def test_pick_code_uses_cached_value_for_formula():
    assert convert.pick_code("=VLOOKUP(1,A1:B2,1,0)", 113103) == "113103"
    assert convert.pick_code("=VLOOKUP(1,A1:B2,1,0)", None) == ""
    assert convert.pick_code("113103", None) == "113103"


def test_pack_same_company_over_five_one_voucher(tmp_path):
    rows = [_ok_row(name="同一家", inv=str(i)) for i in range(8)]
    result = _run(tmp_path, rows)
    nums = _voucher_nums(result.kingdee_path, 8)
    assert nums.count(1) == 24
    assert 2 not in nums


def test_booking_date_not_invoice_date(tmp_path):
    result = _run(tmp_path, [_ok_row(day="2026-08-01")], booking="2026-08-14")
    kd = load_workbook(result.kingdee_path)
    ws = kd["sheet1（名称勿改）"]
    assert str(ws.cell(4, 1).value) == "2026-08-14"
    kd.close()


def test_red_invoice_negative_and_summary(tmp_path):
    rows = [_ok_row(typ="普票", tot=-1060, amt=-1000, tax=-60)]
    result = _run(tmp_path, rows)
    kd = load_workbook(result.kingdee_path)
    ws = kd["sheet1（名称勿改）"]
    assert ws.cell(4, 6).value == "红字普票：甲科技有限公司"
    assert Decimal(str(ws.cell(4, 16).value)) == Decimal("-1060.00")
    kd.close()


def test_detail_covers_every_source_row(tmp_path):
    rows = [
        _ok_row(inv="ok"),
        _ok_row(ar=None, rev=None, inv="hold"),
    ]
    result = _run(tmp_path, rows)
    det = load_workbook(result.detail_path)
    ws = det.active
    statuses = [ws.cell(i, 1).value for i in range(2, ws.max_row + 1)]
    assert len(statuses) == 2
    assert set(statuses) == {"可入账", "待确认"}
    det.close()


def test_missing_master_still_converts(tmp_path):
    inv = tmp_path / "发票.xlsx"
    out = tmp_path / "out"
    _write_invoice(inv, [_ok_row()])
    result = convert.convert(
        invoice_path=inv,
        out_dir=out,
        booking_date="2026-08-14",
        master_path=None,
    )
    assert result.bookable_count == 1
    assert result.kingdee_path.is_file()
    kd = load_workbook(result.kingdee_path)
    ws = kd[convert.KINGDEE_SHEET]
    assert ws.cell(4, 19).value == "甲科技有限公司"
    assert str(ws.cell(4, 22).value) == "15"
    assert ws.cell(4, 25).value == "于占国"
    kd.close()


def test_result_filename_beside_invoice(tmp_path):
    inv = tmp_path / "发票明昊.xlsx"
    leftover = tmp_path / "数据模板_凭证引入引出模板_0813.xlsx"
    _write_invoice(inv, [_ok_row()])
    _copy_template(leftover)
    result = convert.convert(
        invoice_path=inv,
        out_dir=tmp_path,
        booking_date="2026-08-14",
    )
    assert result.kingdee_path == tmp_path / "凭证引入_结果.xlsx"
    assert result.detail_path == tmp_path / "发票明昊_明细结果.xlsx"
    assert leftover.is_file()
    assert result.kingdee_path.is_file()
    assert result.kingdee_path != leftover


def test_inspect_invoice_only(tmp_path):
    inv = tmp_path / "发票.xlsx"
    leftover = tmp_path / "数据模板.xlsx"
    _write_invoice(inv, [_ok_row()])
    _copy_template(leftover)
    found = convert.inspect_dir(tmp_path)
    assert Path(found["invoice"]).name == "发票.xlsx"
    assert found.get("template") in (None, "")
    assert found.get("master") in (None, "")


def test_cli_writes_result_in_same_folder(tmp_path):
    inv = tmp_path / "发票.xlsx"
    leftover = tmp_path / "数据模板.xlsx"
    _write_invoice(inv, [_ok_row()])
    _copy_template(leftover)
    rc = convert.main(["--input-dir", str(tmp_path), "--date", "2026-08-14", "--out-dir", str(tmp_path)])
    assert rc == 0
    assert (tmp_path / "凭证引入_结果.xlsx").is_file()
    assert leftover.is_file()
    assert not (tmp_path / "数据模板_结果.xlsx").is_file()


@pytest.mark.skipif(not GOLD_INVOICE.is_file(), reason="本机金标文件不在")
def test_gold_minghao_30(tmp_path):
    out = tmp_path / "gold"
    out.mkdir()
    master = GOLD_MASTER if GOLD_MASTER.is_file() else None
    result = convert.convert(
        invoice_path=GOLD_INVOICE,
        out_dir=out,
        booking_date="2026-08-14",
        master_path=master,
    )
    assert result.source_count == 30
    assert result.hold_count == 0
    assert result.bookable_count == 30
    kd = load_workbook(result.kingdee_path)
    ws = kd["sheet1（名称勿改）"]
    n_lines = 0
    debit = Decimal("0")
    credit = Decimal("0")
    for row in ws.iter_rows(min_row=4, max_col=17, values_only=True):
        if not row[6]:
            continue
        n_lines += 1
        debit += Decimal(str(row[15] or 0))
        credit += Decimal(str(row[16] or 0))
    assert n_lines == 90
    assert debit == credit
    kd.close()
