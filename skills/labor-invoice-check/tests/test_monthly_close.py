#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""月度劳务发票做账 · 合成用例（先红后绿）。真表不进仓。"""
import os
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(SKILL, "scripts"))
import monthly_close as mc  # noqa


def _pay_xlsx(path, rows, sheet="202608"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(["姓名", "身份证", "应发金额"])
    for r in rows:
        ws.append(list(r))
    wb.save(path)


def _inv_xlsx(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["销售方名称", "纳税人识别号", "合计金额"])
    for r in rows:
        ws.append(list(r))
    # 底稿 sheet，脚本必须忽略（真表里发票汇总也会夹做账底稿）
    ws2 = wb.create_sheet("Sheet2")
    ws2.append(["开户名", "应付金额", "发票"])
    ws2.append(["诱饵", 99999, 1])
    ws3 = wb.create_sheet("Sheet3")
    ws3.append(["日期", "姓名", "应发金额", "有票", "800以下不提供发票", "无票"])
    ws3.append(["8月", "诱饵完成版", 1, 1, None, None])
    wb.save(path)


def test_sum_two_invoices_and_columns():
    with tempfile.TemporaryDirectory() as td:
        pay = os.path.join(td, "pay.xlsx")
        inv = os.path.join(td, "inv.xlsx")
        out = os.path.join(td, "out.xlsx")
        _pay_xlsx(
            pay,
            [
                ("张多张", "110101199001011234", 1000),
                ("李八百", "110101199001011235", 500),
                ("李八百有票", "110101199001011236", 500),
                ("JOHN SMITH", "E04731572", 5000),
                ("MINI JANE", "C6624835", 600),
                ("王缺票", "110101199001011237", 2000),
                ("赵少开", "110101199001011238", 1000.6),
                ("钱多开", "110101199001011239", 1000),
            ],
        )
        _inv_xlsx(
            inv,
            [
                ("张多张", "110101199001011234", 400),
                ("张多张", "110101199001011234", 600),
                ("李八百有票", "110101199001011236", 500),
                ("赵少开", "110101199001011238", 1000),
                ("钱多开", "110101199001011239", 1200),
            ],
        )
        rc = mc.main(
            ["--pay", pay, "--invoice", inv, "--month", "202608", "--out", out]
        )
        assert rc == 0, "main should return 0"
        people = mc.load_pay(pay, "202608", mc.merge_alias(mc.PAY_ALIAS, {}))
        sums, sheet = mc.load_invoices(inv, mc.merge_alias(mc.INV_ALIAS, {}))
        assert sheet == "Sheet1"
        assert abs(sums["张多张"] - 1000) < 1e-6
        assert "诱饵" not in sums
        rows, makeup, special = mc.classify(people, sums, mc.DEFAULTS)
        by = {r["姓名"]: r for r in rows}
        assert by["张多张"]["有票"] == 1000
        assert by["张多张"]["无票"] is None
        assert by["李八百"]["800以下不提供发票"] == 500
        assert by["李八百"]["有票"] is None
        assert by["李八百有票"]["有票"] == 500
        assert by["李八百有票"]["800以下不提供发票"] is None
        assert by["JOHN SMITH"]["无票"] == 5000
        assert by["MINI JANE"]["800以下不提供发票"] == 600
        assert by["王缺票"]["待人工原因"]
        assert makeup == 1
        assert abs(by["赵少开"]["无票"] - 0.6) < 1e-6
        assert by["钱多开"]["有票"] == 1200
        assert by["钱多开"]["无票"] is None
        assert Path(out).is_file()


def test_special_same_amount():
    people = [
        {"name": f"甲{i}", "pay": 37400.0, "idno": f"11010119900101121{i}"}
        for i in range(3)
    ]
    people.append({"name": "乙单独", "pay": 9000.0, "idno": "110101199001011299"})
    rows, makeup, special = mc.classify(people, {}, mc.DEFAULTS)
    assert special == 3
    assert makeup == 1


def _gold_xlsx(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "一般劳务明细"
    ws.append(["日期", "姓名", "应发金额", "有票", "800以下不提供发票", "无票"])
    ws.append(["8月", "甲", 100, 100, None, None])
    ws2 = wb.create_sheet("汇总")
    ws2.append(["行标签", "应发金额", "有票", "800以下不提供发票", "无票"])
    ws2.append(["8月", 100, 100, 0, 0])
    wb.save(path)


def test_old_check_py_refuses():
    import subprocess

    rc = subprocess.call(
        [sys.executable, os.path.join(SKILL, "scripts", "check.py")],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    assert rc == 2


def test_inspect_sniff():
    with tempfile.TemporaryDirectory() as td:
        pay = os.path.join(td, "劳务费用明细.xlsx")
        inv = os.path.join(td, "202608个人发票汇总.xlsx")
        _pay_xlsx(pay, [("甲", "110101199001011234", 100)])
        _inv_xlsx(inv, [("甲", "x", 100)])
        rc = mc.main(["--inspect", "--input-dir", td])
        assert rc == 0


def test_messy_folder_skips_gold_and_other_month():
    """桌面四张表：完成版不是应发，7 月票不是 8 月。"""
    with tempfile.TemporaryDirectory() as td:
        pay = os.path.join(td, "劳务费用明细.xlsx")
        inv7 = os.path.join(td, "202607个人发票汇总.xlsx")
        inv8 = os.path.join(td, "202608个人发票汇总.xlsx")
        gold = os.path.join(td, "2026劳务发票统计-完成版.xlsx")
        _pay_xlsx(pay, [("甲", "110101199001011234", 900)])
        _inv_xlsx(inv7, [("甲", "x", 1)])
        _inv_xlsx(inv8, [("甲", "x", 900)])
        _gold_xlsx(gold)
        aliases_p = mc.merge_alias(mc.PAY_ALIAS, {})
        aliases_i = mc.merge_alias(mc.INV_ALIAS, {})
        assert mc.sniff_role(gold, aliases_p, aliases_i) == "gold"
        assert mc.sniff_role(pay, aliases_p, aliases_i) == "pay"
        assert mc.sniff_role(inv8, aliases_p, aliases_i) == "invoice"
        got_pay, got_inv, got_gold = mc.find_inputs(td, aliases_p, aliases_i, "202608")
        assert got_pay == pay
        assert got_inv == inv8
        assert got_gold == gold
        out = os.path.join(td, "out.xlsx")
        rc = mc.main(["--input-dir", td, "--month", "202608", "--out", out])
        assert rc == 0
        people = mc.load_pay(got_pay, "202608", aliases_p)
        sums, _ = mc.load_invoices(got_inv, aliases_i)
        rows, makeup, _ = mc.classify(people, sums, mc.DEFAULTS)
        assert rows[0]["有票"] == 900
        assert makeup == 0


def test_choose_month_prefers_last_closed_when_two_invoices():
    from datetime import date

    assert mc.choose_month({"202607"}, today=date(2026, 9, 10)) == "202607"
    assert mc.choose_month({"202607", "202608"}, today=date(2026, 9, 10)) == "202608"
    assert mc.choose_month({"202607", "202608"}, today=date(2026, 10, 1)) is None


def test_name_and_amount_format_noise():
    with tempfile.TemporaryDirectory() as td:
        pay = os.path.join(td, "pay.xlsx")
        inv = os.path.join(td, "inv.xlsx")
        _pay_xlsx(
            pay,
            [
                ("张 三", "110101199001011234", 1000),
                ("JOHN SMITH", "E04731572", 2000),
                ("李四", "110101199001011235", 900),
            ],
        )
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(["说明", "别用这行"])
        ws.append([None, None])
        ws.append(["销售方名称*", "纳税人识别号", "合计金额（元）"])
        ws.append(["张三", "x", "￥400.00"])
        ws.append(["张三", "x", "600"])
        ws.append(["John  Smith", "x", "2,000"])
        ws.append(["李四（个人）", "x", "900"])
        wb.save(inv)
        people = mc.load_pay(pay, "202608", mc.merge_alias(mc.PAY_ALIAS, {}))
        sums, sheet = mc.load_invoices(inv, mc.merge_alias(mc.INV_ALIAS, {}))
        assert sheet == "Sheet1"
        rows, makeup, _ = mc.classify(people, sums, mc.DEFAULTS)
        by = {r["姓名"]: r for r in rows}
        assert by["张 三"]["有票"] == 1000
        assert by["JOHN SMITH"]["有票"] == 2000
        assert by["李四"]["有票"] == 900
        assert makeup == 0


def test_scan_only_domestic_near_month():
    with tempfile.TemporaryDirectory() as td:
        old = Path(td) / "2018" / "国内个人发票"
        pub = Path(td) / "202607" / "对公发票"
        foreign = Path(td) / "202607" / "国外个人发票"
        personal = Path(td) / "202607" / "国内个人发票"
        extra = personal / "新增发票"
        for p in (old, pub, foreign, extra):
            p.mkdir(parents=True)
        (old / "发票-张三-2018-1000.pdf").write_bytes(b"%PDF")
        (pub / "发票-张三-202607-1000.pdf").write_bytes(b"%PDF")
        (foreign / "发票-张三-202607-1000.pdf").write_bytes(b"%PDF")
        (personal / "发票-张三-202607-1000.pdf").write_bytes(b"%PDF")
        (extra / "发票-李四-202606-900.pdf").write_bytes(b"%PDF")
        roots = mc.invoice_scan_roots(td, "202608")
        assert any(str(personal) == r or str(personal) in r for r in roots)
        assert not any("对公发票" in r or "国外个人" in r or "2018" in r for r in roots)
        people = [
            {"name": "张三", "pay": 1000.0, "idno": "110101199001011234"},
            {"name": "李四", "pay": 900.0, "idno": "110101199001011235"},
        ]
        found = mc.scan_invoice_dir(td, people, "202608", 0.02)
        assert mc.amount_for(found, "张三") == 1000
        assert mc.amount_for(found, "李四") == 900


if __name__ == "__main__":
    test_sum_two_invoices_and_columns()
    test_special_same_amount()
    test_old_check_py_refuses()
    test_inspect_sniff()
    test_messy_folder_skips_gold_and_other_month()
    test_choose_month_prefers_last_closed_when_two_invoices()
    test_name_and_amount_format_noise()
    test_scan_only_domestic_near_month()
    print("OK")
