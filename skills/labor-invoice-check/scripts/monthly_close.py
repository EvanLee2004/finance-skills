#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""月初劳务发票做账：应发明细 × 发票汇总 Sheet1 → 完成版三列。"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL_DIR = os.path.dirname(HERE)
CONFIG_DIR = os.path.join(SKILL_DIR, "config")

DEFAULTS = {
    "THRESHOLD": 800.0,
    "TOLERANCE": 0.02,
    "SPECIAL_MIN_AMOUNT": 10000.0,
    "SPECIAL_MIN_COUNT": 3,
}


def log(msg: str) -> None:
    print(msg, file=sys.stderr)


def kv(k: str, v) -> None:
    print(f"{k}={v}")


def fail(ask: str, code: int = 2) -> int:
    kv("status", "error")
    kv("ask", ask)
    return code


def load_config() -> dict:
    cfg = dict(DEFAULTS)
    p = os.path.join(CONFIG_DIR, "业务规则.md")
    if not os.path.isfile(p):
        return cfg
    text = Path(p).read_text(encoding="utf-8")
    pairs = {
        "开票门槛": "THRESHOLD",
        "容差": "TOLERANCE",
        "单独写单金额": "SPECIAL_MIN_AMOUNT",
        "单独写单重复人数": "SPECIAL_MIN_COUNT",
    }
    for line in text.splitlines():
        if "|" not in line:
            continue
        cells = [c.strip() for c in line.split("|")]
        if len(cells) < 3:
            continue
        key, val = cells[1], cells[2]
        if key in pairs:
            try:
                cfg[pairs[key]] = float(val)
            except ValueError:
                pass
    return cfg


def load_aliases() -> tuple[dict, dict]:
    p = os.path.join(CONFIG_DIR, "列名别名.json")
    if os.path.isfile(p):
        try:
            d = json.loads(Path(p).read_text(encoding="utf-8"))
            return d.get("应发明细_列别名", {}), d.get("发票汇总_列别名", {})
        except Exception as e:
            log(f"读列名别名失败({e})，用内置。")
    return {}, {}


PAY_ALIAS = {
    "姓名": ["姓名", "供应商姓名", "销售方名称", "开户名"],
    "应发金额": ["应发金额", "应发", "应付金额", "应付", "应发合计"],
    "身份证号": ["身份证", "身份证号", "身份证号/护照号", "护照号", "证件号"],
}
INV_ALIAS = {
    "销售方名称": ["销售方信息名称", "销售方信息-名称", "销售方名称", "开票人", "开票名称", "姓名"],
    "合计金额": ["合计金额（元）", "合计金额(元)", "合计金额", "价税合计（元）", "价税合计(元)", "价税合计", "开票金额", "发票金额"],
}
GOLD_ALIAS = {
    "姓名": ["姓名"],
    "有票": ["有票"],
    "无票": ["无票"],
    "800以下": ["800以下不提供发票", "800以下"],
}


def merge_alias(base: dict, extra: dict) -> dict:
    out = {k: list(v) for k, v in base.items()}
    for k, vs in extra.items():
        if k.startswith("_"):
            continue
        out.setdefault(k, [])
        for x in vs:
            if x not in out[k]:
                out[k].append(x)
    return out


def norm_name(v) -> str:
    if v is None:
        return ""
    s = unicodedata.normalize("NFKC", str(v)).strip()
    s = s.replace("\u3000", " ")
    s = re.sub(r"\s+", " ", s)
    return s


def norm_id(v) -> str:
    if v is None:
        return ""
    s = str(v).strip().upper()
    s = re.sub(r"\s+", "", s)
    if re.fullmatch(r"\d+\.0", s):
        s = s[:-2]
    return s


def to_number(v):
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        if v != v:  # NaN
            return None
        return float(v)
    s = unicodedata.normalize("NFKC", str(v)).strip()
    s = s.replace(",", "").replace("，", "")
    s = s.replace("¥", "").replace("￥", "").replace("元", "").strip()
    if s in ("", "-", "#N/A", "#N/A", "N/A", "NA", "nan", "None"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def strip_name_noise(name: str) -> str:
    s = norm_name(name)
    s = re.sub(r"[（(][^）)]*[）)]", "", s)
    return s.strip()


def name_key(name: str) -> str:
    """空格、全角、大小写、括号昵称、间隔号不算两个人。"""
    s = unicodedata.normalize("NFKC", strip_name_noise(name))
    s = re.sub(r"[\s\u3000·・.\-_'’]+", "", s)
    return s.casefold()


def amount_for(mapping: dict, name: str) -> float:
    if not mapping or not name:
        return 0.0
    kn = name_key(name)
    if kn in mapping:
        return float(mapping[kn] or 0)
    if name in mapping:
        return float(mapping[name] or 0)
    total = 0.0
    hit = False
    for k, v in mapping.items():
        if name_key(str(k)) == kn:
            total += float(v or 0)
            hit = True
    return total if hit else 0.0


def norm_header(h) -> str:
    s = unicodedata.normalize("NFKC", str(h) if h is not None else "").strip()
    s = s.replace("*", "").replace("：", "").replace(":", "")
    s = re.sub(r"\s+", "", s)
    return s


def has_cjk(text: str) -> bool:
    return bool(text) and bool(re.search(r"[一-鿿]", text))


def is_foreigner(name: str, idno: str) -> bool:
    cleaned = strip_name_noise(name)
    if cleaned and (not has_cjk(cleaned)) and re.search(r"[A-Za-z]", cleaned):
        return True
    s = norm_id(idno)
    if s and not re.fullmatch(r"\d{17}[\dX]", s):
        if re.fullmatch(r"[A-Z0-9]{5,15}", s) and re.search(r"[A-Z]", s):
            return True
    return False


def header_map(headers: list[str], aliases: dict) -> dict[str, int]:
    found = {}
    cells = [norm_header(h) for h in headers]
    for logical, names in aliases.items():
        want = {norm_header(n) for n in names}
        for i, h in enumerate(cells):
            if h and h in want:
                found[logical] = i
                break
    return found


def worksheet_rows(ws):
    return [[c.value for c in row] for row in ws.iter_rows()]


def pick_header_row(rows: list[list], aliases: dict) -> int | None:
    best_i, best_n = None, 0
    for i, row in enumerate(rows[:12]):
        m = header_map(row, aliases)
        if len(m) > best_n:
            best_i, best_n = i, len(m)
    if best_n >= 2:
        return best_i
    return None


def month_tokens(month: str) -> list[str]:
    m = re.fullmatch(r"(\d{4})(\d{2})", month)
    if not m:
        return [month]
    y, mo = m.group(1), m.group(2)
    n = int(mo)
    return [
        month,
        f"{y}-{mo}",
        f"{y}/{mo}",
        f"{y}年{n}月",
        f"{n}月",
        f"{y}{n}",
    ]


def month_from_sheet_name(name: str, month: str) -> bool:
    n = str(name or "").replace(" ", "")
    return any(tok in n for tok in month_tokens(month) if tok)


def date_is_month(val, month: str) -> bool:
    if val is None:
        return False
    if isinstance(val, datetime):
        return f"{val.year:04d}{val.month:02d}" == month
    s = str(val).strip()
    return any(tok in s.replace(" ", "") for tok in month_tokens(month) if len(tok) >= 2)


def load_pay(path: str, month: str, aliases: dict) -> list[dict]:
    wb = load_workbook(path, data_only=True)
    try:
        chosen = None
        for name in wb.sheetnames:
            if month_from_sheet_name(name, month):
                chosen = name
                break
        if chosen is None:
            for name in wb.sheetnames:
                ws = wb[name]
                rows = worksheet_rows(ws)
                hi = pick_header_row(rows, aliases)
                if hi is not None and "姓名" in header_map(rows[hi], aliases) and "应发金额" in header_map(rows[hi], aliases):
                    chosen = name
                    break
        if chosen is None:
            raise ValueError("应发明细里找不到当月 sheet（姓名+应发）")
        ws = wb[chosen]
        rows = worksheet_rows(ws)
        hi = pick_header_row(rows, aliases)
        if hi is None:
            raise ValueError(f"{chosen} 认不出姓名/应发列")
        hmap = header_map(rows[hi], aliases)
        people = []
        seen = set()
        for row in rows[hi + 1 :]:
            name = norm_name(row[hmap["姓名"]]) if "姓名" in hmap else ""
            if not name or name in ("合计", "总计", "小计"):
                continue
            pay = to_number(row[hmap["应发金额"]]) if "应发金额" in hmap else None
            if pay is None:
                continue
            idno = ""
            if "身份证号" in hmap:
                idno = norm_id(row[hmap["身份证号"]])
            key = name_key(name) or name.casefold()
            if key in seen:
                for p in people:
                    if (name_key(p["name"]) or p["name"].casefold()) == key:
                        p["pay"] = round(p["pay"] + pay, 2)
                        break
                continue
            seen.add(key)
            people.append({"name": name, "pay": float(pay), "idno": idno, "sheet": chosen})
        return people
    finally:
        wb.close()


def load_invoices(path: str, aliases: dict) -> tuple[dict[str, float], str]:
    wb = load_workbook(path, data_only=True)
    try:
        chosen = None
        for name in wb.sheetnames:
            ws = wb[name]
            rows = worksheet_rows(ws)
            hi = pick_header_row(rows, aliases)
            if hi is None:
                continue
            hmap = header_map(rows[hi], aliases)
            if "销售方名称" in hmap and "合计金额" in hmap:
                chosen = name
                break
        if chosen is None:
            raise ValueError("发票文件里没有同时含销售方名称+合计金额的 sheet")
        ws = wb[chosen]
        rows = worksheet_rows(ws)
        hi = pick_header_row(rows, aliases)
        hmap = header_map(rows[hi], aliases)
        sums: dict[str, float] = defaultdict(float)
        for row in rows[hi + 1 :]:
            seller = norm_name(row[hmap["销售方名称"]])
            amt = to_number(row[hmap["合计金额"]])
            if not seller or amt is None:
                continue
            sums[name_key(seller) or seller] += float(amt)
        return {k: round(v, 2) for k, v in sums.items()}, chosen
    finally:
        wb.close()


def adjacent_yyyymm(month: str) -> list[str]:
    y, mo = int(month[:4]), int(month[4:6])
    triples = [(y, mo)]
    if mo == 1:
        triples.append((y - 1, 12))
    else:
        triples.append((y, mo - 1))
    if mo == 12:
        triples.append((y + 1, 1))
    else:
        triples.append((y, mo + 1))
    return [f"{yy}{mm:02d}" for yy, mm in triples]


def invoice_scan_roots(root: str, month: str) -> list[str]:
    """译员对账清单按月分夹。只扫近月「国内个人发票」「新增发票」，不扫对公/国外/老年度。"""
    if not root or not os.path.isdir(root):
        return []
    month_dirs = [
        name
        for name in os.listdir(root)
        if re.fullmatch(r"20\d{4}", name) and os.path.isdir(os.path.join(root, name))
    ]
    if not month_dirs:
        return [root]
    roots: list[str] = []
    want = set(adjacent_yyyymm(month))
    for name in sorted(month_dirs):
        if name not in want:
            continue
        base = os.path.join(root, name)
        personal = os.path.join(base, "国内个人发票")
        if os.path.isdir(personal):
            roots.append(personal)
            for extra_name in ("新增发票", "新增发票2"):
                extra = os.path.join(personal, extra_name)
                if os.path.isdir(extra):
                    roots.append(extra)
        top_extra = os.path.join(base, "新增发票")
        if os.path.isdir(top_extra) and top_extra not in roots:
            roots.append(top_extra)
        if not os.path.isdir(personal):
            roots.append(base)
    return roots or [root]


def scan_invoice_dir(root: str, people: list[dict], month: str, tol: float) -> dict[str, float]:
    """按文件名里的姓名+金额补票。日期看文件 mtime 是否近月。"""
    scan_from = invoice_scan_roots(root, month)
    if not scan_from:
        return {}
    ok_months = set()
    for yyyymm in adjacent_yyyymm(month):
        ok_months.add((int(yyyymm[:4]), int(yyyymm[4:6])))

    by_name = {p["name"]: p["pay"] for p in people}
    found: dict[str, float] = defaultdict(float)
    skip_bits = ("对公发票", "国外个人发票", "个人转对公")
    for start in scan_from:
        for dirpath, dirnames, files in os.walk(start):
            dirnames[:] = [d for d in dirnames if d not in skip_bits]
            if any(bit in dirpath.replace("\\", "/") for bit in skip_bits):
                continue
            for fn in files:
                if not re.search(r"\.(pdf|jpg|jpeg|png|webp)$", fn, re.I):
                    continue
                path = os.path.join(dirpath, fn)
                try:
                    mt = datetime.fromtimestamp(os.path.getmtime(path))
                except OSError:
                    continue
                if (mt.year, mt.month) not in ok_months:
                    continue
                stem = Path(fn).stem
                stem_k = name_key(stem)
                matched_people = []
                for name, pay in by_name.items():
                    kn = name_key(name)
                    if name and (name in stem or (kn and kn in stem_k)):
                        nums = [float(x) for x in re.findall(r"(\d+(?:\.\d+)?)", stem)]
                        hits = [n for n in nums if abs(n - pay) <= max(tol, 0.05)]
                        if hits:
                            matched_people.append((name, max(hits)))
                if len(matched_people) != 1:
                    continue
                name, amt = matched_people[0]
                key = name_key(name) or name
                found[key] = max(found[key], amt)
    return {k: round(v, 2) for k, v in found.items()}


def classify(people, inv_sums, cfg, extras=None):
    extras = extras or {}
    th = cfg["THRESHOLD"]
    tol = cfg["TOLERANCE"]
    special_amt = cfg["SPECIAL_MIN_AMOUNT"]
    special_n = int(cfg["SPECIAL_MIN_COUNT"])

    missing_dom = []
    for p in people:
        name = p["name"]
        inv = amount_for(inv_sums, name) + amount_for(extras, name)
        p["invoice"] = round(inv, 2)
        p["foreign"] = is_foreigner(name, p.get("idno", ""))
        if p["invoice"] <= 0 and p["pay"] > th and not p["foreign"]:
            missing_dom.append(p)

    amt_cnt = Counter(round(p["pay"], 2) for p in missing_dom)
    special_amts = {
        a for a, n in amt_cnt.items() if n >= special_n and a >= special_amt
    }

    rows = []
    ask_makeup = 0
    ask_special = 0
    for p in people:
        inv = p["invoice"]
        pay = p["pay"]
        has = under = none = None
        reason = ""
        how = ""
        if inv > 0:
            has = inv
            from_inv = amount_for(inv_sums, p["name"]) > 0
            from_dir = amount_for(extras, p["name"]) > 0
            how = "当月汇总加总" if from_inv else "发票目录文件名"
            if from_inv and from_dir:
                how = "当月汇总+目录"
            gap = round(pay - inv, 2)
            if gap > tol:
                none = gap
            # 多开或对齐到容差内：无票空
            # 多开或对齐：无票空
        elif pay <= th:
            under = pay
            how = "≤800无票"
        elif p["foreign"]:
            none = pay
            how = "外籍无票"
        else:
            how = "待人工"
            if round(pay, 2) in special_amts:
                reason = "单独写单（金额大且多人相同，批量目录不应猜有票）"
                ask_special += 1
            else:
                reason = "当月汇总无此销售方，需补票或人核"
                ask_makeup += 1
        rows.append(
            {
                "姓名": p["name"],
                "应发金额": round(pay, 2),
                "有票": has,
                "800以下不提供发票": under,
                "无票": none,
                "任务线": None,
                "匹配方式": how,
                "待人工原因": reason or None,
            }
        )
    return rows, ask_makeup, ask_special


def write_out(path: str, month: str, rows: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "一般劳务明细"
    headers = [
        "日期",
        "姓名",
        "应发金额",
        "有票",
        "800以下不提供发票",
        "无票",
        "任务线",
        "匹配方式",
        "待人工原因",
    ]
    ws.append(headers)
    label = f"{int(month[4:6])}月"
    for r in rows:
        ws.append(
            [
                label,
                r["姓名"],
                r["应发金额"],
                r["有票"],
                r["800以下不提供发票"],
                r["无票"],
                r["任务线"],
                r["匹配方式"],
                r["待人工原因"],
            ]
        )

    ask_rows = [r for r in rows if r.get("待人工原因")]
    ws2 = wb.create_sheet("待人工")
    ws2.append(["姓名", "应发金额", "待人工原因"])
    for r in ask_rows:
        ws2.append([r["姓名"], r["应发金额"], r["待人工原因"]])

    def _sum(key):
        return round(sum((r[key] or 0) for r in rows), 2)

    ws3 = wb.create_sheet("汇总_本月")
    ws3.append(["行标签", "应发金额", "有票", "800以下不提供发票", "无票"])
    ws3.append([label, _sum("应发金额"), _sum("有票"), _sum("800以下不提供发票"), _sum("无票")])
    ws3.append([])
    ws3.append(["公司一般的要求：800以下可以不提供发票直接支付"])
    ws3.append(["实际情况：800以下也有人会提供发票，按有票来统计"])
    ws3.append(["外国人无论是否超过800，均无票（指不开票；≤800仍进800以下列）"])
    ws3.append(["有人会多开发票，有人会少开发票"])

    ws4 = wb.create_sheet("运行报告")
    ws4.append(["字段", "值"])
    ws4.append(["月份", month])
    ws4.append(["人数", len(rows)])
    ws4.append(["有票人数", sum(1 for r in rows if r["有票"])])
    ws4.append(["800以下人数", sum(1 for r in rows if r["800以下不提供发票"])])
    ws4.append(["无票人数", sum(1 for r in rows if r["无票"])])
    ws4.append(["待人工人数", len(ask_rows)])
    wb.save(path)


def gold_counts(path: str, month: str) -> dict:
    wb = load_workbook(path, data_only=True)
    try:
        if "一般劳务明细" not in wb.sheetnames:
            return {}
        ws = wb["一般劳务明细"]
        rows = worksheet_rows(ws)
        if not rows:
            return {}
        headers = [str(c).strip() if c is not None else "" for c in rows[0]]
        def col(*names):
            for n in names:
                if n in headers:
                    return headers.index(n)
            return None
        i_date = col("日期")
        i_has = col("有票")
        i_under = col("800以下不提供发票", "800以下")
        i_none = col("无票")
        n = n_has = n_under = n_none = 0
        for row in rows[1:]:
            if i_date is not None and not date_is_month(row[i_date], month):
                continue
            n += 1
            if i_has is not None and to_number(row[i_has]):
                n_has += 1
            if i_under is not None and to_number(row[i_under]):
                n_under += 1
            if i_none is not None and to_number(row[i_none]):
                n_none += 1
        return {"gold_n": n, "gold_has": n_has, "gold_under": n_under, "gold_none": n_none}
    finally:
        wb.close()


def last_closed_month(today: date | None = None) -> str:
    d = today or date.today()
    if d.month == 1:
        return f"{d.year - 1}12"
    return f"{d.year}{d.month - 1:02d}"


def filename_has_month(path: str, month: str) -> bool:
    fn = os.path.basename(path)
    y, mo = month[:4], month[4:]
    n = int(mo)
    tokens = [month, f"{y}-{mo}", f"{y}/{mo}", f"{y}_{mo}", f"{y}年{n}月", f"{y}年{mo}月"]
    return any(t in fn for t in tokens)


def month_from_filename(path: str) -> str | None:
    m = re.search(r"(20\d{2})[-_/]?([01]\d)", os.path.basename(path))
    if not m:
        return None
    y, mo = m.group(1), m.group(2)
    if 1 <= int(mo) <= 12:
        return f"{y}{mo}"
    return None


def choose_month(candidates: set[str], today: date | None = None) -> str | None:
    cand = {c for c in candidates if c and re.fullmatch(r"\d{6}", c)}
    if len(cand) == 1:
        return next(iter(cand))
    last = last_closed_month(today)
    if last in cand:
        return last
    return None


def is_gold_header(headers: list) -> bool:
    h = header_map(headers, GOLD_ALIAS)
    return "姓名" in h and "有票" in h and ("无票" in h or "800以下" in h)


def sniff_role(path: str, pay_alias, inv_alias) -> str:
    try:
        wb = load_workbook(path, data_only=True)
    except Exception:
        return "unknown"
    try:
        saw_invoice = saw_pay = saw_gold = False
        for name in wb.sheetnames:
            rows = worksheet_rows(wb[name])
            if not rows:
                continue
            for row in rows[:12]:
                if is_gold_header(row):
                    saw_gold = True
            hi_inv = pick_header_row(rows, inv_alias)
            if hi_inv is not None:
                h = header_map(rows[hi_inv], inv_alias)
                if "销售方名称" in h and "合计金额" in h and not is_gold_header(rows[hi_inv]):
                    saw_invoice = True
            hi_pay = pick_header_row(rows, pay_alias)
            if hi_pay is not None:
                h = header_map(rows[hi_pay], pay_alias)
                if "姓名" in h and "应发金额" in h:
                    if is_gold_header(rows[hi_pay]):
                        saw_gold = True
                    else:
                        saw_pay = True
        # 发票汇总里常夹着她的做账底稿（有票/无票 sheet），整本仍是发票。
        if saw_invoice:
            return "invoice"
        if saw_gold:
            return "gold"
        if saw_pay:
            return "pay"
        return "unknown"
    finally:
        wb.close()


def list_role_files(input_dir: str, pay_alias, inv_alias) -> dict[str, list[str]]:
    out = {"pay": [], "invoice": [], "gold": []}
    for fn in sorted(os.listdir(input_dir)):
        if not fn.lower().endswith((".xlsx", ".xlsm")) or fn.startswith("~$"):
            continue
        path = os.path.join(input_dir, fn)
        role = sniff_role(path, pay_alias, inv_alias)
        if role in out:
            out[role].append(path)
    return out


def pick_one(paths: list[str], month: str | None) -> str | None:
    if not paths:
        return None
    if month:
        hits = [p for p in paths if filename_has_month(p, month)]
        if len(hits) == 1:
            return hits[0]
        if len(hits) > 1:
            return None
        if len(paths) == 1:
            return paths[0]
        return None
    if len(paths) == 1:
        return paths[0]
    return None


def find_inputs(
    input_dir: str, pay_alias, inv_alias, month: str | None = None
) -> tuple[str | None, str | None, str | None]:
    roles = list_role_files(input_dir, pay_alias, inv_alias)
    gold = roles["gold"][0] if len(roles["gold"]) == 1 else None
    return pick_one(roles["pay"], month), pick_one(roles["invoice"], month), gold


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="月初劳务发票做账：应发 × 发票汇总 → 完成版三列")
    ap.add_argument("--pay", help="劳务费用明细.xlsx")
    ap.add_argument("--invoice", help="个人发票汇总.xlsx")
    ap.add_argument("--month", help="月份 YYYYMM，如 202608")
    ap.add_argument("--out", help="输出完成版 xlsx")
    ap.add_argument("--gold", help="她的完成版，只比对人数（可选）")
    ap.add_argument("--input-dir", help="放两张表的目录，配合 --inspect 或省略路径时认文件")
    ap.add_argument("--inspect", action="store_true", help="只认文件，不写结果")
    args = ap.parse_args(argv)

    cfg = load_config()
    extra_pay, extra_inv = load_aliases()
    pay_alias = merge_alias(PAY_ALIAS, extra_pay)
    inv_alias = merge_alias(INV_ALIAS, extra_inv)

    pay_path, inv_path, gold_path = args.pay, args.invoice, args.gold
    month = args.month if args.month and re.fullmatch(r"\d{6}", args.month) else None
    month_source = "flag" if month else ""
    roles = {"pay": [], "invoice": [], "gold": []}
    if args.input_dir:
        if not os.path.isdir(args.input_dir):
            return fail("ask=找不到 input-dir，把应发明细和发票汇总放进一个文件夹再跑")
        roles = list_role_files(args.input_dir, pay_alias, inv_alias)
        if not month:
            guessed = choose_month({month_from_filename(p) for p in roles["invoice"] if month_from_filename(p)})
            if guessed:
                month = guessed
                month_source = "last_closed" if len(roles["invoice"]) > 1 else "filename"
        a, b, c = find_inputs(args.input_dir, pay_alias, inv_alias, month)
        pay_path = pay_path or a
        inv_path = inv_path or b
        gold_path = gold_path or c

    if args.inspect:
        ok = bool(pay_path and inv_path)
        kv("status", "ok" if ok else "incomplete")
        kv("pay_file", pay_path or "")
        kv("invoice_file", inv_path or "")
        kv("gold_file", gold_path or "")
        kv("month", month or "")
        kv("month_source", month_source)
        kv("candidates_pay", len(roles["pay"]))
        kv("candidates_invoice", len(roles["invoice"]))
        kv("candidates_gold", len(roles["gold"]))
        if not ok:
            kv("ask", "文件夹里有完成版或上个月发票时，告诉我 YYYYMM，或指出哪张是应发、哪张是当月发票。完成版只能对人数，不能当应发")
            return 2
        kv("ask", "")
        return 0

    if not pay_path or not os.path.isfile(pay_path):
        return fail("ask=缺劳务费用明细（应发已付那张）。完成版不能当应发")
    if not inv_path or not os.path.isfile(inv_path):
        return fail("ask=缺个人发票汇总（资源部当月台账，不要拿上个月）")
    if not month:
        return fail("ask=缺月份，告诉我 YYYYMM，例如 202608。说「做这个月」就按上一个已过完的公历月")

    try:
        people = load_pay(pay_path, month, pay_alias)
        inv_sums, inv_sheet = load_invoices(inv_path, inv_alias)
    except ValueError as e:
        return fail(f"ask={e}")

    rows, ask_makeup, ask_special = classify(people, inv_sums, cfg)
    matched = sum(1 for r in rows if r["有票"])

    out = args.out
    if not out:
        out = os.path.join(os.path.dirname(os.path.abspath(pay_path)), f"劳务发票统计_{month}.xlsx")
    os.makedirs(os.path.dirname(os.path.abspath(out)) or ".", exist_ok=True)
    write_out(out, month, rows)

    kv("status", "ok")
    kv("month", month)
    kv("month_source", month_source or "flag")
    kv("source_pay", len(people))
    kv("invoice_sheet", inv_sheet)
    kv("matched", matched)
    kv("under800", sum(1 for r in rows if r["800以下不提供发票"]))
    kv("none", sum(1 for r in rows if r["无票"]))
    kv("ask_makeup", ask_makeup)
    kv("ask_special", ask_special)
    kv("out", out)
    if gold_path and os.path.isfile(gold_path):
        g = gold_counts(gold_path, month)
        for k, v in g.items():
            kv(k, v)
    ask_bits = []
    if ask_makeup:
        ask_bits.append(f"有{ask_makeup}人当月汇总没有票，请打开「待人工」核对，不要写成有票")
    if ask_special:
        ask_bits.append(f"有{ask_special}人像单独写单，请人核，不要当成普通有票")
    kv("ask", "；".join(ask_bits))
    return 0


if __name__ == "__main__":
    sys.exit(main())
