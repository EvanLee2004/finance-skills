#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from __future__ import annotations

import os
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
SCRIPTS = SKILL / "scripts"
CONVERT = SCRIPTS / "convert.py"
TEMPLATE = SKILL / "config" / "凭证引入空模.xlsx"
KINGDEE_SHEET = "sheet1（名称勿改）"

HUNAN8_SRC = Path(
    "/Users/evanlee/Documents/甲骨易实习/项目/长期项目/自动化记账（金蝶）/原始素材/实操批次/20260902_湖南分公司_8期凭证入账/原始素材/凭证导入模板.xlsx"
)
HUNAN8_GOLD = Path(
    "/Users/evanlee/Documents/甲骨易实习/项目/长期项目/自动化记账（金蝶）/原始素材/实操批次/20260902_湖南分公司_8期凭证入账/待上传/当前上传_湖南分公司_记1至9_2026年8期.xlsx"
)


def _run(args: list[str], env: dict | None = None) -> subprocess.CompletedProcess:
    run_env = os.environ.copy()
    if env:
        run_env.update(env)
    return subprocess.run(
        [sys.executable, str(CONVERT), *args],
        capture_output=True,
        text=True,
    )


def _write_import_template(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "凭证导入模板"
    ws.append(
        [
            "日期",
            "凭证字",
            "凭证号",
            "附件数",
            "分录序号",
            "摘要",
            "科目代码",
            "科目名称",
            "借方金额",
            "贷方金额",
            "客户",
            "供应商",
            "职员",
            "项目",
            "部门",
            "存货",
            "是否限定",
            "自定义辅助核算类别",
            "自定义辅助核算编码",
            "数量",
            "单价",
            "原币金额",
            "币别",
            "汇率",
        ]
    )
    rows = [
        ["2026-08-31", "记", 1, 1, 1, "收到款项", "100201", "银行存款", 1000, None, "", "", "", "", "", "", "", "", "", "", "", 1000, "RMB", 1],
        ["2026-08-31", "记", 1, 1, 2, "收到款项", "224103", "其他应付款", None, 1000, "", "", "", "", "", "", "", "", "", "", "", 1000, "RMB", 1],
        ["2026-08-31", "记", 2, "", 1, "计提工资", "560201", "管理费用_工资", 300, None, "", "", "", "", "", "", "", "", "", "", "", 300, "RMB", 1],
        ["2026-08-31", "记", 2, "", 2, "计提工资", "560201", "管理费用_工资", 50, None, "", "", "", "", "", "", "", "", "", "", "", 50, "RMB", 1],
        ["2026-08-31", "记", 2, "", 3, "计提工资", "221101", "应付职工薪酬", None, 350, "", "", "", "", "", "", "", "", "", "", "", 350, "RMB", 1],
        ["2026-08-31", "记", 3, 1, 1, "退回工资", "100201", "银行存款", 80, None, "", "", "", "", "", "", "", "", "", "", "", 80, "RMB", 1],
        ["2026-08-31", "记", 3, 1, 2, "退回工资", "221101", "应付职工薪酬", -80, None, "", "", "", "", "", "", "", "", "", "", "", -80, "RMB", 1],
    ]
    for row in rows:
        ws.append(row)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()


def _write_voucher_list(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "凭证列表#2026年第7期"
    ws["A1"] = "凭证列表"
    ws["A2"] = "公司名称：甲骨易（北京）语言科技股份有限公司湖南分公司"
    ws["D2"] = "期间：2026年第7期"
    ws.append([])
    ws["A3"] = "日期"
    ws["B3"] = "凭证字号"
    ws["C3"] = "附单据"
    ws["D3"] = "摘要"
    ws["E3"] = "科目"
    ws["F3"] = "借方金额"
    ws["G3"] = "贷方金额"
    ws["H3"] = "制单人"
    data = [
        ["2026-07-31", "记-1", "1", "办公费", "560203 管理费用_办公费", 144, None, "欧莉"],
        [None, None, None, "办公费", "220201 应付账款_某公司", None, 144, None],
        ["2026-07-31", "记-2", "1", "收到款项", "100201 银行存款_中国银行", 500, None, "欧莉"],
        [None, None, None, "收到款项", "224103 其他应付款_总部", None, 500, None],
    ]
    for i, row in enumerate(data, start=4):
        for c, val in enumerate(row, start=1):
            ws.cell(i, c).value = val
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()


def _attach_col(ws):
    for cell in ws[3]:
        if "附件" in str(cell.value or ""):
            return cell.column
    raise KeyError("附件")


def _col(ws, needle):
    for cell in ws[3]:
        if needle in str(cell.value or ""):
            return cell.column
    raise KeyError(needle)


def test_import_template_drops_attachment_keeps_red_debit_and_two_vouchers(tmp_path: Path):
    src = tmp_path / "src.xlsx"
    _write_import_template(src)
    out_dir = tmp_path / "out"
    r = _run(["--input", str(src), "--out-dir", str(out_dir)])
    assert r.returncode == 0, r.stdout + r.stderr
    kingdee = out_dir / "凭证引入_结果.xlsx"
    assert kingdee.is_file()
    wb = load_workbook(kingdee)
    assert KINGDEE_SHEET in wb.sheetnames
    assert "模板说明" in wb.sheetnames
    ws = wb[KINGDEE_SHEET]
    attach = _attach_col(ws)
    debit = _col(ws, "借方 #")
    credit = _col(ws, "贷方 #")
    number = _col(ws, "凭证号 #")
    account = _col(ws, "*科目.编码")
    nums = []
    red = None
    for row in range(4, (ws.max_row or 3) + 1):
        if not ws.cell(row, account).value:
            continue
        assert ws.cell(row, attach).value in (None, "")
        nums.append(ws.cell(row, number).value)
        if ws.cell(row, account).value == "221101" and ws.cell(row, debit).value not in (None, ""):
            red = ws.cell(row, debit).value
            assert ws.cell(row, credit).value in (None, "")
    wb.close()
    assert nums[:2] == [1, 1]
    assert 2 in nums and 3 in nums
    assert red == -80
    report = (out_dir / "运行报告.txt").read_text(encoding="utf-8")
    assert "附件列=空" in report
    assert "3 张" in r.stdout or "凭证张数=3" in report


def test_start_voucher_no_remaps_in_order(tmp_path: Path):
    src = tmp_path / "src.xlsx"
    _write_import_template(src)
    out_dir = tmp_path / "out11"
    r = _run(["--input", str(src), "--out-dir", str(out_dir), "--start-voucher-no", "11"])
    assert r.returncode == 0, r.stdout + r.stderr
    wb = load_workbook(out_dir / "凭证引入_结果.xlsx")
    ws = wb[KINGDEE_SHEET]
    number = _col(ws, "凭证号 #")
    account = _col(ws, "*科目.编码")
    nums = []
    for row in range(4, (ws.max_row or 3) + 1):
        if ws.cell(row, account).value:
            nums.append(ws.cell(row, number).value)
    wb.close()
    assert set(nums) == {11, 12, 13}
    assert nums[0] == 11
    assert "11～13" in r.stdout or "11～13" in (out_dir / "运行报告.txt").read_text(encoding="utf-8")


def test_voucher_list_format_and_company(tmp_path: Path):
    src = tmp_path / "list.xlsx"
    _write_voucher_list(src)
    out_dir = tmp_path / "out7"
    r = _run(["--input", str(src), "--out-dir", str(out_dir)])
    assert r.returncode == 0, r.stdout + r.stderr
    report = (out_dir / "运行报告.txt").read_text(encoding="utf-8")
    assert "湖南分公司" in report
    wb = load_workbook(out_dir / "凭证引入_结果.xlsx")
    ws = wb[KINGDEE_SHEET]
    number = _col(ws, "凭证号 #")
    account = _col(ws, "*科目.编码")
    attach = _attach_col(ws)
    codes = []
    for row in range(4, (ws.max_row or 3) + 1):
        if ws.cell(row, account).value:
            assert ws.cell(row, attach).value in (None, "")
            codes.append((ws.cell(row, number).value, str(ws.cell(row, account).value)))
    wb.close()
    assert codes[0] == (1, "560203")
    assert (2, "100201") in codes


def test_does_not_merge_different_voucher_numbers(tmp_path: Path):
    src = tmp_path / "src.xlsx"
    _write_import_template(src)
    out_dir = tmp_path / "out"
    _run(["--input", str(src), "--out-dir", str(out_dir)])
    wb = load_workbook(out_dir / "凭证引入_结果.xlsx")
    ws = wb[KINGDEE_SHEET]
    number = _col(ws, "凭证号 #")
    account = _col(ws, "*科目.编码")
    pairs = []
    for row in range(4, (ws.max_row or 3) + 1):
        if ws.cell(row, account).value:
            pairs.append(ws.cell(row, number).value)
    wb.close()
    assert pairs.count(1) == 2
    assert pairs.count(2) == 3
    assert pairs.count(3) == 2


def test_default_out_is_desktop(tmp_path: Path):
    sys.path.insert(0, str(SCRIPTS))
    import convert as gl

    desktop = tmp_path / "Desktop"
    desktop.mkdir()
    from datetime import date

    got = gl.default_desktop_dir("序时账入金蝶", today=date(2026, 9, 2), home=tmp_path)
    assert got == desktop / "序时账入金蝶_20260902"
    assert got.is_dir()


def test_official_template_is_rejected(tmp_path: Path):
    dest = tmp_path / "already.xlsx"
    dest.write_bytes(TEMPLATE.read_bytes())
    r = _run(["--input", str(dest), "--out-dir", str(tmp_path / "x")])
    assert r.returncode != 0
    assert "官方" in r.stdout or "引入表" in r.stdout


def test_hunan8_local_matches_gold_structure():
    if not HUNAN8_SRC.is_file() or not HUNAN8_GOLD.is_file():
        return
    import tempfile

    with tempfile.TemporaryDirectory() as td:
        out_dir = Path(td)
        r = _run(["--input", str(HUNAN8_SRC), "--out-dir", str(out_dir)])
        assert r.returncode == 0, r.stdout + r.stderr
        got = load_workbook(out_dir / "凭证引入_结果.xlsx")
        gold = load_workbook(HUNAN8_GOLD)
        gws = got[KINGDEE_SHEET]
        gold_ws = gold[KINGDEE_SHEET]
        attach = _attach_col(gws)
        number = _col(gws, "凭证号 #")
        account = _col(gws, "*科目.编码")
        debit = _col(gws, "借方 #")
        credit = _col(gws, "贷方 #")
        g_attach = _attach_col(gold_ws)
        g_number = _col(gold_ws, "凭证号 #")
        g_account = _col(gold_ws, "*科目.编码")
        g_debit = _col(gold_ws, "借方 #")
        g_credit = _col(gold_ws, "贷方 #")

        def rows(ws, a, n, d, c, att):
            out = []
            for row in range(4, (ws.max_row or 3) + 1):
                if not ws.cell(row, a).value:
                    continue
                out.append(
                    (
                        ws.cell(row, n).value,
                        str(ws.cell(row, a).value),
                        ws.cell(row, d).value,
                        ws.cell(row, c).value,
                        ws.cell(row, att).value in (None, ""),
                    )
                )
            return out

        got_rows = rows(gws, account, number, debit, credit, attach)
        gold_rows = rows(gold_ws, g_account, g_number, g_debit, g_credit, g_attach)
        assert [x[:2] for x in got_rows] == [x[:2] for x in gold_rows]
        assert all(x[4] for x in got_rows)
        assert any(x[2] is not None and float(x[2]) < 0 for x in got_rows)
        got.close()
        gold.close()
