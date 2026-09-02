#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""序时账 / 凭证列表 → 金蝶官方凭证引入表。金额只由本脚本抄，附件列强制空。"""
from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import sys
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
CONFIG = SKILL / "config"
TEMPLATE = CONFIG / "凭证引入空模.xlsx"
KINGDEE_SHEET = "sheet1（名称勿改）"
ALIASES_PATH = CONFIG / "列名别名.json"
MONEY_Q = Decimal("0.01")


def repo_venv_python() -> Path | None:
    root = Path(__file__).resolve().parents[3]
    for rel in (Path(".venv") / "bin" / "python", Path(".venv") / "Scripts" / "python.exe"):
        cand = root / rel
        if cand.is_file():
            return cand
    return None


def _reexec_repo_venv_if_needed() -> None:
    if not sys.argv or Path(sys.argv[0]).name.lower() not in {"convert.py", "convert"}:
        return
    venv_py = repo_venv_python()
    if venv_py is None:
        return
    try:
        if Path(sys.prefix).resolve() == venv_py.parent.parent.resolve():
            return
    except OSError:
        return
    os.execv(str(venv_py), [str(venv_py), *sys.argv])


_reexec_repo_venv_if_needed()

from openpyxl import load_workbook


def load_aliases() -> dict:
    return json.loads(ALIASES_PATH.read_text(encoding="utf-8"))


def default_desktop_dir(prefix: str, today: date | None = None, home: Path | None = None) -> Path:
    day = (today or date.today()).strftime("%Y%m%d")
    root = Path(home) if home else Path.home()
    desktop = root / "Desktop"
    base = desktop if desktop.is_dir() else Path.cwd()
    path = base / f"{prefix}_{day}"
    path.mkdir(parents=True, exist_ok=True)
    return path


def money(value) -> Decimal | None:
    if value is None or str(value).strip() in {"", "-", "—", "None"}:
        return None
    try:
        return Decimal(str(value).replace(",", "")).quantize(MONEY_Q, rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError):
        return None


def as_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).replace("\xa0", " ").strip()
    if text.endswith(".0") and text[:-2].replace("-", "").isdigit():
        text = text[:-2]
    return text


def as_int(value) -> int | None:
    text = as_text(value)
    if not text:
        return None
    m = re.search(r"(\d+)", text)
    if not m:
        return None
    return int(m.group(1))


def parse_word_no(word_cell, number_cell=None) -> tuple[str, int | None]:
    raw = as_text(word_cell)
    number = as_int(number_cell) if number_cell is not None else None
    if raw:
        m = re.match(r"^([记收付转])\s*[-—–]?\s*(\d+)$", raw)
        if m:
            return m.group(1), int(m.group(2))
        if re.fullmatch(r"[记收付转]", raw):
            return raw, number
        m2 = re.search(r"(\d+)", raw)
        if m2 and number is None:
            word = re.sub(r"[\d\s\-—–]+$", "", raw) or "记"
            return word, int(m2.group(1))
    return (raw or "记"), number


def split_account(value) -> tuple[str, str]:
    text = as_text(value)
    if not text:
        return "", ""
    m = re.match(r"^(\d{4,})\s*(.*)$", text)
    if m:
        return m.group(1), m.group(2).strip()
    return text, ""


def header_index(headers: list, aliases: dict) -> dict[str, int]:
    norm = [as_text(h) for h in headers]
    found: dict[str, int] = {}
    for key, names in aliases.items():
        if key.startswith("_"):
            continue
        for name in names:
            if name in norm and key not in found:
                found[key] = norm.index(name)
    return found


def looks_like_official(ws) -> bool:
    title = as_text(ws.title)
    if "名称勿改" in title:
        return True
    blob = " ".join(as_text(c.value) for c in (ws[1] if ws.max_row else [])[:8])
    return "录凭证" in blob and "gl_voucher" in blob


def looks_like_import_template(headers: list[str]) -> bool:
    joined = " ".join(headers)
    return "科目代码" in joined and "凭证号" in joined and "摘要" in joined


def looks_like_voucher_list(headers: list[str]) -> bool:
    joined = " ".join(headers)
    return "凭证字号" in joined and "科目" in joined and "摘要" in joined


@dataclass
class Entry:
    date: str
    word: str
    number: int
    expl: str
    account: str
    account_name: str = ""
    debit: Decimal | None = None
    credit: Decimal | None = None
    amountfor: Decimal | None = None
    currency: str = "RMB"
    rate: Decimal = Decimal("1")
    customer: str = ""
    supplier: str = ""
    employee: str = ""
    department: str = ""
    source_row: int = 0
    source_attachment: str = ""


@dataclass
class ParsedSource:
    kind: str
    path: Path
    company: str
    entries: list[Entry] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


def sniff_workbook(path: Path, aliases: dict) -> dict:
    wb = load_workbook(path, data_only=False, read_only=True)
    try:
        for title in wb.sheetnames:
            ws = wb[title]
            rows = []
            for i, row in enumerate(ws.iter_rows(max_row=8, values_only=True), 1):
                rows.append([as_text(v) for v in row])
            if not rows:
                continue
            if looks_like_official(ws):
                return {
                    "kind": "official",
                    "sheet": title,
                    "ask": "这已经是金蝶官方引入表，不用再转。请直接拿去引入，或换她发的序时账。",
                    "company": _company_from_rows(rows),
                }
            header_row = None
            header_idx = 0
            for i, row in enumerate(rows):
                if looks_like_import_template(row) or looks_like_voucher_list(row):
                    header_row = row
                    header_idx = i + 1
                    break
            if not header_row:
                continue
            kind = "import_template" if looks_like_import_template(header_row) else "voucher_list"
            return {
                "kind": kind,
                "sheet": title,
                "header_row": header_idx,
                "company": _company_from_rows(rows),
                "headers": header_row,
            }
    finally:
        wb.close()
    return {"kind": None}


def _company_from_rows(rows: list[list[str]]) -> str:
    for row in rows:
        for cell in row:
            if "公司名称" in cell:
                return re.sub(r"^公司名称[:：]\s*", "", cell).strip()
    return ""


def parse_source(path: Path, aliases: dict | None = None) -> ParsedSource:
    aliases = aliases or load_aliases()
    sniff = sniff_workbook(path, aliases)
    if not sniff.get("kind") or sniff["kind"] == "official":
        raise SystemExit(sniff.get("ask") or "认不出序时账（要有日期/凭证号/科目/借贷）")
    wb = load_workbook(path, data_only=False)
    ws = wb[sniff["sheet"]]
    header_row = sniff["header_row"]
    headers = [as_text(c.value) for c in ws[header_row]]
    idx = header_index(headers, aliases)
    need = ["expl", "account"]
    missing = [k for k in need if k not in idx]
    if missing or ("debit" not in idx and "credit" not in idx):
        wb.close()
        raise SystemExit("源表缺列：" + "、".join(missing or ["借方/贷方"]))
    entries: list[Entry] = []
    last_date = ""
    last_word = "记"
    last_no: int | None = None
    for r in range(header_row + 1, (ws.max_row or header_row) + 1):
        row = [ws.cell(r, c + 1).value for c in range(len(headers))]

        def cell(key: str):
            i = idx.get(key)
            return row[i] if i is not None and i < len(row) else None

        account_raw = cell("account")
        expl = as_text(cell("expl"))
        debit = money(cell("debit"))
        credit = money(cell("credit"))
        account, acc_name = split_account(account_raw)
        if not account and not expl and debit is None and credit is None:
            continue
        if not account:
            continue
        if not acc_name:
            acc_name = as_text(cell("account_name"))
        word, number = parse_word_no(cell("word"), cell("number"))
        day = as_text(cell("date"))
        if day:
            last_date = day
        else:
            day = last_date
        if number is None:
            number = last_no
        else:
            last_no = number
            last_word = word or last_word
        word = word or last_word
        if number is None:
            wb.close()
            raise SystemExit(f"第 {r} 行没有凭证号")
        amt = money(cell("amountfor"))
        if amt is None:
            amt = debit if debit is not None else credit
        entries.append(
            Entry(
                date=day,
                word=word or "记",
                number=int(number),
                expl=expl,
                account=account,
                account_name=acc_name,
                debit=debit,
                credit=credit,
                amountfor=amt,
                currency=as_text(cell("currency")) or "RMB",
                rate=money(cell("rate")) or Decimal("1"),
                customer=as_text(cell("customer")),
                supplier=as_text(cell("supplier")),
                employee=as_text(cell("employee")),
                department=as_text(cell("department")),
                source_row=r,
                source_attachment=as_text(cell("attachment")),
            )
        )
    company = sniff.get("company") or ""
    wb.close()
    if not entries:
        raise SystemExit("源表没有分录")
    return ParsedSource(kind=sniff["kind"], path=path, company=company, entries=entries)


def remap_numbers(entries: list[Entry], start: int | None) -> dict[int, int]:
    if not start or int(start) == 0:
        return {}
    start_no = int(start)
    if start_no < 1:
        raise SystemExit("凭证号起始必须是正整数")
    order: list[int] = []
    for ent in entries:
        if ent.number not in order:
            order.append(ent.number)
    mapping = {old: start_no + i for i, old in enumerate(order)}
    for ent in entries:
        ent.number = mapping[ent.number]
    return mapping


def col_by_label(ws, needle: str) -> int:
    for cell in ws[3]:
        if needle in str(cell.value or ""):
            return cell.column
    raise KeyError(needle)


def write_official(path: Path, entries: list[Entry], template: Path) -> Path:
    shutil.copy2(template, path)
    wb = load_workbook(path)
    if KINGDEE_SHEET not in wb.sheetnames:
        wb.close()
        raise SystemExit("金蝶空模缺 sheet1（名称勿改）")
    ws = wb[KINGDEE_SHEET]
    cols = {
        "date": col_by_label(ws, "*记账日期"),
        "word": col_by_label(ws, "*凭证字号"),
        "number": col_by_label(ws, "凭证号 #"),
        "attach": col_by_label(ws, "附件"),
        "expl": col_by_label(ws, "摘要 #"),
        "account": col_by_label(ws, "*科目.编码"),
        "currency": col_by_label(ws, "*币别.编码"),
        "rate": col_by_label(ws, "*汇率"),
        "amountfor": col_by_label(ws, "原币金额"),
        "debit": col_by_label(ws, "借方 #"),
        "credit": col_by_label(ws, "贷方 #"),
        "cus_name": col_by_label(ws, "辅助核算.客户.名称"),
        "sup_name": col_by_label(ws, "辅助核算.供应商.名称"),
        "dep_name": col_by_label(ws, "辅助核算.部门.名称"),
        "emp_name": col_by_label(ws, "辅助核算.职员.名称"),
    }
    row_i = 4
    for ent in entries:
        ws.cell(row_i, cols["date"], ent.date)
        ws.cell(row_i, cols["word"], ent.word)
        ws.cell(row_i, cols["number"], ent.number)
        ws.cell(row_i, cols["attach"], None)
        ws.cell(row_i, cols["expl"], ent.expl)
        ws.cell(row_i, cols["account"], ent.account)
        ws.cell(row_i, cols["currency"], ent.currency or "RMB")
        ws.cell(row_i, cols["rate"], float(ent.rate))
        if ent.amountfor is not None:
            ws.cell(row_i, cols["amountfor"], float(ent.amountfor))
        if ent.debit is not None:
            ws.cell(row_i, cols["debit"], float(ent.debit))
        if ent.credit is not None:
            ws.cell(row_i, cols["credit"], float(ent.credit))
        if ent.customer:
            ws.cell(row_i, cols["cus_name"], ent.customer)
        if ent.supplier:
            ws.cell(row_i, cols["sup_name"], ent.supplier)
        if ent.department:
            ws.cell(row_i, cols["dep_name"], ent.department)
        if ent.employee:
            ws.cell(row_i, cols["emp_name"], ent.employee)
        row_i += 1
    wb.save(path)
    wb.close()
    return path


def self_check(entries: list[Entry], out_path: Path) -> dict:
    issues: list[str] = []
    by_no: dict[int, list[Entry]] = {}
    for ent in entries:
        by_no.setdefault(ent.number, []).append(ent)
    for no, rows in by_no.items():
        debit = sum((r.debit or Decimal("0")) for r in rows)
        credit = sum((r.credit or Decimal("0")) for r in rows)
        if debit != credit:
            issues.append(f"记-{no}借贷不平")
        if len(rows) < 2:
            issues.append(f"记-{no}只有一条分录")
    wb = load_workbook(out_path, data_only=False)
    ws = wb[KINGDEE_SHEET]
    attach_col = col_by_label(ws, "附件")
    debit_col = col_by_label(ws, "借方 #")
    red_kept = False
    for r in range(4, (ws.max_row or 3) + 1):
        if ws.cell(r, attach_col).value not in (None, ""):
            issues.append("附件列不是空")
            break
        dv = ws.cell(r, debit_col).value
        if isinstance(dv, (int, float, Decimal)) and dv < 0:
            red_kept = True
    hidden = wb.sheetnames
    wb.close()
    if KINGDEE_SHEET not in hidden:
        issues.append("缺官方 sheet 名")
    return {
        "voucher_count": len(by_no),
        "line_count": len(entries),
        "numbers": [min(by_no), max(by_no)] if by_no else [],
        "red_debit_kept": red_kept,
        "issues": issues,
        "sheets": hidden,
    }


def discover_source(input_dir: str = "", explicit: str = "") -> Path | None:
    if explicit:
        path = Path(explicit).expanduser()
        return path if path.is_file() else None
    homes: list[Path] = []
    if input_dir:
        homes.append(Path(input_dir).expanduser())
    desktop = Path.home() / "Desktop"
    downloads = Path.home() / "Downloads"
    if desktop.is_dir():
        homes.append(desktop)
    if downloads.is_dir():
        homes.append(downloads)
    homes.append(Path.cwd())
    aliases = load_aliases()
    for folder in homes:
        if not folder.is_dir():
            continue
        for path in sorted(folder.glob("*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True):
            if path.name.startswith("~$") or "结果" in path.name or "凭证引入" in path.name:
                continue
            sniff = sniff_workbook(path, aliases)
            if sniff.get("kind") in {"import_template", "voucher_list"}:
                return path
    return None


def inspect_dir(input_dir: str = "", explicit: str = "") -> dict:
    path = discover_source(input_dir, explicit)
    if not path:
        return {"ready": False, "ask": "找不到序时账 Excel。把她发的表放桌面/Downloads，或告诉我路径。"}
    sniff = sniff_workbook(path, load_aliases())
    if sniff.get("kind") == "official":
        return {"ready": False, "path": str(path), "ask": sniff.get("ask")}
    if not sniff.get("kind"):
        return {"ready": False, "path": str(path), "ask": "这份表认不出序时账表头。"}
    return {
        "ready": True,
        "path": str(path.resolve()),
        "kind": sniff["kind"],
        "company": sniff.get("company") or "",
        "sheet": sniff.get("sheet"),
    }


def write_report(path: Path, payload: dict) -> None:
    lines = [
        f"源表={payload.get('source')}",
        f"种类={payload.get('kind')}",
        f"公司名称={payload.get('company') or '源表没写'}",
        f"凭证张数={payload.get('voucher_count')}",
        f"分录行数={payload.get('line_count')}",
        f"凭证号={payload.get('number_from')}～{payload.get('number_to')}",
        "附件列=空",
        f"产物={payload.get('out')}",
        f"issues={','.join(payload.get('issues') or []) or '无'}",
    ]
    if payload.get("company"):
        lines.append(f"引入时请切到账套：{payload['company']}")
    else:
        lines.append("源表没写公司名称，引入前请自己确认账套（湖南分抄作业=湖南分公司，不要进总部/湖南子）。")
    lines.append("本技能只出引入表，没有点引入/审核/过账。")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def run(
    source: Path,
    out_dir: Path,
    start_voucher_no: int | None = None,
    template: Path | None = None,
) -> dict:
    parsed = parse_source(source)
    remap_numbers(parsed.entries, start_voucher_no)
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / "凭证引入_结果.xlsx"
    write_official(out, parsed.entries, template or TEMPLATE)
    check = self_check(parsed.entries, out)
    payload = {
        "source": str(source.resolve()),
        "kind": parsed.kind,
        "company": parsed.company,
        "voucher_count": check["voucher_count"],
        "line_count": check["line_count"],
        "number_from": check["numbers"][0] if check["numbers"] else "",
        "number_to": check["numbers"][1] if check["numbers"] else "",
        "out": str(out.resolve()),
        "issues": check["issues"],
        "red_debit_kept": check["red_debit_kept"],
        "report": str((out_dir / "运行报告.txt").resolve()),
    }
    write_report(out_dir / "运行报告.txt", payload)
    if check["issues"]:
        payload["status"] = "fail"
        return payload
    payload["status"] = "ok"
    return payload


def log(text: str) -> None:
    print(text, flush=True)


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="序时账入金蝶")
    parser.add_argument("--inspect", action="store_true")
    parser.add_argument("--input-dir", default="")
    parser.add_argument("--input", default="")
    parser.add_argument("--out-dir", "--out", dest="out_dir", default="")
    parser.add_argument("--start-voucher-no", type=int, default=0)
    parser.add_argument("--template", default="")
    args = parser.parse_args(argv)
    if args.inspect:
        report = inspect_dir(args.input_dir, args.input)
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return 0 if report.get("ready") else 2
    source = discover_source(args.input_dir, args.input)
    if not source:
        log("找不到序时账 Excel。把她发的表放桌面/Downloads，或告诉我路径。")
        return 2
    out_dir = Path(args.out_dir).expanduser() if args.out_dir else default_desktop_dir("序时账入金蝶")
    template = Path(args.template).expanduser() if args.template else TEMPLATE
    try:
        payload = run(source, out_dir, args.start_voucher_no or None, template)
    except SystemExit as e:
        log(str(e))
        return 2
    log(
        f"序时账已转成金蝶引入表：{payload['voucher_count']} 张 / {payload['line_count']} 行，"
        f"凭证号 {payload['number_from']}～{payload['number_to']}，附件已清空。"
        f"表在 {payload['out']}。"
        f"{'引入时请切到：' + payload['company'] if payload.get('company') else '源表没写公司名称，引入前确认账套。'}"
        "我没有点引入。"
    )
    print(json.dumps({k: payload[k] for k in payload if k != "entries"}, ensure_ascii=False, indent=2))
    return 0 if payload.get("status") == "ok" else 2


if __name__ == "__main__":
    raise SystemExit(main())
