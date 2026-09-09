#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""三模块合一：月底稿认表、1131 公共余额表、付款表头、凭证号失败。打真实 convert.py。"""
from __future__ import annotations

import json
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook

from test_kingdee_posting import (
    SCRIPTS,
    _dummy_pdf,
    _lookups,
    _master,
    _ok_sales,
    _run,
    _write_assist,
    _write_pay,
    _write_sales,
    convert,
    inspect_inputs,
)


def _write_month_draft(path: Path, receipt_rows, org=None, with_flow=True, with_boc_pay=True):
    wb = Workbook()
    org_ws = wb.active
    org_ws.title = "组织架构"
    org_ws.append(["姓名", "部门编码"])
    for pair in org or [("于占国", "15"), ("陈霞", "15")]:
        org_ws.append(list(pair))
    sales = wb.create_sheet("数电票-专票")
    sales.append(["日期", "发票类型", "单位名称", "价税合计", "申请人"])
    sales.append(["2026-09-01", "专票", "甲科技有限公司", 1060, "于占国"])
    if with_flow:
        flow = wb.create_sheet("中行流水")
        flow.append(["日期", "凭证号", "借方（增加）", "贷方（减少）", "余额"])
        flow.append(["2026-09-01", "x", 99999, 1, 1])
    if with_boc_pay:
        pay = wb.create_sheet("中行付款")
        pay.append(["日期", "客户名称", "贷方（减少）"])
        pay.append(["2026-09-01", "不要当付款", 1])
    rec = wb.create_sheet("中行收款")
    rec.append(["日期", "客户名称", "借方（增加）", "销售", "类型"])
    for row in receipt_rows:
        rec.append(row)
    wb.save(path)
    wb.close()


def test_inspect_month_draft_mixed_asks_without_scene(tmp_path):
    _write_month_draft(tmp_path / "2026年9月底稿.xlsx", [["2026-09-01", "甲科技有限公司", 10, "", ""]])
    report = inspect_inputs.inspect_dir(tmp_path)
    assert report["ready"] is False
    assert report["mixed"] is True
    assert "销项发票" in (report.get("ask") or "")


def test_inspect_scene_receipt_reads_boc_receipt_not_flow(tmp_path):
    _write_month_draft(tmp_path / "2026年9月底稿.xlsx", [["2026-09-01", "甲科技有限公司", 10, "", ""]])
    report = inspect_inputs.inspect_dir(tmp_path, "收款")
    assert report["ready"] is True
    assert report["scene"] == "收款"
    assert "部门编码" not in (report.get("ask") or "")
    assert Path(report["files"]["receipt"]).name == "2026年9月底稿.xlsx"


def test_inspect_scene_payment_skips_boc_pay_and_flow(tmp_path):
    _write_month_draft(tmp_path / "2026年9月底稿.xlsx", [["2026-09-01", "甲科技有限公司", 10, "", ""]])
    report = inspect_inputs.inspect_dir(tmp_path, "付款")
    assert report["ready"] is False
    assert any("三列表" in m or "付款" in m for m in report["missing"])


def test_inspect_receipt_does_not_require_dept_column(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "中行收款"
    ws.append(["日期", "客户名称", "借方（增加）", "销售"])
    ws.append(["2026-09-01", "甲科技有限公司", 10, ""])
    wb.save(tmp_path / "稿.xlsx")
    wb.close()
    report = inspect_inputs.inspect_dir(tmp_path, "收款")
    assert report["ready"] is True


def test_inspect_standalone_six_col_receipt_ready(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "收款"
    ws.append(["日期", "客户名称", "借方（增加）", "销售", "部门编码", "应收账款编码"])
    ws.append(["2026-08-01", "甲科技有限公司", 10, "于占国", "15", "113101"])
    wb.save(tmp_path / "收款.xlsx")
    wb.close()
    report = inspect_inputs.inspect_dir(tmp_path, "收款")
    assert report["ready"] is True
    assert report["scene"] == "收款"


def test_convert_receipt_named_sheet_not_flow_or_org(tmp_path):
    _write_month_draft(
        tmp_path / "稿.xlsx",
        [
            ["2026-09-01", "甲科技有限公司", 10, "于占国", ""],
            ["2026-09-01", "甲科技有限公司", 20, "于占国", ""],
        ],
    )
    result = _run(tmp_path, "收款")
    assert result["source_count"] == 2
    assert result["sheet"] == "中行收款"
    assert result["bookable_count"] == 2


def test_convert_receipt_without_dept_column_uses_org(tmp_path):
    _write_month_draft(tmp_path / "稿.xlsx", [["2026-09-01", "甲科技有限公司", 10, "于占国", ""]])
    result = _run(tmp_path, "收款")
    assert result["bookable_count"] == 1
    kd = load_workbook(result["kingdee_path"])
    ws = kd[convert.KINGDEE_SHEET]
    depts = [str(ws.cell(r, 22).value or "") for r in range(4, 8)]
    kd.close()
    assert "15" in depts


def test_table_sales_used_before_zhiyun(tmp_path):
    _write_month_draft(tmp_path / "稿.xlsx", [["2026-09-01", "甲科技有限公司", 10, "陈霞", ""]])
    result = _run(
        tmp_path,
        "收款",
        lookups=_lookups(
            receipt_sales=[{"customer": "甲科技有限公司", "date": "2026-09-01", "amount": "10.00", "sales": ["于占国"]}],
        ),
    )
    assert result["bookable_count"] == 1
    kd = load_workbook(result["kingdee_path"])
    ws = kd[convert.KINGDEE_SHEET]
    emps = [str(ws.cell(r, 24).value or "") for r in range(4, 8)]
    kd.close()
    assert "011" in emps
    assert "103" not in emps


def test_empty_sales_uses_zhiyun_lookups(tmp_path):
    _write_month_draft(tmp_path / "稿.xlsx", [["2026-09-01", "甲科技有限公司", 10, "", ""]])
    result = _run(
        tmp_path,
        "收款",
        lookups=_lookups(
            receipt_sales=[{"customer": "甲科技有限公司", "date": "2026-09-01", "amount": "10.00", "sales": ["于占国"]}],
        ),
    )
    assert result["bookable_count"] == 1
    kd = load_workbook(result["kingdee_path"])
    ws = kd[convert.KINGDEE_SHEET]
    emps = [str(ws.cell(r, 24).value or "") for r in range(4, 8)]
    kd.close()
    assert "103" in emps


def test_haidian_receipt_not_merged_to_0386(tmp_path):
    _write_month_draft(tmp_path / "稿.xlsx", [["2026-09-01", "北京市公安局海淀分局", 10, "于占国", ""]])
    result = _run(tmp_path, "收款")
    assert result["bookable_count"] == 1
    kd = load_workbook(result["kingdee_path"])
    ws = kd[convert.KINGDEE_SHEET]
    names = [ws.cell(r, 19).value for r in range(4, 8)]
    codes = [str(ws.cell(r, 18).value or "") for r in range(4, 8)]
    kd.close()
    assert "北京市公安局海淀分局" in names
    assert "公安部" not in names
    assert "1940" in codes
    assert "0386" not in codes


def test_assist_xlsx_one_account_copied(tmp_path):
    assist = tmp_path / "核算项目余额表_客户_1131_本年.xlsx"
    _write_assist(
        assist,
        [
            {
                "period": "202608",
                "customer_code": "1001",
                "customer_name": "甲科技有限公司",
                "account": "113107",
                "ending_debit": "80",
                "ytd_debit": "80",
            }
        ],
    )
    _write_month_draft(tmp_path / "稿.xlsx", [["2026-09-01", "甲科技有限公司", 10, "于占国", ""]])
    result = convert.run_dir(
        tmp_path,
        "收款",
        "2026-09-07",
        _master(),
        _lookups(include_assist=False),
        start_voucher_no=1,
        out_dir=tmp_path,
        ar_xlsx=assist,
    )
    assert result["bookable_count"] == 1
    kd = load_workbook(result["kingdee_path"])
    ws = kd[convert.KINGDEE_SHEET]
    accounts = [str(ws.cell(r, 7).value or "") for r in range(4, 8)]
    kd.close()
    assert "113107" in accounts
    assert "113103" not in accounts


def test_assist_xlsx_two_ending_holds(tmp_path):
    assist = tmp_path / "核算项目余额表_客户_1131_本年.xlsx"
    _write_assist(
        assist,
        [
            {
                "period": "202608",
                "customer_code": "1001",
                "customer_name": "甲科技有限公司",
                "account": "113103",
                "ending_debit": "80",
            },
            {
                "period": "202608",
                "customer_code": "1001",
                "customer_name": "甲科技有限公司",
                "account": "113107",
                "ending_debit": "20",
            },
        ],
    )
    _write_month_draft(tmp_path / "稿.xlsx", [["2026-09-01", "甲科技有限公司", 10, "于占国", ""]])
    result = convert.run_dir(
        tmp_path,
        "收款",
        "2026-09-07",
        _master(),
        _lookups(include_assist=False),
        out_dir=tmp_path,
        ar_xlsx=assist,
    )
    assert result["bookable_count"] == 0
    assert result["hold_count"] == 1
    detail = load_workbook(result["detail_path"])
    reason = str(detail.active.cell(2, 2).value or "")
    extra = json.loads(str(detail.active.cell(2, 7).value or "{}"))
    detail.close()
    assert "多条" in reason
    assert extra.get("候选1131") == ["113103", "113107"]


def test_table_ar_code_ignored_on_six_col(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "收款"
    ws.append(["日期", "客户名称", "借方（增加）", "销售", "部门编码", "应收账款编码"])
    ws.append(["2026-08-01", "甲科技有限公司", 10, "于占国", "15", "113101"])
    wb.save(tmp_path / "收款.xlsx")
    wb.close()
    result = convert.run_dir(
        tmp_path,
        "收款",
        "2026-08-27",
        _master(),
        _lookups(
            receipt_sales=[{"customer": "甲科技有限公司", "date": "2026-08-01", "amount": "10.00", "sales": ["陈霞"]}],
        ),
        out_dir=tmp_path,
    )
    assert result["bookable_count"] == 1
    kd = load_workbook(result["kingdee_path"])
    ws = kd[convert.KINGDEE_SHEET]
    emps = [str(ws.cell(r, 24).value or "") for r in range(4, 8)]
    accounts = [str(ws.cell(r, 7).value or "") for r in range(4, 8)]
    kd.close()
    assert "103" in emps
    assert "113103" in accounts
    assert "113101" not in accounts


def test_booking_date_from_receipt_row(tmp_path):
    _write_month_draft(tmp_path / "稿.xlsx", [["2026-09-03", "甲科技有限公司", 10, "于占国", ""]])
    result = _run(
        tmp_path,
        "收款",
        lookups=_lookups(
            receipt_sales=[{"customer": "甲科技有限公司", "date": "2026-09-03", "amount": "10.00", "sales": ["于占国"]}],
        ),
    )
    kd = load_workbook(result["kingdee_path"])
    ws = kd[convert.KINGDEE_SHEET]
    dates = [str(ws.cell(r, 1).value or "") for r in range(4, 8)]
    kd.close()
    assert any("2026-09-03" in d for d in dates)
    assert not any("2026-08-27" in d for d in dates)


def test_note_has_counts_not_names(tmp_path):
    _write_month_draft(tmp_path / "稿.xlsx", [["2026-09-01", "甲科技有限公司", 10, "于占国", ""]])
    result = _run(tmp_path, "收款")
    text = Path(result["note_path"]).read_text(encoding="utf-8")
    assert "源：1" in text
    assert "可入账：1" in text
    assert "中行收款" in text
    assert "甲科技" not in text
    assert "10.00" not in text


def test_aux_columns_have_no_formula(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "收款"
    ws.append(["日期", "客户名称", "借方（增加）", "销售", "部门编码"])
    ws.append(["2026-08-01", "甲科技有限公司", 10, "于占国", "=VLOOKUP(A2,组织架构!A:B,2,0)"])
    wb.save(tmp_path / "收款.xlsx")
    wb.close()
    result = _run(tmp_path, "收款")
    assert result["bookable_count"] == 1
    kd = load_workbook(result["kingdee_path"])
    ws = kd[convert.KINGDEE_SHEET]
    for r in range(4, 8):
        for c in range(18, 26):
            val = ws.cell(r, c).value
            if val is not None:
                assert not str(val).startswith("=")
    kd.close()


def test_payment_first_sheet_not_ledger(tmp_path, monkeypatch):
    wb = Workbook()
    cover = wb.active
    cover.title = "说明"
    cover.append(["备注"])
    cover.append(["这不是三列表"])
    pay = wb.create_sheet("应付")
    pay.append(["供应商", "应付金额本币", "开户名"])
    pay.append(["北京某翻译店", 100, "北京某翻译店"])
    wb.save(tmp_path / "付款.xlsx")
    wb.close()
    folder = tmp_path / "北京某翻译店"
    folder.mkdir()
    _dummy_pdf(folder / "a.pdf")
    monkeypatch.setattr(
        convert,
        "parse_invoice_pdf",
        lambda p: {"kind": "普票", "seller": "北京某翻译店", "total": Decimal("100.00"), "tax": None},
    )
    result = _run(tmp_path, "付款")
    assert result["bookable_count"] == 1
    assert result["sheet"] == "应付"


def test_payment_many_uses_9999(tmp_path, monkeypatch):
    master = _master()
    master["supplier"] = [
        {"code": "8001", "name": "重复供应商"},
        {"code": "8002", "name": "重复供应商"},
        {"code": "9999", "name": "其他供应商"},
    ]
    _write_pay(tmp_path / "付款.xlsx", [["重复供应商", 100, "重复供应商"]])
    folder = tmp_path / "重复供应商"
    folder.mkdir()
    _dummy_pdf(folder / "a.pdf")
    monkeypatch.setattr(
        convert,
        "parse_invoice_pdf",
        lambda p: {"kind": "普票", "seller": "重复供应商", "total": Decimal("100.00"), "tax": None},
    )
    result = convert.run_dir(tmp_path, "付款", "2026-08-27", master, _lookups(), out_dir=tmp_path)
    assert result["bookable_count"] == 1
    kd = load_workbook(result["kingdee_path"])
    ws = kd[convert.KINGDEE_SHEET]
    codes = [str(ws.cell(r, 20).value or "") for r in range(4, 8)]
    kd.close()
    assert "9999" in codes
    assert "8001" not in codes


def test_payment_pdf_missing_seller_holds(tmp_path, monkeypatch):
    _write_pay(tmp_path / "付款.xlsx", [["北京某翻译店", 100, "北京某翻译店"]])
    folder = tmp_path / "北京某翻译店"
    folder.mkdir()
    _dummy_pdf(folder / "a.pdf")
    monkeypatch.setattr(
        convert,
        "parse_invoice_pdf",
        lambda p: {"kind": "普票", "seller": "", "total": Decimal("100.00"), "tax": None},
    )
    result = _run(tmp_path, "付款")
    assert result["bookable_count"] == 0
    assert result["hold_count"] == 1
    detail = load_workbook(result["detail_path"])
    reason = str(detail.active.cell(2, 2).value or "")
    detail.close()
    assert "销售方" in reason or "票种" in reason


def test_cli_voucher_query_fail_exits_nonzero(tmp_path, monkeypatch):
    _write_sales(tmp_path / "发票.xlsx", [_ok_sales()], with_org=False)
    master = tmp_path / "master.json"
    lookups = tmp_path / "lookups.json"
    master.write_text(json.dumps(_master()), encoding="utf-8")
    lookups.write_text(json.dumps(_lookups()), encoding="utf-8")
    monkeypatch.setattr(
        convert.kingdee_api,
        "try_fetch_next_voucher_no",
        lambda period=None: {"ok": False, "missing_credentials": False, "error": "voucher list http 500"},
    )
    assert (
        convert.main(
            [
                "--input-dir",
                str(tmp_path),
                "--scene",
                "销项发票",
                "--master",
                str(master),
                "--lookups",
                str(lookups),
                "--date",
                "2026-08-27",
                "--out-dir",
                str(tmp_path),
            ]
        )
        == 2
    )
    assert not (tmp_path / "凭证引入_结果.xlsx").exists()


def test_sales_org_overrides_applicant_json(tmp_path):
    _write_sales(tmp_path / "发票.xlsx", [_ok_sales()], org=[("于占国", "0308")], with_org=True)
    result = _run(tmp_path, "销项发票")
    assert result["bookable_count"] == 1
    kd = load_workbook(result["kingdee_path"])
    ws = kd[convert.KINGDEE_SHEET]
    assert str(ws.cell(4, 22).value) == "0308"
    kd.close()


def test_sales_path_uses_pick_assist_not_voucher_scan():
    text = (SCRIPTS / "convert.py").read_text(encoding="utf-8")
    start = text.index("def resolve_sales_party")
    end = text.index("def parse_invoice_text")
    body = text[start:end]
    assert "pick_assist_account" in body
    assert "resolve_ar" not in body
    rec_start = text.index("def convert_receipt")
    rec_end = text.index("def default_desktop_dir")
    rec = text[rec_start:rec_end]
    assert "pick_assist_account" in rec
    assert "resolve_ar" not in rec
    assert "fill_period" not in rec
    assert "try_fetch_customer_ar" not in text[text.index("def main") :]


def test_default_desktop_prefix_by_scene(tmp_path):
    desktop = tmp_path / "Desktop"
    desktop.mkdir()
    from datetime import date

    sales = convert.default_desktop_dir(convert.SCENE_DESKTOP["销项发票"], today=date(2026, 9, 8), home=tmp_path)
    pay = convert.default_desktop_dir(convert.SCENE_DESKTOP["付款"], today=date(2026, 9, 8), home=tmp_path)
    rec = convert.default_desktop_dir(convert.SCENE_DESKTOP["收款"], today=date(2026, 9, 8), home=tmp_path)
    assert sales == desktop / "金蝶入账_销项_20260908"
    assert pay == desktop / "金蝶入账_付款_20260908"
    assert rec == desktop / "金蝶入账_收款_20260908"


def test_bookable_tieout_zero(tmp_path):
    _write_month_draft(tmp_path / "稿.xlsx", [["2026-09-01", "甲科技有限公司", 10, "于占国", ""]])
    result = _run(tmp_path, "收款")
    assert result["tieout_source"] == result["tieout_debit"] == result["tieout_credit"]
    assert Decimal(result["tieout_debit"]) - Decimal(result["tieout_credit"]) == 0


def test_receipt_employee_dept_fallback(tmp_path):
    master = _master()
    master["employee"].append({"code": "399", "name": "李测试", "dept": "15"})
    _write_month_draft(tmp_path / "稿.xlsx", [["2026-09-01", "甲科技有限公司", 10, "李测试", ""]], org=None)
    result = convert.run_dir(
        tmp_path,
        "收款",
        "2026-09-07",
        master,
        _lookups(receipt_sales=[{"customer": "甲科技有限公司", "date": "2026-09-01", "amount": "10.00", "sales": ["李测试"]}]),
        start_voucher_no=1,
        out_dir=tmp_path,
    )
    assert result["bookable_count"] == 1
    assert result["hold_count"] == 0


def test_receipt_peel_name_matches_zhiyun(tmp_path):
    _write_month_draft(tmp_path / "稿.xlsx", [["2026-09-01", "甲科技有限公司", 10, "", ""]])
    result = convert.run_dir(
        tmp_path,
        "收款",
        "2026-09-07",
        _master(),
        _lookups(receipt_sales=[{"customer": "甲科技", "date": "2026-09-01", "amount": "10.00", "sales": ["于占国"]}]),
        start_voucher_no=1,
        out_dir=tmp_path,
    )
    assert result["bookable_count"] == 1


def test_sales_multi_assist_extra_lists_accounts(tmp_path):
    _write_sales(tmp_path / "发票.xlsx", [_ok_sales()], with_org=True)
    result = convert.run_dir(
        tmp_path,
        "销项发票",
        "2026-09-07",
        _master(),
        _lookups(
            assist_rows=[
                {
                    "period": "202608",
                    "customer_code": "1001",
                    "customer_name": "甲科技有限公司",
                    "account": "113103",
                    "ending_debit": "80",
                },
                {
                    "period": "202608",
                    "customer_code": "1001",
                    "customer_name": "甲科技有限公司",
                    "account": "113107",
                    "ending_debit": "20",
                },
            ]
        ),
        start_voucher_no=1,
        out_dir=tmp_path,
    )
    assert result["hold_count"] == 1
    detail = load_workbook(result["detail_path"])
    extra = json.loads(str(detail.active.cell(2, 7).value or "{}"))
    detail.close()
    assert extra.get("候选1131") == ["113103", "113107"]
