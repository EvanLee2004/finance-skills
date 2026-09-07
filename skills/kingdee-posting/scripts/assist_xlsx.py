#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""客户核算项目余额表：认已有 xlsx，没有则按现网过滤网页引出。OpenAPI 没有这张表。"""
from __future__ import annotations

import os
import sys
import time
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
if str(HERE) not in sys.path:
    sys.path.insert(0, str(HERE))

import lookups as lookup_mod  # noqa: E402

ASSIST_FORM = "https://tf.jdy.com/ierp/index.html?formId=gl_rpt_assistbalance"
EXPORT_LOG = "https://tf.jdy.com/ierp/index.html?formId=bos_exportlog_list"
WORKBENCH = "https://service.jdy.com/workbench/web/index.html"
XINGCHEN_HOME = "https://tf.jdy.com/ierp/index.html?formId=home_page"
HQ_NAME = "甲骨易（北京）语言科技股份有限公司"
STATE = Path.home() / ".config" / "finance" / "xingchen.playwright-state.json"
PROFILE = Path.home() / ".cache" / "chrome-devtools-mcp" / "chrome-profile"
DOWNLOADS = Path.home() / "Downloads"

# 与 技能/金蝶/工作区/20260907_客户核算项目余额表/现网探测.md 同一套过滤。
EXPORT_FILTERS = {
    "form_id": "gl_rpt_assistbalance",
    "assist_type": "客户",
    "account": "1131",
    "period": "本年",
    "hide_zero_balance": False,
    "show_detail_accounts": True,
}


def workspace_assist_dirs() -> list[Path]:
    # scripts/ -> kingdee-posting -> skills -> finance-skills -> 财务部skills
    root = HERE.parents[3]
    work = root / "技能" / "金蝶" / "工作区"
    out = []
    if work.is_dir():
        for child in sorted(work.iterdir()):
            if child.is_dir() and "客户核算项目余额表" in child.name:
                out.append(child)
        extra = work / "引出"
        if extra.is_dir():
            out.append(extra)
    return out


def _is_assist_name(name: str) -> bool:
    if not name.endswith(".xlsx") or name.startswith("~$"):
        return False
    if "结果" in name:
        return False
    return "核算项目余额表" in name


def _collect(folder: Path) -> list[Path]:
    if not folder.is_dir():
        return []
    return [p for p in folder.glob("*.xlsx") if _is_assist_name(p.name)]


def find_assist_xlsx(extra_dirs: list[Path] | None = None) -> Path | None:
    env = os.environ.get("KINGDEE_ASSIST_XLSX", "").strip()
    if env:
        path = Path(env).expanduser()
        if path.is_file():
            return path
    groups: list[list[Path]] = []
    first: list[Path] = []
    for raw in extra_dirs or []:
        first.extend(_collect(Path(raw)))
    if first:
        groups.append(first)
    work_hits: list[Path] = []
    for folder in workspace_assist_dirs():
        work_hits.extend(_collect(folder))
    if work_hits:
        groups.append(work_hits)
    dl = _collect(DOWNLOADS)
    if dl:
        groups.append(dl)
    for group in groups:
        preferred = [p for p in group if "客户" in p.name and "1131" in p.name]
        pool = preferred or group
        return max(pool, key=lambda p: p.stat().st_mtime)
    return None


def parse_assist_xlsx(path: Path) -> list[lookup_mod.AssistRow]:
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        col: dict[str, int] = {}
        ytd_i = None
        end_i = None
        out: list[lookup_mod.AssistRow] = []
        for row in ws.iter_rows(values_only=True):
            vals = list(row)
            texts = [str(c).strip() if c is not None else "" for c in vals]
            if "期间" in texts and "客户编码" in texts and "科目编码" in texts:
                if texts and texts[0] == "期间" and "借方" in texts:
                    continue
                for i, t in enumerate(texts):
                    if t == "期间" and "period" not in col:
                        col["period"] = i
                    elif t == "客户编码":
                        col["code"] = i
                    elif t == "客户名称":
                        col["name"] = i
                    elif t == "科目编码":
                        col["account"] = i
                    elif t == "本年累计" and ytd_i is None:
                        ytd_i = i
                    elif t == "期末" and end_i is None:
                        end_i = i
                continue
            if not col:
                continue
            period_i = col.get("period")
            code_i = col.get("code")
            name_i = col.get("name")
            acc_i = col.get("account")
            if period_i is None or code_i is None or acc_i is None:
                continue
            if acc_i >= len(vals):
                continue
            code = str(vals[code_i] or "").strip()
            acc = str(vals[acc_i] or "").strip()
            if acc.endswith(".0") and acc[:-2].isdigit():
                acc = acc[:-2]
            if code.endswith(".0") and code[:-2].isdigit():
                code = code[:-2]
            if not code or not acc.startswith("1131"):
                continue
            name = str(vals[name_i] or "").strip() if name_i is not None and name_i < len(vals) else ""
            ytd_debit = ytd_credit = end_debit = end_credit = None
            if ytd_i is not None and ytd_i + 1 < len(vals):
                ytd_debit = lookup_mod.money(vals[ytd_i])
                ytd_credit = lookup_mod.money(vals[ytd_i + 1])
            if end_i is not None and end_i + 1 < len(vals):
                end_debit = lookup_mod.money(vals[end_i])
                end_credit = lookup_mod.money(vals[end_i + 1])
            out.append(
                lookup_mod.AssistRow(
                    period=lookup_mod._period_key(vals[period_i] if period_i < len(vals) else ""),
                    customer_code=code,
                    customer_name=name,
                    account=acc,
                    ending_debit=end_debit,
                    ending_credit=end_credit,
                    ytd_debit=ytd_debit,
                    ytd_credit=ytd_credit,
                )
            )
        return out
    finally:
        wb.close()


def _newest_xlsx(folder: Path, after: float) -> Path | None:
    newest = None
    newest_mtime = after
    if not folder.is_dir():
        return None
    for path in folder.glob("*.xlsx"):
        if not _is_assist_name(path.name) and "核算项目余额表" not in path.name:
            continue
        if path.name.startswith("~$"):
            continue
        try:
            mtime = path.stat().st_mtime
        except OSError:
            continue
        if mtime > newest_mtime:
            newest = path
            newest_mtime = mtime
    return newest


async def _click_text(page, label: str) -> bool:
    return bool(
        await page.evaluate(
            """(label) => {
              const walker = document.createTreeWalker(document.body, NodeFilter.SHOW_TEXT);
              let n;
              while ((n = walker.nextNode())) {
                if ((n.textContent || '').trim() === label) {
                  const el = n.parentElement;
                  (el?.closest('label,li,button,div,span') || el)?.click();
                  return true;
                }
              }
              return false;
            }""",
            label,
        )
    )


async def _dismiss_overlays(page) -> None:
    await page.evaluate(
        """() => {
          document.querySelectorAll('iframe').forEach(f => {
            const src = f.src || '';
            if (/piaozone|fpdk|etax/.test(src)) (f.closest('div') || f).remove();
          });
        }"""
    )


async def _open_context(p):
    if PROFILE.is_dir():
        try:
            ctx = await p.chromium.launch_persistent_context(
                user_data_dir=str(PROFILE),
                channel="chrome",
                headless=False,
                accept_downloads=True,
            )
            return ctx, None, "profile"
        except Exception:
            pass
    browser = await p.chromium.launch(headless=False, channel="chrome")
    ctx = await browser.new_context(
        accept_downloads=True,
        storage_state=str(STATE) if STATE.is_file() else None,
    )
    return ctx, browser, "fresh"


async def _ensure_xingchen(page) -> bool:
    await page.goto(WORKBENCH, wait_until="domcontentloaded")
    await page.wait_for_timeout(1500)
    if await page.get_by_text("进入使用").count():
        await page.get_by_text("进入使用").first.click()
        await page.wait_for_timeout(3000)
    body = await page.locator("body").inner_text()
    if "登录" in body and "进入使用" not in body and "云星辰" not in body:
        return False
    await page.goto(XINGCHEN_HOME, wait_until="domcontentloaded")
    await page.wait_for_timeout(1500)
    await _dismiss_overlays(page)
    current = await page.locator("body").inner_text()
    if HQ_NAME in current:
        return True
    handle = page.get_by_text(HQ_NAME).first
    if await handle.count():
        await handle.click()
        await page.wait_for_timeout(800)
        target = page.get_by_text(HQ_NAME, exact=False).first
        if await target.count():
            await target.click()
            await page.wait_for_timeout(2500)
    return HQ_NAME in await page.locator("body").inner_text()


async def _set_customer_ar_filters(page) -> None:
    await _click_text(page, "展开过滤")
    await page.wait_for_timeout(400)
    await _click_text(page, "本年")
    await page.wait_for_timeout(300)
    box = page.locator(".kd-table-cell-basedata-container").first
    if await box.count():
        await box.click()
        await page.keyboard.type(EXPORT_FILTERS["assist_type"])
        await page.wait_for_timeout(800)
        if await page.get_by_text("客户", exact=True).count():
            await page.get_by_text("客户", exact=True).last.click()
            await page.keyboard.press("Enter")
            await page.wait_for_timeout(400)
    await _click_text(page, "1131 应收账款") or await _click_text(page, "应收账款")
    if await page.get_by_text("显示最明细科目").count():
        await _click_text(page, "显示最明细科目")
    checked = page.locator("label").filter(has_text="余额为 0 不显示")
    if await checked.count():
        # 不要勾。若已勾则点掉。
        await page.evaluate(
            """() => {
              const labels = Array.from(document.querySelectorAll('label,span,div'));
              for (const el of labels) {
                if ((el.textContent || '').trim() === '余额为 0 不显示' ||
                    (el.textContent || '').trim() === '余额为0不显示') {
                  const input = el.querySelector('input[type=checkbox]') ||
                    el.parentElement?.querySelector('input[type=checkbox]');
                  if (input && input.checked) el.click();
                  break;
                }
              }
            }"""
        )
    await _click_text(page, "查询")
    await page.wait_for_timeout(2500)


async def _export_via_log(page, dest: Path) -> bool:
    dest.parent.mkdir(parents=True, exist_ok=True)
    mark = time.time()
    await _click_text(page, "引出")
    await page.wait_for_timeout(800)
    if await page.get_by_text("到引出结果界面下载").count():
        await page.get_by_text("到引出结果界面下载").first.click()
        await page.wait_for_timeout(1500)
    else:
        await page.goto(EXPORT_LOG, wait_until="domcontentloaded")
        await page.wait_for_timeout(2000)
    for _ in range(20):
        body = await page.locator("body").inner_text()
        if "成功" in body or "完成" in body:
            break
        await page.wait_for_timeout(1000)
    try:
        async with page.expect_download(timeout=20000) as dl_info:
            if await page.get_by_text("下载").count():
                await page.get_by_text("下载").first.click()
            elif await page.get_by_text("核算项目余额表").count():
                await page.get_by_text("核算项目余额表").first.click()
        download = await dl_info.value
        await download.save_as(str(dest))
        if dest.is_file() and dest.stat().st_size > 400:
            return True
    except Exception:
        pass
    found = _newest_xlsx(DOWNLOADS, mark - 1)
    if found:
        dest.write_bytes(found.read_bytes())
        return dest.is_file() and dest.stat().st_size > 400
    return False


def export_customer_assist_xlsx(dest: Path | None = None) -> dict:
    """网页引出客户+1131 核算项目余额表。失败不得假装 OpenAPI 通了。"""
    dest = dest or (DOWNLOADS / "核算项目余额表_客户_1131_本年.xlsx")
    try:
        import playwright  # noqa: F401
    except Exception as e:
        return {"ok": False, "error": f"playwright_missing:{type(e).__name__}", "path": None}

    async def _run():
        from playwright.async_api import async_playwright

        async with async_playwright() as p:
            ctx, browser, mode = await _open_context(p)
            page = ctx.pages[0] if ctx.pages else await ctx.new_page()
            try:
                if not await _ensure_xingchen(page):
                    return {"ok": False, "error": "星辰未登录，请先在本机 Chrome 登录总部后再引出客户核算项目余额表", "path": None, "mode": mode}
                await _dismiss_overlays(page)
                await page.goto(ASSIST_FORM, wait_until="domcontentloaded")
                await page.wait_for_timeout(2000)
                await _dismiss_overlays(page)
                await _set_customer_ar_filters(page)
                ok = await _export_via_log(page, dest)
                if ok:
                    return {"ok": True, "path": str(dest), "mode": mode}
                return {"ok": False, "error": "引出未拿到核算项目余额表", "path": None, "mode": mode}
            finally:
                await ctx.close()
                if browser:
                    await browser.close()

    import asyncio

    try:
        return asyncio.run(_run())
    except Exception as e:
        return {"ok": False, "error": f"{type(e).__name__}: {str(e)[:200]}", "path": None}


def ensure_assist_xlsx(extra_dirs: list[Path] | None = None) -> Path:
    found = find_assist_xlsx(extra_dirs)
    if found:
        return found
    dest = DOWNLOADS / "核算项目余额表_客户_1131_本年.xlsx"
    got = export_customer_assist_xlsx(dest)
    if got.get("ok") and got.get("path"):
        path = Path(got["path"])
        if path.is_file():
            return path
    reason = got.get("error") or "没有客户核算项目余额表"
    raise SystemExit(f"{reason}。OpenAPI 没有这张表，请把引出的 xlsx 放到文件夹或 Downloads。")
