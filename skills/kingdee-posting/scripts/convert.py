#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""金蝶入账：销项发票 / 付款 / 收款 → 凭证引入表。金额只由本脚本算。"""
from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import sys
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path


def repo_venv_python() -> Path | None:
    root = Path(__file__).resolve().parents[3]
    for rel in (Path(".venv") / "bin" / "python", Path(".venv") / "Scripts" / "python.exe"):
        cand = root / rel
        if cand.is_file():
            return cand
    return None


def _reexec_repo_venv_if_needed() -> None:
    """opencode 常调系统 python3；仓内 .venv 才有 requests/playwright。只在当脚本跑时切，import 进 pytest 不切。"""
    if not sys.argv or Path(sys.argv[0]).name.lower() not in {"convert.py", "convert"}:
        return
    venv_py = repo_venv_python()
    if venv_py is None:
        return
    try:
        # macOS 上 .venv/bin/python 常是指向 Homebrew 的同一条二进制，不能比 executable。
        if Path(sys.prefix).resolve() == venv_py.parent.parent.resolve():
            return
    except OSError:
        return
    os.execv(str(venv_py), [str(venv_py), *sys.argv])


_reexec_repo_venv_if_needed()

from openpyxl import Workbook, load_workbook

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
CONFIG = SKILL / "config"
TEMPLATE = CONFIG / "凭证引入空模.xlsx"
KINGDEE_SHEET = "sheet1（名称勿改）"
TWOPLACES = Decimal("0.01")

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

sys.path.insert(0, str(HERE))
import assist_xlsx as assist_mod  # noqa: E402
import inspect_inputs as inspect_mod  # noqa: E402
import kingdee_api  # noqa: E402
import lookups as lookup_mod  # noqa: E402
import match_name as names  # noqa: E402
import zhiyun_api  # noqa: E402


def log(msg: str) -> None:
    print(msg, file=sys.stderr)


def money(v):
    if v is None or v == "":
        return None
    if isinstance(v, Decimal):
        return v.quantize(TWOPLACES, rounding=ROUND_HALF_UP)
    if isinstance(v, bool):
        return None
    if isinstance(v, int):
        return Decimal(v).quantize(TWOPLACES, rounding=ROUND_HALF_UP)
    if isinstance(v, float):
        return Decimal(str(round(v, 2))).quantize(TWOPLACES, rounding=ROUND_HALF_UP)
    s = str(v).strip().replace(",", "")
    if not s or s.startswith("="):
        return None
    try:
        return Decimal(s).quantize(TWOPLACES, rounding=ROUND_HALF_UP)
    except Exception:
        return None


def code_str(v) -> str:
    if v is None or v == "":
        return ""
    if isinstance(v, bool):
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, int):
        return str(v)
    s = str(v).strip()
    if s.endswith(".0") and s[:-2].isdigit():
        return s[:-2]
    return s


def norm_name(name: str) -> str:
    s = (name or "").strip().replace("(", "（").replace(")", "）")
    return "".join(s.split())


def strip_ge(name: str) -> str:
    s = (name or "").replace("（个体工商户）", "").replace("(个体工商户)", "")
    return s.strip()


def as_day(value, fallback: str = "") -> str:
    if value is None or value == "":
        return str(fallback or "")[:10]
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip().replace("/", "-").replace(".", "-")
    if len(text) >= 10 and text[4] == "-" and text[7] == "-":
        return text[:10]
    return str(fallback or "")[:10]


def load_json(path: Path, default):
    if not path.is_file():
        return default
    return json.loads(path.read_text(encoding="utf-8"))


def load_rules() -> dict:
    return load_json(CONFIG / "rules.json", {})


def load_aliases() -> dict:
    return load_json(CONFIG / "列名别名.json", {})


def load_customer_alias() -> dict:
    raw = load_json(CONFIG / "客户别名.json", {})
    return {k: v for k, v in raw.items() if not str(k).startswith("_") and v}


def load_ar_accounts() -> list[str]:
    raw = load_rules().get("ar_accounts") or []
    return [str(x).strip() for x in raw if str(x).strip()]


def load_applicant_dept() -> dict:
    raw = load_json(CONFIG / "申请人部门.json", {})
    return {k: str(v).strip() for k, v in raw.items() if not str(k).startswith("_") and v}


def load_emp_hang() -> dict:
    raw = load_json(CONFIG / "职员挂靠.json", {})
    return {k: str(v).strip() for k, v in raw.items() if not str(k).startswith("_") and v}


def load_box(raw: dict | None):
    return lookup_mod.box_from_dict(raw, load_ar_accounts(), load_applicant_dept())


def header_index(headers: list, aliases: dict) -> dict:
    idx = {}
    cleaned = [(i, str(h).strip()) for i, h in enumerate(headers) if h is not None and str(h).strip()]
    for field, names in aliases.items():
        for i, h in cleaned:
            if h in names:
                idx[field] = i
                break
    return idx


def cell_at(row, idx: dict, field: str):
    i = idx.get(field)
    if i is None or i >= len(row):
        return None
    return row[i]


def pick_code(formula_cell, value_cell) -> str:
    if isinstance(formula_cell, str) and formula_cell.startswith("="):
        return code_str(value_cell)
    return code_str(formula_cell) or code_str(value_cell)


def pick_amount(formula_cell, value_cell, fallback=None):
    got = money(value_cell)
    if got is not None:
        return got
    got = money(formula_cell)
    if got is not None:
        return got
    return fallback


@dataclass
class VoucherLine:
    status: str
    reason: str = ""
    source_row: int = 0
    key: str = ""
    expl: str = ""
    voucher_no: int | None = None
    booking_date: str = ""
    extra: dict = field(default_factory=dict)
    entries: list = field(default_factory=list)


SCENE_DESKTOP = {
    "销项发票": "金蝶入账_销项",
    "付款": "金蝶入账_付款",
    "收款": "金蝶入账_收款",
}


def unique_lookup(index: dict[str, list[tuple[str, str]]], name: str):
    hits = index.get(norm_name(name)) or []
    if len(hits) == 1:
        return hits[0], None
    if not hits:
        return None, "none"
    return None, "many"


class Master:
    def __init__(self, data: dict | None):
        data = data or {}
        self.emp: dict[str, list[tuple[str, str]]] = {}
        self.emp_dept: dict[str, list[str]] = {}
        self.dept: dict[str, str] = {}
        self.cus: dict[str, list[tuple[str, str]]] = {}
        self.sup: dict[str, list[tuple[str, str]]] = {}
        for e in data.get("employee") or []:
            name = str(e.get("name") or "").strip()
            code = code_str(e.get("code"))
            if name and code:
                self.emp.setdefault(norm_name(name), []).append((code, name))
            dept = code_str(e.get("dept") or e.get("dept_code") or e.get("dept_number"))
            if name and dept:
                bucket = self.emp_dept.setdefault(norm_name(name), [])
                if dept not in bucket:
                    bucket.append(dept)
        for d in data.get("department") or []:
            code = code_str(d.get("code"))
            name = str(d.get("name") or "").strip()
            if code:
                self.dept[code] = name
        for c in data.get("customer") or []:
            name = str(c.get("name") or "").strip()
            code = code_str(c.get("code"))
            if name and code:
                self.cus.setdefault(norm_name(name), []).append((code, name))
        for s in data.get("supplier") or []:
            name = str(s.get("name") or "").strip()
            code = code_str(s.get("code"))
            if name and code:
                self.sup.setdefault(norm_name(name), []).append((code, name))
                stripped = strip_ge(name)
                if stripped != name:
                    self.sup.setdefault(norm_name(stripped), []).append((code, name))
        self.enabled = bool(self.emp or self.dept or self.cus or self.sup)

    def all_customers(self) -> list[tuple[str, str]]:
        out = []
        for pairs in self.cus.values():
            for item in pairs:
                if item not in out:
                    out.append(item)
        return out

    def all_suppliers(self) -> list[tuple[str, str]]:
        out = []
        for pairs in self.sup.values():
            for item in pairs:
                if item not in out:
                    out.append(item)
        return out

    def customer(self, name: str):
        return unique_lookup(self.cus, name)

    def match_customer(self, name: str, applicant: str = "", alias_map: dict | None = None):
        records = self.all_customers()
        if names.maps_to_police_ministry(name, applicant):
            hit = names.by_code(records, names.POLICE_CODE)
            return (hit, None) if hit else (None, "none")
        if names.is_person_heading(name):
            hit = names.by_code(records, names.PERSON_CODE)
            return (hit, None) if hit else (None, "none")
        alias_map = alias_map or {}
        query = alias_map.get(name, name)
        got = names.match_records(query, records)
        if got.status == "ok":
            return got.hit, None
        if query != name:
            got = names.match_records(name, records)
            if got.status == "ok":
                return got.hit, None
        if got.status == "many":
            return None, "many"
        return None, "none"

    def employee(self, name: str):
        return unique_lookup(self.emp, name)

    def match_employee(self, name: str, hang: dict | None = None):
        mapped = names.hang_employee(name, hang or {})
        return unique_lookup(self.emp, mapped)

    def employee_dept_code(self, name: str, hang: dict | None = None) -> str:
        mapped = names.hang_employee(name, hang or {})
        ehit, eerr = unique_lookup(self.emp, mapped)
        if not ehit or eerr == "many":
            return ""
        depts = list(self.emp_dept.get(norm_name(mapped)) or [])
        if len(depts) != 1:
            return ""
        return depts[0]

    def department(self, code: str):
        code = code_str(code)
        if not code:
            return None, "缺部门编码"
        name = self.dept.get(code)
        if not name:
            return None, "部门档案没有该编码"
        return (code, name), None

    def supplier_fuzzy(self, name: str):
        got = names.match_records(name, self.all_suppliers())
        if got.status == "ok":
            return got.hit, None
        stripped = strip_ge(name)
        if stripped != name:
            got = names.match_records(stripped, self.all_suppliers())
            if got.status == "ok":
                return got.hit, None
        if got.status == "many":
            return None, "many"
        keys = [name, stripped]
        seen = []
        for k in keys:
            hits = self.sup.get(norm_name(k)) or []
            for h in hits:
                if h not in seen:
                    seen.append(h)
        if len(seen) == 1:
            return seen[0], None
        if not seen:
            return None, "none"
        return None, "many"


def assign_vouchers(lines: list[VoucherLine], pack_size: int, keep_consecutive: bool = True) -> None:
    bookable = [x for x in lines if x.status == "可入账"]
    if not keep_consecutive:
        for batch, start in enumerate(range(0, len(bookable), pack_size), start=1):
            for line in bookable[start : start + pack_size]:
                line.voucher_no = batch
        return
    batch = 0
    current: list[VoucherLine] = []
    for line in bookable:
        if not current:
            current = [line]
            continue
        same = norm_name(line.key) == norm_name(current[-1].key)
        same_day = (line.booking_date or "") == (current[-1].booking_date or "")
        if not same_day:
            batch += 1
            for item in current:
                item.voucher_no = batch
            current = [line]
            continue
        if same or len(current) < pack_size:
            current.append(line)
            continue
        batch += 1
        for item in current:
            item.voucher_no = batch
        current = [line]
    if current:
        batch += 1
        for item in current:
            item.voucher_no = batch


def _write_aux(ws, row: int, col: int, value) -> None:
    if value is None or value == "":
        return
    text = str(value).strip()
    if not text or text.startswith("="):
        return
    ws.cell(row, col, value)


def col_by_label(ws, needle: str) -> int:
    for cell in ws[3]:
        if needle in str(cell.value or ""):
            return cell.column
    raise KeyError(needle)


def write_kingdee(path: Path, lines: list[VoucherLine], rules: dict, booking: str, template: Path) -> Path:
    shutil.copy2(template, path)
    wb = load_workbook(path)
    if KINGDEE_SHEET not in wb.sheetnames:
        wb.close()
        raise SystemExit("金蝶空模缺 sheet1（名称勿改）")
    ws = wb[KINGDEE_SHEET]
    cols = {
        "date": col_by_label(ws, "*记账日期"),
        "word": col_by_label(ws, "*凭证字号"),
        "number": col_by_label(ws, "凭证号 #"),
        "expl": col_by_label(ws, "摘要 #"),
        "account": col_by_label(ws, "*科目.编码"),
        "account_name": col_by_label(ws, "科目.名称"),
        "currency": col_by_label(ws, "*币别.编码"),
        "currency_name": col_by_label(ws, "币别.名称"),
        "rate": col_by_label(ws, "*汇率"),
        "amountfor": col_by_label(ws, "原币金额"),
        "debit": col_by_label(ws, "借方 #"),
        "credit": col_by_label(ws, "贷方 #"),
        "cus_code": col_by_label(ws, "辅助核算.客户.编码"),
        "cus_name": col_by_label(ws, "辅助核算.客户.名称"),
        "dep_code": col_by_label(ws, "辅助核算.部门.编码"),
        "dep_name": col_by_label(ws, "辅助核算.部门.名称"),
        "emp_code": col_by_label(ws, "辅助核算.职员.编码"),
        "emp_name": col_by_label(ws, "辅助核算.职员.名称"),
        "sup_code": col_by_label(ws, "辅助核算.供应商.编码"),
        "sup_name": col_by_label(ws, "辅助核算.供应商.名称"),
    }
    names = rules.get("account_names") or {}
    row_i = 4
    for line in lines:
        if line.status != "可入账":
            continue
        for ent in line.entries:
            ws.cell(row_i, cols["date"], line.booking_date or booking)
            ws.cell(row_i, cols["word"], rules.get("voucher_word") or "记")
            ws.cell(row_i, cols["number"], line.voucher_no)
            ws.cell(row_i, cols["expl"], line.expl)
            ws.cell(row_i, cols["account"], ent["account"])
            acc_name = names.get(str(ent["account"]), "")
            if acc_name:
                ws.cell(row_i, cols["account_name"], acc_name)
            ws.cell(row_i, cols["currency"], rules.get("currency") or "RMB")
            if rules.get("currency_name"):
                ws.cell(row_i, cols["currency_name"], rules.get("currency_name"))
            ws.cell(row_i, cols["rate"], float(rules.get("exchange_rate") or 1))
            debit = ent.get("debit")
            credit = ent.get("credit")
            amt = debit if debit is not None else credit
            if amt is not None:
                ws.cell(row_i, cols["amountfor"], float(amt))
            if debit is not None:
                ws.cell(row_i, cols["debit"], float(debit))
            if credit is not None:
                ws.cell(row_i, cols["credit"], float(credit))
            if ent.get("aux"):
                _write_aux(ws, row_i, cols["cus_code"], ent.get("cus_code"))
                _write_aux(ws, row_i, cols["cus_name"], ent.get("cus_name"))
                _write_aux(ws, row_i, cols["dep_code"], ent.get("dep_code"))
                _write_aux(ws, row_i, cols["dep_name"], ent.get("dep_name"))
                _write_aux(ws, row_i, cols["emp_code"], ent.get("emp_code"))
                _write_aux(ws, row_i, cols["emp_name"], ent.get("emp_name"))
                _write_aux(ws, row_i, cols["sup_code"], ent.get("sup_code"))
                _write_aux(ws, row_i, cols["sup_name"], ent.get("sup_name"))
            row_i += 1
    wb.save(path)
    wb.close()
    return path


def write_detail(path: Path, lines: list[VoucherLine], scene: str) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "明细"
    ws.append(["状态", "原因", "源行号", "摘要", "凭证批次", "模块", "附加"])
    for line in lines:
        extra = json.dumps(line.extra, ensure_ascii=False) if line.extra else ""
        ws.append([line.status, line.reason, line.source_row, line.expl, line.voucher_no, scene, extra])
    wb.save(path)
    wb.close()
    return path


def sales_sheet(wb) -> str | None:
    for name in wb.sheetnames:
        if name in ("发票", "数电票-专票"):
            return name
    for name in wb.sheetnames:
        ws = wb[name]
        headers = [str(c.value).strip() for c in next(ws.iter_rows(min_row=1, max_row=1)) if c.value]
        if "单位名称" in headers and "价税合计" in headers:
            return name
    return None


def resolve_sales_party(unit: str, applicant: str, inv_day: str, alias_map: dict, master: Master, assist_rows) -> tuple:
    """特殊规则先走，再抄客户核算项目余额表。档案只核验编码存在。"""
    records = lookup_mod.assist_records(assist_rows)
    if names.maps_to_police_ministry(unit, applicant):
        cus_code = names.POLICE_CODE
    elif names.is_person_heading(unit):
        cus_code = names.PERSON_CODE
    else:
        query = (alias_map or {}).get(unit, unit)
        got = names.match_records(query, records)
        if got.status == "none" and query != unit:
            got = names.match_records(unit, records)
        if got.status == "many":
            return None, "客户核算项目余额表名称不唯一，请斯佳确认", {}
        if got.status != "ok" or not got.hit:
            return None, lookup_mod.HOLD_ASSIST_MISSING, {}
        cus_code = got.hit[0]
    extra = {"候选1131": lookup_mod.list_assist_accounts(cus_code, inv_day, assist_rows)}
    ar, rev, why = lookup_mod.pick_assist_account(cus_code, inv_day, assist_rows)
    if not ar:
        return None, why or lookup_mod.HOLD_ASSIST_MISSING, extra
    verified = names.by_code(master.all_customers(), cus_code)
    if not verified:
        return None, "总部档案未核验" if not master.enabled else "客户档案没有此编码，请斯佳确认", extra
    return (verified[0], verified[1], ar, rev), None, extra


def convert_sales(path: Path, master: Master, rules: dict, aliases: dict, box, booking: str, period_fetch=None, assist_rows=None) -> list[VoucherLine]:
    cfg = rules.get("sales") or {}
    tax_account = str(cfg.get("tax_account") or "21710105")
    divisor = Decimal(str(rules.get("tax_rate_divisor") or "1.06"))
    alias_map = load_customer_alias()
    hang = load_emp_hang()
    wb_f = load_workbook(path, data_only=False)
    wb_v = load_workbook(path, data_only=True)
    sheet = sales_sheet(wb_f)
    if not sheet:
        wb_f.close()
        wb_v.close()
        raise SystemExit("找不到发票 sheet")
    ws_f = wb_f[sheet]
    ws_v = wb_v[sheet]
    org_map = inspect_mod.load_org_map(wb_f, aliases)
    headers = [c.value for c in next(ws_f.iter_rows(min_row=1, max_row=1))]
    idx = header_index(headers, aliases.get("销项发票_列别名") or {})
    needed = ["单位名称", "价税合计", "申请人"]
    missing = [k for k in needed if k not in idx]
    if missing:
        wb_f.close()
        wb_v.close()
        raise SystemExit("发票表缺列：" + "、".join(missing))
    lines: list[VoucherLine] = []
    max_row = ws_f.max_row or 1
    for r in range(2, max_row + 1):
        row_f = [ws_f.cell(r, c + 1).value for c in range(len(headers))]
        row_v = [ws_v.cell(r, c + 1).value for c in range(len(headers))]
        unit = str(cell_at(row_f, idx, "单位名称") or "").strip()
        if not unit:
            continue
        total = pick_amount(cell_at(row_f, idx, "价税合计"), cell_at(row_v, idx, "价税合计"), None)
        amt = pick_amount(cell_at(row_f, idx, "金额"), cell_at(row_v, idx, "金额"), None)
        tax = pick_amount(cell_at(row_f, idx, "税额"), cell_at(row_v, idx, "税额"), None)
        if total is not None and amt is None:
            amt = (total / divisor).quantize(TWOPLACES, rounding=ROUND_HALF_UP)
        if total is not None and amt is not None and tax is None:
            tax = (total - amt).quantize(TWOPLACES, rounding=ROUND_HALF_UP)
        typ = str(cell_at(row_f, idx, "发票类型") or "").strip()
        applicant = str(cell_at(row_f, idx, "申请人") or "").strip()
        inv_day = as_day(cell_at(row_v, idx, "日期") or cell_at(row_f, idx, "日期"), booking)
        prefix = typ
        if total is not None and total < 0 and "红字" not in typ:
            prefix = "红字" + typ
        line = VoucherLine(
            status="可入账",
            source_row=r,
            key=unit,
            expl=f"{prefix}：{unit}" if typ else f"{unit}",
            extra={"单位名称": unit, "申请人": applicant, "sheet": sheet, "source_amt": str(total) if total is not None else ""},
        )
        if total is None or amt is None or tax is None:
            line.status, line.reason = "待确认", "缺金额"
            lines.append(line)
            continue
        if total != amt + tax:
            line.status, line.reason = "待确认", "价税合计不等于金额加税额"
            lines.append(line)
            continue
        if not typ:
            line.status, line.reason = "待确认", "缺发票类型"
            lines.append(line)
            continue
        if not applicant:
            line.status, line.reason = "待确认", "缺申请人"
            lines.append(line)
            continue
        dept = inspect_mod.dept_for_sales(
            applicant, org_map, box.applicant_dept, master.employee_dept_code(applicant, hang)
        )
        if not dept:
            line.status, line.reason = "待确认", "申请人不在部门表"
            lines.append(line)
            continue
        dhit, derr = master.department(dept)
        if not dhit:
            line.status, line.reason = "待确认", derr or "部门档案未核验"
            lines.append(line)
            continue
        dept, dep_name = dhit
        ehit, eerr = master.match_employee(applicant, hang)
        if not ehit:
            line.status, line.reason = "待确认", "职员档案一对多" if eerr == "many" else "职员档案没有此人"
            lines.append(line)
            continue
        emp_code, emp_name = ehit
        rows = assist_rows if assist_rows is not None else list(getattr(box, "assist_rows", None) or [])
        party, perr, meta = resolve_sales_party(unit, applicant, inv_day, alias_map, master, rows)
        if meta.get("候选1131"):
            line.extra["候选1131"] = meta["候选1131"]
        if not party:
            line.status = "待确认"
            line.reason = perr or lookup_mod.HOLD_ASSIST_MISSING
            lines.append(line)
            continue
        cus_code, cus_name, ar, rev = party
        aux = {
            "aux": True,
            "cus_code": cus_code,
            "cus_name": cus_name,
            "dep_code": dept,
            "dep_name": dep_name,
            "emp_code": emp_code,
            "emp_name": emp_name,
        }
        line.entries = [
            {"account": ar, "debit": total, **aux},
            {"account": rev, "credit": amt, **aux},
            {"account": tax_account, "credit": tax},
        ]
        lines.append(line)
    wb_f.close()
    wb_v.close()
    assign_vouchers(lines, int(cfg.get("pack_size") or 5), bool(cfg.get("pack_keep_consecutive", True)))
    return lines


def parse_invoice_text(text: str) -> dict:
    kind = ""
    if "专用发票" in (text or ""):
        kind = "专票"
    elif "普通发票" in (text or ""):
        kind = "普票"
    seller = ""
    for key in ("销售方名称：", "销售方名称:", "名称：", "名称:"):
        if key in (text or ""):
            seller = text.split(key, 1)[1].splitlines()[0].strip()
            if seller:
                break
    total = None
    tax = None
    m = re.search(r"价税合计.*?[¥￥]\s*([-0-9,]+\.\d{2})", text or "", re.S)
    if m:
        total = money(m.group(1))
    m = re.search(r"(?:合计)?税额[:：]?\s*[¥￥]?\s*([-0-9,]+\.\d{2})", text or "")
    if m:
        tax = money(m.group(1))
    return {"kind": kind, "seller": seller, "total": total, "tax": tax}


def parse_invoice_pdf(path: Path) -> dict:
    empty = {"kind": "", "seller": "", "total": None, "tax": None}
    try:
        import pdfplumber

        with pdfplumber.open(str(path)) as pdf:
            text = "\n".join((page.extract_text() or "") for page in pdf.pages[:4])
    except Exception:
        return empty
    return parse_invoice_text(text)


def convert_payment(root: Path, ledger: Path, master: Master, rules: dict, aliases: dict) -> list[VoucherLine]:
    cfg = rules.get("payment") or {}
    fallback = rules.get("fallback_supplier") or {"code": "9999", "name": "其他供应商"}
    wb_f = load_workbook(ledger, data_only=False)
    wb_v = load_workbook(ledger, data_only=True)
    sheet = inspect_mod.find_payment_sheet(wb_f, aliases)
    if not sheet:
        wb_f.close()
        wb_v.close()
        raise SystemExit("找不到付款三列表")
    ws_f = wb_f[sheet]
    ws_v = wb_v[sheet] if sheet in wb_v.sheetnames else wb_v[wb_v.sheetnames[0]]
    headers = [c.value for c in next(ws_f.iter_rows(min_row=1, max_row=1))]
    idx = header_index(headers, aliases.get("付款_列别名") or {})
    if "供应商" not in idx or "应付金额本币" not in idx:
        wb_f.close()
        wb_v.close()
        raise SystemExit("付款表缺供应商或应付金额本币")
    lines: list[VoucherLine] = []
    max_row = ws_f.max_row or 1
    dirs = {p.name: p for p in root.iterdir() if p.is_dir()}
    for r in range(2, max_row + 1):
        row_f = [ws_f.cell(r, c + 1).value for c in range(len(headers))]
        row_v = [ws_v.cell(r, c + 1).value for c in range(len(headers))]
        vendor = str(cell_at(row_f, idx, "供应商") or "").strip()
        if not vendor:
            continue
        payable = pick_amount(cell_at(row_f, idx, "应付金额本币"), cell_at(row_v, idx, "应付金额本币"), None)
        line = VoucherLine(
            status="可入账",
            source_row=r,
            key=vendor,
            extra={"供应商": vendor, "sheet": sheet, "source_amt": str(payable) if payable is not None else ""},
        )
        folder = dirs.get(vendor)
        if folder is None:
            line.status, line.reason = "待确认", "缺这家发票夹"
            lines.append(line)
            continue
        pdfs = list(folder.glob("*.pdf")) + list(folder.glob("*.PDF"))
        if not pdfs:
            line.status, line.reason = "待确认", "夹里没有发票 PDF"
            lines.append(line)
            continue
        metas = [parse_invoice_pdf(p) for p in pdfs]
        if any(
            (not m.get("kind")) or (not str(m.get("seller") or "").strip()) or m.get("total") is None
            for m in metas
        ):
            line.status, line.reason = "待确认", "发票缺票种或销售方或金额"
            lines.append(line)
            continue
        kinds = {m.get("kind") for m in metas if m.get("kind")}
        if len(kinds) != 1:
            line.status, line.reason = "待确认", "认不清专票还是普票"
            lines.append(line)
            continue
        kind = kinds.pop()
        sellers = [m.get("seller") or "" for m in metas if m.get("seller")]
        seller = sellers[0] if sellers else vendor
        if payable is None:
            line.status, line.reason = "待确认", "缺应付金额"
            lines.append(line)
            continue
        ticket_total = Decimal("0")
        ticket_tax = Decimal("0")
        for m in metas:
            if m.get("total") is not None:
                ticket_total += m["total"]
            if m.get("tax") is not None:
                ticket_tax += m["tax"]
        if ticket_total and ticket_total < payable:
            line.status, line.reason = "待确认", "票合计小于应付"
            lines.append(line)
            continue
        hit, why = master.supplier_fuzzy(vendor)
        if why == "none":
            hit, why = master.supplier_fuzzy(strip_ge(vendor))
        if why == "none":
            hit, why = master.supplier_fuzzy(seller)
        if hit:
            sup_code, sup_name = hit
        else:
            fallback_hit, fallback_why = master.supplier_fuzzy(str(fallback.get("name") or ""))
            if not fallback_hit or fallback_why:
                line.status, line.reason = "待确认", "其他供应商档案未核验"
                lines.append(line)
                continue
            sup_code, sup_name = fallback_hit
        dhit, derr = master.department(str(cfg.get("dept_code") or ""))
        if not dhit:
            line.status, line.reason = "待确认", derr or "付款部门档案未核验"
            lines.append(line)
            continue
        dep_code, dep_name = dhit
        ehit, eerr = master.employee(str(cfg.get("emp_name") or ""))
        if not ehit:
            line.status, line.reason = "待确认", "职员档案一对多" if eerr == "many" else "付款职员档案未核验"
            lines.append(line)
            continue
        emp_code, emp_name = ehit
        aux = {
            "aux": True,
            "sup_code": sup_code,
            "sup_name": sup_name,
            "dep_code": dep_code,
            "dep_name": dep_name,
            "emp_code": emp_code,
            "emp_name": emp_name,
        }
        cost_acc = str(cfg.get("cost_account") or "540103")
        bank_acc = str(cfg.get("bank_account") or "100206")
        tax_acc = str(cfg.get("input_tax_account") or "21710101")
        line.expl = f"付：{seller}"
        if kind == "普票":
            line.entries = [
                {"account": cost_acc, "debit": payable, **aux},
                {"account": bank_acc, "credit": payable},
            ]
        else:
            tax = ticket_tax if ticket_tax else None
            if tax is None or tax == 0:
                line.status, line.reason = "待确认", "认不清进项税额"
                lines.append(line)
                continue
            if tax >= payable:
                line.status, line.reason = "待确认", "进项不小于应付"
                lines.append(line)
                continue
            cost = (payable - tax).quantize(TWOPLACES, rounding=ROUND_HALF_UP)
            line.entries = [
                {"account": cost_acc, "debit": cost, **aux},
                {"account": tax_acc, "debit": tax},
                {"account": bank_acc, "credit": payable},
            ]
        lines.append(line)
    wb_f.close()
    wb_v.close()
    assign_vouchers(lines, 1, False)
    for i, line in enumerate([x for x in lines if x.status == "可入账"], start=1):
        line.voucher_no = i
    return lines


def match_receipt_customer(master: Master, name: str, alias_map: dict, assist_rows) -> tuple:
    records = master.all_customers()
    query = alias_map.get(name, name) if alias_map else name
    if names.is_haidian_police(name):
        got = names.match_records(query, records)
        if got.status == "ok":
            return got.hit, None
        if query != name:
            got = names.match_records(name, records)
            if got.status == "ok":
                return got.hit, None
        return None, "many" if got.status == "many" else "none"
    chit, cerr = master.match_customer(name, "", alias_map)
    if chit or cerr == "many":
        return chit, cerr
    assist_got = names.match_records(query, lookup_mod.assist_records(assist_rows))
    if assist_got.status == "ok" and assist_got.hit:
        hit = names.by_code(records, assist_got.hit[0])
        return (hit, None) if hit else (None, "none")
    if query != name:
        assist_got = names.match_records(name, lookup_mod.assist_records(assist_rows))
        if assist_got.status == "ok" and assist_got.hit:
            hit = names.by_code(records, assist_got.hit[0])
            return (hit, None) if hit else (None, "none")
    return None, cerr


def convert_receipt(
    path: Path,
    master: Master,
    rules: dict,
    aliases: dict,
    box,
    booking: str,
    period_fetch=None,
    assist_rows=None,
) -> list[VoucherLine]:
    cfg = rules.get("receipt") or {}
    alias_map = load_customer_alias()
    hang = load_emp_hang()
    applicant_dept = load_applicant_dept()
    wb_f = load_workbook(path, data_only=False)
    wb_v = load_workbook(path, data_only=True)
    sheet = inspect_mod.find_receipt_sheet(wb_f, aliases)
    if not sheet:
        wb_f.close()
        wb_v.close()
        raise SystemExit("找不到收款 sheet")
    ws_f = wb_f[sheet]
    ws_v = wb_v[sheet] if sheet in wb_v.sheetnames else wb_v[wb_v.sheetnames[0]]
    headers = [c.value for c in next(ws_f.iter_rows(min_row=1, max_row=1))]
    idx = header_index(headers, aliases.get("收款_列别名") or {})
    needed = ["日期", "客户名称", "借方（增加）"]
    missing = [k for k in needed if k not in idx]
    if missing:
        wb_f.close()
        wb_v.close()
        raise SystemExit("收款表缺列：" + "、".join(missing))
    org_map = inspect_mod.load_org_map(wb_f, aliases)
    rows = list(assist_rows) if assist_rows is not None else lookup_mod.rows_from_balance_box(box)
    bank = str(cfg.get("bank_account") or "100201")
    lines: list[VoucherLine] = []
    max_row = ws_f.max_row or 1
    for r in range(2, max_row + 1):
        row_f = [ws_f.cell(r, c + 1).value for c in range(len(headers))]
        row_v = [ws_v.cell(r, c + 1).value for c in range(len(headers))]
        cust = str(cell_at(row_f, idx, "客户名称") or "").strip()
        if not cust:
            continue
        amt = pick_amount(cell_at(row_f, idx, "借方（增加）"), cell_at(row_v, idx, "借方（增加）"), None)
        rec_day = as_day(cell_at(row_v, idx, "日期") or cell_at(row_f, idx, "日期"), booking)
        table_dept = pick_code(cell_at(row_f, idx, "部门编码"), cell_at(row_v, idx, "部门编码")) if "部门编码" in idx else ""
        table_sales = str(cell_at(row_f, idx, "销售") or "").strip() if "销售" in idx else ""
        line = VoucherLine(
            status="可入账",
            source_row=r,
            key=cust,
            expl=f"收：{cust}",
            booking_date=rec_day,
            extra={"客户名称": cust, "sheet": sheet, "source_amt": str(amt) if amt is not None else ""},
        )
        if amt is None:
            line.status, line.reason = "待确认", "缺金额"
            lines.append(line)
            continue
        chit, cerr = match_receipt_customer(master, cust, alias_map, rows)
        if not chit:
            line.status = "待确认"
            line.reason = "客户档案一对多" if cerr == "many" else "客户档案没有此抬头，请斯佳确认是否新建"
            lines.append(line)
            continue
        cus_code, cus_name = chit
        ar_cands = lookup_mod.list_assist_accounts(cus_code, rec_day, rows)
        if ar_cands:
            line.extra["候选1131"] = ar_cands
        ar, _rev, why = lookup_mod.pick_assist_account(cus_code, rec_day, rows)
        if not ar:
            line.status, line.reason = "待确认", why or lookup_mod.HOLD_ASSIST_MISSING
            lines.append(line)
            continue
        sales_cands: list[str] = []
        if table_sales:
            sales, swhy = table_sales, ""
        else:
            sales, swhy, sales_cands = lookup_mod.resolve_sales_any_detail(
                [cus_name, cust, alias_map.get(cust, cust)], rec_day, amt, box
            )
        if sales_cands:
            line.extra["候选销售"] = sales_cands
        if not sales:
            line.status, line.reason = "待确认", swhy or "找不到销售"
            lines.append(line)
            continue
        line.extra["销售"] = sales
        emp_dept = master.employee_dept_code(sales, hang)
        dept = table_dept if table_dept and not table_dept.startswith("=") else inspect_mod.dept_for_sales(
            sales, org_map, applicant_dept, emp_dept
        )
        if not dept:
            line.status, line.reason = "待确认", "没有销售就没有部门" if not sales else "销售不在部门表"
            lines.append(line)
            continue
        dhit, derr = master.department(dept)
        if not dhit:
            line.status, line.reason = "待确认", derr or "部门档案未核验"
            lines.append(line)
            continue
        dept, dep_name = dhit
        ehit, eerr = master.match_employee(sales, hang)
        emp_code = emp_name = ""
        if ehit:
            emp_code, emp_name = ehit
        elif eerr == "many":
            line.status, line.reason = "待确认", "职员档案一对多"
            lines.append(line)
            continue
        aux = {
            "aux": True,
            "cus_code": cus_code,
            "cus_name": cus_name,
            "dep_code": dept,
            "dep_name": dep_name,
            "emp_code": emp_code,
            "emp_name": emp_name,
        }
        line.entries = [
            {"account": bank, "debit": amt},
            {"account": ar, "credit": amt, **aux},
        ]
        lines.append(line)
    wb_f.close()
    wb_v.close()
    assign_vouchers(lines, int(cfg.get("pack_size") or 10), bool(cfg.get("pack_keep_consecutive", True)))
    return lines


def default_desktop_dir(prefix: str = "金蝶入账", today: date | None = None, home: Path | None = None) -> Path:
    day = (today or date.today()).strftime("%Y%m%d")
    root = Path(home) if home else Path.home()
    desktop = root / "Desktop"
    base = desktop if desktop.is_dir() else Path.cwd()
    path = base / f"{prefix}_{day}"
    path.mkdir(parents=True, exist_ok=True)
    return path


def shift_voucher_numbers(lines: list[VoucherLine], start_no: int | None) -> None:
    if not start_no or int(start_no) == 1:
        return
    start = int(start_no)
    if start < 1:
        raise SystemExit("凭证号起始必须是正整数")
    delta = start - 1
    for line in lines:
        if line.status == "可入账" and line.voucher_no is not None:
            line.voucher_no = int(line.voucher_no) + delta


def reason_counts(lines) -> list[tuple[str, int]]:
    c = Counter()
    for line in lines or []:
        status = getattr(line, "status", "")
        reason = str(getattr(line, "reason", "") or "").strip()
        if status != "可入账" and reason:
            c[reason] += 1
    return sorted(c.items(), key=lambda x: (-x[1], x[0]))


def first_sheet_name(lines) -> str:
    for line in lines or []:
        sheet = (getattr(line, "extra", None) or {}).get("sheet")
        if sheet:
            return str(sheet)
    return ""


def tieout_amounts(lines) -> tuple[Decimal, Decimal, Decimal]:
    src = Decimal("0")
    debit = Decimal("0")
    credit = Decimal("0")
    for line in lines or []:
        if getattr(line, "status", "") != "可入账":
            continue
        amt = money((getattr(line, "extra", None) or {}).get("source_amt"))
        if amt is not None:
            src += amt
        for ent in getattr(line, "entries", None) or []:
            d = money(ent.get("debit"))
            c = money(ent.get("credit"))
            if d is not None:
                debit += d
            if c is not None:
                credit += c
    return src, debit, credit


def write_note(
    path: Path,
    *,
    scene: str,
    source_count: int,
    bookable_count: int,
    hold_count: int,
    sheet: str,
    start_voucher_no: int | None,
    assist_name: str,
    reasons: list[tuple[str, int]],
    extras: list[str] | None = None,
) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    title = {"销项发票": "销项发票", "付款": "付款", "收款": "收款"}.get(scene, scene)
    lines = [
        f"# {title}入金蝶对照说明",
        "",
        f"- 源：{source_count}",
        f"- 可入账：{bookable_count}",
        f"- 待确认：{hold_count}",
        f"- 读了哪个 sheet：{sheet or '（未记）'}",
        f"- 起始凭证号：{start_voucher_no if start_voucher_no else '（未取）'}",
        f"- 余额表：{assist_name or '（本批未用文件名）'}",
        "- 原因类别：",
    ]
    if reasons:
        for reason, n in reasons:
            lines.append(f"  - {reason}：{n}")
    else:
        lines.append("  - （无）")
    for extra in extras or []:
        lines.append(f"- {extra}")
    lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def write_outputs(
    out_dir: Path,
    scene: str,
    lines: list[VoucherLine],
    rules: dict,
    booking: str,
    stem: str,
    start_voucher_no: int | None = None,
    assist_name: str = "",
    extras: list[str] | None = None,
):
    out_dir.mkdir(parents=True, exist_ok=True)
    detail = out_dir / f"{stem}_明细结果.xlsx"
    kingdee = out_dir / "凭证引入_结果.xlsx"
    write_detail(detail, lines, scene)
    write_kingdee(kingdee, lines, rules, booking, TEMPLATE)
    bookable = sum(1 for x in lines if x.status == "可入账")
    hold = sum(1 for x in lines if x.status != "可入账")
    sheet = first_sheet_name(lines)
    note = write_note(
        out_dir / "对照说明.md",
        scene=scene,
        source_count=len(lines),
        bookable_count=bookable,
        hold_count=hold,
        sheet=sheet,
        start_voucher_no=start_voucher_no,
        assist_name=assist_name,
        reasons=reason_counts(lines),
        extras=extras,
    )
    src_amt, debit, credit = tieout_amounts(lines)
    return {
        "scene": scene,
        "source_count": len(lines),
        "bookable_count": bookable,
        "hold_count": hold,
        "kingdee_path": str(kingdee),
        "detail_path": str(detail),
        "note_path": str(note),
        "out_dir": str(out_dir),
        "sheet": sheet,
        "start_voucher_no": start_voucher_no,
        "assist_name": assist_name,
        "tieout_source": str(src_amt),
        "tieout_debit": str(debit),
        "tieout_credit": str(credit),
    }


def _assist_rows_for_run(input_dir: Path, box, required: bool) -> tuple[list, str]:
    if box.assist_supplied:
        return list(box.assist_rows or []), ""
    found = assist_mod.find_assist_xlsx([input_dir])
    if found:
        return assist_mod.parse_assist_xlsx(found), found.name
    if required:
        path = assist_mod.ensure_assist_xlsx([input_dir])
        return assist_mod.parse_assist_xlsx(path), path.name
    try:
        path = assist_mod.ensure_assist_xlsx([input_dir])
        return assist_mod.parse_assist_xlsx(path), path.name
    except SystemExit as e:
        log(str(e))
        return lookup_mod.rows_from_balance_box(box), ""


def run_dir(
    input_dir: Path,
    scene: str | None,
    booking: str | None,
    master_data: dict | None,
    lookups: dict | None = None,
    period_fetch=None,
    start_voucher_no: int = 1,
    out_dir: Path | None = None,
    ar_xlsx: Path | None = None,
) -> dict:
    report = inspect_mod.inspect_dir(input_dir, scene)
    if not report.get("ready"):
        raise SystemExit(report.get("ask") or "材料不齐")
    scene = str(report["scene"])
    rules = load_rules()
    aliases = load_aliases()
    master = Master(master_data or {})
    box = load_box(lookups)
    day = booking or date.today().isoformat()
    files = report["files"]
    assist_name = ""
    extras: list[str] = []
    if scene == "销项发票":
        src = Path(files["invoice"])
        assist_rows, assist_name = _assist_rows_for_run(input_dir, box, required=True)
        extras.append("销项部门：组织架构优先，否则申请人部门.json，再职员档案唯一部门")
        extras.append("记账日：跑批当天")
        lines = convert_sales(src, master, rules, aliases, box, day, assist_rows=assist_rows)
    elif scene == "付款":
        src = Path(files["ledger"])
        extras.append("记账日：跑批当天")
        extras.append("付款按表头认三列表，不读中行付款")
        lines = convert_payment(input_dir, src, master, rules, aliases)
    else:
        src = Path(files["receipt"])
        if ar_xlsx:
            assist_rows = assist_mod.parse_assist_xlsx(Path(ar_xlsx))
            assist_name = Path(ar_xlsx).name
        else:
            assist_rows, assist_name = _assist_rows_for_run(input_dir, box, required=False)
        extras.append("记账日：表上收款日")
        extras.append("表上销售有则用，空则智云回款→下单；部门可回退职员档案")
        lines = convert_receipt(
            src, master, rules, aliases, box, day, period_fetch=period_fetch, assist_rows=assist_rows
        )
    shift_voucher_numbers(lines, start_voucher_no)
    prefix = SCENE_DESKTOP.get(scene, "金蝶入账")
    dest = Path(out_dir) if out_dir else default_desktop_dir(prefix)
    return write_outputs(
        dest,
        scene,
        lines,
        rules,
        day,
        src.stem,
        start_voucher_no=start_voucher_no,
        assist_name=assist_name,
        extras=extras,
    )


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="金蝶入账")
    parser.add_argument("--inspect", action="store_true")
    parser.add_argument("--input-dir", required=True)
    parser.add_argument("--scene", choices=["销项发票", "付款", "收款"])
    parser.add_argument("--date")
    parser.add_argument("--master")
    parser.add_argument("--lookups")
    parser.add_argument("--no-api", action="store_true")
    parser.add_argument("--start-voucher-no", type=int, default=None)
    parser.add_argument("--out-dir", "--out", dest="out_dir")
    args = parser.parse_args(argv)
    root = Path(args.input_dir)
    if args.inspect:
        return inspect_mod.main(["--input-dir", str(root)] + (["--scene", args.scene] if args.scene else []))
    master_data = None
    if args.master:
        master_data = json.loads(Path(args.master).read_text(encoding="utf-8"))
    elif args.no_api:
        log("本次入账必须先读取总部当前档案；--no-api 只可用于开发排查，未生成引入表。")
        return 2
    else:
        loaded = kingdee_api.try_load_master()
        if loaded.get("ok"):
            master_data = loaded["data"]
        else:
            reason = "本机没有金蝶应用号" if loaded.get("missing_credentials") else loaded.get("error") or "读取失败"
            log(f"总部档案未核验（{reason}）；未生成引入表。请检查本机应用号和只读权限后重试。")
            return 2
    lookups = {}
    scene = args.scene
    if not scene:
        inspected = inspect_mod.inspect_dir(root, None)
        scene = inspected.get("scene") if inspected.get("ready") else None
    if args.lookups:
        lookups = json.loads(Path(args.lookups).read_text(encoding="utf-8"))
    elif scene == "收款":
        loaded_zy = zhiyun_api.try_load_lookups()
        if loaded_zy.get("ok"):
            lookups = loaded_zy.get("data") or {}
        else:
            reason = "本机没有智云账号" if loaded_zy.get("missing_credentials") else loaded_zy.get("error") or "读取失败"
            log(f"智云查找未核验（{reason}）；缺销售的行将待确认，其余仍出表。")
    start_no = args.start_voucher_no
    if start_no is None:
        period = (args.date or date.today().isoformat())[:7]
        fetched = kingdee_api.try_fetch_next_voucher_no(period)
        if not fetched.get("ok"):
            reason = "本机没有金蝶应用号" if fetched.get("missing_credentials") else fetched.get("error") or "读取失败"
            log(f"当前月凭证号未核验（{reason}）；未生成引入表。请检查本机应用号后重试。")
            return 2
        start_no = int(fetched["next_number"])

    try:
        result = run_dir(
            root,
            args.scene,
            args.date,
            master_data,
            lookups,
            start_voucher_no=start_no,
            out_dir=Path(args.out_dir).expanduser() if args.out_dir else None,
        )
    except SystemExit as e:
        log(str(e))
        return 2
    log(
        f"这批 {result['source_count']}：可入账 {result['bookable_count']}，待确认 {result['hold_count']}。"
        f"填好的金蝶表在 {result['kingdee_path']}。请您看待确认，再自己去金蝶引入。我没有点引入。"
    )
    print(json.dumps({k: result[k] for k in result if k != "lines"}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
