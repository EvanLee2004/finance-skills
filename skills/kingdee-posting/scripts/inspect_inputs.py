#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""按内容认文件夹：销项发票 / 付款 / 收款。缺材料列出人话。"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
CONFIG = SKILL / "config"

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass


def load_aliases() -> dict:
    p = CONFIG / "列名别名.json"
    return json.loads(p.read_text(encoding="utf-8")) if p.is_file() else {}


def code_str(v) -> str:
    if v is None or v == "":
        return ""
    if isinstance(v, bool):
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, int):
        return str(v)
    s = str(v).strip()
    if s.endswith(".0") and s[:-2].isdigit():
        return s[:-2]
    return s


def norm_name(name: str) -> str:
    s = (name or "").strip().replace("(", "（").replace(")", "）")
    return "".join(s.split())


def header_names(path: Path) -> dict[str, list[str]]:
    out: dict[str, list[str]] = {}
    try:
        wb = load_workbook(path, read_only=True, data_only=False)
    except Exception:
        return out
    try:
        for name in wb.sheetnames:
            ws = wb[name]
            row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
            out[name] = [str(c).strip() for c in (row or []) if c is not None and str(c).strip()]
    finally:
        wb.close()
    return out


def field_hit(headers: list[str], aliases: dict, field: str) -> bool:
    names = aliases.get(field) or [field]
    return any(h in names for h in headers)


def _cleaned(headers) -> list[str]:
    return [str(h).strip() for h in (headers or []) if h is not None and str(h).strip()]


def is_sales_sheet(name: str, headers: list[str], aliases: dict | None = None) -> bool:
    title = str(name or "")
    if "组织架构" in title or "流水" in title:
        return False
    if "收款" in title and "发票" not in title:
        return False
    sales_a = (aliases or {}).get("销项发票_列别名") or {}
    cleaned = _cleaned(headers)
    return field_hit(cleaned, sales_a, "单位名称") and field_hit(cleaned, sales_a, "价税合计")


def is_receipt_sheet(name: str, headers: list[str], aliases: dict | None = None) -> bool:
    title = str(name or "")
    if "流水" in title:
        return False
    if "收款" in title and "付款" not in title:
        return True
    rec_a = (aliases or {}).get("收款_列别名") or {}
    cleaned = _cleaned(headers)
    if any("价税合计" in h for h in cleaned):
        return False
    if any(h == "贷方（减少）" or h.startswith("贷方（减少）") for h in cleaned):
        return False
    has_cust = "客户名称" in cleaned or field_hit(cleaned, rec_a, "客户名称")
    has_debit = "借方（增加）" in cleaned or field_hit(cleaned, rec_a, "借方（增加）")
    return has_cust and has_debit


def is_payment_sheet(name: str, headers: list[str], aliases: dict | None = None) -> bool:
    title = str(name or "")
    if "流水" in title or "组织架构" in title:
        return False
    if "中行付款" in title or ("中行" in title and "付款" in title):
        return False
    if "收款" in title:
        return False
    pay_a = (aliases or {}).get("付款_列别名") or {}
    cleaned = _cleaned(headers)
    if any("价税合计" in h for h in cleaned):
        return False
    return field_hit(cleaned, pay_a, "供应商") and field_hit(cleaned, pay_a, "应付金额本币")


def workbook_has_receipt(path: Path, aliases: dict | None = None) -> bool:
    return any(is_receipt_sheet(name, headers, aliases) for name, headers in header_names(path).items())


def _pick_named(named: list[str], headered: list[str]) -> str | None:
    if named:
        return named[0]
    if headered:
        return headered[0]
    return None


def find_receipt_sheet(wb, aliases: dict | None = None) -> str | None:
    named = []
    headered = []
    for name in wb.sheetnames:
        ws = wb[name]
        row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        headers = _cleaned(row)
        if not is_receipt_sheet(name, headers, aliases):
            continue
        if "收款" in str(name) and "流水" not in str(name) and "付款" not in str(name):
            named.append(name)
        else:
            headered.append(name)
    return _pick_named(named, headered)


def find_payment_sheet(wb, aliases: dict | None = None) -> str | None:
    named = []
    headered = []
    for name in wb.sheetnames:
        ws = wb[name]
        row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        headers = _cleaned(row)
        if not is_payment_sheet(name, headers, aliases):
            continue
        if "付款" in str(name) and "中行" not in str(name):
            named.append(name)
        else:
            headered.append(name)
    return _pick_named(named, headered)


def load_org_map(wb, aliases: dict | None = None) -> dict[str, str]:
    org_a = (aliases or {}).get("组织架构_列别名") or {"姓名": ["姓名"], "部门编码": ["部门编码"]}
    chosen = None
    for name in wb.sheetnames:
        if name == "组织架构":
            chosen = name
            break
    if chosen is None:
        for name in wb.sheetnames:
            if "组织架构" in str(name) and "营销" not in str(name):
                chosen = name
                break
    if chosen is None:
        return {}
    ws = wb[chosen]
    rows = list(ws.iter_rows(min_row=1, values_only=True))
    if not rows:
        return {}
    headers = [str(c).strip() if c is not None else "" for c in rows[0]]
    name_i = dept_i = None
    for i, h in enumerate(headers):
        if h in (org_a.get("姓名") or ["姓名"]) and name_i is None:
            name_i = i
        if h in (org_a.get("部门编码") or ["部门编码"]) and dept_i is None:
            dept_i = i
    if name_i is None or dept_i is None:
        return {}
    out: dict[str, str] = {}
    for row in rows[1:]:
        if not row or name_i >= len(row) or dept_i >= len(row):
            continue
        person = str(row[name_i] or "").strip()
        dept = code_str(row[dept_i])
        if person and dept and not dept.startswith("="):
            out[person] = dept
            out[norm_name(person)] = dept
    return out


def dept_for_sales(sales: str, org_map: dict[str, str], applicant_dept: dict[str, str]) -> str:
    raw = str(sales or "").strip()
    if not raw:
        return ""
    return (
        org_map.get(raw)
        or org_map.get(norm_name(raw))
        or applicant_dept.get(raw)
        or applicant_dept.get(norm_name(raw))
        or ""
    )


def classify_xlsx_kinds(path: Path, aliases: dict) -> list[str]:
    if "结果" in path.stem:
        return []
    if "核算项目余额表" in path.name:
        return []
    sheets = header_names(path)
    if not sheets:
        return []
    kinds: list[str] = []
    for name, headers in sheets.items():
        if is_sales_sheet(name, headers, aliases) and "销项发票" not in kinds:
            kinds.append("销项发票")
        if is_payment_sheet(name, headers, aliases) and "付款" not in kinds:
            kinds.append("付款")
        if is_receipt_sheet(name, headers, aliases) and "收款" not in kinds:
            kinds.append("收款")
    return kinds


def classify_xlsx(path: Path, aliases: dict) -> str | None:
    kinds = classify_xlsx_kinds(path, aliases)
    return kinds[0] if kinds else None


def payment_folders(root: Path, ledger: Path | None) -> list[Path]:
    skip = {ledger.resolve()} if ledger else set()
    out = []
    for p in sorted(root.iterdir()):
        if p.is_dir() and p.resolve() not in skip:
            pdfs = list(p.glob("*.pdf")) + list(p.glob("*.PDF"))
            if pdfs:
                out.append(p)
    return out


def inspect_dir(input_dir: Path, scene: str | None = None) -> dict:
    root = Path(input_dir)
    aliases = load_aliases()
    found: dict[str, list[str]] = {"销项发票": [], "付款": [], "收款": []}
    if not root.is_dir():
        return {
            "ready": False,
            "scene": scene,
            "mixed": False,
            "missing": ["材料文件夹"],
            "files": {},
            "ask": "请把当批材料放到一个文件夹里，再告诉我路径。",
        }
    for p in sorted(root.iterdir()):
        if p.suffix.lower() in {".xlsx", ".xlsm"} and not p.name.startswith("~$"):
            for kind in classify_xlsx_kinds(p, aliases):
                if str(p) not in found[kind]:
                    found[kind].append(str(p))
    hits = [k for k, v in found.items() if v]
    mixed = len(hits) > 1
    if scene:
        chosen = scene
    elif len(hits) == 1:
        chosen = hits[0]
    else:
        chosen = None
    missing = []
    files: dict = {}
    ask = ""
    if mixed and not scene:
        ask = "这个文件夹里同时有不止一种表。请说要跑「销项发票入金蝶」「付款入金蝶」还是「收款入金蝶」。"
        return {
            "ready": False,
            "scene": None,
            "mixed": True,
            "missing": ["指定模块"],
            "files": {k: v for k, v in found.items() if v},
            "ask": ask,
        }
    if chosen == "销项发票":
        if not found["销项发票"]:
            missing.append("发票簿（要有单位名称、价税合计、申请人）")
            ask = "还缺发票簿。把表放进这个文件夹即可，不用改文件名，也不用先填科目。"
        else:
            files["invoice"] = found["销项发票"][0]
    elif chosen == "付款":
        if not found["付款"]:
            missing.append("付款三列表（供应商 / 应付金额本币 / 开户名）")
        ledger = Path(found["付款"][0]) if found["付款"] else None
        folders = payment_folders(root, ledger)
        files["ledger"] = str(ledger) if ledger else None
        files["invoice_dirs"] = [str(p) for p in folders]
        if not folders:
            missing.append("各家发票夹（夹里要有 PDF）")
        if missing:
            ask = "付款还缺：" + "；".join(missing) + "。台账放外面，一家一个夹，夹里放发票 PDF。"
    elif chosen == "收款":
        if not found["收款"]:
            missing.append("收款表（日期 / 客户名称 / 借方（增加）；中行收款 sheet 即可）")
            ask = "还缺收款表。把月底稿放进这个文件夹即可，技能读「中行收款」。部门用组织架构，不必另填部门编码列。"
        else:
            files["receipt"] = found["收款"][0]
    else:
        missing.append("能认出来的业务表")
        ask = "这个文件夹里我还没认出销项发票、付款或收款表。请放好对应 Excel，或直接说要跑哪一句。"
    ready = not missing and chosen is not None
    return {
        "ready": ready,
        "scene": chosen,
        "mixed": mixed,
        "missing": missing,
        "files": files,
        "ask": ask if not ready else "",
    }


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="金蝶入账 · 盘点文件夹")
    parser.add_argument("--input-dir", required=True)
    parser.add_argument("--scene", choices=["销项发票", "付款", "收款"])
    args = parser.parse_args(argv)
    report = inspect_dir(Path(args.input_dir), args.scene)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report.get("ready") else 2


if __name__ == "__main__":
    raise SystemExit(main())
