#!/usr/bin/env python3
"""把已经查过的新闻和这期还要搜的客户分开。不搜新闻，不改风险等级。"""

from __future__ import annotations

import argparse
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

REUSE_DAYS = 7
RISK_OK = {"高", "中", "低", "无"}


def ask(text: str) -> int:
    print("status=ask")
    print(f"ask={text}")
    return 2


def parse_day(value):
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return datetime.strptime(text, "%Y-%m-%d").date()
    except ValueError:
        return None


def usable(row: dict) -> bool:
    blob = " ".join(str(row.get(key) or "") for key in ("summary", "url", "note"))
    if "未检索" in blob or "本次未" in blob:
        return False
    url = str(row.get("url") or "").strip()
    risk = str(row.get("risk") or "").strip()
    has_link = url.startswith("http")
    if has_link:
        return risk in {"高", "中", "低"}
    if "未查到" not in blob:
        return False
    return risk in {"", "无"}


def fresh(row: dict, batch, today) -> bool:
    found = parse_day(row.get("retrieved")) or batch
    if found is None:
        return False
    return 0 <= (today - found).days <= REUSE_DAYS


def customers_from_facts(path: Path) -> list[str]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["事实"] if "事实" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        return []
    try:
        column = [str(cell or "").strip() for cell in header].index("客户")
    except ValueError:
        return []
    names = []
    seen = set()
    for row in rows:
        if column >= len(row) or row[column] in (None, ""):
            continue
        name = str(row[column]).strip()
        if name and name not in seen:
            seen.add(name)
            names.append(name)
    return names


def load_prior(directory: Path) -> dict[str, list[dict]]:
    grouped: dict[str, list[dict]] = {}
    for path in sorted(directory.glob("*.jsonl")):
        for line in path.read_text(encoding="utf-8").splitlines():
            if not line.strip():
                continue
            row = json.loads(line)
            name = str(row.get("customer") or "").strip()
            if name:
                grouped.setdefault(name, []).append(row)
    return grouped


def with_retrieved(row: dict, batch: str) -> dict:
    copied = dict(row)
    if not str(copied.get("retrieved") or "").strip() and batch:
        copied["retrieved"] = batch
    return copied


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="分开沿用的新闻和这次必搜的客户")
    parser.add_argument("--facts", type=Path, required=True)
    parser.add_argument("--prior", type=Path)
    parser.add_argument("--prior-retrieved", default="")
    parser.add_argument("--today", required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args(argv)
    today = parse_day(args.today)
    if today is None:
        return ask("今天的日期要写成 YYYY-MM-DD。")
    if not args.facts.exists():
        return ask("还没有事实表，不能分新闻。")
    batch = parse_day(args.prior_retrieved)
    if args.prior_retrieved and batch is None:
        return ask("上次检索日要写成 YYYY-MM-DD。")
    names = customers_from_facts(args.facts)
    if not names:
        return ask("事实表里没有客户列，或一个客户都没有。")
    prior = load_prior(args.prior) if args.prior and args.prior.exists() else {}
    keep = []
    search = []
    for name in names:
        rows = prior.get(name) or []
        if rows and all(usable(row) and fresh(row, batch, today) for row in rows):
            keep.append((name, rows))
        else:
            search.append(name)
    args.out.mkdir(parents=True, exist_ok=True)
    batch_text = batch.isoformat() if batch else ""
    lines = []
    for _name, rows in keep:
        for row in rows:
            lines.append(json.dumps(with_retrieved(row, batch_text), ensure_ascii=False))
    (args.out / "沿用.jsonl").write_text("\n".join(lines) + ("\n" if lines else ""), encoding="utf-8")
    (args.out / "必搜.txt").write_text("\n".join(search) + ("\n" if search else ""), encoding="utf-8")
    print("status=ok")
    print(f"customers={len(names)}")
    print(f"reuse={len(keep)}")
    print(f"search={len(search)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
