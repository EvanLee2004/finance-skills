"""交卷前检查。五页不齐、新闻是占位、智云页缺列或核对空着，就退出。新闻后补只能先出清单。"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

REQUIRED = ["待抽查清单", "建议本次抽", "豁免与已回款", "风险提示", "智云核对"]


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="检查交给亮晶的工作簿是否五页都落地")
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--list-only", action="store_true", help="新闻后补时只确认清单能先给")
    args = parser.parse_args(argv)
    if not args.workbook.exists():
        print("status=ask")
        print("ask=还没有交给她的工作簿。")
        return 2
    wb = load_workbook(args.workbook, data_only=True)
    missing = [name for name in REQUIRED if name not in wb.sheetnames]
    if missing:
        print("status=ask")
        print("ask=工作簿还缺这些 sheet：" + "、".join(missing))
        wb.close()
        return 2
    news_iter = wb["风险提示"].iter_rows(values_only=True)
    zhiyun_iter = wb["智云核对"].iter_rows(values_only=True)
    news_header = next(news_iter)
    zhiyun_header = next(zhiyun_iter)
    news_customer = next((i for i, cell in enumerate(news_header) if cell and "客户" in str(cell)), 0)
    zhiyun_customer = next((i for i, cell in enumerate(zhiyun_header) if cell and "客户" in str(cell)), 0)
    news_rows = 0
    placeholder = 0
    no_link_and_not_missing = 0
    pending = 0
    news_names = set()
    news_tokens = set()
    for row in news_iter:
        if not any(cell not in (None, "") for cell in row):
            continue
        news_rows += 1
        if news_customer < len(row) and row[news_customer]:
            news_names.add(str(row[news_customer]).strip())
        for cell in row:
            news_tokens.update(part.strip() for part in re.split(r"[、；\n]", str(cell or "")) if part.strip())
        blob = " ".join(str(cell) for cell in row if cell not in (None, ""))
        link = ""
        if len(row) >= 3 and row[2]:
            link = str(row[2])
        later = "新闻后补" in blob
        linked = "http" in link or "http" in blob
        if "未检索" in blob or "本次未" in blob:
            placeholder += 1
        elif later and (linked or "未查到" in blob):
            no_link_and_not_missing += 1
        elif later:
            pending += 1
        elif "未查到" not in blob and not linked:
            no_link_and_not_missing += 1
    def column_index(header, needle: str):
        for index, cell in enumerate(header):
            if cell and str(cell).strip() == needle:
                return index
        return None

    order_col = column_index(zhiyun_header, "订单号")
    archive_col = column_index(zhiyun_header, "合同归档号")
    status_col = column_index(zhiyun_header, "订单状态")
    stage_col = column_index(zhiyun_header, "销售结算阶段")
    verdict_col = column_index(zhiyun_header, "智云订单状态核对")
    zhiyun_rows = 0
    zhiyun_names = set()
    messy = 0
    blank_pair = 0
    multi_order = 0
    bad_verdict = 0
    for row in zhiyun_iter:
        if not any(cell not in (None, "") for cell in row):
            continue
        zhiyun_rows += 1
        if zhiyun_customer < len(row) and row[zhiyun_customer]:
            zhiyun_names.add(str(row[zhiyun_customer]).strip())
        blob = " ".join(str(cell) for cell in row if cell not in (None, ""))
        if "看不出来" in blob or "没挂在这份合同" in blob:
            messy += 1
        if order_col is None or archive_col is None or status_col is None:
            continue
        order = str(row[order_col] or "").strip() if order_col < len(row) else ""
        archive = str(row[archive_col] or "").strip() if archive_col < len(row) else ""
        status = str(row[status_col] or "").strip() if status_col < len(row) else ""
        if not archive or not status:
            blank_pair += 1
        if order != "未找到" and any(mark in order for mark in ("；", ";", "、", "\n")):
            multi_order += 1
        if verdict_col is None or stage_col is None or status_col is None:
            continue
        verdict = str(row[verdict_col] or "").strip() if verdict_col < len(row) else ""
        stage = str(row[stage_col] or "").strip() if stage_col < len(row) else ""
        if verdict not in ("", "不冲突", "不一致"):
            bad_verdict += 1
        elif status and status != "未找到" and stage:
            if verdict not in ("不冲突", "不一致"):
                bad_verdict += 1
        elif verdict:
            bad_verdict += 1
    missing_customers = len((news_names - zhiyun_names) | (zhiyun_names - news_names - news_tokens))
    header_gap = None in (order_col, archive_col, status_col, stage_col, verdict_col)
    wb.close()
    finished = (
        placeholder == 0
        and no_link_and_not_missing == 0
        and zhiyun_rows
        and missing_customers == 0
        and news_rows
        and not header_gap
        and messy == 0
        and blank_pair == 0
        and multi_order == 0
        and bad_verdict == 0
    )
    if finished and pending and args.list_only:
        print("status=list_ready")
    elif finished and pending == 0:
        print("status=ok")
    else:
        print("ask")
    print(f"news_rows={news_rows}")
    print(f"news_placeholder={placeholder}")
    print(f"news_pending={pending}")
    print(f"news_without_link_or_missing={no_link_and_not_missing}")
    print(f"zhiyun_rows={zhiyun_rows}")
    print(f"zhiyun_missing_customers={missing_customers}")
    print(f"zhiyun_messy={messy}")
    print(f"zhiyun_blank={blank_pair}")
    print(f"zhiyun_multi_order={multi_order}")
    print(f"zhiyun_bad_verdict={bad_verdict}")
    if placeholder or no_link_and_not_missing:
        print("ask=风险提示还有没检索或没链接的行。查完再写未查到，不要留未检索。")
        return 2
    if zhiyun_rows == 0:
        print("ask=智云核对还是空的。用 Playwright 按销售反馈里的客户找单子和合同，写进这一页。")
        return 2
    if header_gap or messy or blank_pair or multi_order or bad_verdict:
        print("ask=智云核对要按销售反馈的单号一行，写合同归档号、订单状态和智云订单状态核对。没核到就留空，核对过只写不冲突或不一致。")
        return 2
    if missing_customers:
        print("ask=智云核对没有盖住风险提示里的每个客户。对不上就写未找到，不要少人。")
        return 2
    if news_rows == 0:
        print("ask=风险提示还是空的。")
        return 2
    if pending and not args.list_only:
        print("ask=新闻还没搜完。这版只能当清单。补完新闻后再交，不要加 --list-only。")
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
