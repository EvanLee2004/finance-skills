"""认两张表，把台账月份拆开，给本期每一笔挂上台账证据。

不下「抽不抽」的结论，不删行，不搜新闻，不登录。
"""

from __future__ import annotations

import argparse
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from months import clean_name, is_umbrella, parse_month  # noqa: E402


def header_map(ws) -> dict[int, str]:
    found = {}
    for col in range(1, (ws.max_column or 1) + 1):
        value = ws.cell(1, col).value
        if value is None or str(value).strip() == "":
            continue
        found[col] = str(value).replace("\n", "").replace(" ", "")
    return found


def find_col(headers: dict[int, str], needle: str) -> int | None:
    for col, text in headers.items():
        if needle in text:
            return col
    return None


def score_sheet(ws) -> tuple[str, int]:
    headers = header_map(ws)
    texts = "".join(headers.values())
    if "抽查日期" in texts and "正式确认" in texts and "交付月份" in texts:
        return "ledger", 3
    if "账龄" in texts and "结算阶段" in texts and "交付月份" in texts and ("单号" in texts or "智云" in texts):
        return "sales", 4
    return "", 0


def sniff_file(path: Path) -> list[tuple[str, str, int]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    found = []
    for name in wb.sheetnames:
        role, score = score_sheet(wb[name])
        if score:
            found.append((role, name, score))
    wb.close()
    return found


def assign_roles(paths: list[Path]) -> tuple[tuple[Path, str], tuple[Path, str]]:
    ledger = []
    sales = []
    for path in paths:
        for role, sheet, _score in sniff_file(path):
            if role == "ledger":
                ledger.append((path, sheet))
            elif role == "sales":
                sales.append((path, sheet))
    if len(ledger) != 1 or len(sales) != 1:
        print("status=ask")
        print("ask=认不出哪张是抽查台账、哪张是本期销售反馈。按表头认：台账要有抽查日期和正式确认；销售反馈要有账龄、结算阶段和单号。")
        print(f"ledger_hits={len(ledger)} sales_hits={len(sales)}")
        raise SystemExit(2)
    return ledger[0], sales[0]


def cell(ws, row, col):
    if col is None:
        return None
    return ws.cell(row, col).value


def as_month(value):
    if isinstance(value, float) and value == int(value):
        value = int(value)
    if isinstance(value, int):
        text = str(value)
    elif isinstance(value, str) and value.strip().isdigit():
        text = value.strip()
    else:
        return None
    if len(text) == 6 and text.startswith("20"):
        month = int(text)
        if 1 <= month % 100 <= 12:
            return month
    return None


def load_ledger(path: Path, sheet: str):
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet]
    headers = header_map(ws)
    cols = {
        "date": find_col(headers, "抽查日期"),
        "sales": find_col(headers, "营销人员") or find_col(headers, "销售"),
        "name": find_col(headers, "客户名称"),
        "month": find_col(headers, "交付月份"),
        "confirm": find_col(headers, "正式确认"),
    }
    entries = defaultdict(list)
    pending = []
    for row in range(2, (ws.max_row or 1) + 1):
        values = [cell(ws, row, cols[key]) for key in ("date", "sales", "name", "month", "confirm")]
        if not any(v not in (None, "") for v in values):
            continue
        name, sos = clean_name(values[2])
        months, kind = parse_month(values[3], values[0])
        item = {
            "date": values[0],
            "name": name,
            "months": set(months or []),
            "confirm": "" if values[4] is None else str(values[4]).strip(),
            "raw": values[3],
            "sos": sos,
            "kind": kind,
        }
        if months is None:
            pending.append(item)
        if name:
            entries[name].append(item)
    wb.close()
    return entries, pending


def load_sales(path: Path, sheet: str):
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet]
    headers = header_map(ws)
    cols = {
        "sales": find_col(headers, "销售人员") or find_col(headers, "销售"),
        "name": find_col(headers, "客户名称"),
        "so": find_col(headers, "单号"),
        "month": find_col(headers, "交付月份"),
        "age": find_col(headers, "账龄"),
        "stage": find_col(headers, "结算阶段"),
    }
    rows = []
    for row in range(2, (ws.max_row or 1) + 1):
        raw_name = cell(ws, row, cols["name"])
        month = as_month(cell(ws, row, cols["month"]))
        if raw_name in (None, "") and month is None:
            continue
        name, _sos = clean_name(raw_name)
        stage = cell(ws, row, cols["stage"])
        stage = "" if stage is None else str(stage).strip()
        age = cell(ws, row, cols["age"])
        so = cell(ws, row, cols["so"])
        so = "" if so is None else str(so).strip()
        rows.append(
            {
                "sales": "" if cell(ws, row, cols["sales"]) is None else str(cell(ws, row, cols["sales"])).strip(),
                "name": name,
                "so": so,
                "month": month,
                "age": age,
                "stage": stage,
                "paid": "已回款" in stage,
            }
        )
    wb.close()
    return rows


def matching_entries(entries, row):
    own = entries.get(row["name"], [])
    if is_umbrella(row["name"]):
        if not row["so"]:
            return []
        target = row["so"].upper()
        return [
            item
            for item in own
            if item["months"]
            and item["kind"] not in ("keep", "empty")
            and any(token.upper() == target for token in item["sos"])
        ]
    if row["month"] is None:
        return []
    return [item for item in own if row["month"] in item["months"]]


def build(entries, pending, sales_rows):
    groups = {}
    for row in sales_rows:
        grain = "订单" if is_umbrella(row["name"]) else "客户月"
        key = (
            row["sales"],
            row["name"],
            row["so"] if grain == "订单" else "",
            row["month"],
        )
        bucket = groups.setdefault(
            key,
            {
                "sales": row["sales"],
                "name": row["name"],
                "so": row["so"] if grain == "订单" else "",
                "month": row["month"],
                "grain": grain,
                "orders": 0,
                "paid": 0,
                "ages": set(),
                "stages": [],
                "hits": [],
                "hit_ids": set(),
                "paid_sos": [],
                "sos": [],
                "unpaid_sos": set(),
            },
        )
        bucket["orders"] += 1
        if row["so"] and row["so"].upper() not in {item.upper() for item in bucket["sos"]}:
            bucket["sos"].append(row["so"])
        if row["so"] and not row["paid"]:
            bucket["unpaid_sos"].add(row["so"].upper())
        bucket["paid"] += 1 if row["paid"] else 0
        if isinstance(row["age"], (int, float)) and not isinstance(row["age"], bool):
            bucket["ages"].add(int(row["age"]))
        if row["stage"]:
            bucket["stages"].append(row["stage"])
        if row["paid"] and row["so"] and row["so"].upper() not in {item.upper() for item in bucket["paid_sos"]}:
            bucket["paid_sos"].append(row["so"])
        for item in matching_entries(entries, row):
            if id(item) in bucket["hit_ids"]:
                continue
            bucket["hit_ids"].add(id(item))
            bucket["hits"].append(item)
    unrecognized = defaultdict(int)
    for item in pending:
        unrecognized[item["name"]] += 1
    out = []
    age_conflict = 0
    for bucket in groups.values():
        if len(bucket["ages"]) > 1:
            age_conflict += 1
        confirms = []
        dates = []
        raws = []
        for item in bucket["hits"]:
            if item["confirm"] and item["confirm"] not in confirms:
                confirms.append(item["confirm"])
            if item["date"] not in (None, "") and str(item["date"]) not in dates:
                dates.append(str(item["date"]))
            raw = "" if item["raw"] is None else str(item["raw"]).replace("\n", " ")
            if raw and raw not in raws:
                raws.append(raw)
        if bucket["sos"]:
            order_count = len(bucket["sos"])
            paid_count = sum(1 for so in bucket["sos"] if so.upper() not in bucket["unpaid_sos"])
            paid_list = [so for so in bucket["paid_sos"] if so.upper() not in bucket["unpaid_sos"]]
        else:
            order_count = bucket["orders"]
            paid_count = bucket["paid"]
            paid_list = bucket["paid_sos"]
        out.append(
            [
                bucket["sales"],
                bucket["name"],
                "；".join(bucket["sos"]),
                bucket["month"],
                bucket["grain"],
                order_count,
                paid_count,
                "" if len(bucket["ages"]) != 1 else next(iter(bucket["ages"])),
                "是" if len(bucket["ages"]) > 1 else "",
                "；".join(dict.fromkeys(bucket["stages"])),
                len(bucket["hits"]),
                "；".join(confirms),
                "；".join(dates),
                "；".join(raws),
                unrecognized.get(bucket["name"], 0),
                "；".join(paid_list),
            ]
        )
    return out, age_conflict, len(pending)


HEADERS = [
    "销售",
    "客户",
    "订单号",
    "交付月份",
    "粒度",
    "订单数",
    "已回款订单数",
    "账龄",
    "账龄冲突",
    "结算阶段",
    "台账命中条数",
    "台账确认",
    "台账抽查日",
    "台账原月份",
    "同客户无法识别的台账行",
    "已回款订单号",
]


def write_facts(path: Path, rows, pending):
    wb = Workbook()
    ws = wb.active
    if ws is None:
        raise RuntimeError("workbook has no sheet")
    ws.title = "事实"
    ws.append(HEADERS)
    for row in rows:
        ws.append(row)
    wp = wb.create_sheet("待看")
    wp.append(["客户", "抽查日期", "原交付月份", "原因"])
    for item in pending:
        reason = "交付月份为空" if item["kind"] == "empty" else f"留原值：{item['raw']}"
        wp.append([item["name"], item["date"], item["raw"], reason])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()


def main(argv=None):
    parser = argparse.ArgumentParser(description="应收抽查：认表并挂上台账证据")
    parser.add_argument("--ledger", type=Path)
    parser.add_argument("--sales", type=Path)
    parser.add_argument("--inputs", nargs="*", type=Path, help="不指定角色时，按表头认")
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args(argv)
    paths = []
    if args.ledger and args.sales:
        ledger_hit = sniff_file(args.ledger)
        sales_hit = sniff_file(args.sales)
        if not any(role == "ledger" for role, _, _ in ledger_hit):
            print("status=ask")
            print("ask=--ledger 这张里没有抽查日期和正式确认。")
            return 2
        if not any(role == "sales" for role, _, _ in sales_hit):
            print("status=ask")
            print("ask=--sales 这张里没有账龄、结算阶段和单号。")
            return 2
        ledger_sheets = [sheet for role, sheet, _ in ledger_hit if role == "ledger"]
        sales_sheets = [sheet for role, sheet, _ in sales_hit if role == "sales"]
        if len(ledger_sheets) != 1 or len(sales_sheets) != 1:
            print("status=ask")
            print("ask=同一份文件里有多张都能对上表头。不要默认拿第一张，先问要用哪张。")
            print(f"ledger_hits={len(ledger_sheets)} sales_hits={len(sales_sheets)}")
            return 2
        ledger = (args.ledger, ledger_sheets[0])
        sales = (args.sales, sales_sheets[0])
    else:
        paths = list(args.inputs or [])
        if len(paths) < 2:
            print("status=ask")
            print("ask=需要两张表：抽查台账，和当期销售反馈。")
            return 2
        ledger, sales = assign_roles(paths)
    entries, pending = load_ledger(*ledger)
    sales_rows = load_sales(*sales)
    facts, age_conflict, unparsed = build(entries, pending, sales_rows)
    write_facts(args.out, facts, pending)
    print(f"ledger_sheet={ledger[1]}")
    print(f"sales_sheet={sales[1]}")
    print(f"facts={args.out}")
    print(f"groups={len(facts)}")
    print(f"ledger_unparsed={unparsed}")
    print(f"age_conflict={age_conflict}")
    if age_conflict:
        print("status=ask")
        print("ask=有客户月账龄不一致。先看事实表「账龄冲突」列，问亮晶，不要取平均。")
        return 2
    print("status=ok")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
