import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import months  # noqa: E402
import prepare  # noqa: E402


def test_hyphen_fills_the_middle_and_stays_before_the_check():
    assert months.parse_month("6-12月交付订单", 20260205)[0] == [
        202506, 202507, 202508, 202509, 202510, 202511, 202512,
    ]
    assert months.parse_month("1-4月交付订单", 20260122)[0] == [202501, 202502, 202503, 202504]
    assert months.parse_month("7-10月客服订单", 20251113)[0] == [202507, 202508, 202509, 202510]


def test_numeric_months_stay_numeric():
    assert months.parse_month("202509；202510", 20251023)[0] == [202509, 202510]
    assert months.parse_month(20251023, 20251023)[0] == [202510]
    assert months.parse_month("202509", 20251023)[0] == [202509]


def test_list_does_not_fill_the_gap():
    assert months.parse_month("24年12月，25年6月交付订单", 20260205)[0] == [202412, 202506]
    assert months.parse_month("2024年8、12月交付订单", 20251023)[0] == [202408, 202412]
    assert months.parse_month("6、7月交付订单", 20251023)[0] == [202506, 202507]


def test_explicit_span_and_single_month():
    assert months.parse_month("25年11月-26年3月交付订单", 20260301)[0] == [
        202511, 202512, 202601, 202602, 202603,
    ]
    assert months.parse_month("9月交付订单", 20251023)[0] == [202509]
    assert months.parse_month("11月交付订单", 20260115)[0] == [202511]


def test_unparsed_months_cover_nothing():
    assert months.parse_month("全部订单", 20251113) == (None, "keep")
    assert months.parse_month("折扣进展", 20251120)[0] is None
    assert months.parse_month(None, 20260910)[0] is None


def test_names():
    assert months.clean_name("中华人民共和国公安部（SO12345678）") == ("公安部", ["SO12345678"])
    assert months.clean_name("北京市公安局海淀分局")[0] == "北京市公安局海淀分局"
    assert months.is_umbrella("公安部")
    assert not months.is_umbrella("北京市公安局海淀分局")


def test_prepare_joins_evidence_and_does_not_treat_special_text_as_coverage(tmp_path):
    ledger = tmp_path / "ledger.xlsx"
    sales = tmp_path / "sales.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "合规文件抽查汇总"
    ws.append(["抽查日期", "营销人员", "客户名称", "交付月份", "应收金额是否有客户正式确认（盖章/对公邮件）"])
    ws.append([20251023, "甲", "示例客户", "9月交付订单", "未反馈"])
    ws.append([20251023, "甲", "示例客户", "全部订单", ""])
    ws.append([20251113, "甲", "公安部（so10000001）", "8月交付订单", "对公邮件"])
    wb.save(ledger)
    wb = Workbook()
    ws = wb.active
    ws.title = "销售反馈"
    ws.append(["销售人员", "客户名称", "新智云单号", "交付月份", "账龄(月份）", "结算阶段(请筛选分类)"])
    ws.append(["甲", "示例客户", "SO20000001", 202509, 12, "未对账"])
    ws.append(["甲", "示例客户", "SO20000002", 202510, 11, "已回款未核销"])
    ws.append(["甲", "公安部", "SO10000001", 202508, 13, "未对账"])
    ws.append(["甲", "公安部", "SO10000009", 202508, 13, "未对账"])
    wb.save(sales)
    out = tmp_path / "facts.xlsx"
    code = prepare.main(["--ledger", str(ledger), "--sales", str(sales), "--out", str(out)])
    assert code == 0
    wb = load_workbook(out, data_only=True)
    rows = list(wb["事实"].iter_rows(min_row=2, values_only=True))
    by_key = {(row[1], row[2], row[3]): row for row in rows}
    covered = by_key[("示例客户", "SO20000001", 202509)]
    assert covered[10] == 1
    assert covered[11] == "未反馈"
    open_month = by_key[("示例客户", "SO20000002", 202510)]
    assert open_month[10] == 0
    assert open_month[6] == 1
    assert open_month[15] == "SO20000002"
    checked_case = by_key[("公安部", "SO10000001", 202508)]
    other_case = by_key[("公安部", "SO10000009", 202508)]
    assert checked_case[10] == 1
    assert other_case[10] == 0
    assert wb["待看"].max_row == 2
    wb.close()
