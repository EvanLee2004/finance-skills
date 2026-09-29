import sys
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import check_delivery  # noqa: E402


def test_placeholder_news_and_missing_zhiyun_are_not_done(tmp_path):
    path = tmp_path / "out.xlsx"
    wb = Workbook()
    wb.active.title = "待抽查清单"
    wb.create_sheet("豁免与已回款")
    news = wb.create_sheet("风险提示")
    news.append(["客户", "新闻摘要", "链接", "说明"])
    news.append(["甲", "", "", "本次未检索。不编链接。"])
    wb.save(path)
    assert check_delivery.main(["--workbook", str(path)]) == 2


def test_four_sheets_with_link_or_not_found_pass(tmp_path):
    path = tmp_path / "out.xlsx"
    wb = Workbook()
    wb.active.title = "待抽查清单"
    suggest = wb.create_sheet("建议本次抽")
    suggest.append(["销售", "客户", "档", "原因"])
    wb.create_sheet("豁免与已回款")
    news = wb.create_sheet("风险提示")
    news.append(["客户", "新闻摘要", "链接", "说明"])
    news.append(["甲", "未查到", "", "未查到"])
    zy = wb.create_sheet("智云核对")
    zy.append(["销售", "客户", "订单号", "合同归档号", "订单状态", "说明", "销售结算阶段", "智云订单状态核对"])
    zy.append(["甲销", "甲", "SO1", "20260001", "OP4/项目已交付", "", "已对账，待开票", "不冲突"])
    wb.save(path)
    assert check_delivery.main(["--workbook", str(path)]) == 0


def test_zhiyun_must_cover_every_news_customer(tmp_path):
    path = tmp_path / "out.xlsx"
    wb = Workbook()
    wb.active.title = "待抽查清单"
    wb.create_sheet("豁免与已回款")
    news = wb.create_sheet("风险提示")
    news.append(["客户", "新闻摘要", "链接", "说明"])
    news.append(["甲", "未查到", "", "未查到"])
    news.append(["乙", "未查到", "", "未查到"])
    zy = wb.create_sheet("智云核对")
    zy.append(["销售", "客户", "订单号", "合同归档号", "订单状态", "说明"])
    zy.append(["甲销", "甲", "SO1", "未找到", "未找到", "下单里没有这个单号。"])
    wb.save(path)
    assert check_delivery.main(["--workbook", str(path)]) == 2


def test_mixed_zhiyun_labels_are_not_done(tmp_path):
    path = tmp_path / "out.xlsx"
    wb = Workbook()
    wb.active.title = "待抽查清单"
    wb.create_sheet("豁免与已回款")
    news = wb.create_sheet("风险提示")
    news.append(["客户", "新闻摘要", "链接", "说明"])
    news.append(["甲", "未查到", "", "未查到"])
    zy = wb.create_sheet("智云核对")
    zy.append(["销售", "客户", "订单号", "合同归档号", "订单状态", "说明"])
    zy.append(["甲销", "甲", "SO1；SO2", "看不出来", "没挂在这份合同上", ""])
    wb.save(path)
    assert check_delivery.main(["--workbook", str(path)]) == 2


def test_missing_status_check_column_is_not_done(tmp_path):
    path = tmp_path / "out.xlsx"
    wb = Workbook()
    wb.active.title = "待抽查清单"
    wb.create_sheet("建议本次抽")
    wb.create_sheet("豁免与已回款")
    news = wb.create_sheet("风险提示")
    news.append(["客户", "新闻摘要", "链接", "说明"])
    news.append(["甲", "未查到", "", "未查到"])
    zy = wb.create_sheet("智云核对")
    zy.append(["销售", "客户", "订单号", "合同归档号", "订单状态", "说明", "回款核对"])
    zy.append(["甲销", "甲", "SO1", "20260001", "OP4/项目已交付", "", ""])
    wb.save(path)
    assert check_delivery.main(["--workbook", str(path)]) == 2


def finished_book(path: Path, summary: str, link: str, note: str) -> None:
    wb = Workbook()
    wb.active.title = "待抽查清单"
    wb.create_sheet("建议本次抽")
    wb.create_sheet("豁免与已回款")
    news = wb.create_sheet("风险提示")
    news.append(["客户", "新闻摘要", "链接", "说明"])
    news.append(["甲", summary, link, note])
    zy = wb.create_sheet("智云核对")
    zy.append(["销售", "客户", "订单号", "合同归档号", "订单状态", "说明", "销售结算阶段", "智云订单状态核对"])
    zy.append(["甲销", "甲", "SO1", "20260001", "OP4/项目已交付", "", "已对账，待开票", "不冲突"])
    wb.save(path)


def test_pending_news_is_a_list_until_requested(tmp_path, capsys):
    path = tmp_path / "out.xlsx"
    finished_book(path, "新闻后补", "", "新闻后补")
    assert check_delivery.main(["--workbook", str(path)]) == 2
    assert "新闻还没搜完" in capsys.readouterr().out
    assert check_delivery.main(["--workbook", str(path), "--list-only"]) == 0
    assert "status=list_ready" in capsys.readouterr().out


def test_pending_news_with_a_link_is_not_a_list(tmp_path, capsys):
    path = tmp_path / "out.xlsx"
    finished_book(path, "新闻后补", "https://example.com/a", "新闻后补")
    assert check_delivery.main(["--workbook", str(path), "--list-only"]) == 2
    assert "status=list_ready" not in capsys.readouterr().out


def test_checkable_order_without_verdict_is_not_done(tmp_path):
    path = tmp_path / "out.xlsx"
    wb = Workbook()
    wb.active.title = "待抽查清单"
    wb.create_sheet("建议本次抽")
    wb.create_sheet("豁免与已回款")
    news = wb.create_sheet("风险提示")
    news.append(["客户", "新闻摘要", "链接", "说明"])
    news.append(["甲", "未查到", "", "未查到"])
    zy = wb.create_sheet("智云核对")
    zy.append(["销售", "客户", "订单号", "合同归档号", "订单状态", "说明", "销售结算阶段", "智云订单状态核对"])
    zy.append(["甲销", "甲", "SO1", "20260001", "SP4/已回款", "", "未对账", ""])
    wb.save(path)
    assert check_delivery.main(["--workbook", str(path)]) == 2
