import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import prepare  # noqa: E402

LEDGER_HEADER = ["抽查日期", "营销人员", "客户名称", "交付月份", "正式确认"]
SALES_HEADER = ["销售人员", "客户名称", "新智云单号", "交付月份", "账龄", "结算阶段"]


def save(path, sheets: dict):
    wb = Workbook()
    first = True
    for title, rows in sheets.items():
        ws = wb.active if first else wb.create_sheet(title)
        first = False
        ws.title = title
        for row in rows:
            ws.append(row)
    wb.save(path)


def test_two_matching_sheets_ask(tmp_path):
    ledger = tmp_path / "ledger.xlsx"
    sales = tmp_path / "sales.xlsx"
    save(ledger, {"汇总": [LEDGER_HEADER, [20260101, "甲", "示例", "9月交付订单", "未反馈"]], "另一张": [LEDGER_HEADER]})
    save(sales, {"销售反馈": [SALES_HEADER, ["甲", "示例", "SO10000001", 202509, 1, "未对账"]]})
    code = prepare.main(["--ledger", str(ledger), "--sales", str(sales), "--out", str(tmp_path / "facts.xlsx")])
    assert code == 2


def test_age_conflict_exits_ask(tmp_path):
    ledger = tmp_path / "ledger.xlsx"
    sales = tmp_path / "sales.xlsx"
    save(ledger, {"汇总": [LEDGER_HEADER]})
    save(
        sales,
        {
            "销售反馈": [
                SALES_HEADER,
                ["甲", "示例", "SO10000001", 202509, 3, "未对账"],
                ["甲", "示例", "SO10000002", 202509, 8, "未对账"],
            ]
        },
    )
    out = tmp_path / "facts.xlsx"
    assert prepare.main(["--ledger", str(ledger), "--sales", str(sales), "--out", str(out)]) == 2
    wb = load_workbook(out, data_only=True)
    row = list(wb["事实"].iter_rows(min_row=2, values_only=True))[0]
    assert row[8] == "是"
    wb.close()


def test_umbrella_special_month_and_short_so_do_not_count(tmp_path):
    ledger = tmp_path / "ledger.xlsx"
    sales = tmp_path / "sales.xlsx"
    save(
        ledger,
        {
            "汇总": [
                LEDGER_HEADER,
                [20251113, "甲", "公安部（SO10000001）", "全部订单", "对公邮件"],
                [20251023, "甲", "公安部（SO100000019）", "8月交付订单", "对公邮件"],
                [20251023, "甲", "公安部", "SO10000001的9月交付订单", "对公邮件"],
            ]
        },
    )
    save(
        sales,
        {
            "销售反馈": [
                SALES_HEADER,
                ["甲", "公安部", "SO10000001", 202508, 13, "未对账"],
                ["甲", "公安部", "SO10000001", 202509, 12, "未对账"],
            ]
        },
    )
    out = tmp_path / "facts.xlsx"
    assert prepare.main(["--ledger", str(ledger), "--sales", str(sales), "--out", str(out)]) == 0
    wb = load_workbook(out, data_only=True)
    rows = list(wb["事实"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2
    assert {row[3] for row in rows} == {202508, 202509}
    assert all(row[10] == 0 for row in rows)
    wb.close()


def test_customer_month_keeps_every_order_number(tmp_path):
    ledger = tmp_path / "ledger.xlsx"
    sales = tmp_path / "sales.xlsx"
    save(ledger, {"汇总": [LEDGER_HEADER]})
    save(
        sales,
        {
            "销售反馈": [
                SALES_HEADER,
                ["甲", "示例", "SO10000001", 202509, 1, "未对账"],
                ["甲", "示例", "SO10000002", 202509, 1, "未对账"],
            ]
        },
    )
    out = tmp_path / "facts.xlsx"
    assert prepare.main(["--ledger", str(ledger), "--sales", str(sales), "--out", str(out)]) == 0
    wb = load_workbook(out, data_only=True)
    row = list(wb["事实"].iter_rows(min_row=2, values_only=True))[0]
    assert row[2] == "SO10000001；SO10000002"
    wb.close()


def test_duplicate_order_rows_count_once(tmp_path):
    ledger = tmp_path / "ledger.xlsx"
    sales = tmp_path / "sales.xlsx"
    save(ledger, {"汇总": [LEDGER_HEADER]})
    save(
        sales,
        {
            "销售反馈": [
                SALES_HEADER,
                ["甲", "示例", "SO10000001", 202509, 1, "已回款"],
                ["甲", "示例", "SO10000001", 202509, 1, "已回款"],
                ["甲", "示例", "so10000002", 202509, 1, "已回款"],
                ["甲", "示例", "SO10000002", 202509, 1, "未对账"],
            ]
        },
    )
    out = tmp_path / "facts.xlsx"
    assert prepare.main(["--ledger", str(ledger), "--sales", str(sales), "--out", str(out)]) == 0
    wb = load_workbook(out, data_only=True)
    row = list(wb["事实"].iter_rows(min_row=2, values_only=True))[0]
    assert row[2] == "SO10000001；so10000002"
    assert row[5] == 2
    assert row[6] == 1
    assert row[15] == "SO10000001"
    wb.close()
