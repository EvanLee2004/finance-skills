import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import compose  # noqa: E402


FACTS = [
    "销售", "客户", "订单号", "交付月份", "粒度", "订单数", "已回款订单数", "账龄", "账龄冲突",
    "结算阶段", "台账命中条数", "台账确认", "台账抽查日", "台账原月份", "同客户无法识别的台账行", "已回款订单号",
    "应收金额", "预计回款日", "销售解释", "标记", "单号阶段",
]


def write_facts(path: Path, rows: list[list]):
    wb = Workbook()
    ws = wb.active
    ws.title = "事实"
    ws.append(FACTS)
    for row in rows:
        ws.append(row)
    wb.save(path)


def write_zhiyun(path: Path, names: list[str]):
    wb = Workbook()
    ws = wb.active
    ws.title = "智云核对"
    ws.append(["客户", "订单号", "合同编号", "合同类型", "合同金额", "合同审批", "说明"])
    for name in names:
        ws.append([name, "未找到", "未找到", "未找到", "看不出来", "看不出来", "合同和订单都未找到。未对照本月应收。"])
    wb.save(path)


def write_news(directory: Path, rows: list[dict]):
    directory.mkdir(exist_ok=True)
    (directory / "news.jsonl").write_text(
        "\n".join(json.dumps(row, ensure_ascii=False) for row in rows) + "\n",
        encoding="utf-8",
    )


def fact(sales, name, month, age, hits, confirm, dates="", orders=1, paid=0, so="", paid_so="", grain="客户月", conflict="", unrecognized=0, amount=0, expect="", explain="", marks="", stages=""):
    return [sales, name, so, month, grain, orders, paid, age, conflict, "未对账", hits, confirm, dates, "", unrecognized, paid_so, amount, expect, explain, marks, stages]


def names_of(ws):
    return [row[1] for row in ws.iter_rows(min_row=2, values_only=True) if row[1]]


def write_judgment(path: Path, suggest, keywords=None, people=None, orders=None):
    path.write_text(
        json.dumps(
            {"同一人": people or [], "关键词": keywords or [], "建议": suggest, "订单状态核对": orders or []},
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )


def test_sort_paid_pass_and_soft(tmp_path):
    facts = tmp_path / "facts.xlsx"
    write_facts(
        facts,
        [
            fact("甲", "老单", 202501, 8, 0, "", "20260801"),
            fact("甲", "新单", 202608, 1, 0, "", "20260902"),
            fact("甲", "刚交付", 202609, 0, 0, ""),
            fact("甲", "微信", 202501, 8, 1, "微信对账", "20260101"),
            fact("甲", "已回", 202401, 20, 0, "", orders=2, paid=2, paid_so="SO1；SO2"),
            fact("甲", "拿到", 202501, 8, 2, "未反馈；对公邮件"),
            fact("甲", "盖章", 202501, 8, 1, "已盖章"),
            fact("甲", "没盖成", 202501, 8, 1, "盖章未完成"),
            fact("甲", "空确认", 202502, 7, 1, ""),
            fact("甲", "再抽", 202503, 6, 2, "未反馈；未提供"),
            fact("乙", "更老", 202401, 12, 0, ""),
            fact("甲", "部分回", 202504, 5, 0, "", orders=2, paid=1, unrecognized=1),
        ],
    )
    news = tmp_path / "news"
    customers = ["老单", "新单", "刚交付", "微信", "已回", "拿到", "盖章", "没盖成", "空确认", "再抽", "更老", "部分回"]
    write_news(news, [{"customer": name, "summary": "未查到", "url": "", "note": "未查到"} for name in customers])
    zhiyun = tmp_path / "zhiyun.xlsx"
    write_zhiyun(zhiyun, customers)
    judgment = tmp_path / "判断.json"
    write_judgment(
        judgment,
        [
            {"销售": "甲", "客户": "新单", "交付月份": 202608, "档": "补位", "原因": "这个销售还没有别的可抽"},
            {"销售": "乙", "客户": "更老", "交付月份": 202401, "档": "次危", "原因": "账龄已满6个月"},
        ],
    )
    out = tmp_path / "out.xlsx"
    code = compose.main(["--facts", str(facts), "--news", str(news), "--zhiyun", str(zhiyun), "--check-month", "202609", "--judgment", str(judgment), "--out", str(out)])
    assert code == 0
    wb = load_workbook(out, data_only=True)
    pool = names_of(wb["待抽查清单"])
    assert pool == ["已回", "老单", "空确认", "再抽", "部分回", "新单", "刚交付", "更老", "微信", "没盖成"]
    rules = {row[1]: row[5] for row in wb["待抽查清单"].iter_rows(min_row=2, values_only=True)}
    assert rules["老单"] == "没查过"
    assert rules["再抽"] == "未提供或未反馈"
    assert rules["空确认"] == "确认为空"
    assert rules["微信"] == "待你定"
    assert rules["没盖成"] == "待你定"
    partial = next(row for row in wb["待抽查清单"].iter_rows(min_row=2, values_only=True) if row[1] == "部分回")
    assert partial[6] == 1 and partial[7] == 2
    assert "已回款1/2" in partial[9]
    assert "写不成月份" in partial[9]
    exempt = list(wb["豁免与已回款"].iter_rows(min_row=2, values_only=True))
    assert all(row[1] != "已回" for row in exempt)
    suggest_sheet = wb["建议本次抽"]
    assert "补位" in str(suggest_sheet.cell(1, 1).value)
    assert "高危" in str(suggest_sheet.cell(1, 1).value)
    assert suggest_sheet.cell(2, 1).value == "销售"
    assert suggest_sheet.freeze_panes == "C3"
    suggest = {row[1]: row[6] for row in suggest_sheet.iter_rows(min_row=3, values_only=True)}
    assert suggest == {"新单": "补位", "更老": "次危"}
    assert "刚交付" not in suggest
    assert "已回" not in suggest
    assert "拿到" not in pool and "盖章" not in pool
    assert [row[0] for row in wb["风险提示"].iter_rows(min_row=2, values_only=True)] == customers
    wb.close()


def test_alias_merges_news_but_not_the_pool(tmp_path):
    facts = tmp_path / "facts.xlsx"
    write_facts(facts, [fact("甲", "甲公司", 202601, 3, 0, ""), fact("甲", "甲", 202601, 3, 0, "")])
    news = tmp_path / "news"
    write_news(
        news,
        [
            {"customer": "甲公司", "summary": "未查到", "url": "", "note": "未查到"},
            {"customer": "甲", "summary": "有报道", "url": "https://example.com/a", "note": "媒体", "risk": "低"},
        ],
    )
    zhiyun = tmp_path / "zhiyun.xlsx"
    write_zhiyun(zhiyun, ["甲公司", "甲"])
    config = tmp_path / "config.md"
    config.write_text("# 豁免与别名\n\n## 已确认豁免\n\n客户名称\n\n## 已确认别名\n\n写法 | 合成后的客户\n甲公司 | 甲\n\n## 草稿\n\n还没人点头。待抽清单里仍然留着。\n丙草稿\n", encoding="utf-8")
    judgment = tmp_path / "判断.json"
    write_judgment(judgment, [{"销售": "甲", "客户": "甲", "交付月份": 202601, "档": "补位", "原因": "这个销售还没有别的可抽"}])
    out = tmp_path / "out.xlsx"
    code = compose.main(["--facts", str(facts), "--news", str(news), "--zhiyun", str(zhiyun), "--check-month", "202609", "--config", str(config), "--judgment", str(judgment), "--out", str(out)])
    assert code == 0
    wb = load_workbook(out, data_only=True)
    assert names_of(wb["待抽查清单"]) == ["甲", "甲公司"]
    news_rows = list(wb["风险提示"].iter_rows(min_row=2, values_only=True))
    assert len(news_rows) == 1 and news_rows[0][0] == "甲"
    assert "https://example.com/a" in news_rows[0][2]
    assert "甲公司" in news_rows[0][3]
    exempt = list(wb["豁免与已回款"].iter_rows(min_row=2, values_only=True))
    assert exempt == [(None, "丙草稿", None, None, None, "待确认", None)]
    wb.close()


def test_confirmed_exempt_leaves_the_pool(tmp_path):
    facts = tmp_path / "facts.xlsx"
    write_facts(facts, [fact("甲", "免", 202601, 9, 0, ""), fact("甲", "不免", 202601, 2, 0, "")])
    news = tmp_path / "news"
    write_news(news, [{"customer": name, "summary": "未查到", "url": "", "note": "未查到"} for name in ("免", "不免")])
    zhiyun = tmp_path / "zhiyun.xlsx"
    write_zhiyun(zhiyun, ["免", "不免"])
    config = tmp_path / "config.md"
    config.write_text("## 已确认豁免\n\n免\n\n## 已确认别名\n\n## 草稿\n\n", encoding="utf-8")
    judgment = tmp_path / "判断.json"
    write_judgment(judgment, [{"销售": "甲", "客户": "不免", "交付月份": 202601, "档": "补位", "原因": "这个销售还没有别的可抽"}])
    out = tmp_path / "out.xlsx"
    assert compose.main(["--facts", str(facts), "--news", str(news), "--zhiyun", str(zhiyun), "--check-month", "202609", "--config", str(config), "--judgment", str(judgment), "--out", str(out)]) == 0
    wb = load_workbook(out, data_only=True)
    assert names_of(wb["待抽查清单"]) == ["不免", "免"]
    reasons = {row[1]: row[5] for row in wb["待抽查清单"].iter_rows(min_row=2, values_only=True)}
    assert reasons["免"] == "已豁免"
    assert reasons["不免"] == "没查过"
    exempt = list(wb["豁免与已回款"].iter_rows(min_row=2, values_only=True))
    assert exempt[0][1] == "免" and exempt[0][5] == "已豁免" and exempt[0][6] == "已确认豁免"
    wb.close()


def test_news_risk_sorts_high_to_low_and_colors_column_e(tmp_path):
    facts = tmp_path / "facts.xlsx"
    write_facts(facts, [fact("甲", "无新闻", 202601, 1, 0, ""), fact("甲", "有新闻", 202601, 1, 0, "")])
    news = tmp_path / "news"
    write_news(
        news,
        [
            {"customer": "无新闻", "summary": "未查到", "url": "", "note": "未查到"},
            {"customer": "有新闻", "summary": "被执行", "url": "https://example.com/a", "note": "法院", "risk": "高", "reason": "已被法院执行，付钱会受影响。"},
        ],
    )
    zhiyun = tmp_path / "zhiyun.xlsx"
    write_zhiyun(zhiyun, ["无新闻", "有新闻"])
    judgment = tmp_path / "判断.json"
    write_judgment(judgment, [{"销售": "甲", "客户": "有新闻", "交付月份": 202601, "档": "补位", "原因": "这个销售还没有别的可抽"}])
    out = tmp_path / "out.xlsx"
    assert compose.main(["--facts", str(facts), "--news", str(news), "--zhiyun", str(zhiyun), "--check-month", "202609", "--retrieved", "2026-09-22", "--judgment", str(judgment), "--out", str(out)]) == 0
    wb = load_workbook(out, data_only=True)
    rows = list(wb["风险提示"].iter_rows(min_row=2, values_only=True))
    assert [row[0] for row in rows] == ["有新闻", "无新闻"]
    assert rows[0][4] == "高" and rows[1][4] == "无"
    assert rows[0][5] == "已被法院执行，付钱会受影响。"
    assert rows[1][5] in (None, "")
    assert rows[0][6] == "2026-09-22"
    colored = wb["风险提示"].cell(2, 5)
    assert colored.fill.fgColor.rgb.endswith("F4C7C3")
    wb.close()


def test_asks_on_conflict_placeholder_and_missing_zhiyun(tmp_path):
    facts = tmp_path / "facts.xlsx"
    write_facts(facts, [fact("甲", "冲突", 202601, "", 0, "", conflict="是")])
    news = tmp_path / "news"
    write_news(news, [{"customer": "冲突", "summary": "未查到", "url": "", "note": "未查到"}])
    zhiyun = tmp_path / "zhiyun.xlsx"
    write_zhiyun(zhiyun, ["冲突"])
    out = tmp_path / "out.xlsx"
    assert compose.main(["--facts", str(facts), "--news", str(news), "--zhiyun", str(zhiyun), "--check-month", "202609", "--out", str(out)]) == 2
    assert not out.exists()

    write_facts(facts, [fact("甲", "甲", 202601, 1, 0, "")])
    write_news(news, [{"customer": "甲", "summary": "本次未检索", "url": "", "note": ""}])
    write_zhiyun(zhiyun, ["甲"])
    assert compose.main(["--facts", str(facts), "--news", str(news), "--zhiyun", str(zhiyun), "--check-month", "202609", "--out", str(out)]) == 2

    write_news(news, [{"customer": "甲", "summary": "未查到", "url": "", "note": "未查到"}])
    write_zhiyun(zhiyun, ["别人"])
    assert compose.main(["--facts", str(facts), "--news", str(news), "--zhiyun", str(zhiyun), "--check-month", "202609", "--out", str(out)]) == 2


def test_zhiyun_sales_follows_each_order(tmp_path):
    facts = tmp_path / "facts.xlsx"
    write_facts(facts, [fact("甲", "客户", 202601, 3, 0, "", so="SO1"), fact("乙", "客户", 202601, 3, 0, "", so="SO2")])
    news = tmp_path / "news"
    write_news(news, [{"customer": "客户", "summary": "未查到", "url": "", "note": "未查到"}])
    zhiyun = tmp_path / "zhiyun.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "智云核对"
    ws.append(["销售", "客户", "订单号", "合同归档号", "订单状态", "说明"])
    ws.append(["", "客户", "SO2", "20260001", "OP4/项目已交付", ""])
    ws.append(["", "客户", "SO1", "未找到", "OP1/项目确认中", "下单有这张单，没有合同归档号。"])
    wb.save(zhiyun)
    judgment = tmp_path / "判断.json"
    write_judgment(
        judgment,
        [
            {"销售": "甲", "客户": "客户", "交付月份": 202601, "档": "补位", "原因": "这个销售还没有别的可抽"},
            {"销售": "乙", "客户": "客户", "交付月份": 202601, "档": "补位", "原因": "这个销售还没有别的可抽"},
        ],
        orders=[
            {"客户": "客户", "订单号": "SO1", "核对": "不冲突"},
            {"客户": "客户", "订单号": "SO2", "核对": "不冲突"},
        ],
    )
    out = tmp_path / "out.xlsx"
    assert compose.main(["--facts", str(facts), "--news", str(news), "--zhiyun", str(zhiyun), "--check-month", "202609", "--judgment", str(judgment), "--out", str(out)]) == 0
    wb = load_workbook(out)
    sheet = wb["智云核对"]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert [(row[0], row[2]) for row in rows] == [("乙", "SO2"), ("甲", "SO1")]
    missing = sheet.cell(3, 4)
    assert missing.value == "未找到"
    assert str(missing.fill.fgColor.rgb).endswith("FCE4B3")
    assert not sheet.cell(2, 3).alignment.wrap_text
    pool = wb["待抽查清单"]
    assert pool.freeze_panes == "C2"
    assert not pool.cell(2, 1).alignment.wrap_text
    assert pool.cell(1, 6).comment is not None
    wb.close()


def test_exempt_and_fully_paid_is_marked_on_both_sheets(tmp_path):
    facts = tmp_path / "facts.xlsx"
    write_facts(facts, [fact("甲", "免", 202601, 9, 0, "", orders=2, paid=2, paid_so="SO1；SO2")])
    news = tmp_path / "news"
    write_news(news, [{"customer": "免", "summary": "未查到", "url": "", "note": "未查到"}])
    zhiyun = tmp_path / "zhiyun.xlsx"
    write_zhiyun(zhiyun, ["免"])
    ledger = tmp_path / "ledger.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "豁免清单"
    sheet.append(["豁免客户关键词", "豁免原因"])
    sheet.append(["免", "集团统一"])
    book.save(ledger)
    judgment = tmp_path / "判断.json"
    write_judgment(judgment, [], [{"词": "免", "含义": "客户名称", "豁免原因": "集团统一"}])
    out = tmp_path / "out.xlsx"
    assert compose.main([
        "--facts", str(facts), "--news", str(news), "--zhiyun", str(zhiyun),
        "--ledger", str(ledger), "--check-month", "202609", "--judgment", str(judgment), "--out", str(out),
    ]) == 0
    wb = load_workbook(out, data_only=True)
    pool = list(wb["待抽查清单"].iter_rows(min_row=2, values_only=True))
    assert pool[0][5] == "已豁免"
    assert "销售标了已回款，智云没有" in str(pool[0][9])
    assert "这个月每笔都已回款" not in str(pool[0][9])
    exempt = list(wb["豁免与已回款"].iter_rows(min_row=2, values_only=True))
    assert exempt[0][5] == "已豁免"
    assert exempt[0][6] == "集团统一"
    wb.close()


def test_suggest_note_matches_pool_when_zhiyun_has_not_confirmed(tmp_path):
    facts = tmp_path / "facts.xlsx"
    write_facts(facts, [fact("甲", "甲客户", 202401, 20, 0, "", orders=1, paid=1, so="SO1", paid_so="SO1", stages="SO1=已回款，未核销")])
    news = tmp_path / "news"
    write_news(news, [{"customer": "甲客户", "summary": "未查到", "url": "", "note": "未查到"}])
    zhiyun = tmp_path / "zhiyun.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "智云核对"
    sheet.append(["销售", "客户", "订单号", "合同归档号", "订单状态", "说明"])
    sheet.append(["甲", "甲客户", "SO1", "20260001", "OP5/销售已验收", ""])
    book.save(zhiyun)
    judgment = tmp_path / "判断.json"
    write_judgment(
        judgment,
        [{"销售": "甲", "客户": "甲客户", "交付月份": 202401, "档": "高危", "原因": "账龄已满24个月"}],
        orders=[{"客户": "甲客户", "订单号": "SO1", "核对": "不一致"}],
    )
    out = tmp_path / "out.xlsx"
    assert compose.main([
        "--facts", str(facts), "--news", str(news), "--zhiyun", str(zhiyun),
        "--check-month", "202609", "--judgment", str(judgment), "--out", str(out),
    ]) == 0
    wb = load_workbook(out, data_only=True)
    pool_note = list(wb["待抽查清单"].iter_rows(min_row=2, values_only=True))[0][9]
    suggest_note = list(wb["建议本次抽"].iter_rows(min_row=3, values_only=True))[0][8]
    assert pool_note == suggest_note == "销售标了已回款，智云没有"
    wb.close()


def test_credit_reason_is_the_row_itself(tmp_path):
    facts = tmp_path / "facts.xlsx"
    write_facts(
        facts,
        [
            fact("甲", "季结户", 202607, 2, 0, "", amount=5000, explain="客户季结"),
            fact("乙", "说不清", 202607, 2, 0, "", amount=5000),
        ],
    )
    news = tmp_path / "news"
    write_news(news, [
        {"customer": "季结户", "summary": "未查到", "url": "", "note": "未查到"},
        {"customer": "说不清", "summary": "未查到", "url": "", "note": "未查到"},
    ])
    zhiyun = tmp_path / "zhiyun.xlsx"
    write_zhiyun(zhiyun, ["季结户", "说不清"])
    canned = tmp_path / "canned.json"
    write_judgment(canned, [
        {"销售": "甲", "客户": "季结户", "交付月份": 202607, "档": "信用期内但要看", "原因": "还在信用期，但有经营风险、没合同、季结或预计回款已过"},
        {"销售": "乙", "客户": "说不清", "交付月份": 202607, "档": "信用期内但要看", "原因": "还在信用期，但有经营风险、没合同、季结或预计回款已过"},
    ])
    out = tmp_path / "out.xlsx"
    assert compose.main([
        "--facts", str(facts), "--news", str(news), "--zhiyun", str(zhiyun),
        "--check-month", "202609", "--judgment", str(canned), "--out", str(out),
    ]) == 2
    clear = tmp_path / "clear.json"
    write_judgment(clear, [
        {"销售": "甲", "客户": "季结户", "交付月份": 202607, "档": "信用期内但要看", "原因": "还在信用期，但有经营风险、没合同、季结或预计回款已过"},
        {"销售": "乙", "客户": "说不清", "交付月份": 202607, "档": "信用期内但要看", "原因": "销售口头说下月付清"},
    ])
    assert compose.main([
        "--facts", str(facts), "--news", str(news), "--zhiyun", str(zhiyun),
        "--check-month", "202609", "--judgment", str(clear), "--out", str(out),
    ]) == 0
    wb = load_workbook(out, data_only=True)
    reasons = {row[1]: row[7] for row in wb["建议本次抽"].iter_rows(min_row=3, values_only=True)}
    assert reasons["季结户"] == "季结"
    assert reasons["说不清"] == "销售口头说下月付清"
    wb.close()


def test_unmatched_keyword_is_marked_on_the_exempt_sheet(tmp_path):
    facts = tmp_path / "facts.xlsx"
    write_facts(facts, [fact("甲", "在册客户", 202601, 3, 0, "")])
    news = tmp_path / "news"
    write_news(news, [{"customer": "在册客户", "summary": "未查到", "url": "", "note": "未查到"}])
    zhiyun = tmp_path / "zhiyun.xlsx"
    write_zhiyun(zhiyun, ["在册客户"])
    ledger = tmp_path / "ledger.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "豁免清单"
    sheet.append(["序号", "豁免客户关键词", "豁免原因"])
    sheet.append([1, "对不上的词", "框架合同"])
    book.save(ledger)
    judgment = tmp_path / "判断.json"
    write_judgment(
        judgment,
        [{"销售": "甲", "客户": "在册客户", "交付月份": 202601, "档": "补位", "原因": "这个销售还没有别的可抽"}],
        [{"词": "对不上的词", "含义": "客户名称", "豁免原因": "框架合同"}],
    )
    out = tmp_path / "out.xlsx"
    assert compose.main([
        "--facts", str(facts), "--news", str(news), "--zhiyun", str(zhiyun),
        "--ledger", str(ledger), "--check-month", "202609", "--judgment", str(judgment), "--out", str(out),
    ]) == 0
    wb = load_workbook(out, data_only=True)
    rows = list(wb["豁免与已回款"].iter_rows(min_row=2, values_only=True))
    missed = [row for row in rows if row[1] == "对不上的词"]
    assert len(missed) == 1
    assert missed[0][0] in (None, "")
    assert missed[0][5] == "已豁免"
    assert missed[0][6] == "框架合同"
    wb.close()


def test_suggest_follows_age_amount_and_bad_debt(tmp_path):
    facts = tmp_path / "facts.xlsx"
    write_facts(
        facts,
        [
            fact("甲", "零", 202609, 0, 0, ""),
            fact("甲", "信用", 202608, 2, 0, ""),
            fact("甲", "没合同", 202607, 2, 0, "", amount=5000, marks="没有合同"),
            fact("甲", "逾期", 202605, 4, 0, "", amount=80000),
            fact("甲", "老", 202001, 30, 0, "", amount=20),
            fact("甲", "更大", 202001, 30, 0, "", amount=8000),
            fact("甲", "次危大", 202401, 10, 0, "", amount=120000),
            fact("甲", "坏", 202001, 40, 0, "", explain="这是坏账", marks="坏账"),
        ],
    )
    names = ["零", "信用", "没合同", "逾期", "老", "更大", "次危大", "坏"]
    news = tmp_path / "news"
    write_news(news, [{"customer": name, "summary": "未查到", "url": "", "note": "未查到"} for name in names])
    zhiyun = tmp_path / "zhiyun.xlsx"
    write_zhiyun(zhiyun, names)
    judgment = tmp_path / "判断.json"
    write_judgment(
        judgment,
        [
            {"销售": "甲", "客户": "更大", "交付月份": 202001, "档": "高危", "原因": "账龄已满24个月"},
            {"销售": "甲", "客户": "次危大", "交付月份": 202401, "档": "次危", "原因": "账龄已满6个月"},
            {"销售": "甲", "客户": "没合同", "交付月份": 202607, "档": "信用期内但要看", "原因": "还在信用期，但有经营风险、没合同、季结或预计回款已过"},
        ],
    )
    out = tmp_path / "out.xlsx"
    assert compose.main(["--facts", str(facts), "--news", str(news), "--zhiyun", str(zhiyun), "--check-month", "202609", "--judgment", str(judgment), "--out", str(out)]) == 0
    wb = load_workbook(out, data_only=True)
    suggest = [(row[1], row[6]) for row in wb["建议本次抽"].iter_rows(min_row=3, values_only=True)]
    assert [name for name, _band in suggest] == ["更大", "次危大", "没合同"]
    assert [band for _name, band in suggest] == ["高危", "次危", "信用期内但要看"]
    pool = names_of(wb["待抽查清单"])
    assert "坏" not in pool and "坏" not in [name for name, _band in suggest]
    reasons = [row[5] for row in wb["豁免与已回款"].iter_rows(min_row=2, values_only=True) if row[1] == "坏"]
    assert reasons == ["坏账"]
    wb.close()


def test_paid_needs_zhiyun_and_patent_news_drops(tmp_path):
    facts = tmp_path / "facts.xlsx"
    write_facts(facts, [fact("甲", "已回", 202401, 20, 0, "", orders=1, paid=1, so="SO9", paid_so="SO9", stages="SO9=已回款，已核销")])
    news = tmp_path / "news"
    write_news(news, [{"customer": "已回", "summary": "申请了一项专利", "url": "https://example.com/p", "note": "公告", "risk": "中"}])
    zhiyun = tmp_path / "zhiyun.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "智云核对"
    ws.append(["销售", "客户", "订单号", "合同归档号", "订单状态", "说明"])
    ws.append(["甲", "已回", "SO9", "20260009", "SP4/已回款", ""])
    wb.save(zhiyun)
    judgment = tmp_path / "判断.json"
    write_judgment(judgment, [], orders=[{"客户": "已回", "订单号": "SO9", "核对": "不冲突"}])
    out = tmp_path / "out.xlsx"
    assert compose.main(["--facts", str(facts), "--news", str(news), "--zhiyun", str(zhiyun), "--check-month", "202609", "--judgment", str(judgment), "--out", str(out)]) == 0
    wb = load_workbook(out, data_only=True)
    exempt = list(wb["豁免与已回款"].iter_rows(min_row=2, values_only=True))
    assert exempt[0][1] == "已回" and exempt[0][5] == "销售和智云都已回款"
    assert "已回" not in names_of(wb["待抽查清单"])
    assert list(wb["风险提示"].iter_rows(min_row=2, values_only=True))[0][4] == "中"
    sheet = wb["智云核对"]
    assert sheet.cell(1, 8).value == "智云订单状态核对"
    verdict = list(sheet.iter_rows(min_row=2, values_only=True))[0]
    assert verdict[6] == "已回款，已核销"
    assert verdict[7] == "不冲突"
    wb.close()


def test_keyword_reads_person_and_stops_before_note(tmp_path):
    facts = tmp_path / "facts.xlsx"
    write_facts(
        facts,
        [
            fact("梁玲玲-高美杰", "在册客户", 202601, 4, 0, ""),
            fact("甲", "别人", 202601, 4, 0, ""),
        ],
    )
    news = tmp_path / "news"
    write_news(news, [{"customer": name, "summary": "未查到", "url": "", "note": "未查到"} for name in ("在册客户", "别人")])
    zhiyun = tmp_path / "zhiyun.xlsx"
    write_zhiyun(zhiyun, ["在册客户", "别人"])
    ledger = tmp_path / "ledger.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "豁免清单"
    sheet.append(["序号", "豁免客户关键词", "豁免原因"])
    sheet.append([1, "在册客户", "单独核对"])
    sheet.append([2, "高美杰", "高美杰"])
    sheet.append([None, None, None])
    sheet.append(["豁免口径（供参考）", None, None])
    sheet.append(["交付月份豁免", "交付月份早于 202401", None])
    sheet.append(["GM单号", "以 GM 为前缀的订单豁免", None])
    book.save(ledger)
    judgment = tmp_path / "判断.json"
    write_judgment(
        judgment,
        [{"销售": "甲", "客户": "别人", "交付月份": 202601, "档": "补位", "原因": "这个销售还没有别的可抽"}],
        [
            {"词": "在册客户", "含义": "客户名称", "豁免原因": "单独核对"},
            {"词": "高美杰", "含义": "销售人员", "豁免原因": "高美杰"},
        ],
    )
    out = tmp_path / "out.xlsx"
    assert compose.main([
        "--facts", str(facts), "--news", str(news), "--zhiyun", str(zhiyun),
        "--ledger", str(ledger), "--check-month", "202609", "--judgment", str(judgment), "--out", str(out),
    ]) == 0
    wb = load_workbook(out, data_only=True)
    exempt = list(wb["豁免与已回款"].iter_rows(min_row=2, values_only=True))
    reasons = {row[1]: row[6] for row in exempt}
    assert reasons["在册客户"] == "单独核对；高美杰"
    assert "别人" not in reasons
    assert "交付月份早于 202401" not in reasons
    assert "以 GM 为前缀的订单豁免" not in reasons
    assert all("口径" not in str(row[1]) for row in exempt)
    pool = {row[1]: row[5] for row in wb["待抽查清单"].iter_rows(min_row=2, values_only=True)}
    assert pool["在册客户"] == "已豁免"
    assert pool["别人"] == "没查过"
    suggest = [row[1] for row in wb["建议本次抽"].iter_rows(min_row=3, values_only=True) if row[1]]
    assert "在册客户" not in suggest
    wb.close()


def test_same_person_counts_as_one_salesperson(tmp_path):
    facts = tmp_path / "facts.xlsx"
    write_facts(
        facts,
        [
            fact("梁玲玲", "甲客户", 202401, 20, 0, "", amount=8000),
            fact("梁玲玲-高美杰", "乙客户", 202401, 20, 0, "", amount=9000),
        ],
    )
    news = tmp_path / "news"
    write_news(news, [{"customer": name, "summary": "未查到", "url": "", "note": "未查到"} for name in ("甲客户", "乙客户")])
    zhiyun = tmp_path / "zhiyun.xlsx"
    write_zhiyun(zhiyun, ["甲客户", "乙客户"])
    judgment = tmp_path / "判断.json"
    write_judgment(
        judgment,
        [{"销售": "梁玲玲", "客户": "甲客户", "交付月份": 202401, "档": "高危", "原因": "账龄已满24个月"}],
        people=[{"算作": "梁玲玲", "写成": ["梁玲玲", "梁玲玲-高美杰"]}],
    )
    out = tmp_path / "out.xlsx"
    assert compose.main([
        "--facts", str(facts), "--news", str(news), "--zhiyun", str(zhiyun),
        "--check-month", "202609", "--judgment", str(judgment), "--out", str(out),
    ]) == 0
    wb = load_workbook(out, data_only=True)
    assert [row[1] for row in wb["建议本次抽"].iter_rows(min_row=3, values_only=True) if row[1]] == ["甲客户"]
    wb.close()


def test_missing_salesperson_is_asked(tmp_path):
    facts = tmp_path / "facts.xlsx"
    write_facts(facts, [fact("甲", "甲客户", 202401, 20, 0, ""), fact("乙", "乙客户", 202401, 20, 0, "")])
    news = tmp_path / "news"
    write_news(news, [{"customer": name, "summary": "未查到", "url": "", "note": "未查到"} for name in ("甲客户", "乙客户")])
    zhiyun = tmp_path / "zhiyun.xlsx"
    write_zhiyun(zhiyun, ["甲客户", "乙客户"])
    judgment = tmp_path / "判断.json"
    write_judgment(judgment, [{"销售": "甲", "客户": "甲客户", "交付月份": 202401, "档": "高危", "原因": "账龄已满24个月"}])
    out = tmp_path / "out.xlsx"
    code = compose.main([
        "--facts", str(facts), "--news", str(news), "--zhiyun", str(zhiyun),
        "--check-month", "202609", "--judgment", str(judgment), "--out", str(out),
    ])
    assert code == 2


def test_status_check_blanks_when_it_cannot_see_both_sides(tmp_path):
    facts = tmp_path / "facts.xlsx"
    write_facts(
        facts,
        [
            fact("甲", "甲客户", 202401, 20, 0, "", so="SO1", stages="SO1=已回款，未核销"),
            fact("甲", "乙客户", 202401, 4, 0, "", so="SO2", stages="SO2=未对账"),
        ],
    )
    news = tmp_path / "news"
    write_news(news, [{"customer": name, "summary": "未查到", "url": "", "note": "未查到"} for name in ("甲客户", "乙客户")])
    zhiyun = tmp_path / "zhiyun.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "智云核对"
    sheet.append(["销售", "客户", "订单号", "合同归档号", "订单状态", "说明"])
    sheet.append(["甲", "甲客户", "SO1", "未找到", "未找到", "下单里没有这个单号。"])
    sheet.append(["甲", "乙客户", "SO2", "20260001", "OP5/销售已验收", ""])
    book.save(zhiyun)
    judgment = tmp_path / "判断.json"
    write_judgment(
        judgment,
        [{"销售": "甲", "客户": "乙客户", "交付月份": 202401, "档": "高危", "原因": "账龄已满24个月"}],
        orders=[{"客户": "乙客户", "订单号": "SO2", "核对": "不冲突"}],
    )
    out = tmp_path / "out.xlsx"
    assert compose.main([
        "--facts", str(facts), "--news", str(news), "--zhiyun", str(zhiyun),
        "--check-month", "202609", "--judgment", str(judgment), "--out", str(out),
    ]) == 0
    wb = load_workbook(out)
    rows = {row[2]: row for row in wb["智云核对"].iter_rows(min_row=2, values_only=True)}
    assert rows["SO1"][7] in (None, "")
    assert rows["SO2"][7] == "不冲突"
    for index in range(2, 4):
        cell = wb["智云核对"].cell(index, 8)
        if cell.value == "不冲突":
            assert cell.fill.fgColor is None or cell.fill.fgColor.rgb in (None, "00000000")
    wb.close()


def test_status_conflict_is_red_and_missing_judgment_asks(tmp_path):
    facts = tmp_path / "facts.xlsx"
    write_facts(facts, [fact("甲", "甲客户", 202401, 20, 0, "", so="SO1", stages="SO1=未对账")])
    news = tmp_path / "news"
    write_news(news, [{"customer": "甲客户", "summary": "未查到", "url": "", "note": "未查到"}])
    zhiyun = tmp_path / "zhiyun.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "智云核对"
    sheet.append(["销售", "客户", "订单号", "合同归档号", "订单状态", "说明"])
    sheet.append(["甲", "甲客户", "SO1", "20260001", "SP4/已回款", ""])
    book.save(zhiyun)
    bare = tmp_path / "bare.json"
    write_judgment(bare, [{"销售": "甲", "客户": "甲客户", "交付月份": 202401, "档": "高危", "原因": "账龄已满24个月"}])
    out = tmp_path / "out.xlsx"
    assert compose.main([
        "--facts", str(facts), "--news", str(news), "--zhiyun", str(zhiyun),
        "--check-month", "202609", "--judgment", str(bare), "--out", str(out),
    ]) == 2
    judged = tmp_path / "judged.json"
    write_judgment(
        judged,
        [{"销售": "甲", "客户": "甲客户", "交付月份": 202401, "档": "高危", "原因": "账龄已满24个月"}],
        orders=[{"客户": "甲客户", "订单号": "SO1", "核对": "不一致"}],
    )
    assert compose.main([
        "--facts", str(facts), "--news", str(news), "--zhiyun", str(zhiyun),
        "--check-month", "202609", "--judgment", str(judged), "--out", str(out),
    ]) == 0
    wb = load_workbook(out)
    cell = wb["智云核对"].cell(2, 8)
    assert cell.value == "不一致"
    assert cell.fill.fgColor.rgb == "00F4C7C3"
    wb.close()


def test_asks_without_judgment(tmp_path):
    facts = tmp_path / "facts.xlsx"
    write_facts(facts, [fact("甲", "甲客户", 202401, 20, 0, "")])
    news = tmp_path / "news"
    write_news(news, [{"customer": "甲客户", "summary": "未查到", "url": "", "note": "未查到"}])
    zhiyun = tmp_path / "zhiyun.xlsx"
    write_zhiyun(zhiyun, ["甲客户"])
    out = tmp_path / "out.xlsx"
    code = compose.main(["--facts", str(facts), "--news", str(news), "--zhiyun", str(zhiyun), "--check-month", "202609", "--out", str(out)])
    assert code == 2
