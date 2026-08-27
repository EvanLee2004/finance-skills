#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把连续多日的日期级报表整理为三份任务范围版静态 Excel。

多日任务的日期级 Excel 只是逐日校验和写入时的临时产物。三份范围版全部成功
生成后，自动删除本次日期范围内的日期级同类报告，只保留范围版交付文件。
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import common  # noqa: E402
import workbook_finalize  # noqa: E402


REPORTS = (
    ("核销日清", "核销日清"),
    ("变更清单", "变更清单"),
    ("订单写入差异", "订单写入差异"),
)


def _dates(start: str, end: str):
    current = dt.date.fromisoformat(start)
    finish = dt.date.fromisoformat(end)
    if current > finish:
        raise ValueError("开始日期不能晚于结束日期")
    while current <= finish:
        yield current
        current += dt.timedelta(days=1)


def _copy_sheet(source, target) -> None:
    for row in source.iter_rows():
        for cell in row:
            out = target[cell.coordinate]
            out.value = cell.value
            # 范围版和多年度合并版都是静态审计报表。源报表中的「公式原文」
            # 以 = 开头，但它只是供人核对的说明文字；openpyxl 赋值到新
            # 单元格时会把它重新识别成可执行公式，导致静态报表校验失败。
            if isinstance(cell.value, str) and cell.value.startswith("="):
                out.data_type = "s"
            if cell.has_style:
                out._style = copy(cell._style)
            if cell.number_format:
                out.number_format = cell.number_format
            out.alignment = copy(cell.alignment)
            out.protection = copy(cell.protection)
    for key, dim in source.column_dimensions.items():
        target.column_dimensions[key].width = dim.width
        target.column_dimensions[key].hidden = dim.hidden
    for key, dim in source.row_dimensions.items():
        target.row_dimensions[key].height = dim.height
        target.row_dimensions[key].hidden = dim.hidden
    for merged in source.merged_cells.ranges:
        target.merge_cells(str(merged))
    target.freeze_panes = source.freeze_panes
    target.sheet_view.showGridLines = source.sheet_view.showGridLines
    target.auto_filter.ref = source.auto_filter.ref


def _daily_counts(out_dir: Path, day: dt.date) -> dict:
    token = day.strftime("%Y%m%d")
    result_path = out_dir / f"判定结果_{token}.json"
    if not result_path.is_file():
        return {}
    payload = json.loads(result_path.read_text(encoding="utf-8"))
    counts = payload.get("counts") or {}
    return {
        "到账笔数": payload.get("payment_count", ""),
        "订单行数": counts.get("total", ""),
        "自动": counts.get("auto", ""),
        "挂账": counts.get("hold", ""),
        "异常": counts.get("exception", ""),
    }


def _is_empty_fetched_day(workspace: Path, day: dt.date) -> bool:
    summary_path = workspace / "01_智云导出" / f"取数摘要_{day.strftime('%Y%m%d')}.json"
    if not summary_path.is_file():
        return False
    try:
        summary = json.loads(summary_path.read_text(encoding="utf-8"))
    except (OSError, ValueError, TypeError):
        return False
    if not isinstance(summary, dict):
        return False
    keys = ("回款记录笔数", "下单行数", "核销明细行数", "订单明细SOD行数")
    values = [summary.get(key) for key in keys]
    return bool(
        all(isinstance(value, int) and not isinstance(value, bool) for value in values)
        and all(value == 0 for value in values)
    )


def _daily_report_files(out_dir: Path, prefix: str, day: dt.date) -> list[Path]:
    """返回某核销日的正式日报及多年度内部报告，不误匹配范围版文件。"""
    token = day.strftime("%Y%m%d")
    name_re = re.compile(
        rf"^{re.escape(prefix)}_{token}(?:_\d{{4}})?\.xlsx$"
    )
    return sorted(
        path for path in out_dir.glob(f"{prefix}_{token}*.xlsx")
        if name_re.fullmatch(path.name)
    )


def build(workspace: Path, start: str, end: str) -> list[Path]:
    workspace = Path(workspace)
    if not (workspace / "04_产出").is_dir():
        workspace = common.resolve_workspace(workspace)
    out_dir = workspace / "04_产出"
    days = list(_dates(start, end))
    empty_days = {day for day in days if _is_empty_fetched_day(workspace, day)}
    outputs = []
    daily_sources: set[Path] = set()
    for prefix, label in REPORTS:
        wb = openpyxl.Workbook()
        summary = wb.active
        summary.title = "任务范围"
        headers = ["核销日期", "日报状态", "到账笔数", "订单行数", "自动", "挂账", "异常"]
        summary.append(headers)
        for cell in summary[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
        for day in days:
            token = day.strftime("%Y%m%d")
            daily = out_dir / f"{prefix}_{token}.xlsx"
            daily_sources.update(_daily_report_files(out_dir, prefix, day))
            counts = _daily_counts(out_dir, day)
            status = (
                "已纳入"
                if daily.is_file()
                else "无核销记录，已跳过"
                if day in empty_days
                else "当日无该报表"
            )
            summary.append([
                day.isoformat(), status,
                counts.get("到账笔数", ""), counts.get("订单行数", ""),
                counts.get("自动", ""), counts.get("挂账", ""), counts.get("异常", ""),
            ])
            if not daily.is_file():
                if prefix == "核销日清" and day not in empty_days:
                    raise FileNotFoundError(f"缺少日期级核销日清：{daily}")
                continue
            source_wb = openpyxl.load_workbook(daily, data_only=False, read_only=False)
            for index, source_ws in enumerate(source_wb.worksheets, 1):
                title = f"{day.strftime('%Y%m%d')}_{index}_{source_ws.title}"[:31]
                _copy_sheet(source_ws, wb.create_sheet(title))
            source_wb.close()
        summary.freeze_panes = "A2"
        summary.auto_filter.ref = summary.dimensions
        for col, width in zip("ABCDEFG", (14, 16, 12, 12, 10, 10, 10)):
            summary.column_dimensions[col].width = width
        target = out_dir / f"{label}_{start.replace('-', '')}_{end.replace('-', '')}.xlsx"
        wb.save(target)
        workbook_finalize.finalize_static_report(target)
        outputs.append(target)

    # 只有多日任务执行清理。必须等三份范围版全部保存并完成静态化后再删除日报，
    # 这样任一步失败时仍保留逐日证据，可安全重试。
    if len(days) > 1:
        final_paths = {path.resolve() for path in outputs}
        for source in sorted(daily_sources):
            if source.resolve() not in final_paths and source.is_file():
                source.unlink()
    return outputs


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(
        description="生成连续多日任务范围版三份报表，并移除日期级同类报告"
    )
    parser.add_argument("--workspace", default=str(common.WORK))
    parser.add_argument("--date-from", required=True)
    parser.add_argument("--date-to", required=True)
    args = parser.parse_args(argv)
    try:
        outputs = build(Path(args.workspace), args.date_from, args.date_to)
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2
    for path in outputs:
        print(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
