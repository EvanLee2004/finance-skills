#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
第 1 步：智云只读取数 → 工作区 01_智云导出/

取数架构（2026-07-23 明妹当面口径 + Claude 接口实调重写，v2）
────────────────────────────────────────────────────────────
**入口只有一个**：财务管理 →「回款记录」，**只筛一个条件：核销日期 = T-1**。
然后把每笔回款记录**点进去的明细页**原样抓下来——这正是明妹手工在做的事：

  回款记录（1 笔到账）
    ├─ 关联「下单」            → 这笔到账关联了哪几个 SO + 每个 SO 的交付额
    │   └─ 下单为空时回查「结算」→ 从结算关联补找 SO，再继续取 SOD/交付额
    ├─ 关联「订单同币种核销明细信息」→ 逐 SO 的本次核销金额
    │                            **0 行 = 全额核销**（明妹原话：没有这张就说明到账=交付）
    └─ 由 SO 去「订单明细」表   → 每个 SO 下的 SOD + 逐 SOD 交付额
                                 （SOD 的唯一来源；她盈亏表一行 = 一个 SOD）

**已废弃**：按核销日期单独筛「回款核销对账」表。原因（2026-07-23 实调）：
  1. 它对预存/预收类返回 0 行，靠"按回款类型走不同通道"来补，而通道路由正是
     2026-07-22 那次**3 笔到账 6.5 万静默消失**的根因（预收类既被踢出对账通道、
     又在明细子表拿不到行，两头落空还不报错）。
  2. 它拿不到 SOD 与逐 SOD 交付额，判不出"一个 SO 拆 N 个 SOD"这个 45% 的主场景。
  3. 它是独立筛的第二张表，与回款记录可能不同步；关联读法天然同步。

红线：只读（GetFilterRows / getRowRelationRows / getWorksheetInfo，无任何写接口）；
  账号密码只从本机 Windows 凭据库或环境变量读取，绝不写进代码 / config / git；
  自动任务不会在取数过程中弹出账号密码询问。

用法：
  export ZHIYUN_USER='你的智云账号'
  export ZHIYUN_PASS   # 在 shell 里 export，勿写进任何文件；用完 unset
  python3 scripts/fetch_zhiyun.py --date 2026-07-22 --workspace 工作区/

  或预先配置 MD_PSS_ID 后运行：
    python3 scripts/fetch_zhiyun.py --date yesterday --workspace 工作区/

依赖：playwright（登录）+ requests（取数）+ openpyxl（写出 xlsx）
  pip install playwright requests openpyxl && playwright install chromium
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
import tempfile
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

# ── 常量（表 ID / 字段 ID 来自 2026-07-09/22/23 勘探，非密钥）────────────────
BASE_DEFAULT = "http://192.168.10.167:18880"


def _assert_platform_network_url(url: str) -> None:
    if os.environ.get("FINANCIAL_NETWORK_POLICY_REQUIRED") != "1":
        return
    from urllib.parse import urlsplit

    host = (urlsplit(url).hostname or "").encode("idna").decode("ascii").lower()
    allowed = {
        item.strip().lower()
        for item in os.environ.get("FINANCIAL_NETWORK_ALLOWLIST", "").split(",")
        if item.strip()
    }
    if os.environ.get("FINANCIAL_NETWORK_ACCESS") != "1" or host not in allowed:
        raise RuntimeError("网络目标不在平台批准的精确域名白名单中。")
APP_ID = "6ff4fb2e-e68c-4ee9-83a0-836de8f72c11"
EXPORT_SCHEMA_VERSION = "2026-08-21-atomic-fetch-v5"
CREDENTIAL_SERVICE = "codex.ar-hexiao-daily.zhiyun"
READ_REQUEST_MAX_ATTEMPTS = 3
MAX_BATCH_FETCH_CONCURRENCY = 8

WS_HUIKUAN = "6555d2b1f9460e517040ba6c"  # 回款记录（唯一入口）

# 回款记录字段（只保留新架构真正用到的）
F_HK = {
    "arrival_date": "6555d2b1f9460e517040ba70",
    "amount_orig": "6555d2b1f9460e517040ba71",
    "currency": "663c495f4a2483fd86249c90",
    "ar": "663c4b8a4a2483fd86249ca8",
    "amount_local": "663c4d354a2483fd86249cbb",
    "status": "663ca2204a2483fd8624a107",
    "hexiao_date": "664c571b4a2483fd86250fa5",
    "huikuan_type": "66b1f6986429811e5f304a0b",
    "fee": "675119ab327314202700730a",
    "customer_rel": "676a3314327314202700e7fc",
    "customer_txt": "676b6d01327314202700ed59",
}

# 关联字段按**中文名**取（controlId 会随配置变，名字不会）
REL_XIADAN = "下单"
REL_JIESUAN = "结算"
REL_HEXIAO_MINGXI = "订单同币种核销明细信息"
REL_SODLINE = "订单明细"  # 只借它的 dataSource 定位「订单明细」表

# 关联子表里要取的列（按中文名，取不到就空，不猜）
XIADAN_COLS = [
    "SO", "订单NUM", "订单号", "新智云单号",
    "订单已核销金额", "订单已核销金额/本币",
    "交付额/原币", "汇率", "结算币种", "订单名称", "项目交付日期",
]
# ⚠「同币种核销明细信息」表里**没有**叫 SO 的字段，SO 藏在关联字段「订单NUM」的 name 里
#   （2026-07-23 实调：该表字段 = 核销记录NUM/回款记录NUM/订单NUM/本次核销金额/…）。
#   旧版靠"按金额跟下单栏配对"猜 SO，金额一撞就配错；这一版直接从订单NUM读，不猜。
MINGXI_COLS = [
    "核销记录NUM", "回款记录NUM", "订单NUM", "本次核销金额", "本次核销金额本币", "核销日期",
    "币种", "汇率", "订单名称", "是否已撤销",
]
SODLINE_COLS = ["SO", "SOD", "交付额/原币", "币种", "项目状态"]


class LoginError(RuntimeError):
    pass


class FetchError(RuntimeError):
    pass


def _repair_mojibake(s: str) -> str:
    """修复智云部分选项标签按 GBK 字节误解成 Latin-1 的文本。"""
    try:
        repaired = s.encode("latin1").decode("gbk")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return s
    has_cjk = any("\u4e00" <= ch <= "\u9fff" for ch in repaired)
    original_has_cjk = any("\u4e00" <= ch <= "\u9fff" for ch in s)
    return repaired if has_cjk and not original_has_cjk else s


def _plain(v: Any, options: Optional[Dict[str, str]] = None) -> str:
    """明道云单元格 → 纯文本。"""
    if v is None:
        return ""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return str(v)
    s = str(v).strip()
    if not s:
        return ""
    if s.startswith("["):
        try:
            arr = json.loads(s)
        except ValueError:
            return s
        out: List[str] = []
        for item in arr:
            if isinstance(item, dict):
                out.append(
                    item.get("name")
                    or item.get("fullname")
                    or item.get("sourcevalue")
                    or item.get("departmentName")
                    or ""
                )
            elif options and str(item) in options:
                out.append(options[str(item)])
            else:
                out.append(str(item))
        return _repair_mojibake("、".join(x for x in out if x))
    if options and s in options:
        return _repair_mojibake(options[s])
    return _repair_mojibake(s)


def extract_so(text: str) -> str:
    """从 'ZG SO26060433' / 'SO26060433' / '[{...name:SO...}]' 里抠出 SO 号。"""
    s = _plain(text)
    for token in s.replace("\n", " ").replace("、", " ").replace(",", " ").split():
        t = token.strip().upper()
        if t.startswith("SO") and not t.startswith("SOD") and len(t) >= 8:
            return token.strip()
    return ""


def extract_ar(text: str) -> str:
    """从关联字段文本中提取 AR 号。"""
    s = _plain(text)
    for token in s.replace("\n", " ").replace("、", " ").replace(",", " ").split():
        t = token.strip().upper()
        if t.startswith("AR") and len(t) >= 8:
            return token.strip()
    return ""


def resolve_date(s: str) -> str:
    """
    把 --date 解析成具体某一天。**这里指的永远是「核销日期」**
    （销售哪天在智云把钱核到订单上），不是钱哪天到银行的「到账日期」——
    明妹口径：两者没有固定隔天关系，天然可能差好几天。

    另外：`yesterday` 是相对**运行那一刻**算的，她晚上跑和第二天早上跑不是同一天。
    主流程在接到核销指令后立即把它解析成绝对日期，并在本批内部固定使用，不再二次询问。
    """
    s = (s or "").strip().lower()
    if s in ("yesterday", "t-1", "昨天"):
        return (date.today() - timedelta(days=1)).isoformat()
    if s in ("today", "今天"):
        return date.today().isoformat()
    if s in ("last-workday", "上个工作日", "上一个工作日"):
        cur = date.today() - timedelta(days=1)
        while cur.weekday() >= 5:
            cur -= timedelta(days=1)
        return cur.isoformat()
    datetime.strptime(s, "%Y-%m-%d")
    return s


def resolve_fetch_dates(
    *,
    single_date: str = "",
    date_from: str = "",
    date_to: str = "",
    include_weekends: bool = False,
) -> List[str]:
    """把单日或日期范围解析为固定、升序的核销日期列表。"""
    start_raw = (date_from or "").strip()
    end_raw = (date_to or "").strip()
    if bool(start_raw) != bool(end_raw):
        raise ValueError("--date-from 与 --date-to 必须同时提供")
    if start_raw:
        if (single_date or "").strip():
            raise ValueError("日期范围不能与 --date/--hexiao-date 同时使用")
        start = date.fromisoformat(resolve_date(start_raw))
        end = date.fromisoformat(resolve_date(end_raw))
        if start > end:
            raise ValueError("--date-from 不能晚于 --date-to")
        days: List[str] = []
        current = start
        while current <= end:
            if include_weekends or current.weekday() < 5:
                days.append(current.isoformat())
            current += timedelta(days=1)
        if not days:
            raise ValueError("日期范围内没有需要取数的核销日；周末任务请加 --all-days")
        return days
    return [resolve_date(single_date or "yesterday")]


def login_with_password(
    base_url: str, username: str, password: str, headless: bool = True
) -> Tuple[str, Optional[str]]:
    """账号密码 → (md_pss_id, account_id)。不落盘。"""
    try:
        from playwright.sync_api import sync_playwright
    except ImportError as e:
        raise LoginError(
            "未安装 playwright。请: pip install playwright && playwright install chromium"
        ) from e

    account_sel = "#txtMobilePhone"
    password_sel = "input[type=password]"
    btn_sels = ["text=登 录", "text=登录", ".loginBtn"]

    try:
        with sync_playwright() as p:
            br = p.chromium.launch(headless=headless)
            try:
                ctx = br.new_context(ignore_https_errors=True)
                pg = ctx.new_page()
                pg.goto(base_url, wait_until="networkidle", timeout=45000)
                pg.fill(account_sel, username)
                pg.fill(password_sel, password)
                clicked = False
                for sel in btn_sels:
                    try:
                        pg.click(sel, timeout=2500)
                        clicked = True
                        break
                    except Exception:
                        continue
                if not clicked:
                    pg.keyboard.press("Enter")
                pg.wait_for_timeout(7000)
                token = None
                for c in ctx.cookies():
                    if c.get("name") == "md_pss_id" and c.get("value"):
                        token = c["value"]
                        break
                if not token:
                    raise LoginError(
                        f"登录后未拿到 md_pss_id（url={pg.url}）。账号密码可能错误或页面结构变了。"
                    )
                acct = None
                try:
                    acct = pg.evaluate(
                        "() => { try { return (md && md.global && md.global.Account && "
                        "md.global.Account.accountId) || null; } catch(e) { return null; } }"
                    )
                except Exception:
                    acct = None
                return token, acct
            finally:
                br.close()
    except LoginError:
        raise
    except Exception as e:
        raise LoginError(f"登录异常 {type(e).__name__}: {e}") from e


class ZhiyunClient:
    def __init__(self, base: str, cookie: str, account_id: str = "", page_size: int = 200):
        import requests

        self.base = base.rstrip("/")
        self.page_size = page_size
        self.session = requests.Session()
        self.headers = {
            "Content-Type": "application/json",
            "Authorization": f"md_pss_id {cookie}",
            "X-Requested-With": "XMLHttpRequest",
        }
        if account_id:
            self.headers["AccountId"] = account_id
        self._tpl_cache: Dict[str, List[dict]] = {}

    def close(self) -> None:
        self.session.close()

    def post(self, path: str, body: dict, timeout: int = 90) -> dict:
        import requests

        url = f"{self.base}/wwwapi/{path.lstrip('/')}"
        _assert_platform_network_url(url)
        r = None
        for attempt in range(1, READ_REQUEST_MAX_ATTEMPTS + 1):
            try:
                r = self.session.post(
                    url,
                    headers=self.headers,
                    json=body,
                    timeout=timeout,
                    allow_redirects=False,
                )
            except (requests.ConnectionError, requests.Timeout) as exc:
                if attempt >= READ_REQUEST_MAX_ATTEMPTS:
                    raise
                delay = float(2 ** (attempt - 1))
                print(
                    "WARN: 智云只读请求暂时不可用"
                    f"（{type(exc).__name__}），{delay:g} 秒后重试"
                    f"（{attempt + 1}/{READ_REQUEST_MAX_ATTEMPTS}）",
                    file=sys.stderr,
                )
                time.sleep(delay)
                continue
            if (
                (r.status_code in {408, 429} or 500 <= r.status_code < 600)
                and attempt < READ_REQUEST_MAX_ATTEMPTS
            ):
                delay = float(2 ** (attempt - 1))
                retry_after = str(r.headers.get("Retry-After") or "").strip()
                if retry_after.isdigit():
                    delay = min(10.0, max(delay, float(retry_after)))
                print(
                    f"WARN: 智云只读请求暂时返回 HTTP {r.status_code}，"
                    f"{delay:g} 秒后重试（{attempt + 1}/{READ_REQUEST_MAX_ATTEMPTS}）",
                    file=sys.stderr,
                )
                time.sleep(delay)
                continue
            r.raise_for_status()
            break
        if r is None:  # pragma: no cover - 循环至少执行一次
            raise RuntimeError("智云只读请求未执行。")
        j = r.json()
        if isinstance(j, dict) and "data" in j:
            return j["data"] if j["data"] is not None else {}
        return j

    # ── 模板 / 字段 ──────────────────────────────────────────────
    def controls(self, worksheet_id: str) -> List[dict]:
        if worksheet_id in self._tpl_cache:
            return self._tpl_cache[worksheet_id]
        info = self.post(
            "worksheet/getWorksheetInfo",
            {"worksheetId": worksheet_id, "appId": APP_ID, "getTemplate": True},
        )
        ctrls = (info.get("template") or {}).get("controls") or info.get("controls") or []
        self._tpl_cache[worksheet_id] = ctrls
        return ctrls

    @staticmethod
    def name_map(controls: Sequence[dict]) -> Dict[str, str]:
        """controlId → 中文名。"""
        return {
            (c.get("controlId") or c.get("id") or ""): (c.get("controlName") or c.get("name") or "")
            for c in controls
        }

    @staticmethod
    def id_by_name(controls: Sequence[dict], name: str) -> str:
        for c in controls:
            if (c.get("controlName") or c.get("name") or "") == name:
                return c.get("controlId") or c.get("id") or ""
        return ""

    @staticmethod
    def option_maps(controls: Sequence[dict]) -> Dict[str, Dict[str, str]]:
        """controlId → {optionKey: 显示值}。"""
        out: Dict[str, Dict[str, str]] = {}
        for c in controls:
            opts = c.get("options") or []
            if opts:
                cid = c.get("controlId") or c.get("id") or ""
                out[cid] = {str(o.get("key")): str(o.get("value")) for o in opts}
        return out

    def datasource_of(self, worksheet_id: str, control_name: str) -> str:
        for c in self.controls(worksheet_id):
            if (c.get("controlName") or c.get("name") or "") == control_name:
                return c.get("dataSource") or ""
        return ""

    # ── 取行 ────────────────────────────────────────────────────
    def filter_rows_by_date(
        self, worksheet_id: str, date_control_id: str, day: str, page_size: Optional[int] = None
    ) -> Tuple[List[dict], int]:
        """按日期字段筛单日（filterType 11 + dateRange 18 已实测）。**只此一个条件。**"""
        ps = page_size or self.page_size
        rows: List[dict] = []
        page = 1
        total = 0
        while True:
            d = self.post(
                "worksheet/getFilterRows",
                {
                    "worksheetId": worksheet_id,
                    "appId": APP_ID,
                    "pageSize": ps,
                    "pageIndex": page,
                    "status": 1,
                    "sortControls": [],
                    "notGetTotal": False,
                    "searchType": 1,
                    "keyWords": "",
                    "filterControls": [
                        {
                            "controlId": date_control_id,
                            "dataType": 30,
                            "spliceType": 1,
                            "filterType": 11,
                            "dateRange": 18,
                            "minValue": day,
                            "maxValue": day,
                            "value": "",
                            "values": [],
                        }
                    ],
                    "fastFilters": [],
                    "navGroupFilters": [],
                },
            )
            batch = d.get("data") if isinstance(d, dict) else d
            batch = batch if isinstance(batch, list) else []
            if isinstance(d, dict):
                total = int(d.get("count") or total or 0)
            rows.extend(batch)
            if not batch or (total and len(rows) >= total) or len(batch) < ps:
                break
            page += 1
            if page > 500:
                raise FetchError(f"翻页超过 500，worksheet={worksheet_id}")
        return rows, total or len(rows)

    def search_rows(self, worksheet_id: str, keyword: str, page_size: int = 200) -> List[dict]:
        """全文检索（用于按 SO 找订单明细）。调用方必须再做精确过滤。"""
        rows: List[dict] = []
        page = 1
        while page <= 20:
            d = self.post(
                "worksheet/getFilterRows",
                {
                    "worksheetId": worksheet_id,
                    "appId": APP_ID,
                    "pageSize": page_size,
                    "pageIndex": page,
                    "status": 1,
                    "sortControls": [],
                    "notGetTotal": False,
                    "searchType": 1,
                    "keyWords": keyword,
                    "filterControls": [],
                    "fastFilters": [],
                    "navGroupFilters": [],
                },
            )
            batch = d.get("data") if isinstance(d, dict) else d
            batch = batch if isinstance(batch, list) else []
            rows.extend(batch)
            if len(batch) < page_size:
                break
            page += 1
        return rows

    def relation_rows(
        self, worksheet_id: str, row_id: str, control_id: str, page_size: int = 100
    ) -> Tuple[List[dict], List[dict]]:
        """返回 (行列表, 目标表 controls)。"""
        rows: List[dict] = []
        controls: List[dict] = []
        page = 1
        while page <= 50:
            d = self.post(
                "worksheet/getRowRelationRows",
                {
                    "worksheetId": worksheet_id,
                    "rowId": row_id,
                    "controlId": control_id,
                    "pageIndex": page,
                    "pageSize": page_size,
                    "appId": APP_ID,
                    "getWorksheet": page == 1,
                },
            )
            batch: Any = []
            if isinstance(d, dict):
                batch = d.get("data") or d.get("rows") or []
                if isinstance(batch, dict):
                    batch = batch.get("data") or []
                if page == 1:
                    w = d.get("worksheet") or {}
                    controls = (w.get("template") or {}).get("controls") or w.get("controls") or []
            batch = batch if isinstance(batch, list) else []
            rows.extend(batch)
            if len(batch) < page_size:
                break
            page += 1
        return rows, controls


def pick_named(
    row: dict, names: Dict[str, str], opts: Dict[str, Dict[str, str]], wanted: Sequence[str]
) -> Dict[str, str]:
    """按中文列名从原始行取值 → {列名: 纯文本}。取不到的列留空，不猜。"""
    by_name = {v: k for k, v in names.items() if v}
    out: Dict[str, str] = {}
    for w in wanted:
        cid = by_name.get(w, "")
        out[w] = _plain(row.get(cid), opts.get(cid)) if cid else ""
    return out


def extract_related_orders(
    rows: Sequence[dict], controls: Sequence[dict], source: str
) -> List[dict]:
    """从“下单”或“结算”关联行提取订单；SO 可位于 SO 或订单NUM。"""
    names = ZhiyunClient.name_map(controls)
    opts = ZhiyunClient.option_maps(controls)
    out: List[dict] = []
    for row in rows:
        values = pick_named(row, names, opts, XIADAN_COLS)
        so = extract_so(
            values.get("SO") or values.get("订单NUM")
            or values.get("订单号") or values.get("新智云单号") or ""
        )
        if not so:
            candidates = {
                extract_so(_plain(row.get(cid), opts.get(cid)))
                for cid in names
            } - {""}
            if len(candidates) == 1:
                so = next(iter(candidates))
        if not so:
            continue
        out.append({
            "so": so,
            "written_off": values.get("订单已核销金额"),
            "written_off_local": values.get("订单已核销金额/本币"),
            "deliver": values.get("交付额/原币"),
            "rate": values.get("汇率"),
            "currency": values.get("结算币种"),
            "name": values.get("订单名称"),
            "delivery_date": values.get("项目交付日期"),
            "delivery_date_status": "关联下单明确值" if values.get("项目交付日期") else "",
            "source": source,
        })
    return out


def lookup_order_delivery_date(
    client: ZhiyunClient,
    worksheet_id: str,
    controls: Sequence[dict],
    so: str,
) -> Tuple[str, str]:
    """从智云“下单”订单详情按 SO 精确回读项目交付日期；缺失或冲突都不猜。"""
    if not worksheet_id:
        return "", "项目交付日期缺失：无法定位下单数据源"
    names = client.name_map(controls)
    opts = client.option_maps(controls)
    try:
        hits = client.search_rows(worksheet_id, so)
    except Exception as e:
        return "", f"项目交付日期回读失败：{type(e).__name__}"
    dates = set()
    for row in hits:
        values = pick_named(row, names, opts, XIADAN_COLS)
        found_so = extract_so(
            values.get("SO") or values.get("订单NUM")
            or values.get("订单号") or values.get("新智云单号") or ""
        )
        if found_so != so:
            continue
        raw = str(values.get("项目交付日期") or "").strip()
        if raw:
            dates.add(raw)
    if len(dates) == 1:
        return next(iter(dates)), "订单详情明确值"
    if len(dates) > 1:
        return "", "项目交付日期冲突：订单详情存在多个不同日期"
    return "", "项目交付日期缺失：订单详情没有明确值"


def write_xlsx(path: Path, headers: List[str], rows: List[List[Any]]) -> None:
    from openpyxl import Workbook

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Sheet1"
    ws.append(headers)
    for r in rows:
        ws.append(r)
    wb.save(str(path))


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def publish_day_exports(
    out_dir: Path,
    day_tag: str,
    datasets: Sequence[Tuple[str, List[str], List[List[Any]]]],
    summary: dict,
) -> dict:
    """Build a complete daily bundle off to the side and publish its marker last."""
    out_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=f".fetch-{day_tag}-", dir=out_dir) as raw_staging:
        staging = Path(raw_staging)
        published_names: List[str] = []
        for name, headers, rows in datasets:
            staged_path = staging / name
            write_xlsx(staged_path, headers, rows)
            published_names.append(name)
        completed_summary = {
            **summary,
            "files": published_names,
            "file_sha256": {
                name: _sha256(staging / name)
                for name in published_names
            },
        }
        summary_name = f"取数摘要_{day_tag}.json"
        staged_summary = staging / summary_name
        staged_summary.write_text(
            json.dumps(completed_summary, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        for name in published_names:
            os.replace(staging / name, out_dir / name)
        # 摘要是完成标记，必须最后发布；中断后的新旧混合文件无法通过哈希校验。
        os.replace(staged_summary, out_dir / summary_name)
    return completed_summary


def _row_contains_identifier(row: dict, identifier: str, extractor) -> bool:
    return any(extractor(_plain(value)) == identifier for value in row.values())


def search_supplement_identifiers(
    client: ZhiyunClient,
    ar_ids: Sequence[str],
    so_ids: Sequence[str],
) -> dict[str, dict[str, List[str]]]:
    """按工作人员提供的编号精确搜索智云，只返回编号存在性。"""
    found_ar_ids: List[str] = []
    for ar_id in sorted(set(ar_ids)):
        try:
            hits = client.search_rows(WS_HUIKUAN, ar_id)
        except Exception:
            hits = []
        if any(_plain(row.get(F_HK["ar"])) == ar_id for row in hits):
            found_ar_ids.append(ar_id)

    worksheet_ids = []
    for relation in (REL_XIADAN, REL_HEXIAO_MINGXI, REL_SODLINE):
        try:
            worksheet_id = client.datasource_of(WS_HUIKUAN, relation)
        except Exception:
            worksheet_id = ""
        if worksheet_id and worksheet_id not in worksheet_ids:
            worksheet_ids.append(worksheet_id)
    found_so_ids: List[str] = []
    for so_id in sorted(set(so_ids)):
        found = False
        for worksheet_id in worksheet_ids:
            try:
                hits = client.search_rows(worksheet_id, so_id)
            except Exception:
                continue
            if any(_row_contains_identifier(row, so_id, extract_so) for row in hits):
                found = True
                break
        if found:
            found_so_ids.append(so_id)
    return {"found_ar_ids": found_ar_ids, "found_so_ids": found_so_ids}


def exported_supplement_identifiers(out_dir: Path, day: str) -> dict[str, set[str]]:
    """读取四张导出表中的 AR/SO 编号，用于判断补取前后变化。"""
    from openpyxl import load_workbook

    tag = day.replace("-", "")
    ar_ids: set[str] = set()
    so_ids: set[str] = set()

    def collect(prefix: str, ar_columns: Sequence[int], so_columns: Sequence[int]) -> None:
        matches = sorted(out_dir.glob(f"{prefix}_{tag}*.xlsx"), key=lambda item: item.name)
        if not matches:
            return
        workbook = load_workbook(matches[-1], read_only=True, data_only=True)
        try:
            worksheet = workbook.active
            for row in worksheet.iter_rows(min_row=2, values_only=True):
                for column in ar_columns:
                    if column < len(row):
                        value = extract_ar(_plain(row[column]))
                        if value:
                            ar_ids.add(value)
                for column in so_columns:
                    if column < len(row):
                        value = extract_so(_plain(row[column]))
                        if value:
                            so_ids.add(value)
        finally:
            workbook.close()

    collect("回款记录", (0,), ())
    collect("订单交付", (0,), (1,))
    collect("核销明细", (2,), (8,))
    collect("订单明细", (), (0,))
    return {"ar_ids": ar_ids, "so_ids": so_ids}


def build_supplement_result(
    ar_ids: Sequence[str],
    so_ids: Sequence[str],
    *,
    before: dict[str, set[str]],
    after: dict[str, set[str]],
    searched: dict[str, List[str]],
) -> dict[str, List[str]]:
    requested_ar = set(ar_ids)
    requested_so = set(so_ids)
    before_ar = set(before.get("ar_ids", set()))
    before_so = set(before.get("so_ids", set()))
    after_ar = set(after.get("ar_ids", set()))
    after_so = set(after.get("so_ids", set()))
    return {
        "requested": {"ar_ids": sorted(requested_ar), "so_ids": sorted(requested_so)},
        "found": {
            "ar_ids": sorted(set(searched.get("found_ar_ids", []))),
            "so_ids": sorted(set(searched.get("found_so_ids", []))),
        },
        "added": {
            "ar_ids": sorted(requested_ar & (after_ar - before_ar)),
            "so_ids": sorted(requested_so & (after_so - before_so)),
        },
        "existing": {
            "ar_ids": sorted(requested_ar & before_ar & after_ar),
            "so_ids": sorted(requested_so & before_so & after_so),
        },
        "unresolved": {
            "ar_ids": sorted(requested_ar - after_ar),
            "so_ids": sorted(requested_so - after_so),
        },
    }


def historical_writeoffs_for_sos(
    client: ZhiyunClient,
    worksheet_id: str,
    sos: Sequence[str],
    target_day: str,
) -> List[List[Any]]:
    """
    从全局“订单同币种核销明细信息”补取本批 SO 在目标日前的历史核销。

    关联父回款只会给出当前父 AR 自己的子明细；同一 SO 的首款、尾款若分别落在
    不同 AR，单靠当前父记录会漏掉历史回款。这里按 SO 精确检索全局明细，且只补
    `核销日期 < target_day` 的原始行（含撤销行供审计）；目标日当前行仍由父回款关联子表提供。
    """
    if not worksheet_id or not sos:
        return []
    ctrls = client.controls(worksheet_id)
    names, opts = client.name_map(ctrls), client.option_maps(ctrls)
    out: List[List[Any]] = []
    seen_record_ids = set()
    for wanted_so in sorted({str(x or "").strip() for x in sos if str(x or "").strip()}):
        try:
            hits = client.search_rows(worksheet_id, wanted_so)
        except Exception as exc:
            raise FetchError(
                f"全局核销明细检索失败 SO={wanted_so}: {type(exc).__name__}"
            ) from exc
        for row in hits:
            v = pick_named(row, names, opts, MINGXI_COLS)
            so = extract_so(v.get("订单NUM") or "")
            if so != wanted_so:
                continue
            hx_day = (v.get("核销日期") or "").strip()[:10]
            if not hx_day or hx_day >= target_day:
                continue
            ar = extract_ar(v.get("回款记录NUM") or "")
            if not ar:
                raise FetchError(
                    f"全局核销明细读不到父回款号：SO={wanted_so} 核销日期={hx_day}"
                )
            record_id = (v.get("核销记录NUM") or "").strip()
            if record_id:
                if record_id in seen_record_ids:
                    continue
                seen_record_ids.add(record_id)
            out.append([
                record_id,
                row.get("rowid") or "",
                ar,
                hx_day,
                v.get("本次核销金额"),
                v.get("本次核销金额本币"),
                v.get("币种"),
                v.get("汇率"),
                so,
                v.get("订单名称"),
                v.get("是否已撤销"),
            ])
    return out


def fetch_day(
    client: ZhiyunClient,
    day: str,
    out_dir: Path,
    supplement_ar_ids: Sequence[str] = (),
    supplement_so_ids: Sequence[str] = (),
) -> dict:
    """拉一天的四张表 + 摘要 json。返回计数摘要（无客户名/金额明细）。"""
    out_dir.mkdir(parents=True, exist_ok=True)

    hk_ctrls = client.controls(WS_HUIKUAN)
    hk_opts = client.option_maps(hk_ctrls)
    sales_cid = client.id_by_name(hk_ctrls, "销售")
    cid_xiadan = client.id_by_name(hk_ctrls, REL_XIADAN)
    cid_jiesuan = client.id_by_name(hk_ctrls, REL_JIESUAN)
    cid_mingxi = client.id_by_name(hk_ctrls, REL_HEXIAO_MINGXI)
    ws_xiadan = client.datasource_of(WS_HUIKUAN, REL_XIADAN)
    ws_mingxi = client.datasource_of(WS_HUIKUAN, REL_HEXIAO_MINGXI)
    ws_sodline = client.datasource_of(WS_HUIKUAN, REL_SODLINE)
    if not cid_xiadan:
        raise FetchError(f"回款记录里找不到「{REL_XIADAN}」关联字段——智云表结构变了，停下别猜")

    hk_rows, hk_total = client.filter_rows_by_date(WS_HUIKUAN, F_HK["hexiao_date"], day)
    known_payment_rows = {str(row.get("rowid") or "") for row in hk_rows}
    for ar_id in supplement_ar_ids:
        exact = [
            row
            for row in client.search_rows(WS_HUIKUAN, ar_id)
            if _plain(row.get(F_HK["ar"])) == ar_id
            and _plain(row.get(F_HK["hexiao_date"]))[:10] == day
        ]
        for row in exact:
            row_id = str(row.get("rowid") or "")
            if row_id and row_id not in known_payment_rows:
                hk_rows.append(row)
                known_payment_rows.add(row_id)

    # ── ① 回款记录 ────────────────────────────────────────────
    hk_headers = [
        "回款记录ID", "核销日期", "到账日期", "到账金额/原币", "到账金额/本币",
        "手续费/原币", "原币币种", "回款类型", "核销状态", "开票客户", "销售名称", "rowid",
        "仅历史累计父记录",
    ]
    hk_out: List[List[Any]] = []
    type_counts: Dict[str, int] = {}
    payments: List[dict] = []
    for row in hk_rows:
        ar = _plain(row.get(F_HK["ar"]))
        htype = _plain(row.get(F_HK["huikuan_type"]), hk_opts.get(F_HK["huikuan_type"]))
        type_counts[htype or "(空)"] = type_counts.get(htype or "(空)", 0) + 1
        rec = {
            "ar": ar,
            "hexiao_date": _plain(row.get(F_HK["hexiao_date"])),
            "arrival_date": _plain(row.get(F_HK["arrival_date"])),
            "amount_orig": _plain(row.get(F_HK["amount_orig"])),
            "amount_local": _plain(row.get(F_HK["amount_local"])),
            "fee": _plain(row.get(F_HK["fee"])),
            "currency": _plain(row.get(F_HK["currency"]), hk_opts.get(F_HK["currency"])),
            "huikuan_type": htype,
            "status": _plain(row.get(F_HK["status"]), hk_opts.get(F_HK["status"])),
            "customer": _plain(row.get(F_HK["customer_txt"])) or _plain(row.get(F_HK["customer_rel"])),
            "sales_name": _plain(row.get(sales_cid), hk_opts.get(sales_cid)) if sales_cid else "",
            "rowid": row.get("rowid") or "",
        }
        payments.append(rec)
        hk_out.append([rec[k] for k in (
            "ar", "hexiao_date", "arrival_date", "amount_orig", "amount_local",
            "fee", "currency", "huikuan_type", "status", "customer", "sales_name", "rowid")] + ["否"])

    day_tag = day.replace("-", "")
    # ── ② 下单栏（每笔 → SO + 交付额）+ ③ 同币种核销明细 ───────
    xd_out: List[List[Any]] = []
    mx_out: List[List[Any]] = []
    mx_seen_record_ids = set()
    all_so: List[str] = list(dict.fromkeys(supplement_so_ids))
    ars_without_orders: List[str] = []
    settlement_recovered_ars: List[str] = []
    settlement_rows_used = 0
    delivery_date_cache: Dict[str, Tuple[str, str]] = {}
    xiadan_ctrls = client.controls(ws_xiadan) if ws_xiadan else []

    for rec in payments:
        rid = str(rec.get("rowid") or "")
        if not rid:
            ars_without_orders.append(rec["ar"])
            continue

        xd_rows, xd_ctrls = client.relation_rows(WS_HUIKUAN, rid, cid_xiadan)
        related = extract_related_orders(xd_rows, xd_ctrls, REL_XIADAN)
        if not related and cid_jiesuan:
            js_rows, js_ctrls = client.relation_rows(WS_HUIKUAN, rid, cid_jiesuan)
            related = extract_related_orders(js_rows, js_ctrls, REL_JIESUAN)
            if related:
                settlement_recovered_ars.append(rec["ar"])
                settlement_rows_used += len(related)
        for v in related:
            so = v["so"]
            if so not in all_so:
                all_so.append(so)
            if not str(v.get("delivery_date") or "").strip():
                if so not in delivery_date_cache:
                    delivery_date_cache[so] = lookup_order_delivery_date(
                        client, ws_xiadan, xiadan_ctrls, so
                    )
                v["delivery_date"], v["delivery_date_status"] = delivery_date_cache[so]
            xd_out.append([
                rec["ar"], so, v.get("written_off"), v.get("written_off_local"),
                v.get("deliver"), v.get("rate"),
                v.get("currency"), v.get("name"), v.get("delivery_date"),
                v.get("delivery_date_status"), v.get("source"),
            ])
        if not related:
            ars_without_orders.append(rec["ar"])

        if cid_mingxi:
            mx_rows, mx_ctrls = client.relation_rows(WS_HUIKUAN, rid, cid_mingxi)
            mx_names, mx_opts = client.name_map(mx_ctrls), client.option_maps(mx_ctrls)
            for r in mx_rows:
                v = pick_named(r, mx_names, mx_opts, MINGXI_COLS)
                so = extract_so(v.get("订单NUM") or "")
                if not so:
                    print(
                        "WARN: 核销明细有一行读不出 SO；判定会把对应父回款挂起，不会猜测",
                        file=sys.stderr,
                    )
                if so and so not in all_so:
                    all_so.append(so)
                out_row = [
                    v.get("核销记录NUM"),
                    r.get("rowid") or "",
                    rec["ar"], v.get("核销日期") or rec["hexiao_date"],
                    v.get("本次核销金额"), v.get("本次核销金额本币"),
                    v.get("币种") or rec["currency"], v.get("汇率"), so,
                    v.get("订单名称"), v.get("是否已撤销"),
                ]
                record_id = str(v.get("核销记录NUM") or "").strip()
                if record_id:
                    if record_id in mx_seen_record_ids:
                        continue
                    mx_seen_record_ids.add(record_id)
                mx_out.append(out_row)

    if all_so and not ws_mingxi:
        raise FetchError(
            "回款记录里找不到「订单同币种核销明细信息」的全局数据源，"
            "无法按 SO 补取跨父回款历史核销，停下别猜累计回款"
        )
    historical_rows = historical_writeoffs_for_sos(
        client, ws_mingxi, all_so, day
    )
    # 跨父AR历史核销也必须能按各自父到账额做超核销审计；把仅用于累计的
    # 历史父记录一并保存，但分类器不会把它当目标日任务。
    current_ars = {str(rec.get("ar") or "").strip() for rec in payments}
    historical_ars = sorted({
        str(row[2] or "").strip() for row in historical_rows
        if str(row[2] or "").strip() and str(row[2] or "").strip() not in current_ars
    })
    for historical_ar in historical_ars:
        hits = client.search_rows(WS_HUIKUAN, historical_ar)
        exact = [
            row for row in hits
            if _plain(row.get(F_HK["ar"])) == historical_ar
        ]
        if len(exact) != 1:
            raise FetchError(
                f"历史核销父回款 {historical_ar} 无法唯一回读到账金额，"
                "不能安全计算系统重复核销差额"
            )
        row = exact[0]
        hk_out.append([
            historical_ar,
            _plain(row.get(F_HK["hexiao_date"])),
            _plain(row.get(F_HK["arrival_date"])),
            _plain(row.get(F_HK["amount_orig"])),
            _plain(row.get(F_HK["amount_local"])),
            _plain(row.get(F_HK["fee"])),
            _plain(row.get(F_HK["currency"]), hk_opts.get(F_HK["currency"])),
            _plain(row.get(F_HK["huikuan_type"]), hk_opts.get(F_HK["huikuan_type"])),
            _plain(row.get(F_HK["status"]), hk_opts.get(F_HK["status"])),
            _plain(row.get(F_HK["customer_txt"])) or _plain(row.get(F_HK["customer_rel"])),
            _plain(row.get(sales_cid), hk_opts.get(sales_cid)) if sales_cid else "",
            row.get("rowid") or "",
            "是",
        ])
    historical_added = 0
    for out_row in historical_rows:
        record_id = str(out_row[0] or "").strip()
        if record_id:
            if record_id in mx_seen_record_ids:
                continue
            mx_seen_record_ids.add(record_id)
        mx_out.append(out_row)
        historical_added += 1

    # ── ④ 订单明细（SO → SOD + 逐 SOD 交付额）─────────────────
    sod_out: List[List[Any]] = []
    so_without_sod: List[str] = []
    if ws_sodline and all_so:
        sl_ctrls = client.controls(ws_sodline)
        sl_names, sl_opts = client.name_map(sl_ctrls), client.option_maps(sl_ctrls)
        for so in all_so:
            # 请求失败必须终止该日取数；只有成功响应且确实为空，才能按无 SOD 处理。
            hits = client.search_rows(ws_sodline, so)
            n = 0
            for r in hits:
                v = pick_named(r, sl_names, sl_opts, SODLINE_COLS)
                # 全文检索会带出订单名里含该串的别的单 → 必须精确过滤
                if extract_so(v.get("SO") or "") != so:
                    continue
                sod = (v.get("SOD") or "").strip()
                if not sod:
                    continue
                n += 1
                sod_out.append([so, sod, v.get("交付额/原币"), v.get("币种"), v.get("项目状态")])
            if n == 0:
                so_without_sod.append(so)
    elif all_so:
        print(f"WARN: 找不到「{REL_SODLINE}」表，SOD 取不到", file=sys.stderr)
        so_without_sod = list(all_so)

    summary = {
        "day": day,
        "export_schema_version": EXPORT_SCHEMA_VERSION,
        "business_amount_policy": "ignore_fee_conditional_duplicate_writeoff_correction",
        "架构": "单入口(回款记录·核销日期=T-1)+关联子表",
        "回款记录笔数": len(hk_rows),
        "回款记录_接口报总数": hk_total,
        "下单行数": len(xd_out),
        "结算回查行数": settlement_rows_used,
        "从结算找回单号的AR数": len(set(settlement_recovered_ars)),
        "涉及SO数": len(all_so),
        "核销明细行数": len(mx_out),
        "跨父回款历史核销补取行数": historical_added,
        "订单明细SOD行数": len(sod_out),
        "缺项目交付日期的SO": sorted({
            str(row[1] or "").strip() for row in xd_out
            if str(row[1] or "").strip() and not str(row[8] or "").strip()
        }),
        "回款类型分布": type_counts,
        "无下单行的AR": ars_without_orders,
        "查不到SOD的SO": so_without_sod,
        "read_only": True,
    }
    return publish_day_exports(
        out_dir,
        day_tag,
        [
            (f"回款记录_{day_tag}.xlsx", hk_headers, hk_out),
            (
                f"订单交付_{day_tag}.xlsx",
                ["回款记录ID", "SO", "订单已核销金额", "订单已核销金额/本币",
                 "交付额/原币", "汇率", "结算币种", "订单名称", "项目交付日期",
                 "交付日期取数状态", "单号来源"],
                xd_out,
            ),
            (
                f"核销明细_{day_tag}.xlsx",
                ["核销记录NUM", "rowid", "回款记录NUM", "核销日期", "本次核销金额",
                 "本次核销金额/本币", "币种", "汇率", "SO", "订单名称", "是否已撤销"],
                mx_out,
            ),
            (
                f"订单明细_{day_tag}.xlsx",
                ["SO", "SOD", "交付额/原币", "币种", "项目状态"],
                sod_out,
            ),
        ],
        summary,
    )


def already_fetched(
    out_dir: Path,
    day: str,
    *,
    accept_unversioned: bool = False,
) -> List[str]:
    """
    这天的智云四件套是否齐全且带当前取数版本标记。

    默认不接受无版本摘要或旧版本摘要，避免代码更新后继续复用旧核销明细。
    明确确认是本次新手导文件时，调用方可设置 accept_unversioned=True。
    """
    stamp = day.replace("-", "")
    need = ("回款记录", "订单交付", "核销明细", "订单明细")
    if not out_dir.is_dir():
        return []
    got = []
    for key in need:
        for p in out_dir.glob("*.xlsx"):
            if p.name.startswith("~$"):
                continue
            if key in p.name and (stamp in p.name or day in p.name):
                got.append(p.name)
                break
    if len(got) != 4:
        return got

    summary_path = out_dir / f"取数摘要_{stamp}.json"
    if not summary_path.is_file():
        return got if accept_unversioned else []
    try:
        summary = json.loads(summary_path.read_text(encoding="utf-8"))
    except (OSError, ValueError, TypeError):
        return got if accept_unversioned else []
    if summary.get("export_schema_version") != EXPORT_SCHEMA_VERSION:
        return got if accept_unversioned else []
    file_hashes = summary.get("file_sha256")
    if not isinstance(file_hashes, dict):
        return got if accept_unversioned else []
    for name in got:
        expected = file_hashes.get(name)
        if not isinstance(expected, str) or _sha256(out_dir / name) != expected:
            return got if accept_unversioned else []
    return got


def _saved_credentials() -> Tuple[str, str]:
    """从 Windows 凭据库读取本机保存的智云测试账号。"""
    try:
        import keyring

        user = (keyring.get_password(CREDENTIAL_SERVICE, "__default_username__") or "").strip()
        pwd = (keyring.get_password(CREDENTIAL_SERVICE, user) or "").strip() if user else ""
        return user, pwd
    except Exception:
        return "", ""


def resolve_credentials(args) -> Tuple[str, str]:
    user = (args.user or os.environ.get("ZHIYUN_USER") or "").strip()
    pwd = (args.password or os.environ.get("ZHIYUN_PASS") or "").strip()
    if args.cookie_only:
        return "", ""
    saved_user, saved_pwd = _saved_credentials()
    if not user:
        user = saved_user
    if not pwd and user == saved_user:
        pwd = saved_pwd
    # 核销指令后的主流程必须无人值守；不在中途弹出账号/密码问题，
    # 也不把凭据写进命令行或文件。认证只能来自 MD_PSS_ID、环境变量或
    # 本机 Windows 凭据库；缺失时明确停止，由编排层一次性报告阻塞原因。
    if not user or not pwd:
        raise SystemExit(
            "ERROR: 没有可用的智云登录凭据。自动任务不会在取数中途询问账号密码；"
            "请先配置本机 Windows 凭据库/环境变量，或提供 MD_PSS_ID 后重新执行。"
        )
    return user, pwd


def report_fetched_day(
    day: str,
    out_dir: Path,
    summary: dict,
    supplement_result: Optional[dict] = None,
) -> None:
    """报告并登记一个已完整导出的核销日；多日取数不得在空批处提前结束。"""
    print("✅ 智云只读取数完成（未写系统）")
    print(f"📁 核销日期: {day}   目录: {out_dir.resolve()}")
    if supplement_result is not None:
        result_path = out_dir / f"补取结果_{day.replace('-', '')}.json"
        result_path.write_text(
            json.dumps(supplement_result, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        print(
            "补取复核完成：新增 "
            f"{len(supplement_result['added']['ar_ids']) + len(supplement_result['added']['so_ids'])} 个，"
            "仍未找到 "
            f"{len(supplement_result['unresolved']['ar_ids']) + len(supplement_result['unresolved']['so_ids'])} 个。"
        )

    if not summary["回款记录笔数"]:
        print(
            f"ℹ️ {day} 这天**一笔核销都没有**（不是出错）。常见于周末、假期、"
            "或销售当天没来得及核。这天就算处理完了，已记进跑批台账。"
        )
        try:
            sys.path.insert(0, str(Path(__file__).resolve().parent))
            import batch_ledger
            import common as _c

            batch_ledger.record(
                out_dir.parent,
                _c.norm_date(day),
                "classified",
                payments=0,
                note="空批：那天没有任何核销",
            )
        except Exception:
            pass
        return

    try:
        sys.path.insert(0, str(Path(__file__).resolve().parent))
        import batch_ledger
        import common as _c

        batch_ledger.record(
            out_dir.parent,
            _c.norm_date(day),
            "fetched",
            payments=summary["回款记录笔数"],
        )
    except Exception:
        pass
    print(
        f"   回款记录 {summary['回款记录笔数']} 笔 · 订单关联 {summary['下单行数']} 行"
        f"（结算回查 {summary.get('结算回查行数', 0)} 行）"
        f"（{summary['涉及SO数']} 个 SO）· 核销明细 {summary['核销明细行数']} 行"
        f" · 订单明细 {summary['订单明细SOD行数']} 个 SOD"
    )
    if summary.get("跨父回款历史核销补取行数"):
        print(f"   历史累计：跨父回款补取 {summary['跨父回款历史核销补取行数']} 行")
    print(f"   回款类型分布: {summary['回款类型分布']}")
    if summary["无下单行的AR"]:
        print(
            f"   ⚠ 有 {len(summary['无下单行的AR'])} 笔到账在下单和结算中都没抓到单号，"
            "判定会报异常不会漏"
        )
    if summary["查不到SOD的SO"]:
        print(
            f"   ⚠ 有 {len(summary['查不到SOD的SO'])} 个 SO 查不到 SOD，将退化成按 SO 匹配"
        )


def fetch_days_concurrently(
    *,
    base_url: str,
    cookie: str,
    account_id: str,
    days: Sequence[str],
    out_dir: Path,
) -> List[Tuple[str, dict]]:
    """同一登录凭据下并发取不同日期；每个日期使用独立 HTTP 会话。"""
    ordered_days = list(days)
    if not ordered_days:
        return []

    out_dir.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory(prefix=".batch-fetch-", dir=out_dir) as raw_staging:
        staging_dir = Path(raw_staging)

        def fetch_one(day: str) -> dict:
            client = ZhiyunClient(base_url, cookie, account_id=account_id)
            try:
                return fetch_day(client, day, staging_dir)
            finally:
                close = getattr(client, "close", None)
                if callable(close):
                    close()

        summaries: Dict[str, dict] = {}
        failures: Dict[str, Exception] = {}
        worker_count = min(MAX_BATCH_FETCH_CONCURRENCY, len(ordered_days))
        with ThreadPoolExecutor(
            max_workers=worker_count,
            thread_name_prefix="zhiyun-date",
        ) as executor:
            futures = {executor.submit(fetch_one, day): day for day in ordered_days}
            for future in as_completed(futures):
                day = futures[future]
                try:
                    summaries[day] = future.result()
                except Exception as exc:  # noqa: BLE001 - 聚合所有日期的只读取数错误
                    failures[day] = exc

        if failures:
            failed_day = next(day for day in ordered_days if day in failures)
            failure = failures[failed_day]
            raise FetchError(
                f"核销日期 {failed_day} 取数失败（{type(failure).__name__}）：{failure}"
            ) from failure
        publish_batch_exports(staging_dir, out_dir, ordered_days, summaries)
        return [(day, summaries[day]) for day in ordered_days]


def publish_batch_exports(
    staging_dir: Path,
    out_dir: Path,
    ordered_days: Sequence[str],
    summaries: Dict[str, dict],
) -> None:
    """全部日期取数成功后统一发布；发布异常时恢复原有完整文件。"""
    publish_entries: List[Tuple[Path, Path]] = []
    seen_names: set[str] = set()
    for day in ordered_days:
        summary = summaries.get(day) or {}
        names = summary.get("files")
        if not isinstance(names, list) or len(names) != 4:
            raise FetchError(f"核销日期 {day} 的取数包不完整，禁止发布。")
        day_names = [*names, f"取数摘要_{day.replace('-', '')}.json"]
        for name in day_names:
            if not isinstance(name, str) or Path(name).name != name or name in seen_names:
                raise FetchError(f"核销日期 {day} 的取数文件名无效，禁止发布。")
            source = staging_dir / name
            if not source.is_file():
                raise FetchError(f"核销日期 {day} 缺少取数文件，禁止发布。")
            seen_names.add(name)
            publish_entries.append((source, out_dir / name))

    backup_dir = staging_dir / ".previous"
    backup_dir.mkdir(exist_ok=False)
    promoted: List[Path] = []
    backups: List[Tuple[Path, Path]] = []
    try:
        for source, target in publish_entries:
            if target.exists():
                backup = backup_dir / target.name
                os.replace(target, backup)
                backups.append((backup, target))
            os.replace(source, target)
            promoted.append(target)
    except Exception as exc:  # noqa: BLE001 - 必须回滚任意文件系统发布错误
        for target in reversed(promoted):
            if target.exists():
                target.unlink()
        for backup, target in reversed(backups):
            if backup.exists():
                os.replace(backup, target)
        raise FetchError("批次取数文件发布失败，原有文件已恢复。") from exc


def main(argv: Optional[List[str]] = None) -> int:
    ap = argparse.ArgumentParser(
        description="智云只读取数（单入口：回款记录按核销日期 + 关联子表）"
    )
    ap.add_argument(
        "--date", "--hexiao-date", dest="date", default="",
        help="单个核销日期 YYYY-MM-DD / yesterday / last-workday；不传时默认 yesterday",
    )
    ap.add_argument("--date-from", default="", help="批量取数开始核销日期（与 --date-to 同用）")
    ap.add_argument("--date-to", default="", help="批量取数结束核销日期（与 --date-from 同用）")
    date_scope = ap.add_mutually_exclusive_group()
    date_scope.add_argument(
        "--all-days",
        dest="include_weekends",
        action="store_true",
        default=True,
        help="批量取数时包含周末（当前默认行为，保留参数兼容旧调用）",
    )
    date_scope.add_argument(
        "--workdays-only",
        dest="include_weekends",
        action="store_false",
        help="批量取数时只取周一至周五",
    )
    ap.add_argument(
        "--skip-gap-check", action="store_true",
        help="不查漏天（默认会查：有从没跑过的核销日就先报出来）",
    )
    ap.add_argument(
        "--force", action="store_true",
        help="这天的四件套已在 01_智云导出/ 里也强制重新取一遍",
    )
    ap.add_argument("--supplement-ar", action="append", default=[], help="按完整 AR 编号补取并复核")
    ap.add_argument("--supplement-so", action="append", default=[], help="按完整 SO 编号补取并复核")
    ap.add_argument(
        "--accept-unversioned-existing",
        action="store_true",
        help="明确接受没有当前取数版本标记的手工四件套；默认禁止复用旧取数文件",
    )
    ap.add_argument("--workspace", default="", help="技能工作区根（含 01_智云导出）")
    ap.add_argument("--out", default="", help="直接指定导出目录（优先于 workspace）")
    ap.add_argument("--base-url", default=os.environ.get("ZHIYUN_BASE", BASE_DEFAULT))
    ap.add_argument("--user", default="", help="账号；也可用环境变量 ZHIYUN_USER")
    ap.add_argument(
        "--password", default="",
        help="密码（不推荐写在命令行历史）；优先用 ZHIYUN_PASS 或本机凭据库，任务中不交互询问",
    )
    ap.add_argument("--cookie-only", action="store_true", help="不登录，只用 MD_PSS_ID")
    ap.add_argument("--account-id", default=os.environ.get("ZHIYUN_ACCOUNT_ID", ""))
    ap.add_argument("--headed", action="store_true", help="有头浏览器登录（调试）")
    args = ap.parse_args(argv)

    ar_ids = list(dict.fromkeys(str(value).strip().upper() for value in args.supplement_ar if str(value).strip()))
    so_ids = list(dict.fromkeys(str(value).strip().upper() for value in args.supplement_so if str(value).strip()))
    if any(not re.fullmatch(r"AR[A-Z0-9_-]{3,30}", value) for value in ar_ids):
        ap.error("补取 AR 编号格式无效")
    if any(not re.fullmatch(r"SO[A-Z0-9_-]{3,30}", value) for value in so_ids):
        ap.error("补取 SO 编号格式无效")
    supplement_mode = bool(ar_ids or so_ids)

    try:
        days = resolve_fetch_dates(
            single_date=args.date,
            date_from=args.date_from,
            date_to=args.date_to,
            include_weekends=args.include_weekends,
        )
    except ValueError as exc:
        ap.error(str(exc))
    if supplement_mode and len(days) != 1:
        ap.error("按 AR/SO 补取只支持单个核销日，不能与日期范围同时使用")
    if args.out:
        out_dir = Path(args.out)
    elif args.workspace:
        out_dir = Path(args.workspace) / "01_智云导出"
    else:
        out_dir = Path(__file__).resolve().parent.parent / "工作区" / "01_智云导出"

    # ① 在入口固定全部核销日；取数可以共用一次登录，判定和写表仍逐日串行。
    if len(days) == 1:
        print(f"★ 本次取的是**核销日期 = {days[0]}** 的到账（不是到账日期）")
    else:
        print(
            f"★ 本次一次登录取 {len(days)} 个核销日：{days[0]} → {days[-1]}；"
            "每天仍生成独立四件套"
        )

    # ② 漏天检查：`--date yesterday` 只看昨天，她请假/周末/系统故障跳过的那几天
    #    没有任何机制发现。漏一天 = 那天的到账永远不会回填，而且事后看不出来。
    workspace = out_dir.parent
    if not args.skip_gap_check:
        try:
            sys.path.insert(0, str(Path(__file__).resolve().parent))
            import batch_ledger
            import common as _c

            selected_days = set(days)
            info = batch_ledger.find_gaps(workspace, through=_c.norm_date(days[-1]))
            gaps = [g for g in info["gaps"] if g.isoformat() not in selected_days]
            if gaps:
                print("⚠ 这几个核销日**从来没跑过**（会漏掉那天的到账）：")
                for g in gaps[:10]:
                    print(f"     · {_c.date_cn(g)}")
                if len(gaps) > 10:
                    print(f"     …另有 {len(gaps) - 10} 天")
                print("   → 把漏日纳入本次日期范围；取数可一次完成，后续仍从早到晚逐日处理。")
        except Exception as e:
            print(f"WARN: 漏天检查跳过（{type(e).__name__}）", file=sys.stderr)

    # ③ 每天独立检查四件套；只要有一天需要重取，就复用同一次登录继续取完。
    pending_days: List[str] = []
    for day in days:
        existing_any_version = already_fetched(
            out_dir,
            day,
            accept_unversioned=True,
        )
        have = already_fetched(
            out_dir,
            day,
            accept_unversioned=args.accept_unversioned_existing,
        )
        if (
            len(existing_any_version) == 4
            and len(have) != 4
            and not args.accept_unversioned_existing
        ):
            print(
                f"⚠ {day} 已有四件套，但没有当前取数版本 "
                f"{EXPORT_SCHEMA_VERSION}；本次禁止复用，立即重新抓取。"
            )
        if len(have) == 4 and not args.force and not supplement_mode:
            print(f"✅ {day} 的当前版本四件套已存在，本次跳过重复取数")
            continue
        if have:
            print(
                f"注意：{out_dir} 里已有 {len(have)}/4 份 {day} 文件，"
                "本次会重新取完整四件套。"
            )
        pending_days.append(day)

    if not pending_days:
        print("✅ 本次所有核销日都已有当前版本四件套，无需登录智云")
        return 0

    cookie = (os.environ.get("MD_PSS_ID") or "").strip()
    account_id = (args.account_id or "").strip()

    if not cookie:
        user, pwd = resolve_credentials(args)
        print(f"正在登录智云 {args.base_url} …（密码不打印）")
        try:
            cookie, acct = login_with_password(args.base_url, user, pwd, headless=not args.headed)
        except LoginError as e:
            print(f"ERROR: 登录失败 — {e}", file=sys.stderr)
            return 2
        if acct and not account_id:
            account_id = acct
        del pwd
        print("登录成功，开始只读取数…")

    try:
        if supplement_mode:
            day = pending_days[0]
            client = ZhiyunClient(args.base_url, cookie, account_id=account_id)
            try:
                before_identifiers = exported_supplement_identifiers(out_dir, day)
                searched = search_supplement_identifiers(client, ar_ids, so_ids)
                try:
                    summary = fetch_day(client, day, out_dir, ar_ids, so_ids)
                except Exception as exc:
                    print(
                        f"ERROR: 核销日期 {day} 取数失败（{type(exc).__name__}）：{exc}",
                        file=sys.stderr,
                    )
                    return 2
                after_identifiers = exported_supplement_identifiers(out_dir, day)
                supplement_result = build_supplement_result(
                    ar_ids,
                    so_ids,
                    before=before_identifiers,
                    after=after_identifiers,
                    searched=searched,
                )
                report_fetched_day(day, out_dir, summary, supplement_result)
            finally:
                close = getattr(client, "close", None)
                if callable(close):
                    close()
        else:
            try:
                fetched_days = fetch_days_concurrently(
                    base_url=args.base_url,
                    cookie=cookie,
                    account_id=account_id,
                    days=pending_days,
                    out_dir=out_dir,
                )
            except FetchError as exc:
                print(f"ERROR: {exc}", file=sys.stderr)
                return 2
            # 日志和跑批台账按核销日期登记，后续业务流程也继续按此顺序串行。
            for day, summary in fetched_days:
                report_fetched_day(day, out_dir, summary)
    finally:
        cookie = ""
        del cookie

    print(
        f"✅ 本次 {len(pending_days)} 个核销日已在同一次登录中取完；"
        "后续判定和写表仍须从早到晚逐日执行"
    )
    print(
        "👉 下一步：按上述核销日期从早到晚逐日执行判定、日清和工作副本写入"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
