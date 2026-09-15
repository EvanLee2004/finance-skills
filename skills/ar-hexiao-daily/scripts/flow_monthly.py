"""Monthly receipt allocation on controlled workbook copies.

The custom XML part travels with the workbook material version. It records only
receipt identities and applied allocations; visible rows are verified before reuse.
"""
from __future__ import annotations

import copy
from contextlib import ExitStack
import datetime as dt
import hashlib
import json
import re
import shutil
import tempfile
import zipfile
from decimal import Decimal, InvalidOperation
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.cell.rich_text import CellRichText
import common
import xlsx_patch
import workbook_finalize

SCHEMA = 'receipt-monthly-v2'
PART = 'customXml/financialReceiptMonthly.xml'
NS = 'urn:besteasy:receipt-monthly:v1'
FIELDS = ('日期', '公司名称', '金额', '收款形式', '单号', '预收', '是否更新应收款')
SO = re.compile(r'(SO[A-Z0-9]+)', re.I)
NUMBER = r'\d+(?:\.\d{1,2})?'


def money(value):
    if isinstance(value, str) and re.fullmatch(r'=\d+(?:\.\d+)?\*\d+(?:\.\d+)?', value):
        left, right = value[1:].split('*')
        value = (Decimal(left) * Decimal(right)).quantize(Decimal('.01'))
    try:
        result = Decimal(str(value).replace(',', ''))
    except (InvalidOperation, ValueError):
        raise ValueError('核销金额或预收余额缺失、不是有效数字') from None
    if isinstance(value, float) and result.is_finite():
        rounded = result.quantize(Decimal('.01'))
        if abs(result - rounded) <= Decimal('0.00000001'):
            result = rounded
    if not result.is_finite() or result != result.quantize(Decimal('.01')):
        raise ValueError('金额必须是有限的两位小数')
    return result


def number(value):
    return format(money(value), 'f').rstrip('0').rstrip('.') if '.' in format(money(value), 'f') else format(money(value), 'f')


def formula(start, entries):
    return number(start) + ''.join('-' + number(x['amount']) for x in entries) + '=' + number(balance(start, entries))


def balance(start, entries):
    return money(start) - sum((money(x['amount']) for x in entries), Decimal(0))


def actual_amount(item):
    """Never substitute delivery, cumulative received or accrual for this payment."""
    audit = item.get('write_currency_audit') or {}
    source = item.get('split_payment_source') or {}
    candidates = [source.get('amount_local'), audit.get('amount_local')]
    if common.is_cny(audit.get('currency') or ''):
        candidates.append(audit.get('amount_orig'))
    for value in candidates:
        if value is not None and value != '':
            return money(value)
    raise ValueError('缺少本次核销本币金额，不能用交付金额代替')


def delivery_amount(item):
    """SO delivery is presentation data and never a receipt allocation."""
    source = item.get('split_payment_source') or {}
    candidates = [source.get('so_delivery_local'), item.get('so_delivery_local')]
    if len(source.get('all_sods') or []) <= 1:
        candidates += [source.get('delivery_local'), item.get('delivery_local')]
    for value in candidates:
        if value is not None and value != '':
            result = money(value)
            if result < 0: raise ValueError('SO交付金额无效')
            return result
    raise ValueError('缺少SO本币交付金额，不能用本次核销额代替展示')


def published_receipt_matches(row):
    audit = row.get('baseline_receipt_audit') or {}
    if audit.get('disposition') != 'skip':
        return False
    try:
        identity = json.loads(audit.get('event_key') or 'null')
        event = (row.get('split_payment_source') or {}).get('writeoff_sequence_key') or []
        five = row.get('five_cols') or {}
        return (len(event) >= 3 and identity == [row.get('ar'),row.get('so'),row.get('sod'),event[1]]
                and money(audit.get('current_received')) == actual_amount(row) == money(five.get('回款明细'))
                and five.get('是否结账') == '是')
    except (ValueError, TypeError):
        return False


def checked_receipt_proof(row, ledger_rows, day, parents):
    """Attach receipt-specific evidence without changing the ledger disposition."""
    if (row.get('_check') or {}).get('verdict') != 'skip':
        return {}
    source=row.get('split_payment_source') or {}
    ar,so,sod=row.get('ar'),row.get('so'),row.get('sod')
    parent=parents.get(ar) or {}
    prior=(parent.get('applied_cases') or {}).get(row.get('case_id'))
    try:
        amount=actual_amount(row)
        if amount <= 0:return {}
        arrival=common.norm_date((row.get('flow_identity') or {}).get('date'))
        posting=common.norm_date(day)
        expected_day=common.receipt_time(arrival,posting)
        expected_way=common.pay_way(row.get('status') or '',arrival,posting)
        exact=[r for r in ledger_rows.values() if r.get('SO')==so and r.get('SOD')==sod
               and r.get('是否结账')=='是' and r.get('回款明细') not in (None,'')
               and money(r['回款明细'])==amount and common.norm_date(r.get('收款时间'))==expected_day
               and r.get('收款方式')==expected_way]
        if (prior and parent.get('reused_successful_allocation') and parent.get('applied_at')
                and common.norm_date(parent.get('hexiao_date')) == common.norm_date(day)
                and prior.get('so') == so and prior.get('sod') == sod
                and money(prior.get('amount_local')) == amount):
            scope=source.get('receivable_group_scope') or {}
            represented=set(scope.get('source_sods') or [sod])
            candidates=[r for r in ledger_rows.values() if r.get('SO')==so and r.get('SOD') in represented]
            received=sum((money(r['回款明细']) for r in candidates if r.get('回款明细') not in (None,'')),Decimal(0))
            if len(exact)==1 or (candidates and received == money(source.get('cumulative_local'))):
                return {'basis':'published_parent_case','ar':ar,'so':so,'sod':sod,
                        'date':common.norm_date(day).isoformat(),'amount':number(amount)}
    except (ValueError, TypeError):
        return {}
    return {}


def valid_receipt_proof(row, date):
    proof=row.get('flow_receipt_proof') or {}
    try:
        return (proof.get('basis') == 'published_parent_case'
                and all(proof.get(k)==row.get(k) for k in ('ar','so','sod'))
                and proof.get('date')==date.isoformat() and money(proof.get('amount'))==actual_amount(row))
    except ValueError:
        return False


def complete_parent_history(ar, checked, rows, entries):
    """Only a fully attributed published parent can recover a blank legacy balance."""
    parent=(checked.get('parent_fallback_allocations') or {}).get(ar) or {}
    if not parent.get('reused_successful_allocation') or not parent.get('applied_at'):
        return None
    cases=parent.get('applied_cases') or {}
    by_case={r.get('case_id'):r for _,r in rows if (r.get('flow_receipt_proof') or {}).get('basis')=='published_parent_case'}
    if not cases or set(cases)!=set(by_case):return None
    try:
        if common.norm_date(parent.get('hexiao_date'))!=common.norm_date(checked.get('hexiao_date')):return None
        opening=money(parent.get('parent_net_amount'))
        allocations=parent.get('allocations') or []
        if len({a.get('so') for a in allocations})!=len(allocations):return None
        if {c.get('so') for c in cases.values()} - {a.get('so') for a in allocations}:return None
        for allocation in allocations:
            expected=money(allocation.get('allocated_local'))
            paid=sum((money(c['amount_local']) for c in cases.values() if c['so']==allocation.get('so')),Decimal(0))
            if expected<0 or paid!=expected:return None
        if money(parent.get('unallocated_parent_amount'))!=0:return None
        if sum((money(e['amount']) for e in entries),Decimal(0))!=opening:return None
        return {'date':parent['hexiao_date'],'opening':number(opening),'entries':copy.deepcopy(entries),
                'known_sos':sorted({str(a['so']).upper() for a in allocations})}
    except (ValueError,KeyError,TypeError):
        return None


def complete_current_history(item, checked, rows, entries, require_existing):
    """Full receipt coverage leaves no earlier allocation amount to guess."""
    if require_existing or not entries:return None
    ar=item.get('ar')
    if any(r.get('ar')==ar for r in checked.get('conflict',[])):return None
    try:
        current=[r for kind in ('write','skip') for r in checked.get(kind,[]) if r.get('ar')==ar]
        positive={r['case_id'] for r in current if actual_amount(r)>0}
        if positive != {e['case_id'] for e in entries}:return None
        all_cases={r['case_id'] for r in current}
        if any(any(b not in ('auto','ready') for b in outcome.get('buckets',[]))
               or not set(outcome.get('case_ids',[])).issubset(all_cases)
               for outcome in item.get('so_outcomes',[])):return None
        opening=money((item.get('identity') or {}).get('amount'))
        if sum((money(e['amount']) for e in entries),Decimal(0))!=opening:return None
        return {'date':checked['hexiao_date'],'opening':number(opening),'entries':copy.deepcopy(entries),
                'known_sos':sorted({str(so).upper() for so in item.get('so_list',[])})}
    except (ValueError,KeyError,TypeError):
        return None


def finalize(items, checked):
    """Bind only verified ledger cases to the fixed business date."""
    date = common.norm_date(checked.get('hexiao_date'))
    if not date:
        raise ValueError('缺少有效核销日期')
    bad = {x.get('case_id') for x in checked.get('conflict', [])}
    by_ar = {}
    for kind in ('write', 'skip'):
        for row in checked.get(kind, []):
            if row.get('case_id') in bad:
                continue
            # An already-settled SO alone is not proof that this receipt was applied.
            if row.get('code') == 'OK_SO_ALREADY_SETTLED' and not valid_receipt_proof(row,date):
                continue
            if kind == 'skip':
                reason = (row.get('_check') or {}).get('reason', '')
                proved = reason in {
                    '已经填过且与本次一致（幂等跳过）',
                    '部分回款拆行已完整写入（已收行+紧邻未收行），幂等跳过',
                    '分笔回款链全部父回款行已完整写入，幂等跳过',
                    '结清尾差已合并写入，幂等跳过',
                    '同 SO 多 SOD 已按 SO 交付金额合并写入，幂等跳过',
                    '该 SOD 已并入同一 SO 的合并核销行，不重复写金额',
                    '本父回款属于1元以内结清尾差，审计保留并入目标行，不单独写入',
                }
                proved = proved or published_receipt_matches(row) or valid_receipt_proof(row,date)
                if not proved:
                    continue
            by_ar.setdefault(row.get('ar'), []).append((kind,row))
    for item in items:
        if item.get('monthly_schema') != SCHEMA or item.get('verdict') != 'write':
            continue
        item['monthly_date'] = date.isoformat()
        entries = {}
        deliveries = {}
        try:
            require_existing=[]
            for kind,row in by_ar.get(item.get('ar'), []):
                if row.get('bucket') not in ('auto', 'ready'):
                    continue
                currency = (row.get('write_currency_audit') or {}).get('currency') or ''
                if currency and not common.is_cny(currency) and '原币公式' not in item.get('matched_by', ''):
                    raise ValueError('外币流转缺少本币金额依据，不能混用原币余额与本币核销额')
                amount = actual_amount(row)
                if amount == 0:
                    continue
                if amount < 0:
                    raise ValueError('负数核销需人工核对结转链')
                case = str(row.get('case_id') or '')
                so = str(row.get('so') or '').strip().upper()
                if not case or not so:
                    raise ValueError('缺少核销事项身份')
                deliveries[so] = number(amount)
                event = (row.get('split_payment_source') or {}).get('writeoff_sequence_key')
                parent_proof = valid_receipt_proof(row,date) and row['flow_receipt_proof']['basis']=='published_parent_case'
                if event and len(event)>=3 and (event[1] or event[2]):
                    if common.norm_date(event[0]) != date:
                        raise ValueError('核销事件日期与本次日期不一致')
                    key = 'event|' + json.dumps([case,event[1],event[2]],ensure_ascii=False,separators=(',',':'))
                else:
                    key = date.isoformat() + '|' + case
                    if kind=='skip' and not parent_proof:require_existing.append(key)
                entry = {'key':key, 'case_id':case, 'date':date.isoformat(), 'so':so, 'amount':number(amount)}
                if key in entries and entries[key] != entry:
                    raise ValueError('同一核销事项金额冲突')
                entries[key] = entry
            item['monthly_entries'] = list(entries.values())
            item['monthly_deliveries'] = deliveries
            item['monthly_require_existing'] = require_existing
            item['monthly_receipt_history'] = (complete_parent_history(item.get('ar'),checked,by_ar.get(item.get('ar'),[]),list(entries.values()))
                or complete_current_history(item,checked,by_ar.get(item.get('ar'),[]),list(entries.values()),require_existing))
            if not entries:
                item.update(verdict='skip', reason='没有经验证的本次核销金额可登记')
        except ValueError as exc:
            item.update(verdict='hand', reason=str(exc))
            item['monthly_entries'] = []
    return items


def columns(ws):
    rows = list(ws.iter_rows(values_only=True))
    aliases = common.load_aliases()
    header, names = common.find_header_row(rows, '到账流转', ['日期','公司名称','金额','单号'], aliases)
    found = common.resolve_columns(names, '到账流转', ['日期','公司名称','金额','单号'], aliases)
    for key in ('收款形式','预收','是否更新应收款'):
        options = {'预收':['预收','预收款','预收金额','预收余额']}.get(key, aliases.get('到账流转',{}).get(key,[key]))
        col = common.fuzzy_find_col(names, options)
        if col is None:
            raise ValueError('流转表缺少' + key + '列，需手填')
        found[key] = col
    if len(set(found.values())) != len(found):
        raise ValueError('流转列映射不唯一')
    return {key:value+1 for key,value in found.items()}


def value(ws, row, cols, field):
    return ws.cell(row,cols[field]).value


def signature(ws, row, cols):
    vals = []
    for field in FIELDS:
        v = value(ws,row,cols,field)
        if field == '日期':
            day = common.norm_date(v); v = day.isoformat() if day else str(v or '')
        elif field == '金额':
            v = number(v)
        else:
            v = str(v or '').strip()
        vals.append(v)
    return vals


def load_state(path):
    with zipfile.ZipFile(path) as z:
        if PART not in z.namelist():
            return {'schema':SCHEMA, 'receipts':{}}
        data = json.loads(ET.fromstring(z.read(PART)).text or '')
        if data.get('schema') not in (SCHEMA, 'receipt-monthly-v1') or not isinstance(data.get('receipts'),dict):
            raise ValueError('流转结转记录格式不兼容')
        data['schema'] = SCHEMA
        for chain in data['receipts'].values():
            for month in chain['months']:
                month.setdefault('display_original',month['signature'][4])
                month.setdefault('display_amounts',{})
                month.setdefault('legacy_sos',[e['so'] for e in month['entries'] if e['key'].startswith('legacy:')])
                month['legacy_sos']=[str(so).strip().upper() for so in month['legacy_sos']]
        return data


def save_state(path, state):
    with zipfile.ZipFile(path) as z:
        infos = z.infolist(); payload = {i.filename:z.read(i.filename) for i in infos}
    body = ET.Element('{' + NS + '}receiptLedger')
    body.text = json.dumps(state,ensure_ascii=False,sort_keys=True,separators=(',',':'))
    payload[PART] = ET.tostring(body,encoding='utf-8',xml_declaration=True)
    rels = payload['_rels/.rels'].decode()
    if 'Target="'+PART+'"' not in rels:
        if 'Id="financialReceiptMonthly"' in rels:
            raise ValueError('工作簿关系标识已被占用')
        rels = rels.replace('</Relationships>','<Relationship Id="financialReceiptMonthly" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml" Target="'+PART+'"/></Relationships>')
        payload['_rels/.rels'] = rels.encode()
    types = payload['[Content_Types].xml'].decode()
    if 'PartName="/'+PART+'"' not in types:
        types = types.replace('</Types>','<Override PartName="/'+PART+'" ContentType="application/xml"/></Types>')
        payload['[Content_Types].xml'] = types.encode()
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
        for info in infos:
            z.writestr(info,payload.pop(info.filename))
        for name,data in payload.items():
            z.writestr(name,data)


def resolve_months(ws, cols, chain):
    used=set()
    for month in chain['months']:
        expected=month['signature']; row=int(month['row'])
        if row > ws.max_row or signature(ws,row,cols)!=expected:
            matches=[r for r in range(1,ws.max_row+1)
                     if str(value(ws,r,cols,'公司名称') or '').strip()==expected[1]
                     and common.norm_date(value(ws,r,cols,'日期'))
                     and signature(ws,r,cols)==expected]
            if len(matches)!=1:
                raise ValueError('已有结转行被修改或不能唯一定位，需手填')
            row=matches[0]
        if row in used:
            raise ValueError('多个月份引用同一流转行')
        used.add(row);month['row']=row
        entries=month['entries']
        if balance(month['start'],entries)!=money(month['remaining']):
            raise ValueError('历史结转金额不守恒')
    months=chain['months']
    if [m['month'] for m in months]!=sorted(set(m['month'] for m in months)):
        raise ValueError('月份结转顺序不唯一')
    for previous,current in zip(months,months[1:]):
        if money(previous['remaining'])!=money(current['start']):
            raise ValueError('跨月承接金额不一致')


BALANCE_NUMBER = r'(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d{1,2})?'

def parsed_balance(raw):
    """Return an explicit numeric chain, a balance, or a nonnumeric remark."""
    from openpyxl.cell.rich_text import CellRichText
    if isinstance(raw,CellRichText):raw=str(raw)
    if raw is None or str(raw).strip()=='':return {'blank':True}
    if not isinstance(raw,str):return {'remaining':money(raw)}
    text=str(raw).strip().replace('－','-').replace('＝','=')
    text=re.sub(r'^\s*(?:还剩|剩余|余额|预付)\s*[：:]?\s*','',text)
    text=re.sub(r'[，,;；\s]*转(?:\d+月)?\s*$','',text).strip()
    if re.fullmatch(BALANCE_NUMBER,text):return {'remaining':money(text)}
    excel=text.startswith('=');text=text[1:] if excel else text
    if not re.fullmatch(BALANCE_NUMBER+r'(?:-'+BALANCE_NUMBER+r')*(?:='+BALANCE_NUMBER+r'(?:-'+BALANCE_NUMBER+r')*)*',text):
        if re.fullmatch(r'核销在\s*\d+\s*月',text):return {'remark':str(raw)}
        raise ValueError('预收内容无法解析为数字扣减算式，需核实原始内容')
    segments=text.split('=');first=[money(n) for n in segments[0].split('-')]
    deductions=first[1:];remaining=first[0]-sum(deductions,Decimal(0))
    for segment in segments[1:]:
        nums=[money(n) for n in segment.split('-')]
        if nums[0]!=remaining:raise ValueError('已有预收算式的中间结果或最终结果不一致')
        deductions.extend(nums[1:]);remaining-=sum(nums[1:],Decimal(0))
    if not excel and len(segments)==1 and deductions:
        raise ValueError('预收数字减法缺少等式结果')
    return {'opening':first[0],'deductions':deductions,'remaining':remaining}


def legacy_month(ws, row, cols, history=None):
    date=common.norm_date(value(ws,row,cols,'日期'))
    start=money(value(ws,row,cols,'金额'))
    if date is None or start < 0:
        raise ValueError('原流转日期或金额无效')
    text=str(value(ws,row,cols,'单号') or '').strip()
    raw=value(ws,row,cols,'预收')
    if isinstance(raw,CellRichText):raw=str(raw)
    deductions=[]
    recovered=False
    relocate=False
    if raw is None or raw=='':
        if SO.search(text):
            if (history and money(history['opening'])==start
                    and common.norm_date(history['date'])>=date
                    and {so.upper() for so in SO.findall(text)}.issubset(history['known_sos'])):
                recovered=common.norm_date(history['date']).strftime('%Y-%m')==date.strftime('%Y-%m')
                relocate=not recovered
                deductions=[money(e['amount']) for e in history['entries']] if recovered else []
            else:
                raise ValueError('预收为空且已有SO，缺少实际历史核销记录；单号交付额不能推算余额')
        remaining=start-sum(deductions,Decimal(0))
    else:
        parsed = parsed_balance(raw)
        if 'remaining' not in parsed:
            raise ValueError('预收仅有备注，缺少可核实的余额')
        if 'opening' in parsed and parsed['opening'] != start:
            raise ValueError('已有预收算式起始金额不一致')
        remaining = parsed['remaining']
        deductions = parsed.get('deductions', [start - remaining])
    if remaining<0 or remaining>start or any(x<0 for x in deductions):
        raise ValueError('预收余额超出该行可用金额')
    # The order cell contains delivery amounts only. Prior deductions come solely
    # from the balance/formula; no SO allocation is invented for these totals.
    entries=[{'key':'legacy:'+str(row)+':'+str(i),'date':date.isoformat(),'so':'',
              'amount':number(amount),'basis':'existing_balance'} for i,amount in enumerate(deductions) if amount]
    if recovered:entries=copy.deepcopy(history['entries'])
    prefix_lines=[line for line in text.splitlines() if not SO.search(line) and not re.search(r'转\d+月',line)]
    prefix='\n'.join(prefix_lines).strip()
    if not prefix:
        match=re.match(r'([A-Za-z]+?)SO',text)
        if match:prefix=match.group(1)
    return {'month':date.strftime('%Y-%m'),'date':date.isoformat(),'row':row,'start':number(start),
            'remaining':number(remaining),'entries':entries,'prefix':prefix,'signature':signature(ws,row,cols),
            'legacy_sos':[] if recovered or relocate else sorted({so.upper() for so in SO.findall(text)}),
            'display_original':prefix if relocate else text, 'display_amounts':{},
            '_repair_balance':recovered, '_relocate_display':relocate}


def adopt_chain(ws, cols, item):
    row=int(item['row_no']); first=legacy_month(ws,row,cols,item.get('monthly_receipt_history'))
    chain={'sheet':ws.title,'months':[first]}
    payer=first['signature'][1]
    # Existing employee splits must be adjacent and uniquely explain the carry.
    while row<ws.max_row:
        next_row=row+1
        if str(value(ws,next_row,cols,'收款形式') or '').strip()!='冲预收': break
        if str(value(ws,next_row,cols,'公司名称') or '').strip()!=payer: break
        following=legacy_month(ws,next_row,cols)
        previous=chain['months'][-1]
        if following['month']<=previous['month'] or money(following['start'])!=money(previous['remaining']):
            raise ValueError('人工拆行的月份或结转余额不一致')
        chain['months'].append(following);row=next_row
    known={m['row'] for m in chain['months']}
    for r in range(1,ws.max_row+1):
        if r not in known and str(value(ws,r,cols,'收款形式') or '').strip()=='冲预收' and str(value(ws,r,cols,'公司名称') or '').strip()==payer:
            if common.norm_date(value(ws,r,cols,'日期')) and money(value(ws,r,cols,'金额')) in {money(m['remaining']) for m in chain['months']}:
                raise ValueError('存在非相邻的同名承接候选，不能自动重复拆行')
    if re.search(r'转\d+月',str(value(ws,row,cols,'单号') or '')):
        raise ValueError('原行标记已转月，但缺少可核实承接行')
    return chain


def order_text(month):
    # Each posting date keeps its own SO payment amount; never display delivery.
    amounts = {}
    for entry in month['entries']:
        if entry.get('so'):
            key = (entry['date'], entry['so'])
            amounts[key] = amounts.get(key, Decimal(0)) + money(entry['amount'])
    known = {so for _, so in amounts}
    existing = month.get('display_original', month.get('prefix', ''))
    lines = [line for line in existing.splitlines()
             if not any(so.upper() in known for so in SO.findall(line))]
    prefix = month.get('prefix', '').strip()
    if prefix and prefix not in lines:
        lines.insert(0, prefix)
    # Legacy values without event attribution are preserved, never deducted.
    lines += [f"{so}  {amount:,.2f}" for (_, so), amount in amounts.items()]
    return '\n'.join(lines).strip()


def references_above_insertion(formula, cell_row, insertion_row):
    """Prove a stationary formula references only explicit cells above insertion."""
    from html import unescape
    from openpyxl.formula.tokenizer import Tokenizer
    if cell_row >= insertion_row:
        return False
    text = unescape(formula)
    if re.search(r'\b(?:INDIRECT|OFFSET)\s*\(', text, re.I):
        return False
    try:
        tokens = Tokenizer('=' + text.lstrip('=')).items
    except Exception:
        return False
    for token in tokens:
        if token.type != 'OPERAND' or token.subtype != 'RANGE':
            continue
        for cell in token.value.split(':'):
            match = re.fullmatch(r'\$?[A-Za-z]+\$?(\d+)', cell)
            if not match or int(match[1]) > insertion_row:
                return False
    return True


def insertion_guard(path, sheet, row):
    """Reject structures the existing lossless row patcher cannot safely relocate."""
    with zipfile.ZipFile(path) as z:
        target=xlsx_patch.sheet_path_for(z,sheet)
        xml=z.read(target).decode()
        merges = [m.group(1) for m in re.finditer(r'<mergeCell\b[^>]*ref="([^"]+)"', xml)]
        affected_merge = any(any(int(n)>row for n in re.findall(r'\d+', ref)) for ref in merges)
        if affected_merge or any(tag in xml for tag in ('<tableParts','<drawing','<legacyDrawing','<conditionalFormatting','<dataValidations')):
            raise ValueError('目标流转页含表对象、图形或合并单元格，拆行需人工处理')
        for cell in re.finditer(r'<c\b[^>]*r="[A-Z]+(\d+)"[^>]*>(.*?)</c>',xml,re.S):
            for f in re.finditer(r'<f\b[^>]*>(.*?)</f>',cell.group(2),re.S):
                grouped_formula = re.search(r'<f\b[^>]*\bt=["\'](?:shared|array)["\']', cell.group(2))
                if not grouped_formula and references_above_insertion(f.group(1), int(cell.group(1)), row):
                    continue
                refs=re.findall(r'([A-Za-z]+)(\$?)(\d+)',f.group(1))
                # Same-row relative references move correctly with the copied row.
                if any(absolute or int(n)!=int(cell.group(1)) for _,absolute,n in refs) or '!' in f.group(1):
                    raise ValueError('目标页公式跨行引用，拆行需核实引用范围')
        for name in z.namelist():
            if name.endswith('.xml') and name!=target and (name.startswith('xl/worksheets/') or name.startswith('xl/charts/')):
                text=z.read(name).decode()
                if sheet+'!' in text or sheet+"'!" in text:
                    raise ValueError('其他工作表引用目标流转页，不能直接拆行')
        for match in re.finditer(r'<definedName\b([^>]*)>(.*?)</definedName>',z.read('xl/workbook.xml').decode(),re.S):
            if '_xlnm._FilterDatabase' not in match.group(1):
                raise ValueError('工作簿有业务命名范围，拆行需核实范围引用')



def update_filter_names(path, sheet, sources):
    if not sources: return
    with zipfile.ZipFile(path) as z:
        infos=z.infolist(); payload={i.filename:z.read(i.filename) for i in infos}
    text=payload['xl/workbook.xml'].decode()
    def replace(match):
        content=match.group(2)
        if '_xlnm._FilterDatabase' not in match.group(1) or '!' not in content:return match.group(0)
        qualifier,refs=content.rsplit('!',1)
        if qualifier.strip("'").replace("''", "'")!=sheet:return match.group(0)
        for source in sorted(sources,reverse=True):
            refs=xlsx_patch._shift_a1_token(refs,source,include_inserted=True)
        return '<definedName'+match.group(1)+'>'+qualifier+'!'+refs+'</definedName>'
    text=re.sub(r'<definedName\b([^>]*)>(.*?)</definedName>',replace,text,flags=re.S)
    payload['xl/workbook.xml']=text.encode()
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
        for info in infos:z.writestr(info,payload[info.filename])


def write_file(src, out, items, *, validate_only=False, workbook=None):
    state=load_state(src)
    wb=workbook if workbook is not None else openpyxl.load_workbook(src,data_only=False,rich_text=True)
    edits_by={}; insert_by={}; changes=[]; targets=set(); cols_by={}; expected_cells=[]
    try:
        for chain in state['receipts'].values():
            resolve_months(wb[chain['sheet']], columns(wb[chain['sheet']]), chain)
        for item in items:
            sheet=item['sheet'];ws=wb[sheet];cols=cols_by.setdefault(sheet,columns(ws))
            ar=item['ar'];chain=state['receipts'].get(ar)
            if chain:
                if chain['sheet']!=sheet: raise ValueError('到账关联工作表发生变化')
                resolve_months(ws,cols,chain)
                root=chain['months'][0]['row']
                if root!=int(item['row_no']): raise ValueError('三键定位与已有到账关联不一致')
            else:
                chain=adopt_chain(ws,cols,item);state['receipts'][ar]=chain
            root=chain['months'][0]['row']
            if any(other!=ar and known['sheet']==sheet and known['months'][0]['row']==root for other,known in state['receipts'].items()):
                raise ValueError('原始行已关联另一笔到账')
            if (sheet,root) in targets: raise ValueError('多笔到账计划引用同一原始行')
            targets.add((sheet,root))
            day=common.norm_date(item['monthly_date']);month_key=day.strftime('%Y-%m')
            recorded={e['key']:e for m in chain['months'] for e in m['entries']}
            if any(key not in recorded for key in item.get('monthly_require_existing', [])):
                raise ValueError('盈亏已跳过但缺少流转事件记录，不能推定本笔已登记')
            new=[]
            for entry in item['monthly_entries']:
                if entry['key'] in recorded:
                    if recorded[entry['key']]!=entry: raise ValueError('已登记核销事项金额发生变化')
                else: new.append(entry)
            repair=any([m.pop('_repair_balance',False) for m in chain['months']])
            if not new and not repair: continue
            legacy_sos={str(so).strip().upper() for m in chain['months'] for so in m.get('legacy_sos',[])}
            if any(str(e['so']).strip().upper() in legacy_sos for e in new):
                raise ValueError('本次SO已有人工核销记录，缺少事项身份，不能重复扣减')
            last=chain['months'][-1]
            if day.isoformat()<max([m['date'] for m in chain['months']]+[e['date'] for e in recorded.values()]):
                raise ValueError('倒序补录会影响已有结转余额，需人工核对')
            cross=month_key!=last['month']
            if month_key<last['month']: raise ValueError('不能倒序修改历史结转月份')
            edits=edits_by.setdefault(sheet,[]);insertions=insert_by.setdefault(sheet,[])
            if cross:
                insertion_guard(src,sheet,last['row'])
                month={'month':month_key,'date':day.isoformat(),'row':last['row']+1,
                    'start':last['remaining'],'remaining':last['remaining'],'entries':[],
                    'prefix':chain['months'][0].get('prefix',''),
                    'display_original':chain['months'][0].get('prefix',''), 'display_amounts':{}, 'legacy_sos':[]}
            else: month=last
            month['entries'].extend(new)
            remaining=balance(month['start'],month['entries'])
            if remaining<0: raise ValueError('本次核销超过可用预收余额')
            month['remaining']=number(remaining)
            text=order_text(month)
            if not cross:
                from apply_flow import _line_colors, _rich_signature
                original=value(ws,month['row'],cols,'单号')
                colors=_line_colors(original)
                original_runs=_rich_signature(original)
                original_text=''.join(text for text,_ in original_runs)
                lines=text.splitlines();runs=[]
                for index,line in enumerate(lines):
                    ids=SO.findall(line);color=colors.get(ids[0].upper(),'') if ids else ''
                    if not ids and line in original_text:
                        start=original_text.index(line);end=start+len(line);offset=0
                        for fragment,fragment_color in original_runs:
                            left=max(start,offset);right=min(end,offset+len(fragment))
                            if left<right:runs.append(xlsx_patch.RichTextRun(fragment[left-offset:right-offset],fragment_color))
                            offset+=len(fragment)
                        if index<len(lines)-1:runs.append(xlsx_patch.RichTextRun('\n'))
                    else:runs.append(xlsx_patch.RichTextRun(line+('\n' if index<len(lines)-1 else ''),color))
                text=xlsx_patch.RichTextValue(tuple(runs))
            overrides={cols['单号']:text,cols['预收']:formula(month['start'],month['entries']),cols['是否更新应收款']:'是'}
            if cross:
                overrides.update({cols['日期']:day,cols['金额']:float(money(month['start'])),cols['收款形式']:'冲预收'})
                # Receipt registration is new for this month; do not copy old status.
                headers=list(ws.iter_rows(values_only=True))
                for key in ('是否已登记系统',):
                    aliases=common.load_aliases();_,names=common.find_header_row(headers,'到账流转',['日期','公司名称','金额','单号'],aliases)
                    idx=common.fuzzy_find_col(names,aliases.get('到账流转',{}).get(key,[key]))
                    if idx is not None:overrides[idx+1]=None
                marker=(chain['months'][0].get('prefix','')+'转'+str(day.month)+'月')
                old=(last.get('prefix','') if last.pop('_relocate_display',False) else str(value(ws,last['row'],cols,'单号') or '').strip())
                if not SO.search(old) and old==chain['months'][0].get('prefix',''):old=marker
                elif marker not in old:old=(old+'\n'+marker).strip()
                from apply_flow import _rich_signature
                original_value=value(ws,last['row'],cols,'单号')
                old_plain=str(original_value or '').strip()
                if old.startswith(old_plain) and old_plain:
                    original_runs = _rich_signature(original_value)
                    original_text = ''.join(text for text, _ in original_runs)
                    left = len(original_text) - len(original_text.lstrip())
                    right = len(original_text.rstrip())
                    runs = []
                    offset = 0
                    for fragment, color in original_runs:
                        begin, end = max(left, offset), min(right, offset + len(fragment))
                        if begin < end:
                            runs.append(xlsx_patch.RichTextRun(fragment[begin-offset:end-offset], color))
                        offset += len(fragment)
                    runs.append(xlsx_patch.RichTextRun(old[len(old_plain):]))
                    marker_value=xlsx_patch.RichTextValue(tuple(runs))
                else:marker_value=old
                edits.append((last['row'],cols['单号'],marker_value))
                last['signature'][4]=old
                insertions.append((last['row'],overrides));month['_insert_after']=last['row'];chain['months'].append(month)
            else:
                edits.extend((month['row'],col,v) for col,v in overrides.items())
            expected_cells.append((sheet, month, dict(overrides)))
            changes.append({'AR':ar,'单号':order_text(month),'日期':month['date'],'是否更新应收款':'是','文件':src.name,'sheet':sheet,'行号':month['row'],'操作':'恢复核销公式' if repair else ('跨月结转' if cross else '同月登记'),
                '预收公式':formula(month['start'],month['entries']),'预收余额':float(remaining), '核销事项':[e['key'] for e in new]})
        if validate_only:return changes
        if not changes:return []
        for ar,chain in state['receipts'].items():
            sources=[r for r,_ in insert_by.get(chain['sheet'],[])]
            for month in chain['months']:
                if '_insert_after' in month:
                    source=month.pop('_insert_after');month['row']=source+1+sum(r<source for r in sources)
                else:month['row']+=sum(r<month['row'] for r in sources)
        for change in changes:
            change['行号']=state['receipts'][change['AR']]['months'][-1]['row']
    finally:
        if workbook is None:wb.close()
    current=src
    with tempfile.TemporaryDirectory(dir=out.parent,prefix='.flow-monthly-') as tmp:
        for index,sheet in enumerate(edits_by):
            target=Path(tmp)/f'{index}.xlsx'
            patch=xlsx_patch.patch_cells(current,target,sheet,edits_by[sheet],insertions=insert_by[sheet],return_result=True)
            update_filter_names(target,sheet,[row for row,_ in insert_by[sheet]])
            workbook_finalize.finalize_workbook(target,{sheet:patch});current=target
        # Re-read formulas AND caches before committing; metadata refers to actual rows.
        check=openpyxl.load_workbook(current,data_only=False,rich_text=True)
        cached=openpyxl.load_workbook(current,data_only=True)
        try:
            for sheet, month, expected in expected_cells:
                row=month['row']
                for col,wanted in expected.items():
                    got=check[sheet].cell(row,col).value
                    if isinstance(wanted,xlsx_patch.FormulaValue):
                        valid=got==wanted.formula and money(cached[sheet].cell(row,col).value)==money(wanted.cached)
                    elif isinstance(wanted,xlsx_patch.RichTextValue):
                        from apply_flow import _rich_signature
                        valid=_rich_signature(got)==_rich_signature(wanted)
                    elif isinstance(wanted,dt.date):valid=common.norm_date(got)==wanted
                    else:valid=(str(got or '')==str(wanted or ''))
                    if not valid:raise ValueError('流转单号、日期、状态或金额回读不一致')
            for chain in state['receipts'].values():
                ws=check[chain['sheet']];cols=columns(ws)
                for month in chain['months']:
                    r=month['row']
                    sig=signature(ws,r,cols)
                    # Every managed month's amount/balance remains conserved.
                    expected=formula(month['start'],month['entries'])
                    touched = any(candidate is month for _,candidate,_ in expected_cells)
                    if not touched and month.get('signature') is not None and sig!=month['signature']:
                        raise ValueError('未授权修改的历史结转行发生变化')
                    if touched:
                        if sig[5]!=expected or value(cached[ws.title],r,cols,'预收')!=expected or balance(month['start'],month['entries'])!=money(month['remaining']):
                            raise ValueError('流转预收算式或余额回读不一致')
                    month['signature']=sig
        finally:check.close();cached.close()
        shutil.copyfile(current,out)
    save_state(out,state)
    if load_state(out)!=state: raise ValueError('结转记录回读不一致')
    return changes


def prepare(workspace, items):
    """Decide each receipt on the fixed baseline before any workbook mutation."""
    from apply_flow import _resolve_flow_path,precheck_flow_identity
    selected=[x for x in items if x.get('verdict')=='write']
    targets={}
    for item in selected:
        key=(item.get('file'),item.get('sheet'),item.get('row_no'))
        targets.setdefault(key,[]).append(item)
    for group in targets.values():
        if len(group)>1:
            for item in group:item.update(verdict='hand',reason='多笔到账计划引用同一原始行，需核实到账归属')
    with ExitStack() as stack:
        books={}
        for item in selected:
            if item.get('verdict')!='write':continue
            if item.get('monthly_schema')!=SCHEMA or 'monthly_entries' not in item:
                item.update(verdict='hand',reason='流转计划尚未绑定本版本的已验证核销事项');continue
            errors=precheck_flow_identity(workspace,[item])
            if errors:
                item.update(verdict='hand',reason='；'.join(errors));continue
            src=_resolve_flow_path(workspace,item.get('file') or '')
            if src is None or not src.resolve().is_relative_to((workspace/'02_我的表副本').resolve()):
                raise ValueError('流转文件超出工作副本')
            try:
                if src not in books:
                    books[src]=openpyxl.load_workbook(src,data_only=False,rich_text=True)
                    stack.callback(books[src].close)
                write_file(src,src,[item],validate_only=True,workbook=books[src])
            except ValueError as exc:
                item.update(verdict='hand',reason=str(exc))
    return items


def write(workspace, items, *, in_place, phase):
    if phase=='prefill':return [],[]
    if phase not in ('status','all'):raise ValueError('未知流转阶段')
    selected=[x for x in items if x.get('verdict')=='write']
    if any('monthly_entries' not in x for x in selected):
        return [],['月度流转必须先绑定已验证的盈亏核销事项']
    from apply_flow import _resolve_flow_path,precheck_flow_identity
    problems=precheck_flow_identity(workspace,selected)
    if problems:return [],problems
    grouped={}
    for item in selected:grouped.setdefault(item['file'],[]).append(item)
    changes=[]
    try:
        # Prepare all files first so one rejected receipt cannot leave partial edits.
        with tempfile.TemporaryDirectory(dir=workspace,prefix='.flow-monthly-') as tmp:
            pending=[]
            for index,(name,group) in enumerate(grouped.items()):
                src=_resolve_flow_path(workspace,name)
                if src is None or not src.resolve().is_relative_to((workspace/'02_我的表副本').resolve()):raise ValueError('流转文件超出工作副本')
                before=hashlib.sha256(src.read_bytes()).hexdigest()
                target=Path(tmp)/f'{index}.xlsx'
                made=write_file(src,target,group)
                if made:pending.append((src,target,before));changes.extend(made)
            for src,target,before in pending:
                if hashlib.sha256(src.read_bytes()).hexdigest()!=before:raise ValueError('写入期间流转副本发生变化')
            for src,target,before in pending:
                backup=workspace/'02_我的表副本'/'备份'/(src.stem+'_月度流转_'+before[:16]+src.suffix)
                backup.parent.mkdir(parents=True,exist_ok=True)
                if not backup.exists():shutil.copy2(src,backup)
                destination=src if in_place else workspace/'04_产出'/src.name
                destination.parent.mkdir(parents=True,exist_ok=True)
                shutil.copyfile(target,destination)
        return changes,[]
    except (ValueError,KeyError,TypeError,OSError,zipfile.BadZipFile) as exc:
        return [],[str(exc)]
