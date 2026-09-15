"""Shared formula edits exercise the production ZIP patcher without business data."""
import hashlib
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape

import openpyxl
import xlsx_patch as patch

NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


class SharedFormulaPatchTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.src, self.out = self.root / "source.xlsx", self.root / "result.xlsx"

    def fixture(self, formula="B2", *, invalid=""):
        wb = openpyxl.Workbook()
        wb.active.title = "明细"
        wb.create_sheet("Reference data")["A1"] = 17
        wb.save(self.src)
        wb.close()
        rows = []
        for row in range(1, 6):
            cells = [f'<c r="A{row}"><v>{row}</v></c>', f'<c r="B{row}"><v>{row * 10}</v></c>']
            if 2 <= row <= 4:
                for col in ("C", "D"):
                    attrs = ' t="shared" si="7"'
                    text = ""
                    if col == "C" and row == 2:
                        if invalid == "orphan":
                            continue
                        attrs += ' ref="C2:D3"' if invalid == "outside" else ' ref="C2:D4"'
                        text = escape(formula)
                    if invalid == "duplicate" and col == "D" and row == 2:
                        attrs += ' ref="C2:D4"'
                        text = "C2"
                    cells.append(f'<c r="{col}{row}" s="0"><f{attrs}>{text}</f><v>{row * 100 + ord(col)}</v></c>')
            # A second formula group must remain byte-for-byte unchanged.
            if row in (2, 3):
                master = ' ref="F2:F3"' if row == 2 else ''
                text = 'A2+$B$1' if row == 2 else ''
                cells.append(f'<c r="F{row}"><f t="shared" si="8"{master}>{text}</f><v>19</v></c>')
            rows.append(f'<row r="{row}">'+''.join(cells)+'</row>')
        xml = '<worksheet xmlns="'+NS['s']+'"><dimension ref="A1:F5"/><sheetData>'+''.join(rows)+'</sheetData></worksheet>'
        with zipfile.ZipFile(self.src) as z:
            payload = {n: z.read(n) for n in z.namelist()}
        payload['xl/worksheets/sheet1.xml'] = xml.encode()
        with zipfile.ZipFile(self.src, 'w', zipfile.ZIP_DEFLATED) as z:
            for name, data in payload.items():
                z.writestr(name, data)
        self.original = self.src.read_bytes()
        return xml

    def write(self, edits, insertions=None):
        result = patch.patch_cells(self.src, self.out, "明细", edits, insertions, return_result=True)
        self.assertEqual(self.src.read_bytes(), self.original)
        with zipfile.ZipFile(self.src) as before, zipfile.ZipFile(self.out) as after:
            for name in before.namelist():
                if name not in {'xl/worksheets/sheet1.xml', 'xl/workbook.xml', 'xl/_rels/workbook.xml.rels', '[Content_Types].xml', 'xl/calcChain.xml'}:
                    self.assertEqual(before.read(name), after.read(name), name)
            self.xml = after.read('xl/worksheets/sheet1.xml').decode()
        patch._validate_shared_formula_integrity(self.xml)
        return result

    def cell(self, ref):
        root = ET.fromstring(self.xml)
        return root.find(f'.//s:c[@r="{ref}"]', NS)

    def formula(self, ref):
        f = self.cell(ref).find('s:f', NS)
        return None if f is None else f.text

    def test_overwrite_master_preserves_followers_and_caches(self):
        original_xml = self.fixture()
        self.write([(2, 3, 42), (2, 4, 42), (4, 3, 84), (4, 4, 84)])
        self.assertEqual(self.formula('C3'), 'B3')
        self.assertEqual(self.formula('D3'), 'C3')
        self.assertEqual(self.cell('C3').find('s:v', NS).text, '367')
        self.assertEqual(self.cell('C3').attrib['s'], '0')
        self.assertNotIn('si="7"', self.xml)
        for ref in ['F2', 'F3']:
            before = ET.fromstring(original_xml).find(f'.//s:c[@r="{ref}"]', NS)
            self.assertEqual(ET.tostring(before), ET.tostring(self.cell(ref)))
        wb = openpyxl.load_workbook(self.out, data_only=False, read_only=True)
        self.assertEqual(wb['明细']['C3'].value, '=B3')
        self.assertEqual(wb['明细']['D3'].value, '=C3')
        wb.close()

    def test_clear_master(self):
        self.fixture()
        self.write([(2, 3, None)])
        self.assertIsNone(self.formula('C2'))
        self.assertEqual(self.formula('D2'), 'C2')
        self.assertEqual(self.formula('C4'), 'B4')

    def test_replace_master_with_new_formula(self):
        self.fixture()
        self.write([(2, 3, patch.FormulaValue('=A2*2', 4))])
        self.assertEqual(self.formula('C2'), 'A2*2')
        self.assertEqual(self.formula('C3'), 'B3')

    def test_edit_follower_preserves_other_formulas(self):
        self.fixture()
        self.write([(3, 4, 77)])
        self.assertEqual(self.formula('C2'), 'B2')
        self.assertEqual(self.formula('D2'), 'C2')
        self.assertEqual(self.formula('C4'), 'B4')
        self.assertIsNone(self.formula('D3'))

    def test_overwrite_entire_group(self):
        self.fixture()
        self.write([(r, c, 0) for r in range(2, 5) for c in [3, 4]])
        self.assertNotIn('si="7"', self.xml)

    def test_mixed_absolute_sheet_and_escaped_references(self):
        self.fixture('$A2+B$1+$B$1+\'Reference data\'!A2+IF(A2<1,1,0)')
        self.write([(2, 3, 0)])
        self.assertEqual(self.formula('D3'), '$A3+C$1+$B$1+\'Reference data\'!B3+IF(B3<1,1,0)')
        self.assertEqual(self.cell('D3').find('s:v', NS).text, '368')

    def test_insert_with_master_and_copy_override(self):
        self.fixture()
        self.write([(2, 3, 99)], [(2, {3: 88})])
        self.assertIsNone(self.formula('C2'))
        self.assertIsNone(self.formula('C3'))
        self.assertEqual(self.formula('D3'), 'C3')
        self.assertEqual(self.formula('C4'), 'B4')
        self.assertEqual(self.formula('D5'), 'C5')

    def test_insert_before_master_maps_edit_coordinates(self):
        self.fixture()
        self.write([(2, 3, 99)], [(1, {1: 7})])
        self.assertIsNone(self.formula('C3'))
        self.assertEqual(self.formula('D3'), 'C3')
        self.assertEqual(self.formula('C4'), 'B4')

    def test_unrelated_edit_keeps_shared_group_bytes(self):
        original_xml = self.fixture()
        self.write([(5, 1, 10)])
        for ref in ['C2', 'D2', 'C3', 'D3', 'C4', 'D4']:
            before = ET.fromstring(original_xml).find(f'.//s:c[@r="{ref}"]', NS)
            self.assertEqual(ET.tostring(before), ET.tostring(self.cell(ref)))

    def test_existing_invalid_groups_still_rejected(self):
        for invalid in ['orphan', 'duplicate', 'outside']:
            with self.subTest(invalid=invalid):
                self.fixture(invalid=invalid)
                with self.assertRaisesRegex(ValueError, '共享公式结构非法'):
                    patch.patch_cells(self.src, self.out, '明细', [(r, c, 0) for r in range(2, 5) for c in [3, 4]])
                self.assertFalse(self.out.exists())
                self.assertEqual(self.src.read_bytes(), self.original)


if __name__ == '__main__':
    unittest.main()
