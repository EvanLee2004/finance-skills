import copy
import unittest
from types import SimpleNamespace
import receipt_correction as R
import baseline_receipts as BR
import validate_plan as V


class CorrectionTest(unittest.TestCase):
    def make(self, amount=50, day='2026-08-01', settled='是'):
        raw={'so':'SO1','sod':'SOD1','yingshou':100,'jiti':100,'huikuan':amount,
             'jiezhang':settled,'shoukuan_time':day,'shoukuan_way':'汇','chayi':0}
        from classify_hexiao import LedgerIndex
        ledger=LedgerIndex(synthetic={'so':{'SO1':[2]},'sod':{'SOD1':[2]},'rows':{2:raw}})
        rec={'ar':'AR1','so':'SO1','sod':'SOD1','amount_local':100,'deliver_local':100,'cumulative_received_local':100,
             'shoukuan_date':'2026-09-09','hexiao_date':'2026-09-09','writeoff_sequence_key':['2026-09-09','HX1','HXD1'], 'status':'正常','amount_orig':100,'currency':'CNY','deliver_orig':100,'all_sods':['SOD1']}
        result={'ar':'AR1','so':'SO1','sod':'SOD1','case_id':'AR1|SO1|SOD1','split_payment_source':{'amount_local':100,'writeoff_sequence_key':rec['writeoff_sequence_key']}}
        item=R.candidate(rec,result,ledger)
        return rec,ledger,item,{2:BR.ledger_rows(ledger,'SO1','SOD1')['2']}
    def test_filled_conflict_becomes_explicit_write_and_repeat_skips(self):
        _,_,item,rows=self.make()
        self.assertEqual(V.check_one(item,rows)['verdict'],'write')
        rows[2]=item['receipt_correction']['after']
        self.assertEqual(V.check_one(item,rows)['verdict'],'skip')
    def test_blank_receipt_even_with_settled_flag_is_filled(self):
        _,_,item,rows=self.make(None,None)
        self.assertEqual(V.check_one(item,rows)['verdict'],'write')
        self.assertEqual(item['five_cols']['回款明细'],100)
    def test_changed_row_after_plan_stays_blocked(self):
        _,_,item,rows=self.make();rows[2]['回款明细']=80
        self.assertEqual(V.check_one(item,rows)['verdict'],'conflict')
    def test_wrong_event_or_amount_stays_blocked(self):
        for field,value in [('amount_local',101),('writeoff_sequence_key',['2026-09-09','OTHER','HXD1'])]:
            _,_,item,rows=self.make();item['split_payment_source'][field]=value
            self.assertEqual(V.check_one(item,rows)['verdict'],'conflict')
    def test_ambiguous_or_partial_receipt_uses_existing_planner(self):
        rec,ledger,item,_=self.make()
        rec['cumulative_received_local']=200
        self.assertIsNone(R.candidate(rec,item,ledger))
        rec['cumulative_received_local']=100;ledger.business_rows=lambda so,sod:[2,3];ledger.row_snapshot[3]=copy.deepcopy(ledger.row_snapshot[2])
        self.assertIsNone(R.candidate(rec,item,ledger))

    def test_real_workbook_write_readback_and_original_preservation(self):
        import tempfile
        from pathlib import Path
        import openpyxl
        import apply_to_copy as writer
        with tempfile.TemporaryDirectory() as temp:
            src=Path(temp)/'before.xlsx';out=Path(temp)/'after.xlsx'
            book=openpyxl.Workbook();sheet=book.active;sheet.title='明细'
            sheet.append(['新智云单号','实收金额','应收金额','计提','回款明细','是否结账','收款时间','收款方式(支/汇/现)','差异'])
            sheet.append(['SO1','SOD1',100,100,50,'是','2026-08-01','汇',0])
            book.create_sheet('保留')['A1']='不可修改';book.save(src);book.close()
            original=src.read_bytes();_,_,item,_=self.make()
            self.assertEqual(V.check_one(item,writer.read_ledger_rows(src))['verdict'],'write')
            changes=writer.write_plan(src,out,[item])
            self.assertTrue(changes)
            self.assertEqual(writer.verify_written(out,[item]),[])
            self.assertEqual(V.check_one(item,writer.read_ledger_rows(out))['verdict'],'skip',V.check_one(item,writer.read_ledger_rows(out)))
            self.assertEqual(src.read_bytes(),original)
            book=openpyxl.load_workbook(out);self.assertEqual(book['保留']['A1'].value,'不可修改');book.close()

    def test_settled_flags_without_receipts_do_not_skip_whole_so(self):
        from classify_hexiao import LedgerIndex
        _,ledger,_,_=self.make(None,None)
        ledger.so_index={'SO1':[2]}
        self.assertFalse(LedgerIndex.so_settlement(ledger,'SO1')['all_settled'])
    def test_classifier_selects_correction_before_settled_skip(self):
        from classify_hexiao import classify_one
        rec,ledger,_,_=self.make()
        actual=classify_one(rec,ledger,{},0.01,2026)
        self.assertEqual(actual['receipt_correction']['policy'],R.POLICY)
        self.assertEqual(actual['five_cols']['回款明细'],100)

    def test_verified_parent_allocation_identity_is_preserved(self):
        rec,ledger,result,rows=self.make()
        rec['writeoff_sequence_key']=[];rec['parent_allocation_audit']={'applied':False}
        result['split_payment_source']['writeoff_sequence_key']=[]
        result['parent_allocation_audit']=rec['parent_allocation_audit']
        item=R.candidate(rec,result,ledger)
        self.assertEqual(V.check_one(item,rows)['verdict'],'write')

    def test_partial_receipt_uses_normal_split_plan(self):
        rec,ledger,result,rows=self.make()
        rec.update(amount_local=40,amount_orig=40,cumulative_received_local=40,cumulative_received_orig=40)
        item=R.candidate(rec,result,ledger)
        self.assertIsNotNone(item)
        self.assertEqual(item['row_operation']['type'],'split_below')
        self.assertEqual(item['row_operation']['paid_receivable'],40)
        self.assertEqual(item['row_operation']['unpaid_receivable'],60)
        rows[2]['_差异列存在']=True
        self.assertEqual(V.check_one(item,rows)['verdict'],'write')

    def test_partial_correction_writes_two_rows_and_is_idempotent(self):
        import tempfile
        from pathlib import Path
        import openpyxl
        import apply_to_copy as writer
        with tempfile.TemporaryDirectory() as temp:
            src=Path(temp)/'before.xlsx';out=Path(temp)/'after.xlsx'
            book=openpyxl.Workbook();sheet=book.active;sheet.title='明细'
            sheet.append(['新智云单号','实收金额','应收金额','计提','回款明细','是否结账','收款时间','收款方式(支/汇/现)','差异'])
            sheet.append(['SO1','SOD1',100,100,50,'是','2026-08-01','汇',0]);book.save(src);book.close()
            rec,ledger,result,_=self.make()
            rec.update(amount_local=40,amount_orig=40,cumulative_received_local=40,cumulative_received_orig=40)
            item=R.candidate(rec,result,ledger)
            self.assertEqual(V.check_one(item,writer.read_ledger_rows(src))['verdict'],'write')
            writer.write_plan(src,out,[item])
            self.assertEqual(writer.verify_written(out,[item]),[])
            self.assertEqual(V.check_one(item,writer.read_ledger_rows(out))['verdict'],'skip',V.check_one(item,writer.read_ledger_rows(out)))
            book=openpyxl.load_workbook(out,data_only=True);sheet=book['明细']
            self.assertEqual(sheet['C2'].value,40);self.assertEqual(sheet['C3'].value,60)
            self.assertEqual(sheet['E2'].value,40);self.assertIsNone(sheet['E3'].value);book.close()

if __name__=='__main__':unittest.main()
