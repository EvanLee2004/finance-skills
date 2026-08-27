# -*- coding: utf-8 -*-
import json
import subprocess
import sys
from pathlib import Path

import openpyxl

import build_task_reports as B
import workbook_finalize as W


def _find_lightweight_checker():
    for parent in Path(__file__).resolve().parents:
        for relative in (
            ("tools", "xlsx_lightweight_audit.py"),
            ("03_tools", "scripts", "xlsx_lightweight_audit.py"),
        ):
            candidate = parent.joinpath(*relative)
            if candidate.is_file():
                return candidate
    return None


def _report(path, title, value):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    ws.append(["项目", "值"])
    ws.append(["日期", value])
    wb.save(path)


def test_build_task_reports_keeps_only_three_static_range_files(tmp_path):
    out = tmp_path / "04_产出"
    out.mkdir()
    for token, date in (("20260808", "2026-08-08"), ("20260809", "2026-08-09")):
        for prefix in ("核销日清", "变更清单", "订单写入差异"):
            _report(out / f"{prefix}_{token}.xlsx", "明细", date)
        (out / f"判定结果_{token}.json").write_text(
            json.dumps({
                "payment_count": 2,
                "counts": {"total": 3, "auto": 2, "hold": 1, "exception": 0},
            }, ensure_ascii=False),
            encoding="utf-8",
        )
    _report(out / "变更清单_20260808_2025.xlsx", "内部", "2026-08-08")
    _report(out / "核销日清_20260810.xlsx", "范围外", "2026-08-10")

    outputs = B.build(tmp_path, "2026-08-08", "2026-08-09")

    assert len(outputs) == 3
    for path in outputs:
        wb = openpyxl.load_workbook(path, data_only=False)
        assert wb.sheetnames[0] == "任务范围"
        assert {row[0].value for row in wb["任务范围"].iter_rows(min_row=2)} == {
            "2026-08-08", "2026-08-09"
        }
        assert len(wb.sheetnames) == 3
        assert wb.sheetnames[1].startswith("20260808_")
        assert wb.sheetnames[2].startswith("20260809_")
        wb.close()
        audit = W.inspect_calculation(path)
        assert audit.formula_cells == 0
        assert audit.full_calc_on_load == "0"
        assert audit.force_full_calc == "0"
        checker = _find_lightweight_checker()
        if checker is not None:
            checked = subprocess.run(
                [sys.executable, str(checker), str(path), "--strict"],
                capture_output=True, text=True, encoding="utf-8",
            )
            assert checked.returncode == 0, checked.stdout + checked.stderr

    assert not list(out.glob("核销日清_2026080[89].xlsx"))
    assert not list(out.glob("变更清单_2026080[89].xlsx"))
    assert not list(out.glob("订单写入差异_2026080[89].xlsx"))
    assert not (out / "变更清单_20260808_2025.xlsx").exists()
    assert (out / "核销日清_20260810.xlsx").is_file()
    assert (out / "判定结果_20260808.json").is_file()
    assert (out / "判定结果_20260809.json").is_file()


def test_build_task_reports_accepts_confirmed_empty_fetch_day(tmp_path):
    out = tmp_path / "04_产出"
    export = tmp_path / "01_智云导出"
    out.mkdir()
    export.mkdir()
    (export / "取数摘要_20260808.json").write_text(
        json.dumps(
            {
                "回款记录笔数": 0,
                "下单行数": 0,
                "核销明细行数": 0,
                "订单明细SOD行数": 0,
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    for prefix in ("核销日清", "变更清单", "订单写入差异"):
        _report(out / f"{prefix}_20260809.xlsx", "明细", "2026-08-09")
    (out / "判定结果_20260809.json").write_text(
        json.dumps(
            {"payment_count": 1, "counts": {"total": 1, "auto": 1, "hold": 0, "exception": 0}},
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    outputs = B.build(tmp_path, "2026-08-08", "2026-08-09")

    assert len(outputs) == 3
    workbook = openpyxl.load_workbook(outputs[0], data_only=True)
    rows = list(workbook["任务范围"].iter_rows(min_row=2, values_only=True))
    assert rows[0][:2] == ("2026-08-08", "无核销记录，已跳过")
    assert rows[1][:2] == ("2026-08-09", "已纳入")
    workbook.close()
