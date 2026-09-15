"""Synthetic workbook acceptance at the approved flow-write boundary."""
import datetime as dt
import tempfile
import unittest
from pathlib import Path
import openpyxl
import apply_flow
import build_flow_plan

class MonthlyFlowTest(unittest.TestCase):
    def test_monthly_carry_and_repeat(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "02_我的表副本").mkdir()
            path = root / "02_我的表副本" / "flow.xlsx"
            wb = openpyxl.Workbook(); ws = wb.active; ws.title = "流水"
            ws.append(["日期", "公司名称", "金额", "收款形式", "单号", "预收", "是否更新应收款"])
            ws.append([dt.date(2026,7,27), "合成客户", 9000, "汇款", "WX", 9000, None])
            wb.save(path); wb.close()
            def run(date, so, amount):
                entry = {"ar":"ARTEST", "so":so, "sod":so+"D", "case_id":"ARTEST|"+so,
                    "bucket":"auto", "flow_hits":1,"flow_matched_by":"三键",
                    "flow_file":"flow.xlsx","flow_sheet":"流水","flow_row_no":2,
                    "flow_identity":{"date":"2026-07-27","payer":"合成客户","amount":9000},
                    "split_payment_source":{"amount_local":amount,"so_delivery_local":10000},
                    "write_currency_audit":{"currency":"CNY","amount_local":amount}}
                plan = build_flow_plan.build_plan({"auto":[entry],"hexiao_date":date})
                checked = {"hexiao_date":date,"write":[entry],"skip":[],"conflict":[]}
                final = build_flow_plan.finalize_plan_after_ledger(plan, checked)
                changes, errors = apply_flow.write_flow_items(root, final["items"], in_place=True, phase="status")
                self.assertEqual(errors, [])
                return changes
            run("2026-07-28","SO1",1000)
            wb = openpyxl.load_workbook(path); self.assertEqual(wb["流水"]["F2"].value,"9000-1000=8000");wb.close()
            run("2026-08-01","SO2",3000)
            run("2026-08-02","SO3",5000)
            before = path.read_bytes()
            self.assertEqual(run("2026-08-02","SO3",5000), [])
            self.assertEqual(path.read_bytes(),before)
            wb=openpyxl.load_workbook(path); ws=wb["流水"]
            self.assertEqual(ws.max_row,3)
            self.assertEqual(ws["C3"].value,8000)
            self.assertEqual(ws["A3"].value,dt.datetime(2026,8,1))
            self.assertEqual(ws["F3"].value,"8000-3000-5000=0")
            self.assertIn("SO2",ws["E3"].value);self.assertIn("SO3",ws["E3"].value)
            self.assertNotIn("SO1",ws["E3"].value)
            self.assertEqual(ws["F2"].value,"9000-1000=8000");wb.close()
            wb=openpyxl.load_workbook(path,data_only=True);self.assertEqual(wb["流水"]["F3"].value,"8000-3000-5000=0");wb.close()



class MonthlySafetyTest(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        (self.root/'02_我的表副本').mkdir();(self.root/'04_产出').mkdir()
        self.path=self.root/'02_我的表副本'/'flow.xlsx'
        wb=openpyxl.Workbook();ws=wb.active;ws.title='流水'
        ws.append(['日期','公司名称','金额','收款形式','单号','预收','是否更新应收款'])
        ws.append([dt.date(2026,7,27),'测试甲',9000,'汇款','WX',9000,None])
        wb.save(self.path);wb.close()
    def tearDown(self):self.temp.cleanup()
    def entry(self,so='SO1',amount=1000,ar='AR1',row=2,payer='测试甲'):
        return {'ar':ar,'so':so,'sod':so+'D','case_id':ar+'|'+so+'|'+so+'D','bucket':'auto',
                'flow_hits':1,'flow_matched_by':'三键','flow_file':'flow.xlsx','flow_sheet':'流水','flow_row_no':row,
                'flow_identity':{'date':'2026-07-27','payer':payer,'amount':9000},
                'split_payment_source':{'amount_local':amount,'so_delivery_local':10000},
                'write_currency_audit':{'currency':'CNY','amount_local':amount}}
    def plans(self,day,entries,kind='write'):
        flow=build_flow_plan.build_plan({'auto':entries,'hexiao_date':day})
        checked={'hexiao_date':day,'write':[],'skip':[],'conflict':[]};checked[kind]=entries
        return flow,checked
    def run_items(self,day,entries,kind='write'):
        flow,checked=self.plans(day,entries,kind)
        final=build_flow_plan.finalize_plan_after_ledger(flow,checked)
        return apply_flow.write_flow_items(self.root,final['items'],in_place=True,phase='status')
    def read(self,cell,cached=False):
        wb=openpyxl.load_workbook(self.path,data_only=cached);result=wb['流水'][cell].value;wb.close();return result
    def modify(self,edit):
        wb=openpyxl.load_workbook(self.path);edit(wb['流水']);wb.save(self.path);wb.close()
    def test_literal_equation_adoption_keeps_all_deductions(self):
        self.modify(lambda ws:setattr(ws['F2'],'value','9000-1000-2000=6000'))
        self.assertEqual(self.run_items('2026-07-28',[self.entry(amount=500)])[1],[])
        self.assertEqual(self.read('F2'),'9000-1000-2000-500=5500')
        self.assertEqual(self.read('F2',cached=True),'9000-1000-2000-500=5500')
    def test_wrong_literal_equation_does_not_write(self):
        self.modify(lambda ws:setattr(ws['F2'],'value','9000-1000=7000'))
        before=self.path.read_bytes()
        changes,errors=self.run_items('2026-07-28',[self.entry(amount=500)])
        self.assertEqual(changes,[]);self.assertTrue(errors)
        self.assertEqual(self.path.read_bytes(),before)
    def test_literal_equation_decimal_deductions(self):
        self.assertEqual(self.run_items('2026-07-28',[self.entry(amount=0.01)])[1],[])
        self.assertEqual(self.read('F2'),'9000-0.01=8999.99')
        self.assertEqual(self.run_items('2026-07-29',[self.entry('SO2',8999.99)])[1],[])
        self.assertEqual(self.read('F2'),'9000-0.01-8999.99=0')
    def test_single_so_cross_month(self):
        changes,errors=self.run_items('2026-09-04',[self.entry(amount=9000)])
        self.assertEqual(errors,[]);self.assertEqual(len(changes),1)
        self.assertEqual(self.read('E2'),'WX转9月');self.assertEqual(self.read('D3'),'冲预收')
        self.assertEqual(self.read('F3'),'9000-9000=0');self.assertEqual(self.read('C2'),9000)
    def test_same_so_separate_events(self):
        self.assertEqual(self.run_items('2026-07-28',[self.entry(amount=1000)])[1],[])
        self.assertEqual(self.run_items('2026-07-29',[self.entry(amount=2000)])[1],[])
        self.assertEqual(self.read('E2'),'WX\nSO1  1,000.00\nSO1  2,000.00')
        self.assertEqual(self.read('F2'),'9000-1000-2000=6000')
    def test_skip_state_does_not_deduct(self):
        before=self.path.read_bytes();entry=self.entry();entry['_check']={'reason':'目标业务行全部已结账'}
        self.assertEqual(self.run_items('2026-08-01',[entry],'skip'),([],[]))
        self.assertEqual(self.path.read_bytes(),before)
    def test_skip_missing_event_needs_existing_registration(self):
        entry=self.entry();entry['_check']={'reason':'已经填过且与本次一致（幂等跳过）'}
        before=self.path.read_bytes();self.assertTrue(self.run_items('2026-08-01',[entry],'skip')[1])
        self.assertEqual(self.path.read_bytes(),before)
    def test_event_date_cannot_be_changed(self):
        e=self.entry();e['split_payment_source']['writeoff_sequence_key']=['2026-07-28','EVENT1','R1','AR1','SO1']
        self.assertEqual(self.run_items('2026-07-28',[e])[1],[])
        before=self.path.read_bytes();self.assertEqual(self.run_items('2026-08-01',[e]),([],[]))
        self.assertEqual(self.path.read_bytes(),before)
    def test_overdraw_and_backdate_leave_file_unchanged(self):
        self.assertEqual(self.run_items('2026-08-01',[self.entry()])[1],[])
        before=self.path.read_bytes()
        self.assertTrue(self.run_items('2026-08-02',[self.entry('SO2',9000)])[1])
        self.assertEqual(self.path.read_bytes(),before)
        self.assertTrue(self.run_items('2026-07-29',[self.entry('SO2',1000)])[1])
        self.assertEqual(self.path.read_bytes(),before)
    def test_legacy_split_adoption(self):
        def old(ws):
            ws['E2']='WX\nSO0  1000\nWX转8月';ws['F2']='=9000-1000'
            ws.append([dt.date(2026,8,1),'测试甲',8000,'冲预收','WX\nSO1  3000','=8000-3000','是'])
        self.modify(old)
        self.assertEqual(self.run_items('2026-08-02',[self.entry('SO2',5000)])[1],[])
        self.assertEqual(self.read('F3'),'8000-3000-5000=0');self.assertEqual(self.read('A3'),dt.datetime(2026,8,1))
    def test_legacy_balance_above_receipt_does_not_write(self):
        self.modify(lambda ws:(setattr(ws['E2'],'value','SO0  2000'),setattr(ws['F2'],'value',10000)))
        before=self.path.read_bytes();self.assertTrue(self.run_items('2026-08-01',[self.entry()])[1]);self.assertEqual(self.path.read_bytes(),before)
    def test_multiple_receipts_shift_rows_once(self):
        self.modify(lambda ws:ws.append([dt.date(2026,7,27),'测试乙',9000,'汇款','WX',9000,None]))
        e1=self.entry();e2=self.entry('SO2',2000,'AR2',3,'测试乙')
        self.assertEqual(self.run_items('2026-08-01',[e1,e2])[1],[])
        self.assertEqual(self.read('B3'),'测试甲');self.assertEqual(self.read('B4'),'测试乙');self.assertEqual(self.read('B5'),'测试乙')
        e2=self.entry('SO3',500,'AR2',4,'测试乙')
        self.assertEqual(self.run_items('2026-08-02',[e2])[1],[]);self.assertEqual(self.read('F5'),'9000-2000-500=6500')
    def test_filters_and_same_row_month_formula(self):
        def fixture(ws):
            ws.auto_filter.ref='A1:H2';ws['H1']='月份';ws['H2']='=MONTH(A2)'
        self.modify(fixture)
        self.assertEqual(self.run_items('2026-08-01',[self.entry()])[1],[])
        self.assertEqual(self.read('H3'),'=MONTH(A3)')
        wb=openpyxl.load_workbook(self.path);self.assertEqual(wb['流水'].auto_filter.ref,'A1:H3');wb.close()
    def test_unrelated_zip_part_preserved(self):
        import zipfile
        with zipfile.ZipFile(self.path,'a') as z:z.writestr('customXml/unrelated.xml',b'<preserved/>')
        self.assertEqual(self.run_items('2026-08-01',[self.entry()])[1],[])
        with zipfile.ZipFile(self.path) as z:self.assertEqual(z.read('customXml/unrelated.xml'),b'<preserved/>')
    def test_full_execution_and_independent_workbook_replay(self):
        import json,shutil,execution_flow_stage,verify_execution_write
        baseline=self.root/'baseline';(baseline/'02_我的表副本').mkdir(parents=True)
        shutil.copy2(self.path,baseline/'02_我的表副本'/'flow.xlsx')
        flow,checked=self.plans('2026-08-01',[self.entry()])
        (self.root/'04_产出'/'流转写入计划_校验后.json').write_text(json.dumps(flow),encoding='utf-8')
        checked_file=self.root/'04_产出'/'checked.json';checked_file.write_text(json.dumps(checked),encoding='utf-8')
        result=execution_flow_stage.run(self.root,checked_file)
        self.assertTrue(result['flow_written'])
        (self.root/'04_产出'/'流转阶段执行结果.json').write_text(json.dumps(result),encoding='utf-8')
        proof=verify_execution_write.verify_flow(self.root,baseline,self.path,checked)
        self.assertEqual(proof['mode'],'approved_patch_all_parts')
    def test_cross_year_creates_new_month(self):
        self.assertEqual(self.run_items('2026-12-31',[self.entry()])[1],[])
        self.assertEqual(self.run_items('2027-01-01',[self.entry('SO2',1000)])[1],[])
        self.assertEqual(self.read('A4'),dt.datetime(2027,1,1));self.assertEqual(self.read('C4'),8000)

    def test_same_event_new_record_is_distinct_and_old_record_not_reapplied(self):
        first=self.entry();first['split_payment_source']['writeoff_sequence_key']=['2026-07-28','EVENT1','R1','AR1','SO1']
        self.assertEqual(self.run_items('2026-07-28',[first])[1],[])
        self.assertEqual(self.run_items('2026-07-28',[first]),([],[]))
        second=self.entry(amount=2000);second['split_payment_source']['writeoff_sequence_key']=['2026-07-29','EVENT2','R2','AR1','SO1']
        self.assertEqual(self.run_items('2026-07-29',[second])[1],[])
        self.assertEqual(self.read('F2'),'9000-1000-2000=6000')
    def test_conflict_does_not_enter_balance(self):
        good=self.entry();bad=self.entry('SO2',2000)
        flow=build_flow_plan.build_plan({'auto':[good,bad],'hexiao_date':'2026-07-28'})
        final=build_flow_plan.finalize_plan_after_ledger(flow,{'hexiao_date':'2026-07-28','write':[good],'conflict':[bad]})
        self.assertEqual(apply_flow.write_flow_items(self.root,final['items'],in_place=True,phase='status')[1],[])
        self.assertEqual(self.read('F2'),'9000-1000=8000');self.assertNotIn('SO2',self.read('E2'))
    def test_unsupported_cross_row_formula_is_rejected_without_change(self):
        self.modify(lambda ws:setattr(ws['H2'],'value','=C2+C3'))
        before=self.path.read_bytes();self.assertTrue(self.run_items('2026-08-01',[self.entry()])[1]);self.assertEqual(self.path.read_bytes(),before)
    def test_original_currency_formula_uses_local_balance(self):
        self.modify(lambda ws:(setattr(ws['C2'],'value','=1000*9'),setattr(ws['F2'],'value',9000)))
        # Excel cached amount is produced with the project's formula-preserving patcher.
        import xlsx_patch,shutil
        patched=self.root/'cached.xlsx';xlsx_patch.patch_cells(self.path,patched,'流水',[(2,3,xlsx_patch.FormulaValue('=1000*9',9000))]);shutil.copy2(patched,self.path)
        e=self.entry();e['flow_matched_by']='三键(原币公式)';e['write_currency_audit']['currency']='USD'
        self.assertEqual(self.run_items('2026-08-01',[e])[1],[]);self.assertEqual(self.read('F3'),'9000-1000=8000')
    def test_sorting_managed_receipts_relocates_all_chains(self):
        self.modify(lambda ws:ws.append([dt.date(2026,7,27),'测试乙',9000,'汇款','WX',9000,None]))
        self.assertEqual(self.run_items('2026-08-01',[self.entry(),self.entry('SO2',1000,'AR2',3,'测试乙')])[1],[])
        import xlsx_patch,shutil
        wb=openpyxl.load_workbook(self.path);ws=wb['流水'];edits=[]
        for origin,destination in [(2,4),(3,5),(4,2),(5,3)]:
            for col in range(1,8):
                val=ws.cell(origin,col).value
                if col==6:val=xlsx_patch.FormulaValue(val,8000) if isinstance(val,str) and val.startswith('=') else val
                edits.append((destination,col,val))
        wb.close();patched=self.root/'sorted.xlsx';xlsx_patch.patch_cells(self.path,patched,'流水',edits);shutil.copy2(patched,self.path)
        self.assertEqual(self.run_items('2026-08-02',[self.entry('SO3',500,'AR1',4,'测试甲')])[1],[])
        self.assertEqual(self.read('B3'),'测试乙');self.assertEqual(self.read('F3'),'9000-1000=8000')
        self.assertEqual(self.read('F5'),'9000-1000-500=7500')
    def test_existing_red_text_preserved(self):
        from openpyxl.cell.rich_text import CellRichText,TextBlock
        from openpyxl.cell.text import InlineFont
        def old(ws):
            ws['E2']=CellRichText('WX\n',TextBlock(InlineFont(color='FFFF0000'),'SO0  1,000.00'))
            ws['F2']='=9000-1000'
        self.modify(old)
        self.assertEqual(self.run_items('2026-07-28',[self.entry('SO1',2000)])[1],[])
        wb=openpyxl.load_workbook(self.path,rich_text=True)
        self.assertEqual(apply_flow._line_colors(wb['流水']['E2'].value)['SO0'],'FFFF0000');wb.close()



    def test_order_and_balance_both_use_actual_receipt(self):
        self.assertEqual(self.run_items('2026-07-28',[self.entry(amount=1000)])[1],[])
        self.assertEqual(self.read('E2'),'WX\nSO1  1,000.00')
        self.assertEqual(self.read('F2'),'9000-1000=8000')
    def test_legacy_delivery_is_not_used_as_paid_amount(self):
        def old(ws):
            ws['E2']='WX\nSO0  29,000.00';ws['F2']='=9000-1000'
        self.modify(old)
        self.assertEqual(self.run_items('2026-07-28',[self.entry('SO1',2000)])[1],[])
        self.assertEqual(self.read('F2'),'9000-1000-2000=6000')
        self.assertIn('SO0  29,000.00',self.read('E2'))
        self.assertIn('SO1  2,000.00',self.read('E2'))



    def test_same_so_actual_receipts_are_summed_across_sods(self):
        a=self.entry(amount=1000);b=self.entry(amount=2000);b['sod']='OTHER';b['case_id']='AR1|SO1|OTHER'
        self.assertEqual(self.run_items('2026-07-28',[a,b])[1],[])
        self.assertEqual(self.read('E2'),'WX\nSO1  3,000.00')
        self.assertEqual(self.read('F2'),'9000-1000-2000=6000')
    def test_receipt_display_does_not_require_delivery(self):
        e=self.entry();e['split_payment_source'].pop('so_delivery_local')
        before=self.path.read_bytes();flow,checked=self.plans('2026-07-28',[e])
        final=build_flow_plan.finalize_plan_after_ledger(flow,checked)
        self.assertEqual(final['items'][0]['verdict'],'write')
        self.assertEqual(self.path.read_bytes(),before)
        self.assertEqual(self.run_items('2026-07-28',[e])[1],[])
        self.assertEqual(self.read('E2'),'WX\nSO1  1,000.00')
    def test_plain_legacy_balance_is_independent_of_delivery(self):
        self.modify(lambda ws:(setattr(ws['E2'],'value','SO0  90000'),setattr(ws['F2'],'value','还剩：8000')))
        self.assertEqual(self.run_items('2026-08-01',[self.entry(amount=2000)])[1],[])
        self.assertEqual(self.read('C3'),8000);self.assertEqual(self.read('F3'),'8000-2000=6000')
    def test_blank_prepaid_is_not_inferred_from_delivery(self):
        self.modify(lambda ws:(setattr(ws['E2'],'value','SO0  90000'),setattr(ws['F2'],'value',None)))
        before=self.path.read_bytes();errors=self.run_items('2026-08-01',[self.entry()])[1]
        self.assertTrue(errors);self.assertIn('交付额不能推算',errors[0]);self.assertEqual(self.path.read_bytes(),before)
    def test_preview_displays_delivery_once(self):
        a=self.entry(amount=1000);b=self.entry(amount=2000);b['case_id']='AR1|SO1|OTHER'
        plan,_=self.plans('2026-07-28',[a,b])
        self.assertEqual(plan['items'][0]['order_suggest'],'SO1  10,000.00')
    def test_v1_metadata_remains_readable_without_reallocating(self):
        import flow_monthly,zipfile,json
        from xml.etree import ElementTree as ET
        self.assertEqual(self.run_items('2026-07-28',[self.entry()])[1],[])
        state=flow_monthly.load_state(self.path);state['schema']='receipt-monthly-v1'
        for chain in state['receipts'].values():
            for month in chain['months']:
                for key in ['display_original','display_amounts','legacy_sos']:month.pop(key,None)
        flow_monthly.save_state(self.path,state)
        self.assertEqual(self.run_items('2026-07-29',[self.entry('SO2',2000)])[1],[])
        self.assertEqual(self.read('F2'),'9000-1000-2000=6000')
        self.assertIn('SO2  2,000.00',self.read('E2'))



    def test_lowercase_legacy_so_cannot_bypass_history_guard(self):
        self.modify(lambda ws:(setattr(ws['E2'],'value','so1  90000'),setattr(ws['F2'],'value','=9000-1000')))
        before=self.path.read_bytes();errors=self.run_items('2026-07-28',[self.entry()])[1]
        self.assertTrue(errors);self.assertIn('已有人工核销记录',errors[0])
        self.assertEqual(self.path.read_bytes(),before)



    def test_delivery_can_exceed_receipt_without_balance_conflict(self):
        self.modify(lambda ws:(setattr(ws['C2'],'value',519),setattr(ws['F2'],'value',519)))
        a=self.entry('SOA',323.23);b=self.entry('SOB',195.77)
        for e in [a,b]:e['flow_identity']['amount']=519
        a['split_payment_source']['so_delivery_local']=987.84
        b['split_payment_source']['so_delivery_local']=29000
        self.assertEqual(self.run_items('2026-07-28',[a,b])[1],[])
        self.assertIn('SOA  323.23',self.read('E2'));self.assertIn('SOB  195.77',self.read('E2'))
        self.assertEqual(self.read('F2'),'519-323.23-195.77=0');self.assertEqual(self.read('F2',True),'519-323.23-195.77=0')


    def test_verified_published_receipt_is_automatically_registered(self):
        entry=self.entry()
        entry['split_payment_source']['writeoff_sequence_key']=['2026-08-01','HX1','ROW1','AR1','SO1']
        entry['code']='OK_BASELINE_RECEIPT_APPLIED'
        entry['_check']={'verdict':'skip','reason':'本次父 AR 和核销记录已发布，实际回款行一致；跳过且不重复累计'}
        entry['baseline_receipt_audit']={'disposition':'skip','current_received':1000,
            'event_key':'["AR1", "SO1", "SO1D", "HX1"]'}
        entry['five_cols']={'回款明细':1000,'是否结账':'是','收款时间':'2026-08-01','收款方式':'冲预收'}
        changes,errors=self.run_items('2026-08-01',[entry],'skip')
        self.assertEqual(errors,[]);self.assertEqual(len(changes),1)
        self.assertEqual(self.read('F3'),'9000-1000=8000')


    def test_complete_published_parent_repairs_blank_prepaid_once(self):
        import validate_plan
        self.modify(lambda ws:(setattr(ws['C2'],'value',519),setattr(ws['F2'],'value',None),
            setattr(ws['E2'],'value','WXSOA  987.84\nSOB  29000')))
        entries=[self.entry('SOA',323.23),self.entry('SOB',195.77)]
        rows={};cases={};allocations=[]
        for i,e in enumerate(entries,2):
            e['code']='OK_SO_ALREADY_SETTLED';e['flow_identity']['amount']=519
            e['split_payment_source']['so_delivery_local']=987.84 if i==2 else 29000
            e['split_payment_source']['cumulative_local']=987.84 if i==2 else 195.77
            cases[e['case_id']]={'so':e['so'],'sod':e['sod'],'amount_local':e['split_payment_source']['amount_local']}
            allocations.append({'so':e['so'],'allocated_local':e['split_payment_source']['amount_local']})
            rows[i]={'SO':e['so'],'SOD':e['sod'],'回款明细':e['split_payment_source']['cumulative_local'],'是否结账':'是'}
        parent={'ar':'AR1','hexiao_date':'2026-07-28','parent_net_amount':519,'unallocated_parent_amount':0,
            'applied_at':'2026-07-28T10:00:00','reused_successful_allocation':True,'applied_cases':cases,'allocations':allocations}
        checked=validate_plan.validate({'hexiao_date':'2026-07-28','auto':entries,'parent_fallback_allocations':{'AR1':parent}},rows)
        flow=build_flow_plan.build_plan({'auto':entries,'hexiao_date':'2026-07-28'})
        final=build_flow_plan.finalize_plan_after_ledger(flow,checked)
        changes,errors=apply_flow.write_flow_items(self.root,final['items'],in_place=True,phase='status')
        self.assertEqual(errors,[]);self.assertEqual(len(changes),1)
        self.assertEqual(self.read('F2'),'519-323.23-195.77=0');self.assertEqual(self.read('F2',True),'519-323.23-195.77=0')
        self.assertIn('SOA  323.23',self.read('E2'));self.assertIn('SOB  195.77',self.read('E2'))
        before=self.path.read_bytes()
        self.assertEqual(apply_flow.write_flow_items(self.root,final['items'],in_place=True,phase='status'),([],[]))
        self.assertEqual(self.path.read_bytes(),before)


    def test_one_unresolved_receipt_does_not_block_other_automatic_rows(self):
        import json,shutil,execution_flow_stage,verify_execution_write
        def add(ws):
            ws['E2']='SOOLD  30000';ws['F2']=None
            ws.append([dt.date(2026,7,27),'测试乙',9000,'汇款','WX',9000,None])
        self.modify(add)
        entries=[self.entry(),self.entry('SO2',2000,'AR2',3,'测试乙')]
        flow,checked=self.plans('2026-08-01',entries)
        baseline=self.root/'baseline';shutil.copytree(self.root/'02_我的表副本',baseline/'02_我的表副本')
        out=self.root/'04_产出';(out/'流转写入计划_校验后.json').write_text(json.dumps(flow),encoding='utf-8')
        cp=out/'checked.json';cp.write_text(json.dumps(checked),encoding='utf-8')
        result=execution_flow_stage.run(self.root,cp)
        self.assertTrue(result['flow_written']);self.assertEqual(result['manual_count'],1)
        self.assertEqual(result['manual_items'][0]['ar'],'AR1')
        self.assertEqual(result['phases']['status']['changed_count'],1)
        self.assertIsNone(self.read('F2'));self.assertEqual(self.read('F4'),'9000-2000=7000')
        (out/'流转阶段执行结果.json').write_text(json.dumps(result),encoding='utf-8')
        proof=verify_execution_write.verify_flow(self.root,baseline,self.path,checked)
        self.assertEqual(proof['mode'],'approved_patch_all_parts')


    def test_settled_so_with_only_same_amount_row_cannot_prove_parent(self):
        import validate_plan
        e=self.entry();e['code']='OK_SO_ALREADY_SETTLED'
        e['split_payment_source']['writeoff_sequence_key']=['2026-08-01','HX1','R1','AR1','SO1']
        rows={2:{'SO':'SO1','SOD':'SO1D','回款明细':1000,'是否结账':'是','收款时间':'2026-08-01','收款方式':'冲预收'}}
        checked=validate_plan.validate({'hexiao_date':'2026-08-01','auto':[e]},rows)
        flow=build_flow_plan.build_plan({'auto':[e],'hexiao_date':'2026-08-01'})
        final=build_flow_plan.finalize_plan_after_ledger(flow,checked,workspace=self.root)
        changes,errors=apply_flow.write_flow_items(self.root,final['items'],in_place=True,phase='status')
        self.assertEqual(errors,[]);self.assertEqual(changes,[]);self.assertEqual(self.read('F2'),9000)
    def test_settled_so_without_matching_receipt_does_not_register(self):
        import validate_plan
        e=self.entry();e['code']='OK_SO_ALREADY_SETTLED'
        e['split_payment_source']['writeoff_sequence_key']=['2026-08-01','HX1','R1','AR1','SO1']
        rows={2:{'SO':'SO1','SOD':'SO1D','回款明细':900,'是否结账':'是','收款时间':'2026-08-01','收款方式':'冲预收'}}
        checked=validate_plan.validate({'hexiao_date':'2026-08-01','auto':[e]},rows)
        flow=build_flow_plan.build_plan({'auto':[e],'hexiao_date':'2026-08-01'})
        final=build_flow_plan.finalize_plan_after_ledger(flow,checked,workspace=self.root)
        before=self.path.read_bytes()
        self.assertEqual(apply_flow.write_flow_items(self.root,final['items'],in_place=True,phase='status'),([],[]))
        self.assertEqual(self.path.read_bytes(),before)
    def test_new_parent_proof_does_not_double_deduct_old_managed_event(self):
        e=self.entry();self.assertEqual(self.run_items('2026-07-28',[e])[1],[])
        before=self.path.read_bytes();e['_check']={'verdict':'skip'}
        e['flow_receipt_proof']={'basis':'published_parent_case','ar':'AR1','so':'SO1','sod':'SO1D','date':'2026-07-28','amount':'1000'}
        self.assertEqual(self.run_items('2026-07-28',[e],'skip'),([],[]));self.assertEqual(self.path.read_bytes(),before)
    def test_duplicate_receipt_targets_are_both_manual_before_any_write(self):
        a=self.entry();b=self.entry('SO2',2000,'AR2')
        flow,checked=self.plans('2026-08-01',[a,b]);before=self.path.read_bytes()
        final=build_flow_plan.finalize_plan_after_ledger(flow,checked,workspace=self.root)
        self.assertEqual(final['counts']['hand'],2)
        self.assertEqual(apply_flow.write_flow_items(self.root,final['items'],in_place=True,phase='status'),([],[]))
        self.assertEqual(self.path.read_bytes(),before)


    def test_fully_proven_current_receipt_moves_legacy_prefill_to_its_month(self):
        self.modify(lambda ws:(setattr(ws['E2'],'value','WXSO1  10000'),setattr(ws['F2'],'value',None)))
        e=self.entry(amount=9000);e['split_payment_source']['writeoff_sequence_key']=['2026-08-01','HX1','R1','AR1','SO1']
        changes,errors=self.run_items('2026-08-01',[e])
        self.assertEqual(errors,[]);self.assertEqual(len(changes),1)
        self.assertEqual(self.read('E2'),'WX转8月');self.assertEqual(self.read('E3'),'WX\nSO1  9,000.00')
        self.assertEqual(self.read('C3'),9000);self.assertEqual(self.read('F3'),'9000-9000=0')
        before=self.path.read_bytes();self.assertEqual(self.run_items('2026-08-01',[e]),([],[]));self.assertEqual(self.path.read_bytes(),before)


    def test_colored_remaining_text_can_be_used_for_monthly_carry(self):
        from openpyxl.cell.rich_text import CellRichText,TextBlock
        from openpyxl.cell.text import InlineFont
        self.modify(lambda ws:(setattr(ws['E2'],'value','SO0  30000'),
            setattr(ws['F2'],'value',CellRichText('还剩：',TextBlock(InlineFont(color='FFFF0000'),'8000')))))
        changes,errors=self.run_items('2026-08-01',[self.entry('SO1',2000)])
        self.assertEqual(errors,[]);self.assertEqual(len(changes),1)
        self.assertEqual(self.read('F3'),'8000-2000=6000')


    def test_published_parent_proof_keeps_existing_real_event_identity(self):
        e=self.entry();e['split_payment_source']['writeoff_sequence_key']=['2026-07-28','HX1','ROW1','AR1','SO1']
        self.assertEqual(self.run_items('2026-07-28',[e])[1],[])
        before=self.path.read_bytes();e['_check']={'verdict':'skip'}
        e['flow_receipt_proof']={'basis':'published_parent_case','ar':'AR1','so':'SO1','sod':'SO1D','date':'2026-07-28','amount':'1000'}
        self.assertEqual(self.run_items('2026-07-28',[e],'skip'),([],[]));self.assertEqual(self.path.read_bytes(),before)

if __name__ == "__main__": unittest.main()
