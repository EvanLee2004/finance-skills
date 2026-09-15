"""Six-sheet workpaper with numeric detail data, subtotal and summary formulas and calculated caches."""
from __future__ import annotations
from copy import copy
from decimal import Decimal
import math
import unicodedata
from io import BytesIO
from pathlib import Path
import json
import zipfile
import xml.etree.ElementTree as ET

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from engine import CATALOG, COMPANIES, KINDS, ZERO, VERSION, classify, validate


# Standard-template line IDs; signed entries subtract. Exclude included detail lines.
TOTAL_RULES = {'bs': {14: [1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13],
        33: [15, 16, 17, 18, 19, 20, 21, 22, 23, 24, 25, 26, 27, 28, 29, 30, 31, 32],
        34: [14, 33],
        48: [35, 36, 37, 38, 39, 40, 41, 42, 43, 44, 45, 46, 47],
        59: [49, 50, 53, 54, 55, 56, 57, 58],
        60: [48, 59],
        71: [61, 62, 65, -66, 67, 68, 69, 70],
        72: [60, 71]},
 'is': {19: [1, -2, -3, -4, -5, -6, -7, -10, -11, 12, 13, 15, 16, 17, 18],
        22: [19, 20, -21],
        24: [22, -23],
        27: [28, 34],
        28: [29, 30, 31, 32],
        34: [35, 36, 37, 38, 39, 40],
        42: [24, 27]},
 'cf': {4: [1, 2, 3],
        9: [5, 6, 7, 8],
        10: [4, -9],
        16: [11, 12, 13, 14, 15],
        21: [17, 18, 19, 20],
        22: [16, -21],
        26: [23, 24, 25],
        30: [27, 28, 29],
        31: [26, -30],
        33: [10, 22, 31, 32],
        35: [34, 33]}}

N = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
FONT=Font(name="宋体",size=18)
BORDER=Border(*( [Side(style="thin",color="777777")]*4 ))


def apply_workpaper_format(workbook):
    """Apply the presentation contract without changing values or formulas."""
    for sheet in workbook:
        # Amount columns share the requested width across the six main sheets.
        for row in sheet.iter_rows(min_row=1, max_row=min(8, sheet.max_row)):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.strip() in {
                    "年初数", "期末数", "年初余额", "期末余额",
                    "本月数", "本年累计", "本月金额", "本年累计金额",
                    "本期金额", "上期金额", "上年同期金额",
                }:
                    sheet.column_dimensions[get_column_letter(cell.column)].width = 30
        merged_widths = {
            (area.min_row, area.min_col): sum(
                sheet.column_dimensions[get_column_letter(col)].width
                for col in range(area.min_col, area.max_col + 1)
            ) for area in sheet.merged_cells.ranges
        }
        for row in sheet:
            height = max(30, sheet.row_dimensions[row[0].row].height or 0)
            for cell in row:
                if cell.value is None:
                    continue
                font = copy(cell.font)
                font.sz = 18
                cell.font = font
                if cell.alignment.wrap_text:
                    width = merged_widths.get(
                        (cell.row, cell.column),
                        sheet.column_dimensions[cell.column_letter].width,
                    )
                    capacity = max(1, (width - 2) * 11 / 18)
                    lines = sum(max(1, math.ceil(sum(
                        2 if unicodedata.east_asian_width(ch) in "WF" else 1
                        for ch in line
                    ) / capacity)) for line in str(cell.value).split("\n"))
                    height = max(height, lines * 27)
            sheet.row_dimensions[row[0].row].height = min(409, height)


def build_workbook(reports, period, path, *, input_issues=None):
    reports=[classify(r) for r in reports]
    order={c:i for i,c in enumerate(COMPANIES)}
    reports.sort(key=lambda r:(order[r.company],list(KINDS).index(r.kind)))
    w=openpyxl.Workbook();w.remove(w.active)
    cache={}
    for kind in KINDS:
        for scope in ["合并","母公司"]:
            name=scope+KINDS[kind]+period
            w.create_sheet(name)
    notes=[]
    notes.append(["编制说明","内容"])
    notes.append(["会计月份",period])
    notes.append(["范围","按可用来源直接汇总；未抵销，未进行其他人工调整"])
    notes.append(["分类规则","sd-management-to-cost-v1：山东管理费用全部计入营业成本"])
    notes.append(["标准模板",VERSION])
    for r in reports:
        notes.append(["来源",COMPANIES[r.company][0]+" / "+KINDS[r.kind]+" / "+r.file_name+" / "+r.sheet])
        notes.append(["来源SHA256",r.file_hash])
        for issue in r.issues:notes.append(["待核实",COMPANIES[r.company][0]+" / "+KINDS[r.kind]+" / "+issue])
        for check in validate(r):
            if not check["passed"]:notes.append(["校验未通过",COMPANIES[r.company][0]+" / "+check["check"]+" / 差额 "+check["difference"]])
    matrix=[]
    for company in COMPANIES:
        row={"company":company,"name":COMPANIES[company][0]}
        for kind in KINDS:
            available=next((r for r in reports if r.company==company and r.kind==kind),None)
            related=[issue for issue in input_issues or [] if issue.get("company")==company and issue.get("kind")==kind]
            missing_state=("解析失败" if any(issue.get("status") in {"rejected","needs_review"} for issue in related)
                           else "取数失败" if related else "未上传" if company in {"SHANDONG","SICHUAN","JINAN"} else "未取数")
            row[kind]="已纳入" if available else missing_state
            if not available:notes.append([missing_state,COMPANIES[company][0]+" / "+KINDS[kind]+(" / "+"；".join(issue["message"] for issue in related) if related else "")])
        matrix.append(row)
    for issue in input_issues or []:notes.append(["来源问题",json.dumps(issue,ensure_ascii=False)])
    coverage_complete=all(row[kind]=="已纳入" for row in matrix for kind in KINDS)
    notes.append(["编制状态","资料齐全，仍需检查口径提示" if coverage_complete else "部分底稿：未包含全部八家公司三张报表"])
    notes.append(["已纳入报表",str(len(reports))+" / 24"])
    validation=[]
    for kind in KINDS:
        catalog=CATALOG["standard"][kind]
        width=catalog["width"];delta=1 if kind=="is" else 0
        for scope in ["合并","母公司"]:
            s=w[scope+KINDS[kind]+period]
            selected=[r for r in reports if r.kind==kind and (scope=="合并" or COMPANIES[r.company][2])]
            if not selected:
                s["A1"]=KINDS[kind];s["A3"]="未取得该范围的有效来源；本表不计算为零"
                s.column_dimensions["A"].width=75
                continue
            for block,r in enumerate(selected+[None]):
                start=block*width
                s.merge_cells(start_row=1,start_column=start+1,end_row=1,end_column=start+width)
                s.cell(1,start+1,KINDS[kind]+("（部分范围）" if len(selected)<sum(1 for c in COMPANIES.values() if scope=="合并" or c[2]) else "")).font=Font(name="宋体",size=18,bold=True)
                s.cell(1,start+1).alignment=Alignment(horizontal="center")
                s.merge_cells(start_row=2,start_column=start+1,end_row=2,end_column=start+width)
                s.cell(2,start+1,COMPANIES[r.company][1] if r else scope+"汇总（可用来源）")
                s.cell(2,start+1).alignment=Alignment(horizontal="left",wrap_text=True)
                s.cell(3,start+1,period[:4]+"年"+period[4:]+"月");s.cell(3,start+width,"单位：元")
                hdr=["资产","行次","年初数","期末数","负债及所有者权益","行次","年初数","期末数"] if kind=="bs" else ["项目","行次","本月数","本年累计"]
                for c,text in enumerate(hdr,1):s.cell(4+delta,start+c,text)
                for item in catalog["rows"]:
                    row=item["row"]+delta;col=start+item["col"]+1
                    s.cell(row,col,item["label"])
                    if not item["id"]:continue
                    item_id=item["id"];s.cell(row,col+1,item_id)
                    for metric in range(2):
                        dest=col+2+(1-metric if kind=="bs" else metric)
                        target=s.cell(row,dest)
                        if kind=="is" and item_id in {33,41,43}:continue
                        if r:
                            value=r.values[item_id,metric]
                            if value.state=="semantic_unknown":
                                target.value=None
                                continue
                            expected=value.amount
                            if item_id in TOTAL_RULES[kind]:
                                terms=TOTAL_RULES[kind][item_id]
                                if any(r.values[abs(i),metric].state=="semantic_unknown" for i in terms):
                                    raise ValueError("合计的来源明细口径不明确，不能生成可靠公式")
                                by_id={x["id"]:x for x in catalog["rows"] if x["id"]}
                                parts=[]
                                for component in terms:
                                    source=by_id[abs(component)]
                                    address="$"+get_column_letter(start+source["col"]+3+(1-metric if kind=="bs" else metric))+"$"+str(source["row"]+delta)
                                    parts.append(("+" if component>0 else "-")+address)
                                target.value="="+"".join(parts).lstrip("+")
                                cache[s.title,target.coordinate]=expected
                            else:
                                target.value=expected
                        else:
                            if kind=="is" and item_id in {44,45}:continue
                            if any(x.values[item_id,metric].state=="semantic_unknown" for x in selected):continue
                            refs=[f"${get_column_letter(i*width+item['col']+3+(1-metric if kind=='bs' else metric))}${row}" for i in range(len(selected))]
                            target.value="="+"+".join(refs)
                            expected=sum((x.values[item_id,metric].amount for x in selected),ZERO)
                            cache[s.title,target.coordinate]=expected
                        validation.append({"sheet":s.title,"cell":target.coordinate,"value":str(expected),"company":r.company if r else scope,"kind":kind,"item":item_id,"metric":metric})
                for offset in range(width):
                    s.column_dimensions[get_column_letter(start+offset+1)].width=(42 if kind=="bs" else 60) if offset%4==0 else 7 if offset%4==1 else 30
            maxrow=catalog["height"]+delta
            for row in s.iter_rows(min_row=4+delta,max_row=maxrow,max_col=(len(selected)+1)*width):
                for cell in row:
                    cell.font=FONT;cell.border=BORDER
                    cell.alignment=Alignment(vertical="center",wrap_text=cell.column%4==1)
                    if cell.column%4 in {0,3}:cell.number_format='#,##0.00;-#,##0.00;–'
            s.row_dimensions[2].height=34
            for row in range(5+delta,maxrow+1):s.row_dimensions[row].height=30
            s.freeze_panes="C"+str(5+delta)
            s.sheet_view.zoomScale=70
            s.sheet_properties.pageSetUpPr.fitToPage=True
            s.page_setup.orientation="landscape";s.page_setup.paperSize=s.PAPERSIZE_A3
            s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=1
            s.print_title_rows=f"1:{4+delta}"
            s.print_area=f"A1:{get_column_letter((len(selected)+1)*width)}{maxrow}"
    # Match cash and balance only for the same company; metric 1 means January opening only.
    cross_checks=[]
    for company in COMPANIES:
        bs=next((r for r in reports if r.company==company and r.kind=="bs"),None)
        cf=next((r for r in reports if r.company==company and r.kind=="cf"),None)
        if bs and cf:
            diff=bs.values[1,0].amount-cf.values[35,0].amount
            cross_checks.append({"company":company,"difference":str(diff),"passed":abs(diff)<Decimal(".01")})
            notes.append(["货币资金与期末现金勾稽",COMPANIES[company][0]+" / 差额 "+str(diff)])
    apply_workpaper_format(w)
    w.calculation.calcMode="auto"
    w.calculation.fullCalcOnLoad=True
    w.calculation.forceFullCalc=True
    verify_formulas(w,cache)
    save_cached_workbook(w,path,cache)
    readback=openpyxl.load_workbook(path,data_only=True,read_only=False)
    try:
        for item in validation:
            expected=Decimal(item["value"])
            actual=readback[item["sheet"]][item["cell"]].value
            if actual is None or abs(Decimal(str(actual))-expected)>=Decimal(".005"):
                raise ValueError("工作簿数值或汇总缓存核对失败")
    finally:readback.close()
    complete=coverage_complete and all(x["passed"] for x in cross_checks) and not input_issues and all(not r.issues and all(c["passed"] for c in validate(r)) for r in reports)
    return {"compilation_notes":notes,"cross_checks":cross_checks,"complete":complete,"coverage_complete":coverage_complete,"coverage":matrix,"report_count":len(reports),"formula_count":len(cache),"values":validation,"rules_version":"sd-management-to-cost-v1","template_version":VERSION}



def save_cached_workbook(w, path, cache):
    buf=BytesIO();w.save(buf)
    result=BytesIO()
    with zipfile.ZipFile(buf) as zin,zipfile.ZipFile(result,"w",zipfile.ZIP_DEFLATED) as zout:
        for part in zin.infolist():
            data=zin.read(part.filename)
            if part.filename.startswith("xl/worksheets/sheet") and part.filename.endswith(".xml"):
                index=int(part.filename.rsplit("sheet",1)[1].split(".")[0])-1
                sheet_name=w.sheetnames[index]
                tree=ET.fromstring(data)
                for c in tree.iter(N+"c"):
                    key=(sheet_name,c.attrib["r"])
                    if key in cache:
                        node=c.find(N+"v")
                        if node is None:node=ET.SubElement(c,N+"v")
                        node.text=str(cache[key])
                data=ET.tostring(tree,encoding="utf-8",xml_declaration=True)
            zout.writestr(part,data)
    Path(path).write_bytes(result.getvalue())


def split_workbook(path):
    """Export each scope from the finished workpaper, retaining formulas and caches."""
    path=Path(path)
    reference=openpyxl.load_workbook(path,data_only=True)
    outputs=[]
    try:
        for scope in ["合并","母公司"]:
            book=openpyxl.load_workbook(path)
            try:
                for sheet in list(book):
                    if not sheet.title.startswith(scope):
                        book.remove(sheet)
                if len(book.sheetnames)!=3:
                    raise ValueError("拆分范围必须包含三张报表")
                cache={(s.title,c.coordinate):Decimal(str(reference[s.title][c.coordinate].value))
                       for s in book for row in s for c in row if c.data_type=="f"}
                verify_formulas(book,cache)
                target=path.with_name(path.name.replace("合并报表底稿",scope+"报表",1))
                if target==path:
                    raise ValueError("拆分文件名不能覆盖完整底稿")
                save_cached_workbook(book,target,cache)
                check=openpyxl.load_workbook(target,data_only=True)
                try:
                    for sheet in book:
                        for row in sheet:
                            for cell in row:
                                actual=check[sheet.title][cell.coordinate].value
                                expected=reference[sheet.title][cell.coordinate].value
                                if actual!=expected:
                                    raise ValueError("拆分结果与完整底稿不一致")
                finally:
                    check.close()
                outputs.append(target)
            finally:
                book.close()
    finally:
        reference.close()
    return outputs


def verify_formulas(workbook, expected):
    """Independently evaluate the generated references against the company amount cells."""
    import re
    from openpyxl.utils.cell import coordinate_to_tuple
    token=re.compile(r"([+-]?)(?:'((?:[^']|'')+)'!)?(\$[A-Z]+\$[1-9][0-9]*)")
    memo={};active=set()
    def value(sheet,address):
        key=(sheet,address.replace('$',''))
        if key in memo:return memo[key]
        if key in active:raise ValueError("公式循环引用")
        active.add(key)
        raw=workbook[key[0]][key[1]].value
        if raw is None or raw=="":result=ZERO
        elif isinstance(raw,str) and raw.startswith('='):
            body=raw[1:]
            if body=='0':result=ZERO
            else:
                result=ZERO;end=0
                for match in token.finditer(body):
                    if match.start()!=end:raise ValueError("不支持的公式表达式")
                    sign,name,cell=match.groups()
                    result+=(-1 if sign=='-' else 1)*value(name.replace("''", "'") if name else sheet,cell)
                    end=match.end()
                if end!=len(body) or not end:raise ValueError("公式无法独立求值")
        else:result=Decimal(str(raw).replace(',',''))
        active.remove(key);memo[key]=result
        return result
    for (sheet,address),amount in expected.items():
        if value(sheet,address)!=amount:raise ValueError("公式引用与标准金额不一致")
