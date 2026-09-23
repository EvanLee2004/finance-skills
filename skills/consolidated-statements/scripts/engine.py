"""Deterministic statement parsing. No network, database access or inferred periods."""
from __future__ import annotations
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
import hashlib
import json
import re
import unicodedata
import zipfile

CATALOG = json.loads((Path(__file__).parents[1] / "templates/catalog.json").read_text("utf-8"))
VERSION = CATALOG["version"]
COMPANIES = {
    "HEAD": ("北京本部", "甲骨易（北京）语言科技股份有限公司", True),
    "CULTURE": ("北京文化传媒", "北京甲骨易文化传媒有限公司", False),
    "SHANGHAI": ("上海智译", "甲骨易智译（上海）科技有限公司", False),
    "SHANDONG": ("山东分公司", "甲骨易（北京）语言科技股份有限公司山东分公司", True),
    "HUNAN_TECH": ("湖南科技", "甲骨易（湖南）科技有限公司", False),
    "HUNAN_BRANCH": ("湖南分公司", "甲骨易（北京）语言科技股份有限公司湖南分公司", True),
    "SICHUAN": ("四川分公司", "甲骨易（北京）语言科技股份有限公司四川分公司", True),
    "JINAN": ("济南子公司", "甲骨易（济南）科技有限公司", False),
}
KINDS = {"bs": "资产负债表", "is": "利润表", "cf": "现金流量表"}
SMALL = {
 "is": {1:1,2:2,3:3,4:11,5:14,7:18,8:19,13:20,19:21,20:22,21:24,22:30,23:31,24:32},
 "bs": {1:1,4:3,5:4,7:5,8:8,9:9,13:14,14:15,18:17,22:20,23:21,24:24,27:25,28:26,30:27,32:28,33:29,34:30,35:31,38:32,39:33,40:34,42:35,43:36,44:39,47:40,48:41,49:42,54:43,56:44,58:45,59:46,60:47,61:48,65:49,69:50,70:51,71:52,72:53},
 "cf": {1:1,3:2,5:3,6:4,7:5,8:6,10:7,11:8,12:9,13:10,17:12,18:11,22:13,23:15,24:14,27:16,28:[17,18],31:19,33:20,34:21,35:22},
}
CF_TOTALS = {4:[1,2,3],9:[5,6,7,8],16:[11,12,13,14,15],21:[17,18,19,20],26:[23,24,25],30:[27,28,29]}
ZERO = Decimal("0")
CENT = Decimal(".01")

class SourceError(ValueError):
    pass

def norm(v):
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", str(v or ""))).replace("：", ":")

def label_norm(v):
    label = norm(v).replace("“", '"').replace("”", '"').replace("－", "-")
    label = label.replace("流动负债合计:", "流动负债合计")
    aliases = {
        "教育费附加、矿产资源补偿费、排污费": "教育费附加、矿产资源补偿税、排污费",
        "其中:商品维护费": "其中:商品维修费",
        '加:投资收益(亏损以"-"号填列)': '加:投资收益(损失以"-"号填列)',
        '四、净利润(净亏损以"-"号填列)': '四:净利润(净亏损以"-"号填列)',
    }
    return aliases.get(label, label)


def amount(v):
    if v is None or v == "":
        return ZERO, "blank"
    if isinstance(v, bool):
        raise SourceError("金额中包含布尔值")
    try:
        n = Decimal(str(v).replace(",", "").strip())
    except InvalidOperation:
        raise SourceError("金额无法解析") from None
    if not n.is_finite() or abs(n) >= Decimal("1e20"):
        raise SourceError("金额超出允许范围")
    if n != n.quantize(CENT):
        raise SourceError("金额超过两位小数，需要确认单位或精度")
    return n, "reported"

def line_id(v):
    try:
        d = Decimal(str(v).strip())
        return int(d) if d.is_finite() and d == int(d) else None
    except (ValueError, InvalidOperation):
        return None

@dataclass
class Value:
    amount: Decimal
    state: str
    refs: list[tuple[int, int, int]] = field(default_factory=list)

@dataclass
class Report:
    company: str
    period: str
    kind: str
    file_name: str
    file_hash: str
    sheet: str
    rows: list[list]
    values: dict[tuple[int,int], Value]
    issues: list[str]
    template: str
    rule_version: str = "sd-management-to-cost-v1"
    source_file_id: str = ""
    metric_basis: str = "explicit_headers"

def worksheets(path):
    path = Path(path)
    if path.stat().st_size > 30 * 1024 * 1024:
        raise SourceError("文件超过30MB")
    if path.suffix.lower() == ".xls":
        import xlrd
        w = xlrd.open_workbook(path)
        try:
            for s in w.sheets():
                if s.nrows > 500 or s.ncols > 30:
                    raise SourceError("报表范围超过500行或30列")
                for r in range(s.nrows):
                    for c in range(s.ncols):
                        if s.cell_type(r,c) == xlrd.XL_CELL_ERROR:
                            raise SourceError("来源包含Excel错误")
                yield s.name, [s.row_values(r) for r in range(s.nrows)]
        finally:
            w.release_resources()
    elif path.suffix.lower() == ".xlsx":
        import openpyxl
        with zipfile.ZipFile(path) as z:
            if sum(i.file_size for i in z.infolist()) > 50 * 1024 * 1024:
                raise SourceError("工作簿解压大小超过限制")
            if any(n.startswith("xl/externalLinks/") for n in z.namelist()):
                raise SourceError("来源含外部链接，请提供已计算的独立导出表")
        wf = openpyxl.load_workbook(path, data_only=False, read_only=True)
        w = openpyxl.load_workbook(path, data_only=True, read_only=True)
        try:
            for s in w:
                if s.max_row > 500 or s.max_column > 30:
                    raise SourceError("报表范围超过500行或30列")
                rows = list(s.values)
                for r in wf[s.title]:
                    for cell in r:
                        if cell.data_type == "e":
                            raise SourceError("来源包含Excel错误")
                        if cell.data_type == "f":
                            rr,cc=cell.row-1,cell.column-1
                            if rows[rr][cc] is None:
                                raise SourceError("来源公式没有计算缓存")
                yield s.title, [list(r) for r in rows]
        finally:
            w.close();wf.close()
    else:
        raise SourceError("仅支持xls和xlsx")

def identify(rows, period):
    if not re.fullmatch(r"20\d{2}(0[1-9]|1[0-2])", period):
        raise SourceError("会计月份格式应为YYYYMM")
    header = "\n".join(str(v or "") for row in rows[:5] for v in row)
    kinds = [k for k,v in KINDS.items() if v in header]
    if len(kinds) != 1:
        raise SourceError("不能从表内识别报表类型")
    text = norm(header)
    # Longer branch names win over the head-office substring.
    found = [c for c,(_,legal,_) in COMPANIES.items() if norm(legal) in text]
    if len(found)>1 and "HEAD" in found:
        found.remove("HEAD")
    if len(found)!=1:
        raise SourceError("不能从编制单位唯一识别公司")
    dates = {(int(y),int(m)) for y,m in re.findall(r"(20\d{2})年\s*(\d{1,2})(?:月|期|账期)", header)}
    dates.update((int(y),int(m)) for y,m in re.findall(r"(20\d{2})-(\d{2})(?:-\d{2})?(?!\d)",header))
    if dates != {(int(period[:4]),int(period[4:]))}:
        raise SourceError("表内会计月份缺失、冲突或不符")
    if "单位:元" not in text and "单位:人民币元" not in text:
        raise SourceError("未确认金额单位为元")
    return found[0], kinds[0]

def match_fields(positions, rows, kind, standard, small, header):
    """Resolve known labels to canonical source IDs; unknown amounts fail closed."""
    candidates = []
    required = {"bs": ({34, 72}, {30, 53}), "is": ({24}, {32}), "cf": ({35}, {22})}
    for index, (template, fields) in enumerate([("standard", standard), ("small", small)]):
        names = {}
        for item, label in fields.items():
            names.setdefault(label_norm(label), []).append(item)
        resolved = {}
        unknown = []
        ambiguous = False
        for source_id, position in positions.items():
            r, offset, label = position
            hits = names.get(label_norm(label), [])
            if len(hits) > 1:
                hits = [source_id] if source_id in hits else []
            if len(hits) == 1:
                if hits[0] in resolved:
                    ambiguous = True
                    break
                resolved[hits[0]] = position
            else:
                columns = range(offset + 2, min(len(rows[r]), offset + 4)) if kind == "bs" else range(2, min(len(rows[r]), len(header)))
                if source_id in fields or any(amount(rows[r][col])[0] for col in columns):
                    unknown.append(label.strip())
        if not ambiguous and not unknown and len(resolved) >= 3 and required[kind][index] <= resolved.keys():
            candidates.append((len(resolved), template, resolved))
    if not candidates:
        raise SourceError("字段签名未识别，存在未确认项目或缺少合计项目")
    candidates.sort(key=lambda candidate: candidate[0], reverse=True)
    if len(candidates) > 1 and candidates[0][0] == candidates[1][0]:
        raise SourceError("字段同时匹配多种模板，需核实")
    _, template, resolved = candidates[0]
    return template, resolved


def parse_sheet(rows, sheet, period, file_name, file_hash, *, bs_reclassified=False):
    company,kind=identify(rows,period)
    header_row=next((r for r,row in enumerate(rows[:6]) if any(norm(v)=="行次" for v in row)),None)
    if header_row is None:
        raise SourceError("未找到行次表头")
    header=rows[header_row]
    positions={}
    for r,row in enumerate(rows[header_row+1:], header_row+1):
        if kind=="cf" and norm(row[0])=="补充资料":
            break
        for offset in ([0,4] if kind=="bs" else [0]):
            if offset+1 >= len(row):
                continue
            n=line_id(row[offset+1])
            if n is not None:
                if n in positions:
                    # Standard small cash-flow supplementary schedules repeat line ids.
                    if kind=="cf" and len(positions)>=22:
                        continue
                    raise SourceError("来源行次重复")
                positions[n]=(r,offset,str(row[offset] or ""))
    standard={v["id"]:v["label"] for v in CATALOG["standard"][kind]["rows"] if v["id"]}
    small={int(k):v for k,v in CATALOG["small"][kind].items()}
    template, positions = match_fields(positions, rows, kind, standard, small, header)
    mapping = {i:i for i in standard} if template == "standard" else SMALL[kind]
    metric_basis="explicit_headers"
    if kind=="bs":
        if norm(header[2]) not in {"期末余额","期末数"} or norm(header[3]) not in {"年初余额","年初数"}:
            raise SourceError("资产负债表金额列口径未确认")
        columns=[2,3]
    else:
        columns=[]
        for metric, keywords in enumerate([("本月金额","本月数"),("本年累计金额","本年累计数","本年累计","本年数")]):
            hits=[i for i,v in enumerate(header) if norm(v) in keywords]
            if len(hits)!=1:
                known_cashflow=(kind=="cf" and template=="standard" and
                    company in {"HEAD","CULTURE","SHANGHAI"} and
                    file_name==COMPANIES[company][1]+"_现金流量表__"+period+"期.xlsx" and
                    [norm(v) for v in header] in [
                        ["项目","行次","本月金额","本期金额"],
                        ["项目","行次","本月金额","本期金额","上期金额"],
                    ])
                if metric==1 and known_cashflow:
                    hits=[3]
                    metric_basis="kingdee-standard-cashflow-ytd-v1"
                elif metric==1 and period.endswith("01") and columns:
                    hits=[columns[0]]
                    metric_basis="january-month-equals-ytd"
                else:
                    raise SourceError("本月或本年累计列不明确")
            columns.append(hits[0])
    issues=[]
    if kind=="bs" and not bs_reclassified:
        issues.append("重分类及不重分类应交税费的来源口径待确认")
    values={}
    for item in standard:
        source_ids=mapping.get(item,[])
        source_ids=source_ids if isinstance(source_ids,list) else [source_ids]
        for metric,col in enumerate(columns):
            refs=[];total=ZERO;states=[]
            for sid in source_ids:
                if sid not in positions:continue
                r,offset,_=positions[sid];c=col+offset if kind=="bs" else col
                value,state=amount(rows[r][c]);total+=value;states.append(state)
                refs.append((r+1,c+1,1))
            values[item,metric]=Value(total,"reported" if "reported" in states else "blank" if states else "missing_zero",refs)
    used={sid for v in mapping.values() for sid in (v if isinstance(v,list) else [v])}
    # Unmapped nonzero detail is never silently discarded, even for familiar templates.
    for sid,(r,offset,label) in positions.items():
        if sid in used:continue
        if any(amount(rows[r][col+offset if kind=="bs" else col])[0] for col in columns):
            issues.append("非零来源明细未映射："+label.strip())
    if template=="small" and kind=="cf":
        for item,parts in CF_TOTALS.items():
            for metric in range(2):
                vs=[values[p,metric] for p in parts]
                values[item,metric]=Value(sum((v.amount for v in vs),ZERO),"derived",sum((v.refs for v in vs),[]))
    if kind=="is":
        for metric in range(2):
            if template=="small":
                values[42,metric]=Value(values[24,metric].amount+values[27,metric].amount,"derived",values[24,metric].refs+values[27,metric].refs)
            if values[25,metric].state!="reported":
                issues.append("持续经营净利润未填，不能解释为经济意义上的零")
            # 小企业准则利润表「其中：利息费用（收入以"-"号填列）」是**净利息**，
            # 而「财务费用」= 净利息 + 手续费等，两者本就不相等（实例：济南 −179.97 = −377.16 + 197.19），
            # 属正常口径而非异常。历史实现曾在此把该格归零并挂 issue，会丢失金蝶原值、
            # 并因 issue 使 complete 永远无法达成，故按"以金蝶为准"保留来源原值。
            # （不再归零：保留来源原值，仅由来源自带状态决定）
    report=Report(company,period,kind,file_name,file_hash,sheet,rows,values,sorted(set(issues)),template)
    report.metric_basis=metric_basis
    return report

def classify(report, *, enabled=True):
    import copy
    target=copy.deepcopy(report)
    target.rule_version="sd-management-to-cost-v1" if enabled else "none"
    if enabled and target.company=="SHANDONG" and target.kind=="is":
        for metric in range(2):
            a,b=target.values[2,metric],target.values[5,metric]
            target.values[2,metric]=Value(a.amount+b.amount,"classified",a.refs+b.refs)
            target.values[5,metric]=Value(ZERO,"classified",b.refs+[(r,c,-sign) for r,c,sign in b.refs])
    return target

def parse_file(path, period, *, file_name=None, bs_reclassified=False):
    path=Path(path)
    digest=hashlib.sha256(path.read_bytes()).hexdigest()
    reports=[];issues=[]
    for sheet,rows in worksheets(path):
        try:
            reports.append(parse_sheet(rows,sheet,period,file_name or path.name,digest,bs_reclassified=bs_reclassified))
        except SourceError as exc:
            issue={"file":file_name or path.name,"sheet":sheet,"message":str(exc),"status":"rejected"}
            try:
                company,kind=identify(rows,period)
                issue.update(company=company,kind=kind)
            except SourceError:
                pass
            issues.append(issue)
    return reports,issues

def select_sources(reports):
    """Exact normalized duplicates collapse; conflicting revisions require an explicit choice."""
    groups={}
    for r in reports:groups.setdefault((r.company,r.period,r.kind),[]).append(r)
    chosen=[];issues=[]
    for key,group in groups.items():
        signatures={(json.dumps(r.rows,ensure_ascii=False,sort_keys=True,default=str),r.template,tuple(r.issues),tuple((i,m,str(v.amount),v.state) for (i,m),v in sorted(r.values.items()))) for r in group}
        if len(signatures)>1:
            issues.append({"company":key[0],"kind":key[2],"message":"同公司同月同表有冲突来源，请仅选择一个权威来源"})
        else:chosen.append(group[0])
    return chosen,issues

def validate(report):
    checks=[]
    v=lambda i,m:report.values[i,m].amount
    def check(name,left,right,m):
        checks.append({"check":name,"metric":m,"difference":str(left-right),"passed":abs(left-right)<CENT})
    for m in range(2):
        if report.kind=="bs":
            check("资产等于负债和权益",v(34,m),v(72,m),m)
            check("资产合计",v(14,m)+v(33,m),v(34,m),m)
            check("负债和权益合计",v(60,m)+v(71,m),v(72,m),m)
        elif report.kind=="is":
            check("营业利润",v(1,m)-sum(v(i,m) for i in [2,3,4,5,6,7,10,11])+sum(v(i,m) for i in [12,13,15,16,17,18]),v(19,m),m)
            check("利润总额",v(19,m)+v(20,m)-v(21,m),v(22,m),m)
            check("净利润",v(22,m)-v(23,m),v(24,m),m)
        else:
            check("现金期末余额",v(34,m)+v(33,m),v(35,m),m)
            check("现金净增加额",v(10,m)+v(22,m)+v(31,m)+v(32,m),v(33,m),m)
    return checks
