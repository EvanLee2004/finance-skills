import copy
import json
import sys
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path
import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from engine import Report, Value, CATALOG, COMPANIES, KINDS
import workbook
from eliminations import amounts, apply_eliminations, validate_evidence, verify_formulas
from entry import run

PERIOD = "202608"

def fixture():
    reports = []
    for company in COMPANIES:
        for kind in KINDS:
            values = {(item["id"], metric): Value(Decimal(0), "reported", [])
                      for item in CATALOG["standard"][kind]["rows"] if item["id"] for metric in (0, 1)}
            if company == "HEAD" and kind == "bs":
                for item in (8, 44, 18, 61):
                    values[item, 0] = Value(Decimal(100), "reported", [])
            if kind == "cf":
                for item in (3, 8):
                    for metric in (0, 1):
                        values[item, metric] = Value(Decimal(100), "reported", [])
            for _ in range(5):
                for item, terms in workbook.TOTAL_RULES[kind].items():
                    for metric in (0, 1):
                        values[item, metric] = Value(sum(((1 if t > 0 else -1) * values[abs(t), metric].amount for t in terms), Decimal(0)), "reported", [])
            reports.append(Report(company, PERIOD, kind, "synthetic.xlsx", "0"*64, kind, [], values, [], "standard"))
    companies = {}
    for company in COMPANIES:
        if company == "HEAD":
            continue
        eligible = company == "CULTURE"
        companies[company] = {
            "bs": {"status": "observed", "ending_debit": "1.00", "refs": ["synthetic#bs"]},
            "is": {"status": "observed", "accounts_complete": True, "refs": ["synthetic#is"], "accounts": [
                {"account": "113311", "has_receipt_month": eligible, "has_receipt_ytd": eligible,
                 "receipt_month": "53.00" if eligible else "0.00", "credit_ytd": "106.00",
                 "refs": ["synthetic#account"]}]},
            "cf": {"status": "observed", "month_in": "7.00", "month_out": "10.00",
                   "ytd_in": "15.00", "ytd_out": "20.00", "refs": ["synthetic#cf"]}
        }
    return reports, dict(period=PERIOD, ledger="HEAD", collected_at="synthetic-test", companies=companies)

def cell(book, kind, scope, item, metric):
    sheet = book[scope+KINDS[kind]+PERIOD]
    row = next(x for x in CATALOG["standard"][kind]["rows"] if x["id"] == item)
    count = 8 if scope == "合并" else 4
    column = count * CATALOG["standard"][kind]["width"] + row["col"] + 3 + (1-metric if kind == "bs" else metric)
    return sheet.cell(row["row"] + int(kind == "is"), column)

class EliminationTests(unittest.TestCase):
    def test_company_receipt_eligibility_and_tax(self):
        _, evidence = fixture()
        self.assertEqual(amounts(evidence, ["CULTURE", "SHANDONG"], "is")[0], (Decimal(50), Decimal(100)))

    def test_year_credit_includes_adjustment(self):
        _, e = fixture()
        e["companies"]["CULTURE"]["is"]["accounts"][0]["credit_ytd"] = "212.00"
        self.assertEqual(amounts(e, ["CULTURE"], "is")[0][1], Decimal(200))

    def test_round_per_company(self):
        _, e = fixture()
        e["companies"]["SHANGHAI"]["is"] = copy.deepcopy(e["companies"]["CULTURE"]["is"])
        for c in ("CULTURE", "SHANGHAI"):
            e["companies"][c]["is"]["accounts"][0]["credit_ytd"] = "0.01"
        self.assertEqual(amounts(e, ["CULTURE","SHANGHAI"], "is")[0][1], Decimal(".02"))

    def test_duplicate_cumulative_account_rejected(self):
        _, e = fixture()
        e["companies"]["CULTURE"]["is"]["accounts"] *= 2
        with self.assertRaises(ValueError):
            validate_evidence(e, PERIOD)

    def test_wrong_period_and_missing_evidence(self):
        _, e = fixture()
        with self.assertRaises(ValueError):
            validate_evidence(e, "202607")
        e["companies"]["CULTURE"]["bs"]["status"] = "no_data"
        self.assertEqual(amounts(e, ["CULTURE"], "bs"), (None, ["CULTURE"]))

    def test_signed_bs_and_outflow_only_cash(self):
        _, e = fixture()
        e["companies"]["CULTURE"]["bs"]["ending_debit"] = "-1.00"
        self.assertEqual(amounts(e, ["CULTURE"], "bs")[0], Decimal(-1))
        self.assertEqual(amounts(e, ["CULTURE"], "cf")[0], [(Decimal(10), Decimal(10)), (Decimal(20), Decimal(20))])

    def test_workbook_scope_formulas_and_split(self):
        reports, e = fixture()
        old = workbook.verify_formulas
        workbook.verify_formulas = verify_formulas
        try:
            with tempfile.TemporaryDirectory() as td:
                p = Path(td) / "甲骨易合并报表底稿.xlsx"
                base = workbook.build_workbook(reports, PERIOD, p)
                before = openpyxl.load_workbook(p, data_only=False)
                original_parent_profit = [[c.value for c in row] for row in before["母公司利润表"+PERIOD]]
                result = apply_eliminations(p, reports, PERIOD, e, base)
                after = openpyxl.load_workbook(p, data_only=True)
                formulas = openpyxl.load_workbook(p, data_only=False)
                self.assertEqual(len(after.sheetnames), 6)
                self.assertEqual(cell(after, "bs", "合并", 8, 0).value, 93)
                self.assertEqual(cell(after, "bs", "母公司", 8, 0).value, 97)
                self.assertEqual(cell(after, "bs", "合并", 18, 0).value, 0)
                self.assertEqual(cell(after, "bs", "母公司", 18, 0).value, 100)
                self.assertEqual(cell(after, "bs", "合并", 34, 0).value, cell(after, "bs", "合并", 72, 0).value)
                self.assertEqual(cell(after, "is", "合并", 1, 1).value, -100)
                self.assertEqual(cell(after, "is", "合并", 4, 1).value, -100)
                self.assertEqual(cell(after, "is", "母公司", 1, 1).value, 0)
                self.assertEqual(cell(after, "cf", "合并", 3, 0).value, 730)
                self.assertEqual(cell(after, "cf", "合并", 8, 0).value, 730)
                self.assertEqual(cell(after, "cf", "母公司", 3, 0).value, 370)
                self.assertTrue(all(x["passed"] for x in result["elimination_checks"]))
                # Source region and amounts are preserved; source totals remain formulas.
                for kind in KINDS:
                    for scope in ("合并", "母公司"):
                        count = 8 if scope == "合并" else 4
                        width = count * CATALOG["standard"][kind]["width"]
                        a,b = before[scope+KINDS[kind]+PERIOD], formulas[scope+KINDS[kind]+PERIOD]
                        for row in a.iter_rows(max_col=width):
                            for c in row:
                                self.assertEqual(c.value, b[c.coordinate].value)
                for p2 in workbook.split_workbook(p):
                    split = openpyxl.load_workbook(p2, data_only=True)
                    self.assertEqual(len(split.sheetnames), 3)
                    for s in split:
                        for row in s:
                            for c in row:
                                self.assertEqual(c.value, after[s.title][c.coordinate].value)
                    split.close()
                before.close();after.close();formulas.close()
        finally:
            workbook.verify_formulas = old

    def test_missing_cash_marks_incomplete(self):
        reports, e = fixture()
        e["companies"]["CULTURE"].pop("cf")
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "甲骨易合并报表底稿.xlsx"
            base = workbook.build_workbook(reports, PERIOD, path)
            result = apply_eliminations(path, reports, PERIOD, e, base)
            self.assertFalse(result["complete"])
            self.assertFalse(result["elimination_complete"])
            self.assertTrue(any(x["kind"] == "cf" for x in result["elimination_issues"]))

    def test_output_reuse_rejected(self):
        _, e = fixture()
        with tempfile.TemporaryDirectory() as td:
            Path(td,"prior.txt").write_text("prior")
            with self.assertRaises(ValueError):
                run(dict(parameters=dict(period=PERIOD), output_dir=td), e)


    def test_integrated_entry_and_source_hash(self):
        from unittest.mock import patch
        reports, e = fixture()
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            original = root/"raw.xlsx"
            original.write_bytes(b"synthetic parser fixture")
            output = root/"result"
            request = dict(parameters=dict(period=PERIOD, fetch_kingdee=False), output_dir=str(output),
                           files=dict(reports=[dict(local_path=str(original),name="raw.xlsx",file_id="test",sha256="")]))
            with patch("base_entry.parse_file", return_value=(reports, [])):
                result = run(request, e)
            self.assertEqual(original.read_bytes(), b"synthetic parser fixture")
            self.assertEqual(result["summary"]["report_count"], 24)
            self.assertTrue(result["summary"]["elimination_complete"])
            self.assertEqual(len(list(output.glob("*.xlsx"))), 3)
            audit = json.loads((output/"报表来源与核验记录.json").read_text("utf-8"))
            self.assertEqual(audit["rules_version"], "jgy-eliminations-20260917.1")
            self.assertTrue(audit["eliminations"])
            self.assertFalse(any("未抵销" in str(x) for x in audit["compilation_notes"]))

if __name__ == "__main__":
    unittest.main()
