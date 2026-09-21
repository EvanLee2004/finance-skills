#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""供应商付款入金蝶 · 合成回归。真表不进仓。"""
from __future__ import annotations

import json
import sys
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
import importlib.util

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
SCRIPTS = SKILL / "scripts"


def _load(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod


convert = _load("kingdee_payment_convert", SCRIPTS / "convert.py")
inspect_inputs = _load("kingdee_payment_inspect", SCRIPTS / "inspect_inputs.py")
parse_invoice = _load("kingdee_payment_parse", SCRIPTS / "parse_invoice.py")
kingdee_live = _load("kingdee_payment_live", SCRIPTS / "kingdee_live.py")

YELLOW = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")


def _master():
    return {
        "employee": [{"code": "113", "name": "项目总监"}],
        "department": [{"code": "0405", "name": "项目总监及助理"}],
        "supplier": [{"code": "8001", "name": "北京某翻译店"}, {"code": "9999", "name": "其他供应商"}],
    }


def _write_pay(path: Path, rows, yellow_rows=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "付款"
    ws.append(["供应商", "应付金额本币", "开户名"])
    yellow_rows = set(yellow_rows or [])
    for i, r in enumerate(rows, start=2):
        ws.append(r)
        if i in yellow_rows:
            ws.cell(i, 1).fill = YELLOW
            ws.cell(i, 2).fill = YELLOW
    wb.save(path)
    wb.close()


def _dummy_pdf(path: Path):
    path.write_bytes(b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n")


def _run(tmp_path, monkeypatch, parse_fn, book_short_pay=False, master=None, start=1, citic_as_other=False, boc_as_other=False, one_voucher_per_payee=False):
    monkeypatch.setattr(convert, "parse_invoice_pdf", parse_fn)
    return convert.run_dir(
        tmp_path,
        "2026-09-17",
        master or _master(),
        start_voucher_no=start,
        out_dir=tmp_path / "out",
        book_short_pay=book_short_pay,
        citic_as_other=citic_as_other,
        boc_as_other=boc_as_other,
        one_voucher_per_payee=one_voucher_per_payee,
    )


def test_inspect_needs_pdfs(tmp_path):
    _write_pay(tmp_path / "付款.xlsx", [["北京某翻译店", 100, "北京某翻译店"]])
    report = inspect_inputs.inspect_dir(tmp_path)
    assert report["ready"] is False
    assert any("PDF" in m for m in report["missing"])


def test_inspect_unwraps_one_wrapper_folder(tmp_path):
    _write_pay(tmp_path / "付款.xlsx", [["北京某翻译店", 100, "北京某翻译店"]])
    wrap = tmp_path / "个人转对公"
    d = wrap / "北京某翻译店"
    d.mkdir(parents=True)
    _dummy_pdf(d / "a.pdf")
    report = inspect_inputs.inspect_dir(tmp_path)
    assert report["ready"] is True


def test_inspect_ready(tmp_path):
    _write_pay(tmp_path / "付款.xlsx", [["北京某翻译店", 100, "北京某翻译店"]])
    d = tmp_path / "北京某翻译店"
    d.mkdir()
    _dummy_pdf(d / "a.pdf")
    report = inspect_inputs.inspect_dir(tmp_path)
    assert report["ready"] is True


def test_inspect_ignores_xlsx_in_folder(tmp_path):
    _write_pay(tmp_path / "付款.xlsx", [["北京某翻译店", 100, "北京某翻译店"]])
    d = tmp_path / "北京某翻译店"
    d.mkdir()
    _dummy_pdf(d / "a.pdf")
    (d / "别人的表.xlsx").write_bytes(b"not-a-real-xlsx")
    report = inspect_inputs.inspect_dir(tmp_path)
    assert report["ready"] is True


def test_parse_spaced_special_and_tax():
    text = "电 子 发 票 （ 增 值税专 用发 票 ）\n合 计 ¥59405.94 ¥594.06\n价税合计（大写）陆万圆整 （小写）¥60000.00\n销售方名称：甲店"
    got = parse_invoice.parse_invoice_text(text)
    assert got["kind"] == "专票"
    assert got["total"] == Decimal("60000.00")
    assert got["tax"] == Decimal("594.06")
    assert got["seller"] == "甲店"


def test_parse_normal_invoice():
    text = "电子发票（普通发票）\n价税合计（小写）¥2719.92\n销售方名称：乙店"
    got = parse_invoice.parse_invoice_text(text)
    assert got["kind"] == "普票"
    assert got["total"] == Decimal("2719.92")


def test_yellow_uses_boc_white_uses_citic(tmp_path, monkeypatch):
    _write_pay(
        tmp_path / "付款.xlsx",
        [["北京某翻译店", 106, "北京某翻译店"], ["无档店", 200, "无档店"]],
        yellow_rows={2},
    )
    (tmp_path / "北京某翻译店").mkdir()
    (tmp_path / "无档店").mkdir()
    _dummy_pdf(tmp_path / "北京某翻译店" / "a.pdf")
    _dummy_pdf(tmp_path / "无档店" / "b.pdf")

    def fake_parse(path: Path):
        if path.parent.name == "北京某翻译店":
            return {"kind": "专票", "seller": "北京某翻译店", "total": Decimal("106.00"), "tax": Decimal("6.00")}
        return {"kind": "普票", "seller": "无档店", "total": Decimal("200.00"), "tax": None}

    result = _run(tmp_path, monkeypatch, fake_parse, citic_as_other=True)
    assert result["bookable_count"] == 2
    kd = load_workbook(result["kingdee_path"])
    ws = kd[convert.KINGDEE_SHEET]
    acc_col = convert.col_by_label(ws, "*科目.编码")
    emp_col = convert.col_by_label(ws, "辅助核算.职员.编码")
    accounts = [ws.cell(r, acc_col).value for r in range(4, 10)]
    assert "100201" in accounts
    assert "100206" in accounts
    assert "21710101" in accounts
    emp = [ws.cell(r, emp_col).value for r in range(4, 10)]
    assert "113" in emp
    kd.close()


def test_white_missing_supplier_holds_then_9999(tmp_path, monkeypatch):
    _write_pay(tmp_path / "付款.xlsx", [["无档店", 200, "无档店"]])
    (tmp_path / "无档店").mkdir()
    _dummy_pdf(tmp_path / "无档店" / "a.pdf")
    fake = lambda p: {"kind": "普票", "seller": "无档店", "total": Decimal("200.00"), "tax": None}
    hold = _run(tmp_path, monkeypatch, fake)
    assert hold["bookable_count"] == 0
    assert hold["new_supplier_names_citic"] == ["无档店"]
    assert hold["new_supplier_names_boc"] == []
    ok = _run(tmp_path, monkeypatch, fake, citic_as_other=True)
    kd = load_workbook(ok["kingdee_path"])
    ws = kd[convert.KINGDEE_SHEET]
    col = convert.col_by_label(ws, "辅助核算.供应商.编码")
    codes = [ws.cell(r, col).value for r in range(4, 8)]
    assert "9999" in codes
    kd.close()


def test_yellow_missing_supplier_asks(tmp_path, monkeypatch):
    _write_pay(tmp_path / "付款.xlsx", [["新黄店", 100, "新黄店"]], yellow_rows={2})
    (tmp_path / "新黄店").mkdir()
    _dummy_pdf(tmp_path / "新黄店" / "a.pdf")
    result = _run(
        tmp_path,
        monkeypatch,
        lambda p: {"kind": "普票", "seller": "新黄店", "total": Decimal("100.00"), "tax": None},
    )
    assert result["bookable_count"] == 0
    assert result["new_supplier_names_boc"] == ["新黄店"]
    assert result["new_supplier_names_citic"] == []


def test_ticket_greater_asks_once_then_flag_books(tmp_path, monkeypatch):
    _write_pay(tmp_path / "付款.xlsx", [["北京某翻译店", 100, "北京某翻译店"]])
    (tmp_path / "北京某翻译店").mkdir()
    _dummy_pdf(tmp_path / "北京某翻译店" / "a.pdf")
    fake = lambda p: {"kind": "专票", "seller": "北京某翻译店", "total": Decimal("106.00"), "tax": Decimal("6.00")}
    hold = _run(tmp_path, monkeypatch, fake, book_short_pay=False)
    assert hold["short_pay_count"] == 1
    assert hold["bookable_count"] == 0
    ok = _run(tmp_path, monkeypatch, fake, book_short_pay=True)
    assert ok["bookable_count"] == 1
    kd = load_workbook(ok["kingdee_path"])
    ws = kd[convert.KINGDEE_SHEET]
    dcol = convert.col_by_label(ws, "借方 #")
    ccol = convert.col_by_label(ws, "贷方 #")
    debits = [ws.cell(r, dcol).value for r in range(4, 7)]
    credits = [ws.cell(r, ccol).value for r in range(4, 7)]
    assert 94.0 in debits
    assert 6.0 in debits
    assert 100.0 in credits
    kd.close()


def test_ticket_less_holds(tmp_path, monkeypatch):
    _write_pay(tmp_path / "付款.xlsx", [["北京某翻译店", 200, "北京某翻译店"]])
    (tmp_path / "北京某翻译店").mkdir()
    _dummy_pdf(tmp_path / "北京某翻译店" / "a.pdf")
    result = _run(
        tmp_path,
        monkeypatch,
        lambda p: {"kind": "普票", "seller": "北京某翻译店", "total": Decimal("100.00"), "tax": None},
    )
    assert result["hold_count"] == 1


def test_many_invoices_sum(tmp_path, monkeypatch):
    _write_pay(tmp_path / "付款.xlsx", [["北京某翻译店", 300, "北京某翻译店"]])
    d = tmp_path / "北京某翻译店"
    d.mkdir()
    _dummy_pdf(d / "a.pdf")
    _dummy_pdf(d / "b.pdf")
    amounts = [Decimal("100.00"), Decimal("200.00")]

    def fake(path: Path):
        return {"kind": "普票", "seller": "北京某翻译店", "total": amounts.pop(0) if amounts else Decimal("0"), "tax": None}

    result = _run(tmp_path, monkeypatch, fake)
    assert result["bookable_count"] == 1


def test_next_supplier_number_skips_9999():
    assert kingdee_live.next_supplier_number([{"code": "8001"}, {"code": "9999"}]) == "8002"


def test_pick_banks_requires_names():
    rules = json.loads((SKILL / "config" / "rules.json").read_text(encoding="utf-8"))
    acc = [
        {"code": "100201", "name": "中行宣武门支行1001"},
        {"code": "100206", "name": "中信银行广安门支行"},
    ]
    assert kingdee_live.pick_payment_banks(acc, rules) == ("100201", "100206")


def test_cli_asks_short_pay(tmp_path, monkeypatch, capsys):
    _write_pay(tmp_path / "付款.xlsx", [["北京某翻译店", 100, "北京某翻译店"]])
    (tmp_path / "北京某翻译店").mkdir()
    _dummy_pdf(tmp_path / "北京某翻译店" / "a.pdf")
    master = tmp_path / "master.json"
    master.write_text(json.dumps(_master(), ensure_ascii=False), encoding="utf-8")
    monkeypatch.setattr(
        convert,
        "parse_invoice_pdf",
        lambda p: {"kind": "专票", "seller": "北京某翻译店", "total": Decimal("106.00"), "tax": Decimal("6.00")},
    )
    rc = convert.main(
        [
            "--input-dir",
            str(tmp_path),
            "--master",
            str(master),
            "--date",
            "2026-09-17",
            "--start-voucher-no",
            "1",
            "--out-dir",
            str(tmp_path / "out"),
            "--skip-browser",
        ]
    )
    assert rc == 2
    out = capsys.readouterr().out
    assert "ask=" in out
    assert "大于应付" in out


def test_cli_asks_boc_and_citic_together(tmp_path, monkeypatch, capsys):
    _write_pay(
        tmp_path / "付款.xlsx",
        [["新黄店", 100, "新黄店"], ["无档店", 200, "无档店"]],
        yellow_rows={2},
    )
    (tmp_path / "新黄店").mkdir()
    (tmp_path / "无档店").mkdir()
    _dummy_pdf(tmp_path / "新黄店" / "a.pdf")
    _dummy_pdf(tmp_path / "无档店" / "b.pdf")
    master = tmp_path / "master.json"
    master.write_text(json.dumps(_master(), ensure_ascii=False), encoding="utf-8")
    monkeypatch.setattr(
        convert,
        "parse_invoice_pdf",
        lambda p: {"kind": "普票", "seller": p.parent.name, "total": Decimal("100.00") if p.parent.name == "新黄店" else Decimal("200.00"), "tax": None},
    )
    rc = convert.main(
        [
            "--input-dir",
            str(tmp_path),
            "--master",
            str(master),
            "--date",
            "2026-09-17",
            "--start-voucher-no",
            "1",
            "--out-dir",
            str(tmp_path / "out"),
            "--skip-browser",
        ]
    )
    assert rc == 2
    out = capsys.readouterr().out
    assert "标黄（中行）" in out
    assert "没标黄（中信）" in out
    assert "新黄店" in out
    assert "无档店" in out
    assert "一次确认" in out


# ---------- 2026-09-21 返工：摘要取开票抬头（销方）+ 全批一张凭证 ----------

def test_seller_gou_xiao_same_line():
    """2026 版电子发票「购/销 双栏同行」：摘要必须取销方（开票抬头），不许取我方购买方。"""
    text = (
        "电子发票（增值税专用发票）\n"
        "发票号码：00000000000000000000\n"
        "开票日期：2026年09月07日\n"
        "购 名称：某（北京）语言科技股份有限公司 销 名称：某翻译店（个体工商户）\n"
        "买 售\n方 方\n"
        "信 统一社会信用代码/纳税人识别号：91110000000000000X 信 统一社会信用代码/纳税人识别号：91120000000000000Y\n"
        "息 息\n"
        "*生产生活服务*翻译费 无 次 1 6897.0099009901 6897.01 1% 68.97\n"
        "合 计 ¥6897.01 ¥68.97\n"
        "价税合计（大写） 陆仟玖佰陆拾伍圆玖角捌分 （小写）¥6965.98\n"
        "开票人：某甲\n"
    )
    got = parse_invoice.parse_invoice_text(text)
    assert got["seller"] == "某翻译店（个体工商户）"
    assert "语言科技" not in got["seller"]
    assert "名称" not in got["seller"]
    assert got["kind"] == "专票"
    assert got["total"] == Decimal("6965.98")
    assert got["tax"] == Decimal("68.97")


def test_seller_with_spaced_buyer_name():
    text = "购 名称：某 （北京）语言科技股份有限公司 销 名称：某翻译工作室\n买 售\n方 方\n"
    assert parse_invoice.parse_invoice_text(text)["seller"] == "某翻译工作室"


def test_seller_despace_broken_name():
    text = "购 名称：某公司 销 名称：某翻译服务有 限公司\n"
    assert parse_invoice.parse_invoice_text(text)["seller"] == "某翻译服务有限公司"


def test_seller_name_outside_table():
    """名称被排版挤到表格外、与买家名同一行：左买右销，取最右。"""
    text = (
        "电子发票（普通发票）\n购 销\n买 名称： 售 名称：\n方 方\n"
        "信 信 统一社会信用代码/纳税人识别号： 统一社会信用代码/纳税人识别号：\n息 息\n"
        "某公司 某信息技术有限公司\n9111 9222\n"
    )
    assert parse_invoice.parse_invoice_text(text)["seller"] == "某信息技术有限公司"


def test_seller_legacy_standalone_field():
    text = "电子发票（普通发票）\n价税合计（小写）¥2719.92\n销售方名称：乙店"
    assert parse_invoice.parse_invoice_text(text)["seller"] == "乙店"


def test_seller_empty_when_no_counterparty():
    """销方抽不到（只剩购买方）→ 留空，由 convert 判待确认；不许拿买方名顶替。"""
    text = "购 名称：某（北京）语言科技股份有限公司 销 名称：\n买 售\n方 方\n"
    got = parse_invoice.parse_invoice_text(text)
    assert got["seller"] == ""


def test_empty_seller_holds_the_row(tmp_path, monkeypatch):
    _write_pay(tmp_path / "付款.xlsx", [["北京某翻译店", 100, "北京某翻译店"]])
    (tmp_path / "北京某翻译店").mkdir()
    _dummy_pdf(tmp_path / "北京某翻译店" / "a.pdf")
    fake = lambda p: {"kind": "普票", "seller": "", "total": Decimal("100.00"), "tax": None}
    result = _run(tmp_path, monkeypatch, fake)
    assert result["bookable_count"] == 0
    assert result["hold_count"] == 1


def _three_bookable(tmp_path, monkeypatch, **kw):
    names = ["甲店", "乙店", "丙店"]
    _write_pay(tmp_path / "付款.xlsx", [[n, (i + 1) * 100, n] for i, n in enumerate(names)])
    for i, n in enumerate(names):
        (tmp_path / n).mkdir()
        _dummy_pdf(tmp_path / n / f"{i}.pdf")
    amt = {n: Decimal(str((i + 1) * 100)) for i, n in enumerate(names)}

    def fake(path: Path):
        n = path.parent.name
        return {"kind": "普票", "seller": n, "total": amt[n], "tax": None}

    return _run(tmp_path, monkeypatch, fake, citic_as_other=True, **kw)


def _read_kingdee(result, label):
    kd = load_workbook(result["kingdee_path"])
    ws = kd[convert.KINGDEE_SHEET]
    col = convert.col_by_label(ws, label)
    vals = [ws.cell(r, col).value for r in range(4, 60)]
    kd.close()
    return [v for v in vals if v not in (None, "")]


def test_payment_one_voucher_by_default(tmp_path, monkeypatch):
    """斯佳 2026-09-21：本批全部合一张凭证。"""
    result = _three_bookable(tmp_path, monkeypatch, start=123)
    assert result["bookable_count"] == 3
    assert result["voucher_count"] == 1
    assert set(_read_kingdee(result, "凭证号 #")) == {123}


def test_payment_one_voucher_per_payee_flag(tmp_path, monkeypatch):
    result = _three_bookable(tmp_path, monkeypatch, start=123, one_voucher_per_payee=True)
    assert result["voucher_count"] == 3
    assert set(_read_kingdee(result, "凭证号 #")) == {123, 124, 125}


def test_payment_one_voucher_balances_and_expl(tmp_path, monkeypatch):
    result = _three_bookable(tmp_path, monkeypatch, start=123)
    debits = _read_kingdee(result, "借方 #")
    credits = _read_kingdee(result, "贷方 #")
    assert abs(sum(debits) - sum(credits)) < 0.005
    assert abs(sum(debits) - 600.0) < 0.005
    expl = _read_kingdee(result, "摘要 #")
    assert set(expl) == {"付：甲店", "付：乙店", "付：丙店"}
    assert not any("语言科技" in str(e) for e in expl)


def test_note_reports_voucher_count(tmp_path, monkeypatch):
    result = _three_bookable(tmp_path, monkeypatch, start=123)
    note = Path(result["note_path"]).read_text(encoding="utf-8")
    assert "凭证张数：1" in note


def test_seller_joined_across_lines():
    """销方名太长被排版折行：名字断在两行，必须接回来。"""
    text = "购 名称：某公司 销 名称：某翻译服务有\n限公司\n"
    assert parse_invoice.parse_invoice_text(text)["seller"] == "某翻译服务有限公司"


def test_seller_does_not_join_unrelated_next_line():
    """下一行是买家/销售方栏的断字（「方 方」「买 售」）时，不许当续字拼进摘要。"""
    text = "购 名称：某公司 销 名称：某翻译服务有\n方 方\n"
    got = parse_invoice.parse_invoice_text(text)["seller"]
    assert "方" not in got
