#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""供应商付款 → 金蝶凭证引入表。金额只由本脚本算。"""
from __future__ import annotations

import argparse
import json
import os
import shutil
import sys
from collections import Counter
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook, load_workbook


def repo_venv_python() -> Path | None:
    root = Path(__file__).resolve().parents[3]
    for rel in (Path(".venv") / "bin" / "python", Path(".venv") / "Scripts" / "python.exe"):
        cand = root / rel
        if cand.is_file():
            return cand
    return None


def _reexec_repo_venv_if_needed() -> None:
    if not sys.argv or Path(sys.argv[0]).name.lower() not in {"convert.py", "convert"}:
        return
    venv_py = repo_venv_python()
    if venv_py is None:
        return
    try:
        if Path(sys.prefix).resolve() == venv_py.parent.parent.resolve():
            return
    except OSError:
        return
    os.execv(str(venv_py), [str(venv_py), *sys.argv])


_reexec_repo_venv_if_needed()

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
CONFIG = SKILL / "config"
TEMPLATE = CONFIG / "凭证引入空模.xlsx"
KINGDEE_SHEET = "sheet1（名称勿改）"
TWOPLACES = Decimal("0.01")
POSTING_SCRIPTS = HERE.parent.parent / "kingdee-posting" / "scripts"
if str(HERE) not in sys.path:
    sys.path.insert(0, str(HERE))


def _load_mod(name: str, path: Path):
    import importlib.util

    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod


inspect_mod = _load_mod("kingdee_payment_inspect_inputs", HERE / "inspect_inputs.py")
kingdee_api = _load_mod("kingdee_payment_kingdee_api", POSTING_SCRIPTS / "kingdee_api.py")
names = _load_mod("kingdee_payment_match_name", POSTING_SCRIPTS / "match_name.py")
kingdee_live = _load_mod("kingdee_payment_live", HERE / "kingdee_live.py")
parse_invoice = _load_mod("kingdee_payment_parse_invoice", HERE / "parse_invoice.py")
NO_PDF_READER_ASK = parse_invoice.NO_PDF_READER_ASK
money = parse_invoice.money
parse_invoice_pdf = parse_invoice.parse_invoice_pdf

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass


def log(msg: str) -> None:
    print(msg, file=sys.stderr)


def ask_and_stop(msg: str) -> int:
    log(msg)
    print(f"ask={msg}", flush=True)
    return 2


def load_rules() -> dict:
    return json.loads((CONFIG / "rules.json").read_text(encoding="utf-8"))


def load_aliases() -> dict:
    return inspect_mod.load_aliases()


def code_str(v) -> str:
    if v is None or v == "":
        return ""
    if isinstance(v, bool):
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, int):
        return str(v)
    s = str(v).strip()
    if s.endswith(".0") and s[:-2].isdigit():
        return s[:-2]
    return s


def norm_name(name: str) -> str:
    s = (name or "").strip().replace("(", "（").replace(")", "）")
    return "".join(s.split())


def strip_ge(name: str) -> str:
    s = (name or "").replace("（个体工商户）", "").replace("(个体工商户)", "")
    return s.strip()


def cell_at(row, idx, field):
    i = idx.get(field)
    if i is None or i >= len(row):
        return None
    return row[i]


def header_index(headers, aliases) -> dict[str, int]:
    pay_a = aliases.get("付款_列别名") or {}
    out = {}
    cleaned = [str(h).strip() if h is not None else "" for h in headers]
    for field in ("供应商", "应付金额本币", "开户名"):
        names_a = pay_a.get(field) or [field]
        for i, h in enumerate(cleaned):
            if h in names_a:
                out[field] = i
                break
    return out


def pick_amount(formula_cell, value_cell, fallback=None):
    got = money(value_cell)
    if got is not None:
        return got
    got = money(formula_cell)
    if got is not None:
        return got
    return fallback


def is_yellow_fill(cell) -> bool:
    fill = getattr(cell, "fill", None)
    if fill is None or fill.fill_type in (None, "none"):
        return False
    fg = getattr(fill, "fgColor", None)
    if fg is None:
        return False
    rgb = str(getattr(fg, "rgb", "") or "").upper()
    return "FFFF00" in rgb


def row_is_yellow(ws, row: int, col_count: int) -> bool:
    for c in range(1, min(col_count, 3) + 1):
        if is_yellow_fill(ws.cell(row, c)):
            return True
    return False


@dataclass
class VoucherLine:
    status: str
    reason: str = ""
    source_row: int = 0
    key: str = ""
    expl: str = ""
    voucher_no: int | None = None
    booking_date: str = ""
    extra: dict = field(default_factory=dict)
    entries: list = field(default_factory=list)


class Master:
    def __init__(self, data: dict | None):
        data = data or {}
        self.emp: dict[str, list[tuple[str, str]]] = {}
        self.dept: dict[str, str] = {}
        self.sup: dict[str, list[tuple[str, str]]] = {}
        self.supplier_rows = list(data.get("supplier") or [])
        for e in data.get("employee") or []:
            name = str(e.get("name") or "").strip()
            code = code_str(e.get("code"))
            if name and code:
                self.emp.setdefault(norm_name(name), []).append((code, name))
        for d in data.get("department") or []:
            code = code_str(d.get("code"))
            name = str(d.get("name") or "").strip()
            if code:
                self.dept[code] = name
        for s in data.get("supplier") or []:
            name = str(s.get("name") or "").strip()
            code = code_str(s.get("code"))
            if name and code:
                self.sup.setdefault(norm_name(name), []).append((code, name))
                stripped = strip_ge(name)
                if stripped != name:
                    self.sup.setdefault(norm_name(stripped), []).append((code, name))

    def all_suppliers(self) -> list[tuple[str, str]]:
        out = []
        for pairs in self.sup.values():
            for item in pairs:
                if item not in out:
                    out.append(item)
        return out

    def employee(self, name: str):
        hits = self.emp.get(norm_name(name)) or []
        if len(hits) == 1:
            return hits[0], None
        if not hits:
            return None, "none"
        return None, "many"

    def department(self, code: str):
        code = code_str(code)
        if not code:
            return None, "缺部门编码"
        name = self.dept.get(code)
        if not name:
            return None, "部门档案没有该编码"
        return (code, name), None

    def supplier_fuzzy(self, name: str):
        got = names.match_records(name, self.all_suppliers())
        if got.status == "ok":
            return got.hit, None
        stripped = strip_ge(name)
        if stripped != name:
            got = names.match_records(stripped, self.all_suppliers())
            if got.status == "ok":
                return got.hit, None
        if got.status == "many":
            return None, "many"
        return None, "none"


def find_vendor_folder(vendor: str, dirs: dict[str, Path]) -> tuple[Path | None, str]:
    folder = dirs.get(vendor)
    if folder is not None:
        return folder, ""
    m = names.match_records(vendor, [(name, name) for name in dirs])
    if m.status == "ok" and m.hit:
        return dirs[m.hit[0]], ""
    if m.status == "many":
        return None, "发票夹名有多个像的，认不出是哪一个"
    return None, "缺这家发票夹"


def find_desktop(root: Path) -> Path | None:
    for rel in ("Desktop", "桌面", "OneDrive/Desktop", "OneDrive/桌面"):
        cand = root / rel
        if cand.is_dir():
            return cand
    return None


def default_desktop_dir(prefix: str = "金蝶入账_付款", today: date | None = None, home: Path | None = None) -> Path:
    day = (today or date.today()).strftime("%Y%m%d")
    root = Path(home) if home else Path.home()
    desktop = find_desktop(root)
    base = desktop if desktop else Path.cwd()
    path = base / f"{prefix}_{day}"
    path.mkdir(parents=True, exist_ok=True)
    return path


def col_by_label(ws, needle: str) -> int:
    for cell in ws[3]:
        if needle in str(cell.value or ""):
            return cell.column
    raise KeyError(needle)


def _write_aux(ws, row: int, col: int, value) -> None:
    if value is None or value == "":
        return
    text = str(value).strip()
    if not text or text.startswith("="):
        return
    ws.cell(row, col, value)


def write_kingdee(path: Path, lines: list[VoucherLine], rules: dict, booking: str, template: Path) -> Path:
    shutil.copy2(template, path)
    wb = load_workbook(path)
    if KINGDEE_SHEET not in wb.sheetnames:
        wb.close()
        raise SystemExit("金蝶空模缺 sheet1（名称勿改）")
    ws = wb[KINGDEE_SHEET]
    cols = {
        "date": col_by_label(ws, "*记账日期"),
        "word": col_by_label(ws, "*凭证字号"),
        "number": col_by_label(ws, "凭证号 #"),
        "expl": col_by_label(ws, "摘要 #"),
        "account": col_by_label(ws, "*科目.编码"),
        "account_name": col_by_label(ws, "科目.名称"),
        "currency": col_by_label(ws, "*币别.编码"),
        "currency_name": col_by_label(ws, "币别.名称"),
        "rate": col_by_label(ws, "*汇率"),
        "amountfor": col_by_label(ws, "原币金额"),
        "debit": col_by_label(ws, "借方 #"),
        "credit": col_by_label(ws, "贷方 #"),
        "cus_code": col_by_label(ws, "辅助核算.客户.编码"),
        "cus_name": col_by_label(ws, "辅助核算.客户.名称"),
        "dep_code": col_by_label(ws, "辅助核算.部门.编码"),
        "dep_name": col_by_label(ws, "辅助核算.部门.名称"),
        "emp_code": col_by_label(ws, "辅助核算.职员.编码"),
        "emp_name": col_by_label(ws, "辅助核算.职员.名称"),
        "sup_code": col_by_label(ws, "辅助核算.供应商.编码"),
        "sup_name": col_by_label(ws, "辅助核算.供应商.名称"),
    }
    acc_names = rules.get("account_names") or {}
    row_i = 4
    for line in lines:
        if line.status != "可入账":
            continue
        for ent in line.entries:
            ws.cell(row_i, cols["date"], line.booking_date or booking)
            ws.cell(row_i, cols["word"], rules.get("voucher_word") or "记")
            ws.cell(row_i, cols["number"], line.voucher_no)
            ws.cell(row_i, cols["expl"], line.expl)
            ws.cell(row_i, cols["account"], ent["account"])
            acc_name = acc_names.get(str(ent["account"]), "")
            if acc_name:
                ws.cell(row_i, cols["account_name"], acc_name)
            ws.cell(row_i, cols["currency"], rules.get("currency") or "RMB")
            if rules.get("currency_name"):
                ws.cell(row_i, cols["currency_name"], rules.get("currency_name"))
            ws.cell(row_i, cols["rate"], float(rules.get("exchange_rate") or 1))
            debit = ent.get("debit")
            credit = ent.get("credit")
            amt = debit if debit is not None else credit
            if amt is not None:
                ws.cell(row_i, cols["amountfor"], float(amt))
            if debit is not None:
                ws.cell(row_i, cols["debit"], float(debit))
            if credit is not None:
                ws.cell(row_i, cols["credit"], float(credit))
            if ent.get("aux"):
                _write_aux(ws, row_i, cols["dep_code"], ent.get("dep_code"))
                _write_aux(ws, row_i, cols["dep_name"], ent.get("dep_name"))
                _write_aux(ws, row_i, cols["emp_code"], ent.get("emp_code"))
                _write_aux(ws, row_i, cols["emp_name"], ent.get("emp_name"))
                _write_aux(ws, row_i, cols["sup_code"], ent.get("sup_code"))
                _write_aux(ws, row_i, cols["sup_name"], ent.get("sup_name"))
            row_i += 1
    wb.save(path)
    wb.close()
    return path


def write_detail(path: Path, lines: list[VoucherLine]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "明细"
    ws.append(["状态", "原因", "源行号", "摘要", "凭证批次", "附加"])
    for line in lines:
        extra = json.dumps(line.extra, ensure_ascii=False) if line.extra else ""
        ws.append([line.status, line.reason, line.source_row, line.expl, line.voucher_no, extra])
    wb.save(path)
    wb.close()
    return path


def reason_counts(lines) -> list[tuple[str, int]]:
    c = Counter()
    for line in lines or []:
        if getattr(line, "status", "") != "可入账" and str(getattr(line, "reason", "") or "").strip():
            c[str(line.reason).strip()] += 1
    return sorted(c.items(), key=lambda x: (-x[1], x[0]))


def write_note(path: Path, *, source_count, bookable_count, hold_count, sheet, start_voucher_no, voucher_count, reasons, extras):
    lines = [
        "# 付款入金蝶对照说明",
        "",
        f"- 源：{source_count}",
        f"- 可入账：{bookable_count}",
        f"- 待确认：{hold_count}",
        f"- 读了哪个 sheet：{sheet or '（未记）'}",
        f"- 起始凭证号：{start_voucher_no if start_voucher_no else '（未取）'}",
        f"- 凭证张数：{voucher_count}",
        "- 原因类别：",
    ]
    if reasons:
        for reason, n in reasons:
            lines.append(f"  - {reason}：{n}")
    else:
        lines.append("  - （无）")
    for extra in extras or []:
        lines.append(f"- {extra}")
    lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def convert_payment(
    root: Path,
    ledger: Path,
    master: Master,
    rules: dict,
    aliases: dict,
    *,
    book_short_pay: bool = False,
    boc: str = "100201",
    citic: str = "100206",
    boc_as_other: bool = False,
    citic_as_other: bool = False,
    one_voucher_per_payee: bool = False,
) -> list[VoucherLine]:
    cfg = rules.get("payment") or {}
    fallback = rules.get("fallback_supplier") or {"code": "9999", "name": "其他供应商"}
    wb_f = load_workbook(ledger, data_only=False)
    wb_v = load_workbook(ledger, data_only=True)
    sheet = inspect_mod.find_payment_sheet(wb_f, aliases)
    if not sheet:
        wb_f.close()
        wb_v.close()
        raise SystemExit("找不到付款三列表")
    ws_f = wb_f[sheet]
    ws_v = wb_v[sheet] if sheet in wb_v.sheetnames else wb_v[wb_v.sheetnames[0]]
    pay_alias = aliases.get("付款_列别名") or {}
    header_r, headers = inspect_mod.find_header_row(ws_f, pay_alias, ["供应商", "应付金额本币"])
    idx = header_index(headers, aliases)
    if "供应商" not in idx or "应付金额本币" not in idx:
        wb_f.close()
        wb_v.close()
        raise SystemExit("付款表缺供应商或应付金额本币")
    lines: list[VoucherLine] = []
    max_row = ws_f.max_row or 1
    dirs = {p.name: p for p in inspect_mod.payment_folders(root, ledger)}
    for r in range(header_r + 1, max_row + 1):
        row_f = [ws_f.cell(r, c + 1).value for c in range(len(headers))]
        row_v = [ws_v.cell(r, c + 1).value for c in range(len(headers))]
        vendor = str(cell_at(row_f, idx, "供应商") or "").strip()
        if not vendor:
            continue
        payable = pick_amount(cell_at(row_f, idx, "应付金额本币"), cell_at(row_v, idx, "应付金额本币"), None)
        yellow = row_is_yellow(ws_f, r, len(headers))
        line = VoucherLine(
            status="可入账",
            source_row=r,
            key=vendor,
            extra={
                "供应商": vendor,
                "sheet": sheet,
                "source_amt": str(payable) if payable is not None else "",
                "yellow": yellow,
            },
        )
        folder, folder_reason = find_vendor_folder(vendor, dirs)
        if folder is None:
            line.status, line.reason = "待确认", folder_reason
            lines.append(line)
            continue
        pdfs = [p for p in folder.iterdir() if p.is_file() and p.suffix.lower() == ".pdf"]
        if not pdfs:
            line.status, line.reason = "待确认", "夹里没有发票 PDF"
            lines.append(line)
            continue
        metas = [parse_invoice_pdf(p) for p in pdfs]
        if any(m.get("error") == "no_pdfplumber" for m in metas):
            wb_f.close()
            wb_v.close()
            raise SystemExit(NO_PDF_READER_ASK)
        if any(m.get("error") == "unreadable" for m in metas):
            line.status, line.reason = "待确认", "发票 PDF 读不出文字（可能是扫描件或坏文件）"
            lines.append(line)
            continue
        if any((not m.get("kind")) or (not str(m.get("seller") or "").strip()) or m.get("total") is None for m in metas):
            line.status, line.reason = "待确认", "发票缺票种或销售方或金额"
            lines.append(line)
            continue
        kinds = {m.get("kind") for m in metas if m.get("kind")}
        if len(kinds) != 1:
            line.status, line.reason = "待确认", "认不清专票还是普票"
            lines.append(line)
            continue
        kind = kinds.pop()
        sellers = [m.get("seller") or "" for m in metas if m.get("seller")]
        seller = sellers[0] if sellers else ""   # 抽不到就留空，绝不静默降级取台账名
        if payable is None:
            line.status, line.reason = "待确认", "缺应付金额"
            lines.append(line)
            continue
        ticket_total = Decimal("0")
        ticket_tax = Decimal("0")
        for m in metas:
            if m.get("total") is not None:
                ticket_total += m["total"]
            if m.get("tax") is not None:
                ticket_tax += m["tax"]
        if ticket_total and ticket_total < payable:
            line.status, line.reason = "待确认", "票合计小于应付"
            lines.append(line)
            continue
        if ticket_total and ticket_total > payable and not book_short_pay:
            line.status, line.reason = "待确认", "票大于应付"
            lines.append(line)
            continue
        hit, why = master.supplier_fuzzy(vendor)
        if why == "none":
            hit, why = master.supplier_fuzzy(strip_ge(vendor))
        if why == "none":
            hit, why = master.supplier_fuzzy(seller)
        if why == "many":
            hit = None
            why = "none"
        if hit:
            sup_code, sup_name = hit
        elif yellow and not boc_as_other:
            line.status, line.reason = "待确认", "无档是否新建（中行标黄）"
            lines.append(line)
            continue
        elif (not yellow) and not citic_as_other:
            line.status, line.reason = "待确认", "无档是否新建（中信）"
            lines.append(line)
            continue
        else:
            fallback_hit, fallback_why = master.supplier_fuzzy(str(fallback.get("name") or ""))
            if not fallback_hit or fallback_why:
                line.status, line.reason = "待确认", "其他供应商档案未核验"
                lines.append(line)
                continue
            sup_code, sup_name = fallback_hit
        dhit, derr = master.department(str(cfg.get("dept_code") or ""))
        if not dhit:
            line.status, line.reason = "待确认", derr or "付款部门档案未核验"
            lines.append(line)
            continue
        dep_code, dep_name = dhit
        ehit, eerr = master.employee(str(cfg.get("emp_name") or ""))
        if not ehit:
            line.status, line.reason = "待确认", "职员档案一对多" if eerr == "many" else "付款职员档案未核验"
            lines.append(line)
            continue
        emp_code, emp_name = ehit
        aux = {
            "aux": True,
            "sup_code": sup_code,
            "sup_name": sup_name,
            "dep_code": dep_code,
            "dep_name": dep_name,
            "emp_code": emp_code,
            "emp_name": emp_name,
        }
        cost_acc = str(cfg.get("cost_account") or "540103")
        bank_acc = boc if yellow else citic
        tax_acc = str(cfg.get("input_tax_account") or "21710101")
        line.expl = f"付：{seller}"
        line.extra["bank"] = bank_acc
        if kind == "普票":
            line.entries = [
                {"account": cost_acc, "debit": payable, **aux},
                {"account": bank_acc, "credit": payable},
            ]
        else:
            tax = ticket_tax if ticket_tax else None
            if tax is None or tax == 0:
                line.status, line.reason = "待确认", "认不清进项税额"
                lines.append(line)
                continue
            if tax >= payable:
                line.status, line.reason = "待确认", "进项不小于应付"
                lines.append(line)
                continue
            cost = (payable - tax).quantize(TWOPLACES, rounding=ROUND_HALF_UP)
            line.entries = [
                {"account": cost_acc, "debit": cost, **aux},
                {"account": tax_acc, "debit": tax},
                {"account": bank_acc, "credit": payable},
            ]
        lines.append(line)
    wb_f.close()
    wb_v.close()
    bookable = [x for x in lines if x.status == "可入账"]
    # 斯佳 2026-09-21：本批全部合一张凭证（同号）。要一家一张加 --one-voucher-per-payee。
    for i, line in enumerate(bookable, start=1):
        line.voucher_no = i if one_voucher_per_payee else 1
    return lines


def names_needing_create(lines, yellow: bool | None = None) -> list[str]:
    out = []
    for line in lines or []:
        if "是否新建" not in str(getattr(line, "reason", "") or ""):
            continue
        extra = getattr(line, "extra", None) or {}
        is_yellow = bool(extra.get("yellow"))
        if yellow is True and not is_yellow:
            continue
        if yellow is False and is_yellow:
            continue
        title = str(extra.get("供应商") or line.key or "").strip()
        if title and title not in out:
            out.append(title)
    return out


def missing_supplier_ask(boc_names: list[str], citic_names: list[str]) -> str:
    bits = ["本批金蝶供应商档案没有的，请一次确认（标黄走中行，没标黄走中信）："]
    if boc_names:
        bits.append(f"标黄（中行）{len(boc_names)} 家：{'、'.join(boc_names)}。要不要新建？")
    else:
        bits.append("标黄（中行）：没有缺档的。")
    if citic_names:
        bits.append(
            f"没标黄（中信）{len(citic_names)} 家：{'、'.join(citic_names)}。"
            "按你上次说的记「其他供应商」9999，还是也要新建？"
        )
    else:
        bits.append("没标黄（中信）：没有缺档的。")
    bits.append("请一次性回：中行建或不建，中信新建或挂9999。")
    return " ".join(bits)


def create_confirmed_suppliers(to_create: list[str], suppliers: list) -> dict:
    created = []
    pool = list(suppliers or [])
    recs = [(str(c.get("code") or ""), str(c.get("name") or "")) for c in pool]
    for title in to_create:
        already = names.match_records(title, recs)
        if already.status == "ok" and already.hit:
            created.append({"number": already.hit[0], "existed": True})
            continue
        number = kingdee_live.next_supplier_number(pool)
        got = kingdee_live.try_create_supplier(title, number)
        if not got.get("ok"):
            return {"ok": False, "error": got.get("error") or "建档失败", "created": created}
        pool.append({"code": number, "name": title})
        recs.append((number, title))
        created.append({"number": number, "existed": False})
    return {"ok": True, "created": created, "suppliers": pool}


def run_dir(
    input_dir: Path,
    booking: str | None,
    master_data: dict | None,
    start_voucher_no: int = 1,
    out_dir: Path | None = None,
    book_short_pay: bool = False,
    boc: str = "100201",
    citic: str = "100206",
    boc_as_other: bool = False,
    citic_as_other: bool = False,
    one_voucher_per_payee: bool = False,
) -> dict:
    report = inspect_mod.inspect_dir(input_dir)
    if not report.get("ready"):
        raise SystemExit(report.get("ask") or "材料不齐")
    rules = load_rules()
    aliases = load_aliases()
    master = Master(master_data or {})
    day = booking or date.today().isoformat()
    src = Path(report["files"]["ledger"])
    extras = [
        "记账日：当前月最后一张凭证的日期（可 --date）",
        "凭证号：本批全部合一张凭证（同号；要一家一张加 --one-voucher-per-payee）",
        "黄=中行，白=中信",
        "夹里非 PDF 忽略",
    ]
    if book_short_pay:
        extras.append("票大于应付：已按应付记")
    if citic_as_other:
        extras.append("没标黄缺档：已按其他供应商 9999")
    if boc_as_other:
        extras.append("标黄缺档：已按其他供应商 9999")
    lines = convert_payment(
        Path(input_dir),
        src,
        master,
        rules,
        aliases,
        book_short_pay=book_short_pay,
        boc=boc,
        citic=citic,
        boc_as_other=boc_as_other,
        citic_as_other=citic_as_other,
        one_voucher_per_payee=one_voucher_per_payee,
    )
    if start_voucher_no and int(start_voucher_no) != 1:
        delta = int(start_voucher_no) - 1
        for line in lines:
            if line.status == "可入账" and line.voucher_no is not None:
                line.voucher_no = int(line.voucher_no) + delta
    dest = Path(out_dir) if out_dir else default_desktop_dir()
    dest.mkdir(parents=True, exist_ok=True)
    for line in lines:
        line.booking_date = line.booking_date or day
    kingdee = dest / "凭证引入_付款_结果.xlsx"
    detail = dest / f"{src.stem}_付款_明细结果.xlsx"
    write_kingdee(kingdee, lines, rules, day, TEMPLATE)
    write_detail(detail, lines)
    bookable = sum(1 for x in lines if x.status == "可入账")
    hold = sum(1 for x in lines if x.status != "可入账")
    sheet = ""
    for line in lines:
        if line.extra.get("sheet"):
            sheet = str(line.extra["sheet"])
            break
    voucher_nos = {int(x.voucher_no) for x in lines if x.status == "可入账" and x.voucher_no is not None}
    write_note(
        dest / "对照说明_付款.md",
        source_count=len(lines),
        bookable_count=bookable,
        hold_count=hold,
        sheet=sheet,
        start_voucher_no=start_voucher_no,
        voucher_count=len(voucher_nos),
        reasons=reason_counts(lines),
        extras=extras,
    )
    short_n = sum(1 for x in lines if x.reason == "票大于应付")
    new_boc = names_needing_create(lines, yellow=True)
    new_citic = names_needing_create(lines, yellow=False)
    new_n = names_needing_create(lines)
    return {
        "scene": "付款",
        "source_count": len(lines),
        "bookable_count": bookable,
        "hold_count": hold,
        "voucher_count": len(voucher_nos),
        "last_voucher_no": max(voucher_nos) if voucher_nos else None,
        "kingdee_path": str(kingdee),
        "detail_path": str(detail),
        "note_path": str(dest / "对照说明_付款.md"),
        "out_dir": str(dest),
        "sheet": sheet,
        "start_voucher_no": start_voucher_no,
        "new_supplier_names": new_n,
        "new_supplier_names_boc": new_boc,
        "new_supplier_names_citic": new_citic,
        "short_pay_count": short_n,
        "boc": boc,
        "citic": citic,
        "one_voucher_per_payee": one_voucher_per_payee,
    }


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="供应商付款入金蝶")
    parser.add_argument("--inspect", action="store_true")
    parser.add_argument("--input-dir", required=True)
    parser.add_argument("--date")
    parser.add_argument("--master")
    parser.add_argument("--no-api", action="store_true")
    parser.add_argument("--skip-browser", action="store_true", help="测试用：不登网页，仍要档案")
    parser.add_argument("--start-voucher-no", type=int, default=None)
    parser.add_argument("--out-dir", "--out", dest="out_dir")
    parser.add_argument("--book-short-pay", action="store_true", help="斯佳点头：票大于应付的一律按应付记")
    parser.add_argument("--create-new-suppliers", action="store_true", help="斯佳点头：标黄（中行）缺档按现网编号新建")
    parser.add_argument("--create-new-suppliers-citic", action="store_true", help="斯佳点头：没标黄（中信）缺档也新建")
    parser.add_argument("--citic-as-other", action="store_true", help="斯佳点头：没标黄缺档记其他供应商 9999")
    parser.add_argument("--boc-as-other", action="store_true", help="斯佳点头：标黄缺档也记其他供应商 9999")
    parser.add_argument("--one-voucher-per-payee", action="store_true", help="一家一张凭证（默认本批全部合一张）")
    args = parser.parse_args(argv)
    root = Path(args.input_dir)
    if args.inspect:
        return inspect_mod.main(["--input-dir", str(root)])

    rules = load_rules()
    boc = str((rules.get("payment") or {}).get("boc_bank_account") or "100201")
    citic = str((rules.get("payment") or {}).get("citic_bank_account") or "100206")
    master_data = None
    if args.master:
        master_data = json.loads(Path(args.master).read_text(encoding="utf-8"))
    elif args.no_api:
        return ask_and_stop("本次入账必须先读取总部当前档案；--no-api 只可用于开发排查，未生成引入表。")
    else:
        if not args.skip_browser:
            try:
                kingdee_live.ensure_hq_session()
            except SystemExit as e:
                return ask_and_stop(str(e) if str(e) else "金蝶网页未登录总部。未生成引入表。")
        try:
            accounts = kingdee_live.fetch_bank_accounts()
            boc, citic = kingdee_live.pick_payment_banks(accounts, rules)
        except SystemExit as e:
            return ask_and_stop(str(e) if str(e) else "银行科目未核验。未生成引入表。")
        loaded = kingdee_api.try_load_master()
        if loaded.get("ok"):
            master_data = loaded["data"]
        else:
            reason = "本机没有金蝶应用号" if loaded.get("missing_credentials") else loaded.get("error") or "读取失败"
            return ask_and_stop(f"总部档案未核验（{reason}）。未生成引入表。请检查本机应用号和只读权限后重试。")

    start_no = args.start_voucher_no
    if start_no is None or not args.date:
        period = (args.date or date.today().isoformat())[:7]
        fetched = kingdee_api.try_fetch_next_voucher_no(period)
        if not fetched.get("ok"):
            reason = "本机没有金蝶应用号" if fetched.get("missing_credentials") else fetched.get("error") or "读取失败"
            return ask_and_stop(f"当前月凭证号未核验（{reason}）。未生成引入表。请检查本机应用号后重试。")
        if start_no is None:
            start_no = int(fetched["next_number"])
        if not args.date:
            args.date = fetched.get("last_date") or date.today().isoformat()
            log(f"记账日取当前月最后一张凭证：{args.date}")

    run_kw = dict(
        start_voucher_no=start_no,
        out_dir=Path(args.out_dir).expanduser() if args.out_dir else None,
        book_short_pay=args.book_short_pay,
        boc=boc,
        citic=citic,
        boc_as_other=args.boc_as_other,
        citic_as_other=args.citic_as_other,
        one_voucher_per_payee=args.one_voucher_per_payee,
    )
    try:
        result = run_dir(root, args.date, master_data, **run_kw)
    except SystemExit as e:
        return ask_and_stop(str(e) if str(e) else "材料不齐，未生成引入表。")

    to_create = []
    if args.create_new_suppliers:
        to_create.extend(result.get("new_supplier_names_boc") or [])
    if args.create_new_suppliers_citic:
        to_create.extend(result.get("new_supplier_names_citic") or [])
    # 去重保序
    seen = set()
    to_create = [n for n in to_create if not (n in seen or seen.add(n))]
    if to_create:
        created = create_confirmed_suppliers(to_create, (master_data or {}).get("supplier") or [])
        if not created.get("ok"):
            return ask_and_stop(f"供应商档案未建成（{created.get('error')}）；未覆盖引入表。")
        master_data = dict(master_data or {})
        master_data["supplier"] = created.get("suppliers") or []
        if not args.master:
            kingdee_api.clear_master_cache()
        log(f"已新建供应商档案 {len(to_create)} 家。")
        result = run_dir(root, args.date, master_data, **run_kw)

    asks = []
    if result.get("short_pay_count"):
        asks.append(
            f"本批有 {result['short_pay_count']} 家发票合计大于应付。"
            "是否一律按应付记（专票：银行=应付、进项=票面税、成本=应付−税；普票：成本=银行=应付）？"
            "请回是或否。是则加 --book-short-pay 重跑。"
        )
    boc_need = result.get("new_supplier_names_boc") or []
    citic_need = result.get("new_supplier_names_citic") or []
    if boc_need or citic_need:
        asks.append(missing_supplier_ask(boc_need, citic_need))

    log(
        f"付款这批 {result['source_count']}：可入账 {result['bookable_count']}，待确认 {result['hold_count']}。"
        f"填好的金蝶表在 {result['kingdee_path']}。请您看待确认，再自己去金蝶引入。我没有点引入。"
    )
    if asks:
        print("ask=" + " ".join(asks), flush=True)
    print(json.dumps({k: v for k, v in result.items() if k != "lines"}, ensure_ascii=False, indent=2))
    return 2 if asks else 0


if __name__ == "__main__":
    raise SystemExit(main())
