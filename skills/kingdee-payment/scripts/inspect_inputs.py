#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""认付款三列表 + 各家发票夹。夹里非 PDF 忽略。"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
CONFIG = SKILL / "config"

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass


def load_aliases() -> dict:
    p = CONFIG / "列名别名.json"
    return json.loads(p.read_text(encoding="utf-8")) if p.is_file() else {}


def _cleaned(headers) -> list[str]:
    return [str(h).strip() for h in (headers or []) if h is not None and str(h).strip()]


def field_hit(headers: list[str], aliases: dict, field: str) -> bool:
    names = aliases.get(field) or [field]
    return any(h in names for h in headers)


def find_header_row(ws, alias_map: dict | None, required: list[str], max_scan: int = 20) -> tuple[int, list]:
    best_i, best, best_score = 1, [], -1
    alias_map = alias_map or {}
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=max_scan, values_only=True), 1):
        headers = list(row or [])
        cleaned = _cleaned(headers)
        score = sum(1 for key in required if field_hit(cleaned, alias_map, key))
        if score > best_score:
            best_score, best_i, best = score, i, headers
        if score >= len(required) and required:
            return i, headers
    return best_i, best


def is_payment_sheet(name: str, headers: list[str], aliases: dict | None = None) -> bool:
    title = str(name or "")
    if "流水" in title or "组织架构" in title:
        return False
    if "中行付款" in title or ("中行" in title and "付款" in title):
        return False
    if "收款" in title:
        return False
    pay_a = (aliases or {}).get("付款_列别名") or {}
    cleaned = _cleaned(headers)
    if any("价税合计" in h for h in cleaned):
        return False
    return field_hit(cleaned, pay_a, "供应商") and field_hit(cleaned, pay_a, "应付金额本币")


def find_payment_sheet(wb, aliases: dict | None = None) -> str | None:
    named = []
    headered = []
    for name in wb.sheetnames:
        ws = wb[name]
        _, raw = find_header_row(ws, (aliases or {}).get("付款_列别名"), ["供应商", "应付金额本币"])
        headers = _cleaned(raw)
        if not is_payment_sheet(name, headers, aliases):
            continue
        if "付款" in str(name) and "中行" not in str(name):
            named.append(name)
        else:
            headered.append(name)
    if named:
        return named[0]
    if headered:
        return headered[0]
    return None


def _pdfs(folder: Path) -> list[Path]:
    return [p for p in folder.iterdir() if p.is_file() and p.suffix.lower() == ".pdf"]


def payment_folders(root: Path, ledger: Path | None) -> list[Path]:
    skip = {ledger.resolve()} if ledger else set()
    out = []
    for p in sorted(root.iterdir()):
        if p.is_dir() and p.resolve() not in skip:
            if _pdfs(p):
                out.append(p)
    if out:
        return out
    wrappers = [p for p in root.iterdir() if p.is_dir() and p.resolve() not in skip]
    if len(wrappers) == 1:
        for p in sorted(wrappers[0].iterdir()):
            if p.is_dir() and _pdfs(p):
                out.append(p)
    return out


def inspect_dir(input_dir: Path) -> dict:
    root = Path(input_dir)
    aliases = load_aliases()
    if not root.is_dir():
        return {
            "ready": False,
            "scene": "付款",
            "missing": ["材料文件夹"],
            "files": {},
            "ask": "请把当批材料放到一个文件夹里，再告诉我路径。",
        }
    ledgers = []
    for p in sorted(root.iterdir()):
        if p.suffix.lower() not in {".xlsx", ".xlsm"} or p.name.startswith("~$"):
            continue
        if "结果" in p.stem or "核算项目余额表" in p.name:
            continue
        try:
            wb = load_workbook(p, read_only=True, data_only=False)
        except Exception:
            continue
        try:
            if find_payment_sheet(wb, aliases):
                ledgers.append(p)
        finally:
            wb.close()
    ledger = ledgers[0] if ledgers else None
    folders = payment_folders(root, ledger)
    missing = []
    asks = []
    if not ledger:
        missing.append("付款三列表（供应商 / 应付金额本币 / 开户名）")
    if not folders:
        missing.append("各家发票夹（夹里要有 PDF）")
    if missing:
        asks.append("付款还缺：" + "；".join(missing) + "。台账放外面，一家一个夹，夹里放发票 PDF。夹里的 Excel 不用管。")
    return {
        "ready": not missing,
        "scene": "付款",
        "missing": missing,
        "files": {
            "ledger": str(ledger) if ledger else None,
            "invoice_dirs": [str(p) for p in folders],
        },
        "ask": " ".join(asks),
    }


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="供应商付款入金蝶 · 盘点文件夹")
    parser.add_argument("--input-dir", required=True)
    args = parser.parse_args(argv)
    report = inspect_dir(Path(args.input_dir))
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if report.get("ask"):
        print(f"ask={report['ask']}", flush=True)
    return 0 if report.get("ready") else 2


if __name__ == "__main__":
    raise SystemExit(main())
