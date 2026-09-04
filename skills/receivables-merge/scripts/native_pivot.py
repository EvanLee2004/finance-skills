"""在 XLSX 中安装可刷新的原生 Excel 数据透视表。

openpyxl 可以保留已有的透视表，但不能创建透视表。本模块先用
openpyxl 写好透视表的初始显示，再补齐 OOXML 的 pivot cache / pivot
table 部件。这样 Python Worker 不需要安装 Excel，也不会把真实工作簿
作为模板带进 Skill 包。
"""

from __future__ import annotations

import math
import numbers
import os
import posixpath
import re
import shutil
import tempfile
import zipfile
from collections import OrderedDict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.styles import Border, Font, PatternFill, Side

try:
    import pandas as pd
except ImportError:  # pragma: no cover - merge.py already requires pandas
    pd = None


MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"

PIVOT_CACHE_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotCacheDefinition"
PIVOT_CACHE_RECORDS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotCacheRecords"
PIVOT_TABLE_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotTable"

PIVOT_CACHE_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.pivotCacheDefinition+xml"
PIVOT_CACHE_RECORDS_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.pivotCacheRecords+xml"
PIVOT_TABLE_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.pivotTable+xml"

PIVOT_NAME = "透视汇总"
STATIC_BACKUP_NAME = "透视汇总_静态备份"
SOURCE_NAME = "ReceivablesPivotSource"
ACCT_FMT = '_-* #,##0.00_-;-* #,##0.00_-;_-* "-"??_-;_-@_-'
ILLEGAL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

ET.register_namespace("", MAIN_NS)
ET.register_namespace("r", REL_NS)
ET.register_namespace("pr", PKG_REL_NS)
ET.register_namespace("ct", CONTENT_TYPES_NS)


def _q(tag: str) -> str:
    return f"{{{MAIN_NS}}}{tag}"


def _r(tag: str) -> str:
    return f"{{{REL_NS}}}{tag}"


def _pkg(tag: str) -> str:
    return f"{{{PKG_REL_NS}}}{tag}"


def _ct(tag: str) -> str:
    return f"{{{CONTENT_TYPES_NS}}}{tag}"


def _xml_bytes(root: ET.Element) -> bytes:
    return ET.tostring(root, encoding="utf-8", xml_declaration=True, short_empty_elements=True)


def _is_blank(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, float):
        return math.isnan(value) or math.isinf(value)
    if pd is not None:
        try:
            missing = pd.isna(value)
            if isinstance(missing, bool):
                return missing
        except (TypeError, ValueError):
            pass
    try:
        result = value != value
        return bool(result) if isinstance(result, (bool, numbers.Number)) else False
    except Exception:
        return False


def _safe_text(value: Any) -> str:
    return ILLEGAL_RE.sub("", str(value))


def _normal_key(value: Any) -> str | None:
    if _is_blank(value):
        return None
    text = _safe_text(value)
    return None if not text.strip() else text


def _number(value: Any) -> float:
    if _is_blank(value):
        return 0.0
    try:
        result = float(value)
    except (TypeError, ValueError):
        return 0.0
    return result if math.isfinite(result) else 0.0


def _display_number(value: float) -> float | int:
    rounded = round(value, 2)
    return int(rounded) if rounded == int(rounded) else rounded


def _excel_scalar(value: Any) -> tuple[str, str] | None:
    """返回 pivot cache record 的 XML 标签和 v 值。"""
    if _is_blank(value):
        return None
    if isinstance(value, bool):
        return "b", "1" if value else "0"
    if isinstance(value, numbers.Number):
        number = float(value)
        if not math.isfinite(number):
            return None
        if number.is_integer():
            return "n", str(int(number))
        return "n", format(number, ".15g")
    if isinstance(value, (datetime, date)):
        return "s", value.isoformat()
    return "s", _safe_text(value)


def _column_letter(index: int) -> str:
    result = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _sales_summary(pivot: Any) -> list[tuple[str | None, float]]:
    totals: OrderedDict[str | None, float] = OrderedDict()
    if pivot is None:
        return []
    for _, row in pivot.iterrows():
        key = _normal_key(row.get("销售人员"))
        totals[key] = totals.get(key, 0.0) + _number(row.get("应收金额"))
    order = list(totals)
    order.sort(key=lambda key: (-totals[key], "" if key is None else key))
    return [(key, totals[key]) for key in order]


def _unique_values(master_out: Any, column: str) -> list[str | None]:
    values: OrderedDict[str | None, None] = OrderedDict()
    for value in master_out[column].tolist():
        values.setdefault(_normal_key(value), None)
    return list(values)


def _write_collapsed_display(ws: Any, pivot: Any) -> int:
    """写入透视表折叠后的初始显示区域；刷新后 Excel 会用 cache 重算。"""
    sub_font = Font(name="等线", size=10.5, bold=True)
    hdr_fill = PatternFill("solid", fgColor="D9E1F2")
    thin = Border(*[Side(style="thin", color="D9D9D9")] * 4)

    for col, value in ((1, "行标签"), (2, "求和项:应收金额")):
        cell = ws.cell(3, col, value)
        cell.font = sub_font
        cell.fill = hdr_fill
        cell.border = thin

    sales = _sales_summary(pivot)
    row = 4
    for person, total in sales:
        first = ws.cell(row, 1, "(空白)" if person is None else person)
        amount = ws.cell(row, 2, _display_number(total))
        first.font = sub_font
        amount.font = sub_font
        first.border = thin
        amount.border = thin
        amount.number_format = ACCT_FMT
        row += 1

    grand = sum(total for _, total in sales)
    first = ws.cell(row, 1, "总计")
    amount = ws.cell(row, 2, _display_number(grand))
    first.font = sub_font
    amount.font = sub_font
    first.border = thin
    amount.border = thin
    amount.number_format = ACCT_FMT

    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 18
    ws.sheet_properties.outlinePr.summaryBelow = False
    return row


def _prepare_workbook(path: Path, pivot: Any) -> None:
    wb = load_workbook(path)
    if PIVOT_NAME not in wb.sheetnames:
        raise RuntimeError(f"工作簿缺少静态透视汇总 sheet：{PIVOT_NAME}")
    if STATIC_BACKUP_NAME in wb.sheetnames:
        raise RuntimeError(f"工作簿已存在备份 sheet：{STATIC_BACKUP_NAME}")

    wb[PIVOT_NAME].title = STATIC_BACKUP_NAME
    pivot_ws = wb.create_sheet(PIVOT_NAME, 1)
    _write_collapsed_display(pivot_ws, pivot)
    wb.save(path)


def _relationship_id(root: ET.Element) -> str:
    used: set[int] = set()
    for node in root.findall(_pkg("Relationship")):
        match = re.fullmatch(r"rId(\d+)", node.get("Id", ""))
        if match:
            used.add(int(match.group(1)))
    candidate = 1
    while candidate in used:
        candidate += 1
    return f"rId{candidate}"


def _part_index(names: Iterable[str], pattern: str) -> int:
    expression = re.compile(pattern)
    used = [int(match.group(1)) for name in names if (match := expression.fullmatch(name))]
    return max(used, default=0) + 1


def _target_part(target: str) -> str:
    target = target.replace("\\", "/")
    return target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))


def _find_sheet_part(entries: dict[str, bytes], name: str) -> str:
    workbook = ET.fromstring(entries["xl/workbook.xml"])
    relationships = ET.fromstring(entries["xl/_rels/workbook.xml.rels"])
    rel_targets = {
        node.get("Id"): node.get("Target")
        for node in relationships.findall(_pkg("Relationship"))
    }
    sheets = workbook.find(_q("sheets"))
    if sheets is None:
        raise RuntimeError("工作簿缺少 sheets 节点")
    for sheet in sheets:
        if sheet.get("name") == name:
            target = rel_targets.get(sheet.get(_r("id")))
            if target:
                return _target_part(target)
    raise RuntimeError(f"找不到 sheet：{name}")


def _find_numfmt_id(styles_xml: bytes) -> str | None:
    root = ET.fromstring(styles_xml)
    numfmts = root.find(_q("numFmts"))
    if numfmts is None:
        return None
    for node in numfmts.findall(_q("numFmt")):
        if node.get("formatCode") == ACCT_FMT:
            return node.get("numFmtId")
    return None


def _shared_items(parent: ET.Element, values: list[str | None]) -> dict[str, int]:
    nonblank = [value for value in values if value is not None]
    attrs: dict[str, str] = {}
    if nonblank:
        attrs["count"] = str(len(values))
    if len(nonblank) != len(values):
        attrs["containsBlank"] = "1"
    shared = ET.SubElement(parent, _q("sharedItems"), attrs)
    mapping: dict[str, int] = {}
    for index, value in enumerate(nonblank):
        ET.SubElement(shared, _q("s"), {"v": value})
        mapping.setdefault(value, index)
    if len(nonblank) != len(values):
        ET.SubElement(shared, _q("m"))
    return mapping


def _cache_field(parent: ET.Element, name: str, values: list[Any], shared: list[str | None] | None) -> dict[str, int] | None:
    field = ET.SubElement(parent, _q("cacheField"), {"name": _safe_text(name), "numFmtId": "0"})
    if shared is not None:
        return _shared_items(field, shared)

    attrs: dict[str, str] = {}
    if any(_is_blank(value) for value in values):
        attrs["containsBlank"] = "1"
    nonblank = [value for value in values if not _is_blank(value)]
    if any(isinstance(value, str) for value in nonblank):
        attrs["containsString"] = "1"
    if any(isinstance(value, numbers.Number) and not isinstance(value, bool) for value in nonblank):
        attrs["containsNumber"] = "1"
    ET.SubElement(field, _q("sharedItems"), attrs)
    return None


def _build_cache_parts(
    master_out: Any,
    headers: list[str],
    sales_values: list[str | None],
    customer_values: list[str | None],
) -> tuple[bytes, bytes]:
    root = ET.Element(
        _q("pivotCacheDefinition"),
        {
            _r("id"): "rId1",
            "refreshedBy": "finance-skill",
            "createdVersion": "8",
            "refreshedVersion": "8",
            "minRefreshableVersion": "3",
            "recordCount": str(len(master_out)),
            "saveData": "1",
        },
    )
    source = ET.SubElement(root, _q("cacheSource"), {"type": "worksheet"})
    ET.SubElement(
        source,
        _q("worksheetSource"),
        {"ref": f"A1:{_column_letter(len(headers))}1048576", "sheet": "主表"},
    )
    cache_fields = ET.SubElement(root, _q("cacheFields"), {"count": str(len(headers))})

    shared_maps: dict[int, dict[str, int]] = {}
    for index, header in enumerate(headers):
        values = master_out[header].tolist()
        if header == "销售人员":
            shared_maps[index] = _cache_field(cache_fields, header, values, sales_values) or {}
        elif header == "客户名称":
            shared_maps[index] = _cache_field(cache_fields, header, values, customer_values) or {}
        else:
            _cache_field(cache_fields, header, values, None)

    records = ET.Element(_q("pivotCacheRecords"), {"count": str(len(master_out))})
    for _, row in master_out.iterrows():
        record = ET.SubElement(records, _q("r"))
        for index, header in enumerate(headers):
            value = row[header]
            if index in shared_maps:
                key = _normal_key(value)
                if key is None:
                    ET.SubElement(record, _q("m"))
                else:
                    ET.SubElement(record, _q("x"), {"v": str(shared_maps[index][key])})
                continue
            scalar = _excel_scalar(value)
            if scalar is None:
                ET.SubElement(record, _q("m"))
            else:
                kind, text = scalar
                ET.SubElement(record, _q(kind), {"v": text})
    return _xml_bytes(root), _xml_bytes(records)


def _pivot_field(parent: ET.Element, *, axis: str | None = None, data: bool = False) -> ET.Element:
    attrs = {"showAll": "0"}
    if axis:
        attrs.update({"axis": axis, "sortType": "descending"})
    if data:
        attrs["dataField"] = "1"
    return ET.SubElement(parent, _q("pivotField"), attrs)


def _pivot_items(field: ET.Element, values: list[str | None], mapping: dict[str, int], *, collapsed: bool) -> None:
    items = ET.SubElement(field, _q("items"), {"count": str(len(values) + 1)})
    for value in values:
        attrs: dict[str, str] = {}
        if collapsed:
            attrs["sd"] = "0"
        # sharedItems 把空值放在所有非空值之后；pivotField 的空值项仍须
        # 用 x 指向这个 <m> 槽位，留空会被 Excel 判定为损坏的透视表。
        attrs["x"] = str(len(mapping) if value is None else mapping[value])
        ET.SubElement(items, _q("item"), attrs)
    default_attrs = {"t": "default"}
    if collapsed:
        default_attrs["sd"] = "0"
    ET.SubElement(items, _q("item"), default_attrs)


def _build_pivot_xml(
    headers: list[str],
    sales_values: list[str | None],
    customer_values: list[str | None],
    sales_mapping: dict[str, int],
    customer_mapping: dict[str, int],
    last_row: int,
    amount_index: int,
    numfmt_id: str | None,
    cache_id: int,
) -> bytes:
    root = ET.Element(
        _q("pivotTableDefinition"),
        {
            "name": "数据透视表1",
            "cacheId": str(cache_id),
            "applyNumberFormats": "0",
            "applyBorderFormats": "0",
            "applyFontFormats": "0",
            "applyPatternFormats": "0",
            "applyAlignmentFormats": "0",
            "applyWidthHeightFormats": "1",
            "dataCaption": "值",
            "updatedVersion": "8",
            "minRefreshableVersion": "3",
            "useAutoFormatting": "1",
            "itemPrintTitles": "1",
            "createdVersion": "6",
            "indent": "0",
            "outline": "1",
            "outlineData": "1",
            "multipleFieldFilters": "0",
        },
    )
    ET.SubElement(root, _q("location"), {"ref": f"A3:B{last_row}", "firstHeaderRow": "1", "firstDataRow": "1", "firstDataCol": "1"})
    fields = ET.SubElement(root, _q("pivotFields"), {"count": str(len(headers))})
    sales_index = headers.index("销售人员")
    customer_index = headers.index("客户名称")
    for index in range(len(headers)):
        if index == sales_index:
            field = _pivot_field(fields, axis="axisRow")
            _pivot_items(field, sales_values, sales_mapping, collapsed=True)
        elif index == customer_index:
            field = _pivot_field(fields, axis="axisRow")
            _pivot_items(field, customer_values, customer_mapping, collapsed=False)
        elif index == amount_index:
            _pivot_field(fields, data=True)
        else:
            _pivot_field(fields)

    row_fields = ET.SubElement(root, _q("rowFields"), {"count": "2"})
    ET.SubElement(row_fields, _q("field"), {"x": str(sales_index)})
    ET.SubElement(row_fields, _q("field"), {"x": str(customer_index)})
    row_items = ET.SubElement(root, _q("rowItems"), {"count": str(len(sales_values) + 1)})
    for value in sales_values:
        item = ET.SubElement(row_items, _q("i"))
        if value is None:
            ET.SubElement(item, _q("x"))
        else:
            ET.SubElement(item, _q("x"), {"v": str(sales_mapping[value])})
    grand = ET.SubElement(row_items, _q("i"), {"t": "grand"})
    ET.SubElement(grand, _q("x"))
    col_items = ET.SubElement(root, _q("colItems"), {"count": "1"})
    ET.SubElement(col_items, _q("i"))
    data_fields = ET.SubElement(root, _q("dataFields"), {"count": "1"})
    data_attrs = {"name": "求和项:应收金额", "fld": str(amount_index), "baseField": "0", "baseItem": "0"}
    if numfmt_id:
        data_attrs["numFmtId"] = numfmt_id
    ET.SubElement(data_fields, _q("dataField"), data_attrs)
    ET.SubElement(
        root,
        _q("pivotTableStyleInfo"),
        {
            "name": "PivotStyleLight16",
            "showRowHeaders": "1",
            "showColHeaders": "1",
            "showRowStripes": "0",
            "showColStripes": "0",
            "showLastColumn": "1",
        },
    )
    return _xml_bytes(root)


def _add_relationship(root: ET.Element, rel_id: str, rel_type: str, target: str) -> None:
    ET.SubElement(root, _pkg("Relationship"), {"Id": rel_id, "Type": rel_type, "Target": target})


def _ensure_override(root: ET.Element, part_name: str, content_type: str) -> None:
    for node in root.findall(_ct("Override")):
        if node.get("PartName") == part_name:
            node.set("ContentType", content_type)
            return
    ET.SubElement(root, _ct("Override"), {"PartName": part_name, "ContentType": content_type})


def _next_cache_id(pivot_cache_nodes: ET.Element | None) -> int:
    if pivot_cache_nodes is None:
        return 0
    ids = []
    for node in pivot_cache_nodes.findall(_q("pivotCache")):
        try:
            ids.append(int(node.get("cacheId", "0")))
        except ValueError:
            pass
    return max(ids, default=-1) + 1


def _rewrite_zip(path: Path, replacements: dict[str, bytes]) -> None:
    fd, temp_name = tempfile.mkstemp(prefix="native-pivot-package-", suffix=".xlsx", dir=str(path.parent))
    os.close(fd)
    temp_path = Path(temp_name)
    try:
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temp_path, "w", zipfile.ZIP_DEFLATED) as target:
            existing: set[str] = set()
            for info in source.infolist():
                existing.add(info.filename)
                data = replacements.get(info.filename)
                if data is None:
                    data = source.read(info.filename)
                target.writestr(info, data)
            for name, data in replacements.items():
                if name not in existing:
                    target.writestr(name, data)
        os.replace(temp_path, path)
    finally:
        temp_path.unlink(missing_ok=True)


def _install_parts(path: Path, master_out: Any, pivot: Any) -> None:
    with zipfile.ZipFile(path, "r") as archive:
        entries = {info.filename: archive.read(info.filename) for info in archive.infolist()}

    pivot_sheet_part = _find_sheet_part(entries, PIVOT_NAME)
    workbook_root = ET.fromstring(entries["xl/workbook.xml"])
    workbook_rels = ET.fromstring(entries["xl/_rels/workbook.xml.rels"])
    sheet_rels_name = posixpath.join(
        posixpath.dirname(pivot_sheet_part),
        "_rels",
        posixpath.basename(pivot_sheet_part) + ".rels",
    )
    sheet_rels = ET.fromstring(entries[sheet_rels_name]) if sheet_rels_name in entries else ET.Element(_pkg("Relationships"))

    names = set(entries)
    cache_index = _part_index(names, r"xl/pivotCache/pivotCacheDefinition(\d+)\.xml")
    table_index = _part_index(names, r"xl/pivotTables/pivotTable(\d+)\.xml")
    cache_part = f"xl/pivotCache/pivotCacheDefinition{cache_index}.xml"
    records_part = f"xl/pivotCache/pivotCacheRecords{cache_index}.xml"
    table_part = f"xl/pivotTables/pivotTable{table_index}.xml"
    cache_rels_part = f"xl/pivotCache/_rels/pivotCacheDefinition{cache_index}.xml.rels"
    table_rels_part = f"xl/pivotTables/_rels/pivotTable{table_index}.xml.rels"

    headers = [str(header) for header in master_out.columns]
    required = {"销售人员", "客户名称", "应收金额"}
    if not required.issubset(headers):
        raise RuntimeError("主表缺少原生透视表所需字段")
    sales_values = [key for key, _ in _sales_summary(pivot)]
    customer_values = _unique_values(master_out, "客户名称")
    sales_mapping = {value: index for index, value in enumerate(value for value in sales_values if value is not None)}
    customer_mapping = {value: index for index, value in enumerate(value for value in customer_values if value is not None)}
    cache_definition, cache_records = _build_cache_parts(master_out, headers, sales_values, customer_values)
    numfmt_id = _find_numfmt_id(entries["xl/styles.xml"])
    pivot_cache_nodes = workbook_root.find(_q("pivotCaches"))
    cache_id = _next_cache_id(pivot_cache_nodes)
    pivot_table = _build_pivot_xml(
        headers,
        sales_values,
        customer_values,
        sales_mapping,
        customer_mapping,
        4 + len(sales_values),
        headers.index("应收金额"),
        numfmt_id,
        cache_id,
    )

    workbook_rel_id = _relationship_id(workbook_rels)
    _add_relationship(workbook_rels, workbook_rel_id, PIVOT_CACHE_REL, f"pivotCache/pivotCacheDefinition{cache_index}.xml")
    if pivot_cache_nodes is None:
        pivot_cache_nodes = ET.Element(_q("pivotCaches"))
        insert_at = len(workbook_root)
        for index, child in enumerate(workbook_root):
            if child.tag == _q("extLst"):
                insert_at = index
                break
        workbook_root.insert(insert_at, pivot_cache_nodes)
    ET.SubElement(pivot_cache_nodes, _q("pivotCache"), {"cacheId": str(cache_id), _r("id"): workbook_rel_id})

    # 旧版本曾为透视源写入 OFFSET 动态名称。Excel 对该 worksheetSource
    # 组合的兼容性不稳定；当前透视缓存直接使用主表的全列范围。若工作簿
    # 中已有旧名称，只清理本技能生成的名称，不影响其他名称。
    defined_names = workbook_root.find(_q("definedNames"))
    if defined_names is not None:
        for node in list(defined_names.findall(_q("definedName"))):
            if node.get("name") == SOURCE_NAME:
                defined_names.remove(node)
        if not list(defined_names):
            workbook_root.remove(defined_names)

    sheet_rel_id = _relationship_id(sheet_rels)
    _add_relationship(sheet_rels, sheet_rel_id, PIVOT_TABLE_REL, f"../pivotTables/pivotTable{table_index}.xml")
    cache_rels = ET.Element(_pkg("Relationships"))
    _add_relationship(cache_rels, "rId1", PIVOT_CACHE_RECORDS_REL, f"pivotCacheRecords{cache_index}.xml")
    table_rels = ET.Element(_pkg("Relationships"))
    _add_relationship(table_rels, "rId1", PIVOT_CACHE_REL, f"../pivotCache/pivotCacheDefinition{cache_index}.xml")

    content_types = ET.fromstring(entries["[Content_Types].xml"])
    _ensure_override(content_types, f"/{cache_part}", PIVOT_CACHE_CONTENT_TYPE)
    _ensure_override(content_types, f"/{records_part}", PIVOT_CACHE_RECORDS_CONTENT_TYPE)
    _ensure_override(content_types, f"/{table_part}", PIVOT_TABLE_CONTENT_TYPE)
    replacements = {
        "xl/workbook.xml": _xml_bytes(workbook_root),
        "xl/_rels/workbook.xml.rels": _xml_bytes(workbook_rels),
        "[Content_Types].xml": _xml_bytes(content_types),
        sheet_rels_name: _xml_bytes(sheet_rels),
        cache_part: cache_definition,
        records_part: cache_records,
        cache_rels_part: _xml_bytes(cache_rels),
        table_part: pivot_table,
        table_rels_part: _xml_bytes(table_rels),
    }
    _rewrite_zip(path, replacements)


def install_native_pivot(workbook_path: str | os.PathLike[str], master_out: Any, pivot: Any) -> None:
    """原子地把静态透视汇总升级为原生、可刷新的透视表。"""
    path = Path(workbook_path).resolve()
    if not path.is_file():
        raise FileNotFoundError(path)
    fd, temp_name = tempfile.mkstemp(prefix="native-pivot-workbook-", suffix=".xlsx", dir=str(path.parent))
    os.close(fd)
    working = Path(temp_name)
    try:
        shutil.copyfile(path, working)
        _prepare_workbook(working, pivot)
        _install_parts(working, master_out, pivot)
        os.replace(working, path)
    finally:
        working.unlink(missing_ok=True)
