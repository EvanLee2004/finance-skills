# -*- coding: utf-8 -*-
r"""
销售反馈合并 · merge_feedback.py
==============================
把一个文件夹里各销售的反馈 Excel（每人一份、单人 sheet、同结构 17 列）
合并成一张总表 销售反馈合并_<日期>.xlsx，并统一样式：
  - 前 8 列表头：蓝色底 FFBDD7EE、粗体、居中
  - 第 9 列起表头：黄色底 FFFFFF00、粗体、居中
  - 数据行：宋体 10.5、整体细边框、冻结首行

用法：
  python merge_feedback.py --input-dir "<反馈文件夹绝对路径>" --date-tag 0820
  (--out 可省，默认落在输入目录下 销售反馈合并_<日期>.xlsx)
"""
import os
import sys
import argparse

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
except ImportError:
    print("✗ 缺 openpyxl → pip install openpyxl")
    sys.exit(1)

BLUE_FILL = PatternFill(start_color="FFBDD7EE", end_color="FFBDD7EE", fill_type="solid")
YELLOW_FILL = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")
THIN = Border(left=Side(style="thin"), right=Side(style="thin"),
              top=Side(style="thin"), bottom=Side(style="thin"))


def list_feedback_files(input_dir):
    files = []
    for fn in sorted(os.listdir(input_dir)):
        if not fn.lower().endswith(".xlsx"):
            continue
        if fn.startswith("~$"):
            continue
        if fn.startswith("销售反馈合并"):
            continue
        files.append(os.path.join(input_dir, fn))
    return files


def norm_col(name):
    """列名归一化：去空白/换行/括号，用于跨文件列名比对（同一列可能带换行或空格差异）"""
    if name is None:
        return ""
    s = str(name)
    import re
    return re.sub(r"[\s\xa0()（）【】：:]+", "", s)


def read_one(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = []
    header = None
    for r in ws.iter_rows(values_only=True):
        vals = list(r)
        if header is None:
            while vals and vals[-1] is None:
                vals.pop()
            header = vals
            continue
        if all(v is None or str(v).strip() == "" for v in vals):
            continue
        vals = (vals + [None] * len(header))[:len(header)]
        rows.append(vals)
    wb.close()
    return header, rows


def main():
    ap = argparse.ArgumentParser(description="合并销售反馈 Excel")
    ap.add_argument("--input-dir", required=True, help="反馈文件夹绝对路径")
    ap.add_argument("--date-tag", required=True, help="日期标签（如0820）")
    ap.add_argument("--out", default="", help="输出路径（默认输入目录下 销售反馈合并_<日期>.xlsx）")
    args = ap.parse_args()

    input_dir = os.path.abspath(args.input_dir)
    files = list_feedback_files(input_dir)
    if not files:
        print(f"✗ 文件夹里没有反馈 Excel：{input_dir}")
        sys.exit(1)

    out_path = os.path.abspath(args.out) if args.out else os.path.join(
        input_dir, f"销售反馈合并_{args.date_tag}.xlsx")
    if os.path.exists(os.path.join(input_dir, "~$" + os.path.basename(out_path)[1:])):
        pass  # read_only 打开不冲突；输出占用会在 save 时报错

    header = None
    all_rows = []
    per_person = []
    for f in files:
        h, rows = read_one(f)
        if header is None:
            header = h
        elif [norm_col(x) for x in h] != [norm_col(x) for x in header]:
            print(f"✗ 列结构与第一份不一致，跳过：{os.path.basename(f)}")
            print(f"  该文件列: {[str(x)[:12] for x in h]}")
            continue
        all_rows.extend(rows)
        name = os.path.basename(f)
        per_person.append((name, len(rows)))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"销售反馈合并{args.date_tag}"

    ncol = len(header)
    for c, h in enumerate(header, 1):
        ws.cell(row=1, column=c, value=h)

    for r, vals in enumerate(all_rows, 2):
        for c, v in enumerate(vals, 1):
            ws.cell(row=r, column=c, value=v)

    # 样式
    for c in range(1, ncol + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = BLUE_FILL if c <= 8 else YELLOW_FILL
        cell.font = Font(name="宋体", size=10.5, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN

    data_font = Font(name="宋体", size=10.5)
    for r in range(2, len(all_rows) + 2):
        for c in range(1, ncol + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = data_font
            cell.border = THIN
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    for c in range(1, ncol + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 14
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["E"].width = 28
    ws.column_dimensions["F"].width = 12
    ws.column_dimensions["J"].width = 18
    ws.column_dimensions["L"].width = 18
    ws.freeze_panes = "A2"

    wb.save(out_path)

    print(f"✓ 合并完成：{out_path}")
    print(f"  共 {len(files)} 份文件、{len(all_rows)} 行数据、{ncol} 列")
    for name, n in per_person:
        print(f"    {n:>4} 行 | {name}")


if __name__ == "__main__":
    main()
