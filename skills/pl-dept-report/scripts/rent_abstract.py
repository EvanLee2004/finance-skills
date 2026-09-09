#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""总部管理费用-房租：摘要含深圳进广东、含长沙进湖南分。无源不拆。"""
from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from common import CHECK_TOL, add_money, load_json, money
from layout import dept_col_letter


def load_rent_rules() -> dict:
    return (load_json("残差归集.json") or {}).get("rent_abstract") or {}


def entry_explanation(entry: dict) -> str:
    if not isinstance(entry, dict):
        return ""
    for key in (
        "explanation",
        "remark",
        "comment",
        "description",
        "expl",
        "摘要",
        "voucher_explanation",
    ):
        val = entry.get(key)
        if val:
            return str(val).replace("\xa0", " ").strip()
    return ""


def classify_rent_expl(text: str, rules: dict | None = None) -> str | None:
    rules = rules or load_rent_rules()
    raw = str(text or "")
    if not raw:
        return None
    for rule in rules.get("rules") or []:
        needle = str(rule.get("contains") or "")
        dest = str(rule.get("excel_dept") or "").strip()
        if needle and dest and needle in raw:
            return dest
    return None


def is_rent_account(code: str, name: str, rules: dict | None = None) -> bool:
    rules = rules or load_rent_rules()
    want = str(rules.get("code") or "550212")
    code = str(code or "").strip()
    if code == want or code.startswith(want):
        return True
    return "房租" in str(name or "") and code.startswith("5502")


def lines_from_entries(entries: list[dict], rules: dict | None = None) -> list[dict]:
    rules = rules or load_rent_rules()
    out: list[dict] = []
    for entry in entries or []:
        if not isinstance(entry, dict):
            continue
        code = str(
            entry.get("account_number")
            or entry.get("accountNumber")
            or entry.get("number")
            or entry.get("account")
            or entry.get("code")
            or ""
        ).strip()
        name = str(entry.get("account_name") or entry.get("name") or "")
        if not is_rent_account(code, name, rules):
            continue
        dest = classify_rent_expl(entry_explanation(entry), rules)
        if not dest:
            continue
        debit = money(entry.get("debit_amount") or entry.get("debitAmount") or entry.get("debit"))
        credit = money(entry.get("credit_amount") or entry.get("creditAmount") or entry.get("credit"))
        amt = (debit or Decimal("0")) - (credit or Decimal("0"))
        if abs(amt) <= CHECK_TOL:
            continue
        out.append({"code": str(rules.get("code") or "550212"), "excel_dept": dest, "amount": amt})
    return out


def _header_map(values: list) -> dict[str, int]:
    found: dict[str, int] = {}
    for i, raw in enumerate(values):
        t = str(raw or "").replace("\xa0", " ").strip()
        if not t:
            continue
        if t in {"摘要", "摘要 #"} and "expl" not in found:
            found["expl"] = i
        if t in {"科目代码", "科目编码"} and "code" not in found:
            found["code"] = i
        if t in {"科目名称", "科目"} and "name" not in found:
            found["name"] = i
        if t in {"借方", "借方金额", "本期发生借方"} and "debit" not in found:
            found["debit"] = i
        if t in {"贷方", "贷方金额", "本期发生贷方"} and "credit" not in found:
            found["credit"] = i
    return found


def looks_like_journal_headers(values: list) -> bool:
    blob = " ".join(str(v or "") for v in values)
    if "摘要" not in blob:
        return False
    return ("科目代码" in blob) or ("科目编码" in blob) or ("凭证字号" in blob and "科目" in blob)


def parse_journal_rent_lines(path: Path, rules: dict | None = None) -> list[dict]:
    rules = rules or load_rent_rules()
    path = Path(path)
    if not path.is_file() or path.suffix.lower() not in {".xlsx", ".xlsm"}:
        return []
    if "损益类部门科目余额表" in path.name or path.name.startswith("月度损益表_"):
        return []
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception:
        return []
    try:
        for title in wb.sheetnames:
            ws = wb[title]
            header = None
            header_r = 0
            for r, row in enumerate(ws.iter_rows(max_row=12, values_only=True), 1):
                vals = list(row or [])
                if looks_like_journal_headers(vals):
                    header = _header_map(vals)
                    header_r = r
                    break
            if not header or "expl" not in header:
                continue
            code_at = header.get("code")
            name_at = header.get("name")
            if code_at is None and name_at is None:
                continue
            entries = []
            for row in ws.iter_rows(min_row=header_r + 1, values_only=True):
                vals = list(row or [])
                if not vals:
                    continue
                expl = vals[header["expl"]] if header["expl"] < len(vals) else None
                code = vals[code_at] if code_at is not None and code_at < len(vals) else ""
                name = vals[name_at] if name_at is not None and name_at < len(vals) else ""
                debit = vals[header["debit"]] if header.get("debit") is not None and header["debit"] < len(vals) else None
                credit = vals[header["credit"]] if header.get("credit") is not None and header["credit"] < len(vals) else None
                entries.append(
                    {
                        "code": code,
                        "account": code,
                        "name": name,
                        "explanation": expl,
                        "debit": debit,
                        "credit": credit,
                    }
                )
            got = lines_from_entries(entries, rules)
            if got:
                return got
    finally:
        wb.close()
    return []


def collect_rent_from_dir(paths: list[Path], rules: dict | None = None) -> list[dict]:
    found: list[dict] = []
    seen: set[str] = set()
    for path in paths:
        key = str(path.resolve()) if path.is_file() else ""
        if not key or key in seen:
            continue
        seen.add(key)
        found.extend(parse_journal_rent_lines(path, rules))
    return found


def apply_rent_abstract_split(dept_amts: dict, lines: list[dict], layout: dict, notes: list[str]) -> None:
    rules = load_rent_rules()
    code = str(rules.get("code") or "550212")
    src_name = str(rules.get("from_dept") or "运营保障中心")
    src = dept_col_letter(layout, src_name, 1)
    if not src:
        return
    if not lines:
        notes.append("房租摘要=无源")
        return
    bucket = dept_amts.setdefault(code, {})
    available = Decimal(str(bucket.get(src) or 0))
    moved: list[str] = []
    dest_names: list[str] = []
    for rule in rules.get("rules") or []:
        dest_name = str(rule.get("excel_dept") or "")
        if dest_name and dest_name not in dest_names:
            dest_names.append(dest_name)
    for dest_name in dest_names:
        dest = dept_col_letter(layout, dest_name, 1)
        if not dest:
            continue
        need = Decimal("0")
        for row in lines:
            if str(row.get("excel_dept") or "") != dest_name:
                continue
            if str(row.get("code") or code) != code:
                continue
            need += Decimal(str(row.get("amount") or 0))
        if abs(need) <= CHECK_TOL or available <= CHECK_TOL:
            continue
        take = need if abs(need) <= available else available
        if take <= CHECK_TOL:
            continue
        bucket[src] = available - take
        bucket[dest] = add_money(bucket.get(dest), take)
        available = Decimal(str(bucket.get(src) or 0))
        moved.append(dest_name)
    if moved:
        notes.append("房租摘要=" + ",".join(moved))
    else:
        notes.append("房租摘要=无匹配")
