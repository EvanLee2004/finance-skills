#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把总部开放平台取数落成和网页引出同结构的源 Excel。stdout 只报文件/占用个数，不含金额。"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
if str(HERE) not in sys.path:
    sys.path.insert(0, str(HERE))

from openpyxl import Workbook

from layout import layout_code_names, load_books, load_layout


def _is_api_path(path: Path | str) -> bool:
    p = Path(path)
    return "API" in p.parts or "_API_" in p.name


def is_api_source(item: dict) -> bool:
    return _is_api_path(str(item.get("path") or ""))


def prefer_web_over_api(inspected: list[dict]) -> list[dict]:
    """同一账套同一类表：有网页引出就不用 API 落盘，避免加两遍。"""
    have_web: set[tuple] = set()
    for item in inspected:
        if not item.get("entity") or not item.get("kind"):
            continue
        if is_api_source(item):
            continue
        have_web.add((item.get("entity"), item.get("kind"), item.get("period")))
    out = []
    for item in inspected:
        key = (item.get("entity"), item.get("kind"), item.get("period"))
        if is_api_source(item) and key in have_web:
            continue
        out.append(item)
    return out


def _legal_hq() -> str:
    for acc in load_books().get("xingchen_accounts") or []:
        if acc.get("key") == "jiagu":
            return str(acc.get("legal_name") or "甲骨易（北京）语言科技股份有限公司")
    return "甲骨易（北京）语言科技股份有限公司"


def _period_banner(period: str) -> str:
    return f"{period[:4]}年{int(period[4:6]):02d}期"


def _cell_num(value):
    if value is None:
        return None
    return float(value)


def write_account_xlsx(path: Path, company: str, period: str, rows: list[tuple]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "科目余额表"
    ws["A1"] = "科目余额表"
    ws["A2"] = f"公司名称：{company}"
    ws["B2"] = f"期间：{period}-{period}"
    ws["C2"] = _period_banner(period)
    ws["A3"] = "科目编码"
    ws["B3"] = "科目名称"
    ws["C3"] = "本期发生借方"
    ws["D3"] = "本期发生贷方"
    for i, (code, name, debit, credit) in enumerate(rows, start=4):
        ws.cell(i, 1).value = code
        ws.cell(i, 2).value = name
        ws.cell(i, 3).value = _cell_num(debit)
        ws.cell(i, 4).value = _cell_num(credit)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_assist_xlsx(path: Path, company: str, period: str, rows: list[tuple]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "核算项目余额表"
    ws["A1"] = "核算项目余额表"
    ws["B1"] = "核算项目类别：部门"
    ws["A2"] = f"公司名称：{company}"
    ws["B2"] = f"期间：{period}-{period}"
    ws["C2"] = _period_banner(period)
    ws["A3"] = "科目编码"
    ws["B3"] = "科目名称"
    ws["C3"] = "核算项目名称"
    ws["D3"] = "本期发生借方"
    ws["E3"] = "本期发生贷方"
    for i, (code, name, dept, debit, credit) in enumerate(rows, start=4):
        ws.cell(i, 1).value = code
        ws.cell(i, 2).value = name
        ws.cell(i, 3).value = dept
        ws.cell(i, 4).value = _cell_num(debit)
        ws.cell(i, 5).value = _cell_num(credit)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_profit_xlsx(path: Path, company: str, period: str, items: list[tuple]) -> None:
    year = period[:4]
    month = int(period[4:6])
    wb = Workbook()
    ws = wb.active
    ws.title = "利润表"
    ws["A1"] = f"{year}年{month}期利润表（月报）"
    ws["A2"] = f"公司名称：{company}"
    ws["A3"] = "项目"
    ws["B3"] = "本月金额"
    for i, (name, amt) in enumerate(items, start=4):
        ws.cell(i, 1).value = name
        ws.cell(i, 2).value = _cell_num(amt)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def dump_hq_to_dir(period: str, out_dir: Path, creds: dict, hq: dict | None = None) -> list[str]:
    if hq is None:
        from fetch_reports import fetch_hq

        hq = fetch_hq(period, creds)
    layout = load_layout()
    names = layout_code_names(layout)
    company = _legal_hq()
    out_dir = Path(out_dir)
    written: list[str] = []
    accounts = hq.get("accounts") or {}
    if accounts:
        dest = out_dir / f"甲骨易_API_科目余额表_{period}.xlsx"
        rows = [
            (code, (pair or {}).get("name") or names.get(code) or "", (pair or {}).get("debit"), (pair or {}).get("credit"))
            for code, pair in sorted(accounts.items())
        ]
        write_account_xlsx(dest, company, period, rows)
        written.append(dest.name)
    depts = hq.get("depts") or []
    if depts:
        dest = out_dir / f"甲骨易_API_核算项目余额表_{period}.xlsx"
        rows = [
            (
                row.get("code") or "",
                row.get("name") or names.get(str(row.get("code") or "")) or "",
                row.get("dept") or "",
                row.get("debit"),
                row.get("credit"),
            )
            for row in depts
        ]
        write_assist_xlsx(dest, company, period, rows)
        written.append(dest.name)
    profit = hq.get("profit") or {}
    items = [(label, amt) for label, amt in profit.items() if amt is not None]
    if items:
        dest = out_dir / f"甲骨易_API_利润表_{period}.xlsx"
        write_profit_xlsx(dest, company, period, items)
        written.append(dest.name)
    return written


def dump_if_needed(period: str, out_dir: Path, notes: list[str] | None = None) -> list[str]:
    from kingdee_client import load_client

    creds = load_client()
    if not creds:
        return []
    try:
        written = dump_hq_to_dir(period, Path(out_dir), creds)
    except Exception as e:
        if notes is not None:
            notes.append(f"api_dump={type(e).__name__}")
        return []
    if written and notes is not None:
        notes.append(f"总部API已落源表={len(written)}")
    return written


def dump_prev_profit_if_needed(period: str, out_dir: Path, notes: list[str] | None = None) -> list[str]:
    """总部上月利润表走开放平台，不网页、不问她。"""
    from kingdee_client import load_client

    from common import prev_period

    creds = load_client()
    if not creds:
        return []
    prev = prev_period(period)
    try:
        from fetch_reports import fetch_hq

        hq = fetch_hq(prev, creds, ledger=False)
        written = dump_hq_to_dir(prev, Path(out_dir), creds, hq=hq)
    except Exception as e:
        if notes is not None:
            notes.append(f"api_dump_prev={type(e).__name__}")
        return []
    if written and notes is not None:
        notes.append(f"总部API上月利润表已落={len(written)}")
    return written


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="总部 API 落源 Excel。只报文件名和占用个数。")
    parser.add_argument("--period", default="202608")
    parser.add_argument("--out-dir", required=True)
    args = parser.parse_args(argv)
    from kingdee_client import load_client

    creds = load_client()
    if not creds:
        print("api=no_creds")
        return 2
    written = dump_hq_to_dir(args.period, Path(args.out_dir).expanduser(), creds)
    print(f"period={args.period} files={len(written)} " + ",".join(written))
    return 0 if written else 2


if __name__ == "__main__":
    raise SystemExit(main())
