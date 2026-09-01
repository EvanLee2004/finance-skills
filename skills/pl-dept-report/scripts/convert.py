#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""月度损益表 / 科目余额表。发生额脚本写，核对/合计/父行/勾稽用 Excel 公式。"""
from __future__ import annotations

import argparse
import json
import os
import sys
from decimal import Decimal
from pathlib import Path

HERE = Path(__file__).resolve().parent
if str(HERE) not in sys.path:
    sys.path.insert(0, str(HERE))


def repo_venv_python() -> Path | None:
    root = Path(__file__).resolve().parents[3]
    for rel in (Path(".venv") / "bin" / "python", Path(".venv") / "Scripts" / "python.exe"):
        cand = root / rel
        if cand.is_file():
            return cand
    return None


def _reexec_repo_venv_if_needed() -> None:
    if not sys.argv or Path(sys.argv[0]).name.lower() not in {"convert.py", "convert"}:
        return
    venv_py = repo_venv_python()
    if venv_py is None:
        return
    try:
        if Path(sys.prefix).resolve() == venv_py.parent.parent.resolve():
            return
    except OSError:
        return
    os.execv(str(venv_py), [str(venv_py), *sys.argv])


_reexec_repo_venv_if_needed()

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, Border, Side
from openpyxl.utils import get_column_letter

from common import (
    CHECK_TOL,
    add_money,
    cell_num,
    default_period,
    load_json,
    money,
    parse_period,
    prev_period,
    stdout_safe,
)
from formula_eval import eval_workbook
from inspect_inputs import inspect_dir
from layout import (
    account_row_map,
    dept_columns,
    direct_children,
    first_dept_col,
    load_books,
    load_dept_map,
    load_layout,
    map_dept_name,
)
from parse_export import parse_inspected

NF = r'_ * #,##0.00_ ;_ * \-#,##0.00_ ;_ * "-"??_ ;_ @_ '
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)


def is_income(code: str, layout: dict) -> bool:
    return any(code.startswith(p) for p in layout.get("income_prefixes") or ["51", "52", "53"])


def nature_amount(code: str, debit, credit, layout: dict):
    if debit is None and credit is None:
        return None
    d = debit or Decimal("0")
    c = credit or Decimal("0")
    net = (c - d) if is_income(code, layout) else (d - c)
    return net


def load_prev_profit(path: Path | None, layout: dict) -> dict[str, dict]:
    if not path or not path.is_file():
        return {}
    try:
        wb = load_workbook(path, data_only=False)
    except Exception:
        return {}
    if "利润表" not in wb.sheetnames:
        return {}
    ws = wb["利润表"]
    entities = layout["entities"]
    labels = [row["label"] for row in layout["profit_rows"]]
    out: dict[str, dict] = {e: {} for e in entities}
    for i, label in enumerate(labels):
        if not label:
            continue
        r = 2 + i
        for j, ent in enumerate(entities):
            val = money(ws.cell(r, 2 + j).value)
            if val is not None:
                out[ent][label] = val
    return out


def merge_dept_rows(rows: list[dict], layout: dict, mapping: dict, report: dict):
    kids = direct_children(layout)
    amounts: dict[str, dict[str, Decimal]] = {}
    unmapped: list[str] = []
    seen_unmapped: set[str] = set()
    for row in rows:
        code = str(row.get("code") or "").strip()
        if not code:
            continue
        excel_dept = map_dept_name(str(row.get("dept") or ""), mapping, layout)
        amt = nature_amount(code, row.get("debit"), row.get("credit"), layout)
        if excel_dept is None:
            name = str(row.get("dept") or "").strip()
            if name and name not in seen_unmapped and amt:
                seen_unmapped.add(name)
                unmapped.append(name)
            continue
        if kids.get(code):
            continue
        if amt is None:
            continue
        bucket = amounts.setdefault(code, {})
        bucket[excel_dept] = (bucket.get(excel_dept) or Decimal("0")) + amt
    report["unmapped_depts"] = unmapped
    return amounts


def write_pl_sheet(wb, layout: dict, entity_amts: dict, dept_amts: dict) -> None:
    ws = wb.create_sheet("损益表", 0)
    ws.freeze_panes = layout.get("freeze") or "C3"
    header_font = Font(name="等线", size=11, bold=True)
    parent_font = Font(name="宋体", size=10, bold=True)
    leaf_font = Font(name="等线", size=11, bold=False)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center")

    ws["A2"] = "科目编码"
    ws["B2"] = "科目名称"
    ws["A2"].font = header_font
    ws["B2"].font = header_font
    for pair in layout["entity_pairs"]:
        ws[f"{pair['debit']}2"] = "本期发生借方"
        ws[f"{pair['credit']}2"] = "本期发生贷方"
        ws[f"{pair['debit']}2"].font = header_font
        ws[f"{pair['credit']}2"].font = header_font
        ws[f"{pair['debit']}2"].alignment = center
        ws[f"{pair['credit']}2"].alignment = center
    ws["S2"] = "本期发生借方"
    ws["T2"] = "本期发生贷方"
    ws["S2"].font = header_font
    ws["T2"].font = header_font
    for name, letter in dept_columns(layout):
        cell = ws[f"{letter}2"]
        cell.value = name
        cell.font = header_font
        cell.alignment = center
    ws["AU2"] = "费用合计"
    ws["AV2"] = "核对"
    ws["AU2"].font = header_font
    ws["AV2"].font = header_font

    for group in layout.get("row1_groups") or []:
        start, end = group["start"], group["end"]
        ws.merge_cells(f"{start}1:{end}1")
        cell = ws[f"{start}1"]
        cell.value = group["label"]
        cell.font = header_font
        cell.alignment = center
    ws.merge_cells("A1:B1")

    for letter, width in (layout.get("column_widths") or {}).items():
        ws.column_dimensions[letter].width = width

    row_of = account_row_map(layout)
    kids = direct_children(layout)
    debit_letters = [p["debit"] for p in layout["entity_pairs"]]
    credit_letters = [p["credit"] for p in layout["entity_pairs"]]
    dept_letters = [letter for _n, letter in dept_columns(layout)]

    for acc in layout["accounts"]:
        code = acc["code"]
        r = row_of[code]
        ws.cell(r, 1).value = code
        ws.cell(r, 2).value = acc["name"]
        font = parent_font if acc.get("bold") or kids.get(code) else leaf_font
        ws.cell(r, 1).font = font
        ws.cell(r, 2).font = font
        ws.cell(r, 2).alignment = left
        for pair in layout["entity_pairs"]:
            header = pair["header"]
            bucket = (entity_amts.get(header) or {}).get(code) or {}
            dcell = ws[f"{pair['debit']}{r}"]
            ccell = ws[f"{pair['credit']}{r}"]
            dcell.value = cell_num(bucket.get("debit"))
            ccell.value = cell_num(bucket.get("credit"))
            dcell.number_format = NF
            ccell.number_format = NF
        ws[f"S{r}"] = "=" + "+".join(f"{col}{r}" for col in debit_letters)
        ws[f"T{r}"] = "=" + "+".join(f"{col}{r}" for col in credit_letters)
        ws[f"S{r}"].number_format = NF
        ws[f"T{r}"].number_format = NF
        child_codes = kids.get(code) or []
        for name, letter in dept_columns(layout):
            cell = ws[f"{letter}{r}"]
            cell.number_format = NF
            if child_codes:
                parts = []
                for child in child_codes:
                    cr = row_of[child]
                    parts.append(f"{letter}{cr}")
                cell.value = "=" + "+".join(parts)
            else:
                amt = (dept_amts.get(code) or {}).get(name)
                cell.value = cell_num(amt)
        ws[f"AU{r}"] = f"=SUM(U{r}:AT{r})"
        ws[f"AU{r}"].number_format = NF
        if is_income(code, layout):
            ws[f"AV{r}"] = f"=T{r}-AU{r}"
        else:
            ws[f"AV{r}"] = f"=S{r}-AU{r}"
        ws[f"AV{r}"].number_format = NF
        for col in range(1, 49):
            ws.cell(r, col).border = THIN


def profit_label_rows(layout: dict) -> dict[str, int]:
    out = {}
    for i, row in enumerate(layout["profit_rows"]):
        if row.get("label"):
            out[row["label"]] = 2 + i
    return out


def write_profit_sheet(wb, layout: dict, period: str, current: dict, previous: dict) -> None:
    ws = wb.create_sheet("利润表", 1)
    header_font = Font(name="等线", size=11, bold=True)
    bold_font = Font(name="等线", size=11, bold=True)
    normal = Font(name="等线", size=11, bold=False)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    y, m = int(period[:4]), int(period[4:6])
    prev = prev_period(period)
    py, pm = int(prev[:4]), int(prev[4:6])
    entities = layout["entities"]
    for j, ent in enumerate(entities):
        ws.cell(1, 2 + j).value = f"{ent}{y % 100}年 {m}月"
        ws.cell(1, 10 + j).value = f"{ent}{py % 100}年 {pm}月"
        ws.cell(1, 2 + j).font = header_font
        ws.cell(1, 10 + j).font = header_font
        ws.cell(1, 2 + j).alignment = center
        ws.cell(1, 10 + j).alignment = center
    for letter, width in (layout.get("profit_column_widths") or {}).items():
        ws.column_dimensions[letter].width = width

    label_row = profit_label_rows(layout)
    row_of = account_row_map(layout)
    cur_cols = list(range(2, 10))  # B-I
    prev_cols = list(range(10, 18))  # J-Q

    def set_nf(cell):
        cell.number_format = NF

    for i, spec in enumerate(layout["profit_rows"]):
        r = 2 + i
        label = spec.get("label") or ""
        ws.cell(r, 1).value = label or None
        kind = spec.get("kind")
        bold = kind in {"operating", "total_profit", "actual_profit", "net_profit"} or label in {
            "收入",
            "成本",
            "利润总额",
            "净利润",
            "应交所得税",
            "实交所得税",
        }
        ws.cell(r, 1).font = bold_font if bold else normal
        if kind == "blank" or not label:
            continue
        au_code = spec.get("pl_au_code")
        if kind == "amount":
            for j, ent in enumerate(entities):
                cur = (current.get(ent) or {}).get(label)
                prev_v = (previous.get(ent) or {}).get(label)
                ws.cell(r, 2 + j).value = cell_num(cur)
                ws.cell(r, 10 + j).value = cell_num(prev_v)
                set_nf(ws.cell(r, 2 + j))
                set_nf(ws.cell(r, 10 + j))
        elif kind == "operating":
            # 收入-成本-销售-管理-研发-财务+其他收益-税金+资产处置+投资+公允价值
            parts = [
                ("收入", 1),
                ("成本", -1),
                ("销售费用", -1),
                ("管理费用", -1),
                ("研发费用", -1),
                ("财务费用", -1),
                ("+其他收益", 1),
                ("税金及附加", -1),
                ("+资产处置收益", 1),
                ("+投资收益", 1),
                ("+公允价值变动收益", 1),
            ]
            for col in cur_cols + prev_cols:
                letter = get_column_letter(col)
                bits = []
                for name, sign in parts:
                    rr = label_row.get(name)
                    if not rr:
                        continue
                    bits.append(("+" if sign > 0 else "-") + f"{letter}{rr}")
                expr = "".join(bits)
                if expr.startswith("+"):
                    expr = expr[1:]
                ws.cell(r, col).value = "=" + expr
                set_nf(ws.cell(r, col))
                ws.cell(r, col).font = bold_font
        elif kind == "total_profit":
            a = label_row["营业利润"]
            b = label_row["+营业外收入"]
            c = label_row["-营业外支出"]
            for col in cur_cols + prev_cols:
                L = get_column_letter(col)
                ws.cell(r, col).value = f"={L}{a}+{L}{b}-{L}{c}"
                set_nf(ws.cell(r, col))
                ws.cell(r, col).font = bold_font
        elif kind == "actual_profit":
            a = label_row["利润总额"]
            b = label_row["以前年度亏损"]
            for col in cur_cols + prev_cols:
                L = get_column_letter(col)
                ws.cell(r, col).value = f"={L}{a}+{L}{b}"
                set_nf(ws.cell(r, col))
                ws.cell(r, col).font = bold_font
        elif kind == "net_profit":
            a = label_row["实际利润额"]
            b = label_row["当年计提所得税"]
            for col in cur_cols + prev_cols:
                L = get_column_letter(col)
                ws.cell(r, col).value = f"={L}{a}-{L}{b}"
                set_nf(ws.cell(r, col))
                ws.cell(r, col).font = bold_font
        elif kind == "tax_paid":
            a = label_row.get("应交所得税")
            b = label_row.get("减免所得税")
            if a and b:
                for col in cur_cols + prev_cols:
                    L = get_column_letter(col)
                    ws.cell(r, col).value = f"={L}{a}-{L}{b}"
                    set_nf(ws.cell(r, col))
        elif kind == "cost_ratio":
            inc, cost = label_row["收入"], label_row["成本"]
            for col in cur_cols + prev_cols:
                L = get_column_letter(col)
                ws.cell(r, col).value = f'=IF(N({L}{inc})=0,"",{L}{cost}/{L}{inc})'
                set_nf(ws.cell(r, col))
        elif kind == "gross_margin":
            inc, cost = label_row["收入"], label_row["成本"]
            for col in cur_cols + prev_cols:
                L = get_column_letter(col)
                ws.cell(r, col).value = f'=IF(N({L}{inc})=0,"",({L}{inc}-{L}{cost})/{L}{inc})'
                set_nf(ws.cell(r, col))
        elif kind == "change":
            base = spec.get("change_of")
            br = label_row.get(base or "")
            if br:
                for j in range(8):
                    cur_l = get_column_letter(2 + j)
                    prev_l = get_column_letter(10 + j)
                    ws.cell(r, 2 + j).value = (
                        f'=IF(N({prev_l}{br})=0,"",({cur_l}{br}-{prev_l}{br})/{prev_l}{br})'
                    )
                    set_nf(ws.cell(r, 2 + j))

        if kind not in {"cost_ratio", "gross_margin", "change", "blank"}:
            ws.cell(r, 18).value = "=" + "+".join(f"{get_column_letter(c)}{r}" for c in cur_cols)
            set_nf(ws.cell(r, 18))
        if au_code and au_code in row_of:
            ws.cell(r, 19).value = f"=损益表!AU{row_of[au_code]}"
            set_nf(ws.cell(r, 19))
        elif kind == "operating":
            # 对照损益表大类：收入-成本-费用+投资等 vs 当月合计
            s_row = label_row
            ws.cell(r, 19).value = (
                f"=S{s_row['收入']}-S{s_row['成本']}-S{s_row['销售费用']}-S{s_row['管理费用']}"
                f"-S{s_row['研发费用']}-S{s_row['财务费用']}+S{s_row['+其他收益']}-S{s_row['税金及附加']}"
                f"+S{s_row['+资产处置收益']}+S{s_row['+投资收益']}+S{s_row['+公允价值变动收益']}"
            )
            set_nf(ws.cell(r, 19))
        elif kind == "total_profit":
            ws.cell(r, 19).value = f"=S{label_row['营业利润']}+S{label_row['+营业外收入']}-S{label_row['-营业外支出']}"
            set_nf(ws.cell(r, 19))
        if ws.cell(r, 19).value:
            ws.cell(r, 20).value = f"=R{r}-S{r}"
            set_nf(ws.cell(r, 20))


def map_profit_dict(raw: dict, layout: dict) -> dict:
    aliases = layout.get("profit_item_aliases") or {}
    out = {}
    inv = {}
    for label, names in aliases.items():
        inv[label] = label
        for n in names:
            inv[str(n).strip()] = label
    for key, val in (raw or {}).items():
        label = inv.get(str(key).strip())
        if label:
            out[label] = val if isinstance(val, Decimal) else money(val)
    return out


def build_workbook(period: str, entity_amts, dept_amts, profit_cur, profit_prev, layout) -> Workbook:
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    write_pl_sheet(wb, layout, entity_amts, dept_amts)
    write_profit_sheet(wb, layout, period, profit_cur, profit_prev)
    return wb


def count_nonzero_checks(wb, layout) -> int:
    book, _ = eval_workbook(wb)
    row_of = account_row_map(layout)
    n = 0
    av = 48  # AV
    for code, r in row_of.items():
        val = book.cell_value("损益表", r, av)
        if val is None or val == "":
            continue
        try:
            from decimal import Decimal

            if abs(Decimal(str(val))) > CHECK_TOL:
                n += 1
        except Exception:
            n += 1
    return n


def write_report(path: Path, payload: dict) -> None:
    lines = [
        f"期间={payload['period']}",
        f"有源账套={','.join(payload['has_source']) or '无'}",
        f"缺源账套={','.join(payload['missing']) or '无'}",
        f"未映射部门个数={payload['unmapped_count']}",
        f"核对非0科目编码个数={payload['nonzero_checks']}",
        f"产物={payload['out']}",
        f"status={payload['status']}",
    ]
    if payload.get("unmapped_names"):
        lines.append("未映射部门=" + "、".join(payload["unmapped_names"]))
    if payload.get("notes"):
        lines.append("说明=" + "；".join(payload["notes"]))
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def run(period: str, input_dir: Path, out: Path, no_api: bool) -> int:
    layout = load_layout()
    books = load_books()
    mapping = load_dept_map()
    notes = []
    secrets: list[str] = []
    inspected = inspect_dir(input_dir) if input_dir.is_dir() else []
    parsed = parse_inspected(inspected) if inspected else {"accounts": {}, "depts": [], "profits": {}}
    entity_amts: dict = {e: {} for e in layout["entities"]}
    for ent, codes in (parsed.get("accounts") or {}).items():
        entity_amts.setdefault(ent, {})
        for code, pair in codes.items():
            if code not in {a["code"] for a in layout["accounts"]}:
                notes.append(f"表外科目={code}")
                continue
            entity_amts[ent][code] = pair

    api_depts = []
    profit_cur = {e: {} for e in layout["entities"]}
    for ent, raw in (parsed.get("profits") or {}).items():
        profit_cur[ent] = map_profit_dict(raw, layout)

    if not no_api:
        try:
            from kingdee_client import load_client

            creds = load_client()
        except Exception:
            creds = None
        if creds:
            for key in ("client_secret", "app_secret"):
                val = str(creds.get(key) or "")
                if val:
                    secrets.append(val)
            try:
                from fetch_reports import fetch_hq

                hq = fetch_hq(period, creds)
                notes.extend(hq.get("notes") or [])
                entity_amts.setdefault("甲骨易", {})
                for code, pair in (hq.get("accounts") or {}).items():
                    if code in {a["code"] for a in layout["accounts"]}:
                        entity_amts["甲骨易"][code] = pair
                api_depts.extend(hq.get("depts") or [])
                mapped = hq.get("profit") or {}
                if mapped:
                    profit_cur["甲骨易"].update(mapped)
                try:
                    prev = prev_period(period)
                    hq_prev = fetch_hq(prev, creds, ledger=False)
                    prev_map = hq_prev.get("profit") or {}
                except Exception:
                    prev_map = {}
            except Exception as e:
                notes.append(f"api={type(e).__name__}")
                prev_map = {}
        else:
            notes.append("无密钥")
            prev_map = {}
    else:
        notes.append("skip-api")
        prev_map = {}

    dept_rows = list(parsed.get("depts") or []) + api_depts
    report_tmp: dict = {}
    dept_amts = merge_dept_rows(dept_rows, layout, mapping, report_tmp)

    prev_xlsx = input_dir / f"月度损益表_{prev_period(period)}.xlsx"
    profit_prev = load_prev_profit(prev_xlsx if prev_xlsx.is_file() else None, layout)
    if prev_map:
        profit_prev.setdefault("甲骨易", {}).update(prev_map)
    if not any(profit_prev.get(e) for e in layout["entities"]):
        notes.append("无上月列")

    xingchen = [a["excel_header"] for a in books.get("xingchen_accounts") or []]
    has_source = []
    missing = []
    for header in layout["entities"]:
        has = bool(entity_amts.get(header)) or bool(profit_cur.get(header))
        if has:
            has_source.append(header)
        else:
            missing.append(header)

    wb = build_workbook(period, entity_amts, dept_amts, profit_cur, profit_prev, layout)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    nonzero = count_nonzero_checks(load_workbook(out, data_only=False), layout)
    incomplete = [h for h in xingchen if h not in has_source]
    status = "ok" if not incomplete else "incomplete"
    payload = {
        "period": period,
        "has_source": has_source,
        "missing": missing,
        "unmapped_count": len(report_tmp.get("unmapped_depts") or []),
        "unmapped_names": report_tmp.get("unmapped_depts") or [],
        "nonzero_checks": nonzero,
        "out": str(out.resolve()),
        "status": status,
        "notes": notes,
    }
    report_path = out.with_name(out.stem + "_运行报告.txt")
    write_report(report_path, payload)
    text = (
        f"期间={period}\n"
        f"有源账套={','.join(has_source) or '无'}\n"
        f"缺源账套={','.join(missing) or '无'}\n"
        f"未映射部门个数={payload['unmapped_count']}\n"
        f"核对非0科目编码个数={nonzero}\n"
        f"产物={out.resolve()}\n"
        f"status={status}\n"
    )
    print(stdout_safe(text, secrets), end="", flush=True)
    return 0 if status == "ok" else 2


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--period", default="")
    parser.add_argument("--input-dir", default="")
    parser.add_argument("--out", default="")
    parser.add_argument("--no-api", action="store_true")
    args = parser.parse_args(argv)
    period = parse_period(args.period)
    input_dir = Path(args.input_dir).expanduser() if args.input_dir else Path.cwd()
    if args.out:
        out = Path(args.out).expanduser()
    else:
        desktop = Path.home() / "Desktop"
        target_dir = desktop if desktop.is_dir() else input_dir
        out = target_dir / f"月度损益表_{period}.xlsx"
    return run(period, input_dir, out, bool(args.no_api) or os.environ.get("PL_DEPT_SKIP_API") == "1")


if __name__ == "__main__":
    raise SystemExit(main())
