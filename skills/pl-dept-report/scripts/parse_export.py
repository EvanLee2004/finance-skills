#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from __future__ import annotations

import re
import sys
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
if str(HERE) not in sys.path:
    sys.path.insert(0, str(HERE))

from common import money, add_money, col_idx
from inspect_inputs import header_map
from layout import load_books, load_export_aliases, load_layout


def _norm_code(value) -> str:
    if value is None:
        return ""
    text = str(value).replace("\xa0", "").strip()
    if text.endswith(".0") and text.replace(".", "", 1)[:-1].isdigit():
        text = text[:-2]
    return text


def parse_account_sheet(ws, headers: dict) -> dict[str, dict]:
    code_at = headers.get("account_code")
    debit_at = headers.get("period_debit")
    credit_at = headers.get("period_credit")
    name_at = headers.get("account_name")
    if not code_at:
        return {}
    start_row = code_at[0] + 1
    out: dict[str, dict] = {}
    for r in range(start_row, (ws.max_row or start_row) + 1):
        code = _norm_code(ws.cell(r, code_at[1]).value)
        if not code or not code[0].isdigit():
            continue
        debit = money(ws.cell(r, debit_at[1]).value) if debit_at else None
        credit = money(ws.cell(r, credit_at[1]).value) if credit_at else None
        name = ""
        if name_at:
            name = str(ws.cell(r, name_at[1]).value or "").replace("\xa0", "").strip()
        if code not in out:
            out[code] = {"debit": debit, "credit": credit, "name": name}
        else:
            from common import add_money

            out[code]["debit"] = add_money(out[code]["debit"], debit)
            out[code]["credit"] = add_money(out[code]["credit"], credit)
            if name and not out[code].get("name"):
                out[code]["name"] = name
    return out


def parse_assist_sheet(ws, headers: dict) -> list[dict]:
    code_at = headers.get("account_code")
    dept_at = headers.get("dept_name")
    debit_at = headers.get("period_debit")
    credit_at = headers.get("period_credit")
    name_at = headers.get("account_name")
    if not code_at or not dept_at:
        return []
    start_row = max(code_at[0], dept_at[0]) + 1
    rows = []
    for r in range(start_row, (ws.max_row or start_row) + 1):
        code = _norm_code(ws.cell(r, code_at[1]).value)
        dept = str(ws.cell(r, dept_at[1]).value or "").replace("\xa0", "").strip()
        if not code or not dept:
            continue
        name = ""
        if name_at:
            name = str(ws.cell(r, name_at[1]).value or "").replace("\xa0", "").strip()
        rows.append(
            {
                "code": code,
                "name": name,
                "dept": dept,
                "debit": money(ws.cell(r, debit_at[1]).value) if debit_at else None,
                "credit": money(ws.cell(r, credit_at[1]).value) if credit_at else None,
            }
        )
    return rows


def parse_profit_sheet(ws, headers: dict, aliases: dict) -> dict[str, Decimal | None]:
    item_at = headers.get("profit_item")
    amt_at = headers.get("profit_month")
    if not amt_at:
        return {}
    start_row = (item_at[0] if item_at else amt_at[0]) + 1
    item_col = item_at[1] if item_at else 1
    out: dict[str, Decimal | None] = {}
    for r in range(start_row, (ws.max_row or start_row) + 1):
        label = str(ws.cell(r, item_col).value or "").replace("\xa0", "").strip()
        if not label:
            continue
        out[label] = money(ws.cell(r, amt_at[1]).value)
    return out


def _header_entity(text: str, entities: list[str], aliases: dict) -> str | None:
    blob = str(text or "")
    for ent in sorted(entities, key=len, reverse=True):
        if ent and ent in blob:
            return ent
    for short, ent in aliases.items():
        if short and short in blob:
            return ent
    return None


def parse_monthly_pl_sheet(ws, layout: dict) -> dict[str, dict[str, dict]]:
    entities = layout["entities"]
    pairs = layout.get("entity_pairs") or []
    out: dict[str, dict[str, dict]] = {}
    for r in range(3, (ws.max_row or 2) + 1):
        code = _norm_code(ws.cell(r, 1).value)
        if not code or not code[0].isdigit():
            continue
        name = str(ws.cell(r, 2).value or "").replace("\xa0", "").strip()
        for pair in pairs:
            header = pair.get("header")
            if header not in entities:
                continue
            dcol = col_idx(pair["debit"])
            ccol = col_idx(pair["credit"])
            debit = money(ws.cell(r, dcol).value)
            credit = money(ws.cell(r, ccol).value)
            if debit is None and credit is None:
                continue
            bucket = out.setdefault(header, {})
            if code not in bucket:
                bucket[code] = {"debit": debit, "credit": credit, "name": name}
            else:
                bucket[code]["debit"] = add_money(bucket[code]["debit"], debit)
                bucket[code]["credit"] = add_money(bucket[code]["credit"], credit)
    return out


def parse_monthly_profit_sheet(ws, layout: dict, books: dict, period: str) -> dict[str, dict[str, dict]]:
    entities = layout["entities"]
    aliases = books.get("profit_header_aliases") or {}
    labels = [row.get("label") for row in layout.get("profit_rows") or [] if row.get("label")]
    by_period: dict[str, dict[str, dict]] = {}
    header_map_ent: dict[int, tuple[str, str]] = {}
    for c in range(2, (ws.max_column or 1) + 1):
        text = str(ws.cell(1, c).value or "")
        ent = _header_entity(text, entities, aliases)
        if not ent:
            continue
        m = None
        found = re.search(r"(20)?(\d{2})年\s*(\d{1,2})\s*月", text)
        if found:
            year = int(found.group(2))
            year = 2000 + year if year < 100 else year
            month = int(found.group(3))
            m = f"{year}{month:02d}"
        header_map_ent[c] = (ent, m or period)
    for r in range(2, (ws.max_row or 1) + 1):
        label = str(ws.cell(r, 1).value or "").replace("\xa0", "").strip()
        if not label or label not in labels:
            continue
        for c, (ent, per) in header_map_ent.items():
            val = money(ws.cell(r, c).value)
            if val is None:
                continue
            by_period.setdefault(per, {}).setdefault(ent, {})[label] = val
    return by_period


def parse_inspected(items: list[dict]) -> dict:
    aliases = load_export_aliases()
    accounts: dict[str, dict[str, dict[str, Decimal | None]]] = {}
    depts: list[dict] = []
    profits: dict[str, dict[str, Decimal | None]] = {}
    profits_by_period: dict[str, dict[str, dict[str, Decimal | None]]] = {}
    extra_codes: set[str] = set()
    monthly_accounts: dict[str, dict[str, dict]] = {}
    monthly_profits: dict[str, dict[str, dict]] = {}
    layout = load_layout()
    books = load_books()
    for item in items:
        path = item["path"]
        data_only = item.get("kind") in {"monthly_overlay", "monthly_profit"}
        wb = load_workbook(path, data_only=data_only)
        try:
            ws = wb[item["sheet"]]
            headers = {k: (v["row"], v["col"]) for k, v in (item.get("headers") or {}).items()}
            if not headers:
                headers = header_map(ws, aliases)
            entity = item.get("entity")
            kind = item["kind"]
            if kind in {"monthly_overlay", "monthly_profit", "agency_profit"}:
                continue
            if kind == "account":
                parsed = parse_account_sheet(ws, headers)
                if entity:
                    bucket = accounts.setdefault(entity, {})
                    for code, pair in parsed.items():
                        if code not in bucket:
                            bucket[code] = pair
                        else:
                            bucket[code] = {
                                "debit": add_money(bucket[code]["debit"], pair["debit"]),
                                "credit": add_money(bucket[code]["credit"], pair["credit"]),
                            }
            elif kind == "assist":
                for row in parse_assist_sheet(ws, headers):
                    row["entity"] = entity
                    depts.append(row)
            elif kind == "profit" and entity:
                parsed = parse_profit_sheet(ws, headers, aliases)
                profits[entity] = parsed
                period = item.get("period")
                if period:
                    profits_by_period.setdefault(period, {})[entity] = parsed
        finally:
            wb.close()
    return {
        "accounts": accounts,
        "depts": depts,
        "profits": profits,
        "profits_by_period": profits_by_period,
        "monthly_accounts": monthly_accounts,
        "monthly_profits": monthly_profits,
        "extra_codes": sorted(extra_codes),
    }
