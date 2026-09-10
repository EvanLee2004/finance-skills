#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""本机星辰引出：切账套 → 科目余额 / 核算项目 / 利润表查询页 → xlsx。

利润表走「财务报表 → 利润表」查询页，选会计期间再查本月金额。
不要走「已生成报表列表」当唯一路径，不要点新增 / 生成 / 引入 / 审核 / 过账。
口令只读 ~/.config/finance/xingchen.local.json，会话写 playwright-state。不打印、不进 git。
"""
from __future__ import annotations

import argparse
import asyncio
import re
import sys
import time
from pathlib import Path
from urllib.parse import parse_qs, urlparse

HERE = Path(__file__).resolve().parent
if str(HERE) not in sys.path:
    sys.path.insert(0, str(HERE))

from common import detect_period_text
from inspect_inputs import inspect_file
from layout import load_books
from xingchen_login import STATE_PATH as STATE
from xingchen_login import ensure_login, load_creds, save_state, state_path

HOME = Path.home()
XINGCHEN_HOME = "https://tf.jdy.com/ierp/index.html?formId=home_page"
ASSIST_FORM = "https://tf.jdy.com/ierp/index.html?formId=gl_rpt_assistbalance"
ACCT_FORM = "https://tf.jdy.com/ierp/index.html?formId=gl_rpt_acctbalance"
PROFIT_FORM = "https://tf.jdy.com/ierp/index.html?formId=gl_profitsheet1"
DOWNLOADS = HOME / "Downloads"
FORBIDDEN_CLICKS = frozenset(
    {
        "新增",
        "生成",
        "生成报表",
        "审核",
        "过账",
        "引入",
        "购买",
        "加购",
        "上报",
        "续费",
    }
)
BOOK_SEARCH = {
    "wenhua": "文化传媒",
    "shanghai": "智译",
    "hunan_fgs": "湖南分公司",
    "hunan_zgs": "湖南子公司",
}


def period_label(period: str) -> str:
    year = int(period[:4])
    month = int(period[4:6])
    return f"{year}年{month}期"


def period_labels(period: str) -> list[str]:
    year = int(period[:4])
    month = int(period[4:6])
    out = [f"{year}年{month}期", f"{year}年{month:02d}期"]
    return list(dict.fromkeys(out))


def profit_dest_name(out_dir: Path, header: str, period: str) -> Path:
    return Path(out_dir) / f"{header}_利润表_{period}.xlsx"


def account_dest_name(out_dir: Path, header: str, period: str) -> Path:
    return Path(out_dir) / f"{header}_科目余额表_{period}.xlsx"


def assist_dest_name(out_dir: Path, header: str, period: str) -> Path:
    return Path(out_dir) / f"{header}_核算项目余额表_{period}.xlsx"


def _export_books() -> list[tuple[str, str, str]]:
    data = load_books()
    out = []
    for acc in data.get("xingchen_accounts") or []:
        if acc.get("api_ready_default"):
            continue
        key = str(acc.get("key") or "")
        legal = str(acc.get("legal_name") or "")
        header = str(acc.get("excel_header") or key)
        if key and legal:
            out.append((key, legal, header))
    return out


def _form_id(url: str) -> str:
    return (parse_qs(urlparse(url).query).get("formId") or [""])[0]


def _denied(body: str) -> bool:
    return "无" in body and "权限" in body and ("个别报表" in body or "查询权限" in body)


async def _ensure_login(page) -> None:
    creds = load_creds()
    if not creds and not state_path().is_file():
        raise SystemExit("missing_creds")
    await ensure_login(page, creds)


async def _dismiss_overlays(page) -> None:
    await page.evaluate(
        """() => {
          document.querySelectorAll('iframe').forEach(f => {
            const src = f.src || '';
            if (/piaozone|fpdk|etax/.test(src)) (f.closest('div') || f).remove();
          });
        }"""
    )


async def _click_text(page, label: str) -> bool:
    if label in FORBIDDEN_CLICKS or any(bad in label for bad in ("新增", "生成报表", "审核", "过账")):
        return False
    return bool(
        await page.evaluate(
            """(label) => {
              const walker = document.createTreeWalker(document.body, NodeFilter.SHOW_TEXT);
              let n;
              while ((n = walker.nextNode())) {
                if ((n.textContent || '').trim() === label) {
                  const el = n.parentElement;
                  (el?.closest('label,li,button,div,span,a') || el)?.click();
                  return true;
                }
              }
              return false;
            }""",
            label,
        )
    )


async def _switch_book(page, legal_name: str, token: str = "") -> bool:
    await _dismiss_overlays(page)
    await page.goto(XINGCHEN_HOME, wait_until="domcontentloaded")
    await page.wait_for_timeout(1600)
    await _dismiss_overlays(page)
    current = await page.locator("body").inner_text()
    if legal_name in current[:5000] and "输入账套名称" not in current:
        return True
    handle = page.locator("span").filter(has_text=re.compile("甲骨易|文化传媒|智译|湖南"))
    if await handle.count():
        await handle.last.click(force=True)
        await page.wait_for_timeout(700)
    box = page.get_by_placeholder("输入账套名称或搜索空间")
    needle = token or legal_name[:8]
    if await box.count():
        await box.first.fill(needle)
        await page.wait_for_timeout(500)
    target = page.get_by_text(legal_name, exact=True)
    if not await target.count():
        target = page.get_by_text(legal_name, exact=False)
    if not await target.count():
        return False
    await target.last.click(force=True)
    await page.wait_for_timeout(2500)
    return legal_name in await page.locator("body").inner_text()


async def _visible_period(page) -> str:
    return str(
        await page.evaluate(
            """() => {
              const re = /20\\d{2}年\\s*\\d{1,2}期/;
              const inputs = [...document.querySelectorAll('input')].filter(el => el.offsetParent !== null);
              for (const el of inputs) {
                const t = (el.value || '').trim();
                const m = t.match(re);
                if (m) return m[0].replace(/\\s+/g, '');
              }
              const walker = document.createTreeWalker(document.body, NodeFilter.SHOW_TEXT);
              let n;
              while ((n = walker.nextNode())) {
                const t = (n.textContent || '').trim();
                const m = t.match(re);
                if (!m) continue;
                const el = n.parentElement;
                if (!el || el.offsetParent === null) continue;
                return m[0].replace(/\\s+/g, '');
              }
              return '';
            }"""
        )
        or ""
    )


def _period_match(shown: str, period: str) -> bool:
    compact = (shown or "").replace(" ", "")
    return any(label.replace(" ", "") in compact for label in period_labels(period))


def period_inputs_match(values: list[str], period: str) -> bool:
    """从–至两格都要是当期。只看第一格会把 08–09 当成选对。"""
    cleaned = [str(v).replace(" ", "") for v in values if v and "年" in str(v) and "期" in str(v)]
    if not cleaned:
        return False
    return all(_period_match(v, period) for v in cleaned)


async def _open_period_picker(page) -> bool:
    return bool(
        await page.evaluate(
            """() => {
              const re = /^20\\d{2}年\\s*\\d{1,2}期$/;
              const walker = document.createTreeWalker(document.body, NodeFilter.SHOW_TEXT);
              let n;
              while ((n = walker.nextNode())) {
                const t = (n.textContent || '').trim();
                if (!re.test(t)) continue;
                const el = n.parentElement;
                if (!el || el.offsetParent === null) continue;
                (el.closest('span,div,label,a,button') || el).click();
                return true;
              }
              return false;
            }"""
        )
    )


async def _period_input_values(page) -> list[str]:
    return list(
        await page.evaluate(
            """() => {
              const re = /20\\d{2}年\\s*\\d{1,2}期/;
              return [...document.querySelectorAll('input')]
                .filter(el => el.offsetParent !== null)
                .map(el => (el.value || '').trim())
                .filter(t => re.test(t));
            }"""
        )
        or []
    )


async def _all_period_inputs_match(page, period: str) -> bool:
    values = await _period_input_values(page)
    if values:
        return period_inputs_match(values, period)
    return _period_match(await _visible_period(page), period)


async def _pick_period_chip(page, period: str) -> None:
    year = period[:4]
    month = int(period[4:6])
    chips = [f"{month}期", f"{month:02d}期", *period_labels(period)]
    year_hit = page.get_by_text(f"{year}年", exact=True)
    if await year_hit.count():
        try:
            if await year_hit.last.is_visible():
                await year_hit.last.click()
                await page.wait_for_timeout(150)
        except Exception:
            pass
    for chip in chips:
        hit = page.get_by_text(chip, exact=True)
        count = await hit.count()
        if not count:
            if await _click_text(page, chip):
                return
            continue
        # 科目余额弹出 2026 | 2027 两列，last 会点到明年。当期年在左边，点第一颗可见的。
        for i in range(count):
            try:
                loc = hit.nth(i)
                if await loc.is_visible():
                    await loc.click()
                    return
            except Exception:
                continue
        try:
            await hit.first.click()
            return
        except Exception:
            continue


async def _input_readonly(box) -> bool:
    try:
        return bool(await box.evaluate("el => !!el.readOnly"))
    except Exception:
        return False


async def _set_one_period_input(box, page, period: str) -> None:
    labels = period_labels(period)
    typed = labels[-1] if len(labels) > 1 else labels[0]
    readonly = await _input_readonly(box)
    try:
        await box.click(force=True)
        await page.wait_for_timeout(350)
    except Exception:
        return
    if not readonly:
        try:
            await box.fill(typed)
            await page.keyboard.press("Enter")
            await page.wait_for_timeout(150)
        except Exception:
            pass
        try:
            current = await box.input_value()
        except Exception:
            current = ""
        if _period_match(current, period):
            return
    await _pick_period_chip(page, period)
    await page.wait_for_timeout(200)


async def _set_period(page, period: str) -> bool:
    if await _all_period_inputs_match(page, period):
        return True
    boxes = page.locator("input:visible")
    n = await boxes.count()
    idxs: list[int] = []
    for i in range(min(n, 8)):
        try:
            current = await boxes.nth(i).input_value()
        except Exception:
            continue
        if "年" in current and "期" in current:
            idxs.append(i)
    if not idxs:
        await _open_period_picker(page)
        await page.wait_for_timeout(300)
        await _pick_period_chip(page, period)
        return await _all_period_inputs_match(page, period)
    for i in idxs:
        try:
            current = await boxes.nth(i).input_value()
        except Exception:
            continue
        if _period_match(current, period):
            continue
        await _set_one_period_input(boxes.nth(i), page, period)
        await page.wait_for_timeout(200)
    if await _all_period_inputs_match(page, period):
        return True
    await _open_period_picker(page)
    await page.wait_for_timeout(300)
    await _pick_period_chip(page, period)
    return await _all_period_inputs_match(page, period)


async def _filter_checked(page, label: str) -> bool | None:
    return await page.evaluate(
        """(label) => {
          const walker = document.createTreeWalker(document.body, NodeFilter.SHOW_TEXT);
          let n;
          while ((n = walker.nextNode())) {
            if ((n.textContent || '').trim() !== label) continue;
            let el = n.parentElement;
            for (let i = 0; i < 3 && el; i += 1) {
              const html = el.innerHTML || '';
              if (html.includes('kdfont-fuxuankuangxuanzhong_fang')) return true;
              if (html.includes('kdfont-fuxuankuangweixuanzhong_fang')) return false;
              el = el.parentElement;
            }
            return null;
          }
          return null;
        }""",
        label,
    )


async def _set_filter_check(page, label: str, want: bool) -> None:
    state = await _filter_checked(page, label)
    if state is want:
        return
    if await page.get_by_text(label, exact=True).count():
        try:
            await page.get_by_text(label, exact=True).first.click(force=True)
            await page.wait_for_timeout(120)
            return
        except Exception:
            pass
    await _click_text(page, label)


async def _set_acct_filters(page, period: str) -> None:
    if not await _set_period(page, period):
        shown = ",".join(await _period_input_values(page) or [await _visible_period(page) or "-"])
        raise RuntimeError(f"period_not_set:{shown}")
    await _click_text(page, "展开过滤")
    await page.wait_for_timeout(400)
    await _set_filter_check(page, "展开所有级次", True)
    await _set_filter_check(page, "显示所有科目", True)
    await _set_filter_check(page, "无发生额且余额为0不显示", False)
    await _set_filter_check(page, "余额为0不显示", False)
    await _set_filter_check(page, "无发生额不显示", False)
    await _click_text(page, "查询")
    await page.wait_for_timeout(2500)


async def _set_assist_dept(page, period: str) -> None:
    if not await _set_period(page, period):
        shown = ",".join(await _period_input_values(page) or [await _visible_period(page) or "-"])
        raise RuntimeError(f"period_not_set:{shown}")
    box = page.locator(".kd-table-cell-basedata-container").first
    if await box.count():
        await box.click()
        await page.keyboard.type("部门")
        await page.wait_for_timeout(800)
        if await page.get_by_text("部门", exact=True).count():
            await page.get_by_text("部门", exact=True).last.click()
            await page.keyboard.press("Enter")
            await page.wait_for_timeout(400)
    await _click_text(page, "查询")
    await page.wait_for_timeout(2500)


def _newest_xlsx(folder: Path, after: float) -> Path | None:
    newest = None
    newest_mtime = after
    if not folder.is_dir():
        return None
    for path in folder.glob("*.xlsx"):
        if path.name.startswith("~$"):
            continue
        try:
            mtime = path.stat().st_mtime
        except OSError:
            continue
        if mtime > newest_mtime:
            newest = path
            newest_mtime = mtime
    return newest


async def _export_current(page, dest: Path) -> bool:
    dest.parent.mkdir(parents=True, exist_ok=True)
    mark = time.time()
    try:
        hit = page.get_by_text("引出", exact=True)
        if await hit.count():
            await hit.first.click()
        else:
            await _click_text(page, "引出")
        await page.wait_for_timeout(500)
        if await page.get_by_text("到引出结果界面下载").count():
            async with page.expect_download(timeout=15000) as dl_info:
                await page.get_by_text("到引出结果界面下载").first.click()
            download = await dl_info.value
        else:
            async with page.expect_download(timeout=15000) as dl_info:
                confirm = page.get_by_text("引出", exact=True)
                if await confirm.count() >= 2:
                    await confirm.last.click()
                elif await page.get_by_text("Excel", exact=True).count():
                    await page.get_by_text("Excel", exact=True).first.click()
                    await confirm.last.click()
                else:
                    await _click_text(page, "引出")
            download = await dl_info.value
        await download.save_as(str(dest))
        if dest.is_file() and dest.stat().st_size > 400:
            return True
    except Exception:
        pass
    found = _newest_xlsx(DOWNLOADS, mark - 1)
    if found:
        dest.write_bytes(found.read_bytes())
        return dest.is_file() and dest.stat().st_size > 400
    return dest.is_file() and dest.stat().st_size > 400


async def _open_profit_query(page) -> str:
    await _dismiss_overlays(page)
    await page.goto(PROFIT_FORM, wait_until="domcontentloaded")
    await page.wait_for_timeout(2200)
    await _dismiss_overlays(page)
    return page.url


def _xlsx_period(path: Path) -> str | None:
    if not path.is_file():
        return None
    for item in inspect_file(path):
        if item.get("period"):
            return str(item["period"])
    return detect_period_text(path.name)


def _profit_file_period(path: Path) -> str | None:
    if not path.is_file():
        return None
    for item in inspect_file(path):
        if item.get("kind") == "profit" and item.get("period"):
            return str(item["period"])
    try:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=False)
        try:
            blob = "\n".join(
                str(ws.cell(r, c).value or "")
                for ws in wb.worksheets
                for r in range(1, min(6, (ws.max_row or 1) + 1))
                for c in range(1, min(6, (ws.max_column or 1) + 1))
            )
        finally:
            wb.close()
        return detect_period_text(blob + "\n" + path.name)
    except Exception:
        return detect_period_text(path.name)


async def _export_profit_query(page, dest: Path, period: str) -> tuple[bool, str]:
    url = await _open_profit_query(page)
    body = await page.locator("body").inner_text()
    if _denied(body):
        return False, "denied"
    if "本月金额" not in body and "本月数" not in body:
        return False, f"no_query_page:{_form_id(url) or '-'}"
    if not await _set_period(page, period):
        return False, f"period_not_set:{await _visible_period(page) or '-'}"
    await _click_text(page, "查询")
    await _click_text(page, "刷新")
    await page.wait_for_timeout(2500)
    body = await page.locator("body").inner_text()
    if _denied(body):
        return False, "denied"
    if not _period_match(await _visible_period(page), period):
        return False, f"wrong_period={await _visible_period(page) or '-'}"
    tmp = dest.with_name(dest.stem + "._tmp.xlsx")
    if not await _export_current(page, tmp):
        return False, "export_failed"
    got = _profit_file_period(tmp)
    if got and got != period:
        alt = dest.with_name(f"{dest.name.split('_利润表')[0]}_利润表_{got}.xlsx")
        if alt != tmp:
            alt.write_bytes(tmp.read_bytes())
        tmp.unlink(missing_ok=True)
        return False, f"wrong_period={got}"
    dest.write_bytes(tmp.read_bytes())
    tmp.unlink(missing_ok=True)
    return True, "ok"


async def export_book(
    page,
    key: str,
    legal: str,
    header: str,
    period: str,
    out_dir: Path,
    kinds: list[str] | None = None,
) -> dict:
    wanted = set(kinds or ["account", "assist", "profit"])
    note = {"book": key, "header": header, "ok": False, "files": [], "kinds": sorted(wanted)}
    token = BOOK_SEARCH.get(key) or header
    if not await _switch_book(page, legal, token):
        note["reason"] = "switch_failed"
        return note
    jobs = []
    if "account" in wanted:
        jobs.append(("account", ACCT_FORM, account_dest_name(out_dir, header, period), _set_acct_filters))
    if "assist" in wanted:
        jobs.append(("assist", ASSIST_FORM, assist_dest_name(out_dir, header, period), _set_assist_dept))
    for kind, url, path, setup in jobs:
        try:
            await _dismiss_overlays(page)
            await page.goto(url, wait_until="domcontentloaded")
            await page.wait_for_timeout(2000)
            await _dismiss_overlays(page)
            await setup(page, period)
            if not await _all_period_inputs_match(page, period):
                shown = ",".join(await _period_input_values(page) or [await _visible_period(page) or "-"])
                raise RuntimeError(f"period_not_set:{shown}")
            body = await page.locator("body").inner_text()
            if "暂无数据" in body:
                note.setdefault("empty", []).append(kind)
                continue
            if await _export_current(page, path):
                got = _xlsx_period(path)
                if got and got != period:
                    note.setdefault("errors", []).append(f"{kind}:wrong_period={got}")
                    path.unlink(missing_ok=True)
                    note.setdefault("empty", []).append(kind)
                else:
                    note["files"].append(path.name)
            else:
                note.setdefault("empty", []).append(kind)
        except Exception as e:
            detail = str(e).split(":")[0][:24]
            note.setdefault("errors", []).append(f"{kind}:{type(e).__name__}:{detail}")
    if "profit" in wanted:
        dest = profit_dest_name(out_dir, header, period)
        try:
            ok, reason = await _export_profit_query(page, dest, period)
            if ok:
                note["files"].append(dest.name)
            else:
                note.setdefault("empty", []).append("profit")
                if reason == "denied":
                    note["denied"] = True
                note.setdefault("errors", []).append(f"profit:{reason}")
        except Exception as e:
            note.setdefault("errors", []).append(f"profit:{type(e).__name__}")
    note["ok"] = bool(note["files"])
    return note


async def _open_context(p):
    browser = await p.chromium.launch(headless=True)
    kwargs = {
        "accept_downloads": True,
        "locale": "zh-CN",
        "extra_http_headers": {"Accept-Language": "zh-CN,zh;q=0.9"},
    }
    if STATE.is_file():
        kwargs["storage_state"] = str(STATE)
    ctx = await browser.new_context(**kwargs)
    return ctx, browser, "fresh"


async def export_missing_async(
    period: str, out_dir: Path, keys: list[str], kinds: dict[str, list[str]] | None = None
) -> dict:
    out_dir.mkdir(parents=True, exist_ok=True)
    STATE.parent.mkdir(parents=True, exist_ok=True)
    wanted = set(keys)
    catalog = [(k, legal, header) for k, legal, header in _export_books() if k in wanted]
    summary = {"ok": [], "denied": [], "errors": [], "books": []}
    from playwright.async_api import async_playwright

    async with async_playwright() as p:
        ctx, browser, mode = await _open_context(p)
        page = ctx.pages[0] if ctx.pages else await ctx.new_page()
        await _ensure_login(page)
        await save_state(ctx)
        for key, legal, header in catalog:
            book_kinds = (kinds or {}).get(key) or ["account", "assist", "profit"]
            result = await export_book(page, key, legal, header, period, out_dir, book_kinds)
            summary["books"].append(result)
            if result.get("ok"):
                summary["ok"].append(key)
            if result.get("denied"):
                summary["denied"].append(header)
            for err in result.get("errors") or []:
                summary["errors"].append(f"{key}:{err}")
        await ctx.close()
        if browser:
            await browser.close()
    summary["mode"] = mode
    return summary


def export_missing(
    period: str, out_dir: Path, keys: list[str], kinds: dict[str, list[str]] | None = None
) -> dict:
    return asyncio.run(export_missing_async(period, Path(out_dir), keys, kinds))


async def main_async(period: str, out_dir: Path, keys: list[str]) -> int:
    summary = await export_missing_async(period, out_dir, keys)
    results = summary.get("books") or []
    ok = len(summary.get("ok") or [])
    print(f"period={period} books_ok={ok}/{len(results)}")
    for r in results:
        extra = r.get("reason") or ",".join(r.get("errors") or r.get("empty") or [])
        print(
            f"book={r['book']} ok={int(bool(r.get('ok')))} files={len(r.get('files') or [])} {extra}".strip()
        )
    return 0 if ok else 2


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--period", default="202608")
    parser.add_argument("--out-dir", required=True)
    parser.add_argument("--books", default="wenhua,shanghai,hunan_fgs,hunan_zgs")
    args = parser.parse_args()
    keys = [k.strip() for k in args.books.split(",") if k.strip()]
    return asyncio.run(main_async(args.period, Path(args.out_dir).expanduser(), keys))


if __name__ == "__main__":
    raise SystemExit(main())
