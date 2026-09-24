#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""材料夹分成三层：01_脏数据、02_清洗后/模型、03_结果。结果只读模型。"""
from __future__ import annotations

import shutil
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from common import cell_num, money

DIRTY_NAME = "01_脏数据"
CLEAN_NAME = "02_清洗后"
RESULT_NAME = "03_结果"
RESERVED = {DIRTY_NAME, CLEAN_NAME, RESULT_NAME}

KIND_LABEL = {
    "account": "科目余额",
    "assist": "部门核算",
    "profit": "利润表",
    "agency_profit": "代账利润表",
    "unreadable": "未认出",
}

_HEADER_FONT = Font(name="等线", size=11, bold=True)
_CELL_FONT = Font(name="宋体", size=10)


def model_path(input_dir: Path, period: str) -> Path:
    return Path(input_dir) / CLEAN_NAME / f"模型_{period}.xlsx"


def result_path(input_dir: Path, period: str) -> Path:
    return Path(input_dir) / RESULT_NAME / f"月度损益表_{period}.xlsx"


def _under(path: Path, root: Path) -> bool:
    try:
        path.resolve().relative_to(root.resolve())
        return True
    except ValueError:
        return False


def _same_bytes(left: Path, right: Path) -> bool:
    if left.stat().st_size != right.stat().st_size:
        return False
    return left.read_bytes() == right.read_bytes()


def _skipped_dir(name: str) -> bool:
    return name in RESERVED or name == "引出" or name.startswith("月度损益表_") or name.startswith("全源_")


def _place(src: Path, dest: Path, input_dir: Path) -> None:
    """夹子里的新表覆盖同名旧脏数据，比的是内容不是字节数。夹子外的原件只复制。"""
    if not src.is_file() or src.name.startswith("~$") or src.name.startswith("月度损益表_"):
        return
    if src.suffix.lower() not in {".xlsx", ".xlsm", ".xls"}:
        return
    if any(_skipped_dir(part) for part in src.resolve().relative_to(input_dir.resolve()).parts[:-1]) if _under(src, input_dir) else False:
        return
    dest.parent.mkdir(parents=True, exist_ok=True)
    if dest.resolve() == src.resolve():
        return
    inside = _under(src, input_dir)
    if dest.exists() and _same_bytes(src, dest):
        if inside:
            src.unlink()
        return
    if inside and not dest.exists():
        shutil.move(str(src), str(dest))
        return
    shutil.copy2(src, dest)
    if inside and src.exists() and src.resolve() != dest.resolve():
        src.unlink()


def find_staged(dirty: Path, raw: str) -> Path:
    """显式指定的文件收进脏数据后，按文件名找回，不要求还在夹子根上。"""
    name = Path(str(raw)).name
    direct = dirty / name
    if direct.is_file():
        return direct
    matches = [path for path in dirty.rglob(name) if path.is_file()]
    return matches[0] if matches else Path(str(raw))


def _rel_under_dirty(src: Path, input_dir: Path, dirty: Path) -> Path:
    if _under(src, input_dir):
        return dirty / src.resolve().relative_to(input_dir.resolve())
    return dirty / src.name


def stage_dirty(input_dir: Path, extra_files: list[str] | None = None) -> Path:
    """材料夹里各层的源表收进 01_脏数据，相对路径保留。夹子外的文件只复制，不挪走。"""
    input_dir = Path(input_dir)
    input_dir.mkdir(parents=True, exist_ok=True)
    dirty = input_dir / DIRTY_NAME
    dirty.mkdir(parents=True, exist_ok=True)
    (input_dir / CLEAN_NAME).mkdir(parents=True, exist_ok=True)
    (input_dir / RESULT_NAME).mkdir(parents=True, exist_ok=True)
    if input_dir.is_dir():
        for path in list(input_dir.rglob("*")):
            if not path.is_file():
                continue
            rel = path.relative_to(input_dir)
            if any(_skipped_dir(part) for part in rel.parts[:-1]):
                continue
            _place(path, dirty / rel, input_dir)
        for folder in sorted(
            (path for path in input_dir.rglob("*") if path.is_dir()),
            key=lambda path: len(path.parts),
            reverse=True,
        ):
            rel_parts = folder.relative_to(input_dir).parts
            if any(_skipped_dir(part) for part in rel_parts):
                continue
            try:
                folder.rmdir()
            except OSError:
                pass
    for raw in extra_files or []:
        path = Path(raw).expanduser()
        if path.is_file():
            _place(path, _rel_under_dirty(path, input_dir, dirty), input_dir)
        elif path.is_dir() and not _under(path, input_dir):
            for child in path.rglob("*"):
                if child.is_file():
                    _place(child, dirty / child.name, input_dir)
    return dirty


def _write_sheet(wb: Workbook, title: str, headers: list[str], rows: list[list]) -> None:
    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title)
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(1, col, header)
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(r, c, value)
            cell.font = _CELL_FONT
            if isinstance(value, float):
                cell.number_format = '#,##0.00'
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(rows) + 1)}"
    for col, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(col)].width = max(12, min(36, len(header) * 2 + 6))


def _blank_book() -> Workbook:
    wb = Workbook()
    default = wb.active
    if default is not None:
        wb.remove(default)
    return wb


def _save(wb: Workbook, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    tmp.replace(path)


def _unused_reason(item: dict, period: str) -> str:
    if item.get("period_conflict"):
        return "期间冲突"
    if item.get("kind") == "unreadable":
        return "表头认不出"
    got = str(item.get("period") or "")
    if got and got != period:
        return "不是本期"
    if not got:
        return "没写期间"
    return "未进入本期"


def write_source_model(
    path: Path,
    period: str,
    inspected_all: list[dict],
    inspected_used: list[dict],
    parsed: dict,
) -> Path:
    used = {
        (str(item.get("path") or ""), str(item.get("sheet") or ""), str(item.get("kind") or ""))
        for item in inspected_used
        if item.get("kind") in {"account", "assist", "profit", "agency_profit"}
    }
    source_rows = []
    for item in inspected_all:
        file_name = Path(str(item.get("path") or "")).name
        kind = str(item.get("kind") or "")
        key = (str(item.get("path") or ""), str(item.get("sheet") or ""), kind)
        if key in used and kind != "unreadable" and not item.get("period_conflict"):
            adopted, reason = "用了", ""
        else:
            adopted, reason = "未用", _unused_reason(item, period)
        source_rows.append(
            [
                file_name,
                KIND_LABEL.get(kind, kind or "未认出"),
                str(item.get("entity") or ""),
                str(item.get("period") or ""),
                adopted,
                reason,
            ]
        )
    account_rows = []
    for ent, codes in sorted((parsed.get("accounts") or {}).items()):
        for code, pair in sorted(codes.items()):
            debit = (pair or {}).get("debit")
            credit = (pair or {}).get("credit")
            if (debit is None or Decimal(str(debit)) == 0) and (credit is None or Decimal(str(credit)) == 0):
                continue
            account_rows.append(
                [
                    period,
                    ent,
                    code,
                    str((pair or {}).get("name") or ""),
                    cell_num((pair or {}).get("debit")),
                    cell_num((pair or {}).get("credit")),
                ]
            )
    dept_rows = []
    for row in parsed.get("depts") or []:
        debit = row.get("debit")
        credit = row.get("credit")
        if (debit is None or Decimal(str(debit)) == 0) and (credit is None or Decimal(str(credit)) == 0):
            continue
        dept_rows.append(
            [
                period,
                str(row.get("entity") or ""),
                str(row.get("code") or ""),
                str(row.get("name") or ""),
                str(row.get("dept") or ""),
                cell_num(row.get("debit")),
                cell_num(row.get("credit")),
            ]
        )
    profit_rows = []
    for per, by_ent in sorted((parsed.get("profits_by_period") or {}).items()):
        for ent, labels in sorted(by_ent.items()):
            for label, amt in labels.items():
                if amt is None or Decimal(str(amt)) == 0:
                    continue
                profit_rows.append([per, ent, label, cell_num(amt)])
    wb = _blank_book()
    _write_sheet(wb, "来源清单", ["文件名", "种类", "账套", "期间", "采用", "原因"], source_rows)
    _write_sheet(wb, "缺口", ["类别", "对象", "说明"], [])
    _write_sheet(wb, "科目发生", ["期间", "账套", "科目编码", "科目名称", "借方", "贷方"], account_rows)
    _write_sheet(wb, "部门核算", ["期间", "账套", "科目编码", "科目名称", "部门", "借方", "贷方"], dept_rows)
    _write_sheet(wb, "利润表", ["期间", "账套", "项目", "本月金额"], profit_rows)
    _write_sheet(wb, "薪酬", ["账套", "科目编码", "项目", "部门", "发生额列", "借方", "填左列"], [])
    _write_sheet(wb, "房租摘要", ["科目编码", "部门", "金额"], [])
    _write_sheet(wb, "代账利润表", ["期间", "账套", "项目", "本月金额"], [])
    _save(wb, path)
    return path


def _rows(path: Path, title: str) -> list[dict]:
    if not path.is_file():
        return []
    wb = load_workbook(path, data_only=True)
    try:
        if title not in wb.sheetnames:
            return []
        ws = wb[title]
        headers = [str(ws.cell(1, c).value or "") for c in range(1, (ws.max_column or 1) + 1)]
        out = []
        for r in range(2, (ws.max_row or 1) + 1):
            values = [ws.cell(r, c).value for c in range(1, len(headers) + 1)]
            if all(v in (None, "") for v in values):
                continue
            out.append({headers[i]: values[i] for i in range(len(headers)) if headers[i]})
        return out
    finally:
        wb.close()


def read_source_model(path: Path) -> dict:
    accounts: dict = {}
    for row in _rows(path, "科目发生"):
        ent = str(row.get("账套") or "")
        code = str(row.get("科目编码") or "").strip()
        if not ent or not code:
            continue
        accounts.setdefault(ent, {})[code] = {
            "name": str(row.get("科目名称") or ""),
            "debit": money(row.get("借方")),
            "credit": money(row.get("贷方")),
        }
    depts = []
    for row in _rows(path, "部门核算"):
        depts.append(
            {
                "entity": str(row.get("账套") or ""),
                "code": str(row.get("科目编码") or "").strip(),
                "name": str(row.get("科目名称") or ""),
                "dept": str(row.get("部门") or ""),
                "debit": money(row.get("借方")),
                "credit": money(row.get("贷方")),
            }
        )
    profits: dict = {}
    for row in _rows(path, "利润表"):
        per = str(row.get("期间") or "")
        ent = str(row.get("账套") or "")
        label = str(row.get("项目") or "")
        if not per or not ent or not label:
            continue
        profits.setdefault(per, {}).setdefault(ent, {})[label] = money(row.get("本月金额"))
    return {"accounts": accounts, "depts": depts, "profits_by_period": profits}


def _replace_sheet(path: Path, title: str, headers: list[str], rows: list[list]) -> None:
    wb = load_workbook(path)
    try:
        _write_sheet(wb, title, headers, rows)
        if title == "缺口" and "来源清单" in wb.sheetnames:
            wb.move_sheet(title, offset=1 - wb.sheetnames.index(title))
        _save(wb, path)
    finally:
        wb.close()


def write_payroll_rows(path: Path, rows: list[dict]) -> None:
    body = []
    for row in rows or []:
        body.append(
            [
                str(row.get("entity") or ""),
                str(row.get("code") or ""),
                str(row.get("name") or ""),
                str(row.get("dept") or ""),
                int(row.get("occurrence") or 1),
                cell_num(row.get("debit")),
                "是" if row.get("fill_left") else "",
            ]
        )
    _replace_sheet(path, "薪酬", ["账套", "科目编码", "项目", "部门", "发生额列", "借方", "填左列"], body)


def read_payroll_rows(path: Path) -> list[dict]:
    out = []
    for row in _rows(path, "薪酬"):
        out.append(
            {
                "entity": str(row.get("账套") or ""),
                "code": str(row.get("科目编码") or ""),
                "name": str(row.get("项目") or ""),
                "dept": str(row.get("部门") or ""),
                "occurrence": int(row.get("发生额列") or 1),
                "debit": money(row.get("借方")),
                "credit": None,
                "fill_left": str(row.get("填左列") or "") == "是",
            }
        )
    return out


def write_agency_records(path: Path, records: list[dict]) -> None:
    body = []
    for rec in records or []:
        ent = str(rec.get("entity") or "")
        per = str(rec.get("period") or "")
        for label, amt in (rec.get("values") or {}).items():
            body.append([per, ent, str(label), cell_num(amt)])
    _replace_sheet(path, "代账利润表", ["期间", "账套", "项目", "本月金额"], body)


def read_agency_records(path: Path, period: str) -> list[dict]:
    by_ent: dict[str, dict] = {}
    for row in _rows(path, "代账利润表"):
        if str(row.get("期间") or "") != period:
            continue
        ent = str(row.get("账套") or "")
        label = str(row.get("项目") or "")
        if not ent or not label:
            continue
        bucket = by_ent.setdefault(ent, {"entity": ent, "period": period, "values": {}})
        bucket["values"][label] = money(row.get("本月金额"))
    return list(by_ent.values())


def write_rent_lines(path: Path, lines: list[dict]) -> None:
    body = [
        [str(row.get("code") or ""), str(row.get("excel_dept") or ""), cell_num(row.get("amount"))]
        for row in lines or []
    ]
    _replace_sheet(path, "房租摘要", ["科目编码", "部门", "金额"], body)


def read_rent_lines(path: Path) -> list[dict]:
    out = []
    for row in _rows(path, "房租摘要"):
        out.append(
            {
                "code": str(row.get("科目编码") or ""),
                "excel_dept": str(row.get("部门") or ""),
                "amount": money(row.get("金额")),
            }
        )
    return out


def write_gap_sheet(path: Path, notes: list[str], profit_missing: list[str]) -> None:
    rows = []
    if profit_missing:
        rows.append(["没取到", "利润表", ",".join(profit_missing)])
    for prefix, label in (
        ("取到本期发生为0=", "取到本期发生为0"),
        ("取到利润表本月为0=", "取到利润表本月为0"),
        ("期间冲突=", "清洗未用"),
        ("期间不符=", "清洗未用"),
        ("期间未写明=", "清洗未用"),
        ("表头认不出=", "清洗未用"),
        ("网页期间未对准=", "网页期间未对准"),
        ("网页引出未保存=", "网页引出未保存"),
        ("缺线下利润表=", "没取到"),
    ):
        for note in notes:
            if note.startswith(prefix):
                rows.append([label, note.split("=", 1)[1], note])
    _replace_sheet(path, "缺口", ["类别", "对象", "说明"], rows)
