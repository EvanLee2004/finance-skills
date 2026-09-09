#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""代账小企业利润表 → 山东/四川/济南。损益表成品不是源。"""
from __future__ import annotations

import shutil
import subprocess
import tempfile
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from common import detect_period_text, load_json, money
from layout import dept_col_letter, direct_children, load_layout


def writable_leaf(code: str, layout: dict) -> str:
    kids = direct_children(layout)
    cur = code
    guard = 0
    while kids.get(cur) and guard < 8:
        cur = kids[cur][0]
        guard += 1
    return cur

CONFIG_NAME = "线下利润表.json"


def load_offline_rules() -> dict:
    return load_json(CONFIG_NAME)


def _text(value) -> str:
    if value is None:
        return ""
    return str(value).replace("\xa0", " ").strip()


def ensure_xlsx(path: Path) -> Path:
    path = Path(path)
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        return path
    if path.suffix.lower() != ".xls":
        raise ValueError(f"不支持的利润表格式: {path.suffix}")
    soffice = shutil.which("soffice")
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    exe = soffice or (str(mac) if mac.is_file() else None)
    if not exe:
        raise SystemExit(f"利润表是 .xls，本机没有 LibreOffice 可转。请另存为 xlsx：{path}")
    out_dir = Path(tempfile.mkdtemp(prefix="pl-xls-"))
    subprocess.run(
        [exe, "--headless", "--convert-to", "xlsx", "--outdir", str(out_dir), str(path)],
        check=True,
        capture_output=True,
    )
    got = out_dir / (path.stem + ".xlsx")
    if not got.is_file():
        hits = list(out_dir.glob("*.xlsx"))
        if not hits:
            raise SystemExit(f"xls 转 xlsx 失败：{path}")
        got = hits[0]
    return got


def looks_like_agency_profit(blob: str, title: str) -> bool:
    rules = load_offline_rules()
    text = f"{title}\n{blob}"
    if "本月金额" not in text and "本月数" not in text:
        return False
    if title.strip() in {"损益表", "确认情况"} or "确认情况" in title:
        return False
    if "公司名称：" in text:
        return False
    hints = rules.get("sheet_hints") or []
    if any(h in text for h in hints):
        return True
    if "本年累计" in text and match_offline_entity(text, "", rules):
        return True
    return False


def match_offline_entity(blob: str, filename: str, rules: dict | None = None) -> str | None:
    rules = rules or load_offline_rules()
    text = f"{blob}\n{filename}"
    table = rules.get("legal_name_to_entity") or {}
    best = None
    best_len = 0
    for legal, header in table.items():
        if legal and legal in text and len(legal) > best_len:
            best = header
            best_len = len(legal)
    if best:
        return best
    for header in rules.get("entities") or []:
        if header and header in text:
            return header
    return None


def _skip_label(label: str, rules: dict) -> bool:
    for needle in rules.get("skip_label_contains") or []:
        if needle and needle in label and needle not in ("管理费用", "销售费用", "财务费用"):
            if needle in {"其中", "开办费", "消费税", "商品维修", "商品维护", "广告费", "业务招待费", "研究费用", "利息费用", "政府补助", "坏账损失", "行次"}:
                if needle == "其中" and label.startswith("其中"):
                    return True
                if needle != "其中" and needle in label and not any(
                    k in label for k in ("营业收入", "营业成本", "管理费用", "销售费用", "财务费用")
                ):
                    return True
    if label.startswith("其中"):
        return True
    return False


def map_profit_label(label: str, rules: dict) -> str | None:
    raw = _text(label)
    if not raw or _skip_label(raw, rules):
        return None
    for spec in rules.get("profit_lines") or []:
        contains = spec.get("contains") or []
        if any(c and c in raw for c in contains):
            if "其中" in raw:
                return None
            return spec.get("profit_label")
    return None


def parse_agency_profit_sheet(ws, rules: dict | None = None) -> dict[str, Decimal]:
    rules = rules or load_offline_rules()
    header_r = None
    month_c = None
    item_c = 1
    last_r = min(ws.max_row or 1, 12)
    last_c = min(ws.max_column or 1, 8)
    for r in range(1, last_r + 1):
        for c in range(1, last_c + 1):
            t = _text(ws.cell(r, c).value)
            if t in {"本月金额", "本月数"} or t.endswith("本月金额"):
                header_r = r
                month_c = c
            if t in {"项目", "项目名称"}:
                item_c = c
    if not header_r or not month_c:
        return {}
    out: dict[str, Decimal] = {}
    for r in range(header_r + 1, (ws.max_row or header_r) + 1):
        lab = _text(ws.cell(r, item_c).value)
        mapped = map_profit_label(lab, rules)
        if not mapped:
            continue
        amt = money(ws.cell(r, month_c).value)
        if amt is None or amt == 0:
            continue
        out[mapped] = amt
    return out


def parse_agency_file(path: Path) -> dict:
    rules = load_offline_rules()
    xlsx = ensure_xlsx(path)
    wb = load_workbook(xlsx, data_only=True)
    try:
        blob_parts = [xlsx.name, path.name]
        chosen = None
        lines: dict[str, Decimal] = {}
        for title in wb.sheetnames:
            ws = wb[title]
            head = []
            for r in range(1, min(ws.max_row or 1, 8) + 1):
                for c in range(1, min(ws.max_column or 1, 8) + 1):
                    t = _text(ws.cell(r, c).value)
                    if t:
                        head.append(t)
            blob = "\n".join(head)
            blob_parts.append(blob)
            if not looks_like_agency_profit(blob, title):
                continue
            parsed = parse_agency_profit_sheet(ws, rules)
            if parsed:
                chosen = title
                lines = parsed
                break
        blob = "\n".join(blob_parts)
        entity = match_offline_entity(blob, path.name, rules)
        period = detect_period_text(blob)
        return {
            "path": str(path),
            "entity": entity,
            "period": period,
            "sheet": chosen,
            "lines": {k: str(v) for k, v in lines.items()},
            "values": lines,
        }
    finally:
        wb.close()


def discover_agency_files(homes: list[Path]) -> list[Path]:
    found: list[Path] = []
    seen: set[str] = set()
    for folder in homes:
        if not folder or not folder.is_dir():
            continue
        for path in sorted(folder.rglob("*"), key=lambda p: p.stat().st_mtime if p.is_file() else 0, reverse=True):
            if not path.is_file():
                continue
            if path.suffix.lower() not in {".xls", ".xlsx", ".xlsm"}:
                continue
            if path.name.startswith("~$"):
                continue
            if "损益类部门科目余额表" in path.name or path.name.startswith("月度损益表_"):
                continue
            key = str(path.resolve())
            if key in seen:
                continue
            seen.add(key)
            try:
                rec = parse_agency_file(path)
            except Exception:
                continue
            if rec.get("entity") and rec.get("values"):
                found.append(path)
    return found


def apply_offline_profit(
    entity_amts: dict,
    dept_amts: dict,
    profit_cur: dict,
    records: list[dict],
    period: str,
    notes: list[str],
    skip_split_entities: set[str] | None = None,
    skip_split_codes: dict[str, set[str]] | None = None,
) -> None:
    rules = load_offline_rules()
    layout = load_layout()
    splits = rules.get("pl_split") or {}
    defaults = rules.get("default_split") or {}
    for rec in records:
        ent = rec.get("entity")
        if not ent:
            continue
        per = rec.get("period") or period
        if per != period:
            notes.append(f"线下利润表期间不符={ent}:{per}")
            continue
        values: dict[str, Decimal] = rec.get("values") or {}
        if not values:
            continue
        notes.append(f"线下利润表={ent}")
        pay_codes = (skip_split_codes or {}).get(ent) or set()
        skip_split = ent in (skip_split_entities or set())
        for label, amt in values.items():
            profit_cur.setdefault(ent, {})[label] = amt
            spec = ((splits.get(ent) or {}).get(label)) or defaults.get(label)
            leaf = str((spec or {}).get("leaf") or (spec or {}).get("parent") or "")
            if pay_codes and label in {"管理费用", "销售费用", "成本"} and leaf:
                if any(c == leaf or str(c).startswith(leaf) for c in pay_codes):
                    continue
            elif skip_split and label in {"管理费用", "销售费用", "成本"}:
                continue
            if not spec:
                continue
            if spec.get("profit_only"):
                continue
            parent = spec.get("parent")
            leaf = spec.get("leaf") or parent
            dump = writable_leaf(leaf, layout) if leaf else leaf
            side = spec.get("side") or "debit"
            dept_name = spec.get("excel_dept")
            bucket = entity_amts.setdefault(ent, {})
            for code in {parent, leaf, dump}:
                if not code:
                    continue
                cell = bucket.setdefault(code, {"debit": None, "credit": None})
                if side == "credit":
                    cell["credit"] = amt
                else:
                    cell["debit"] = amt
            if dept_name and dump:
                letter = dept_col_letter(layout, dept_name, 1)
                if letter:
                    dept_amts.setdefault(dump, {})
                    dept_amts[dump][letter] = (dept_amts[dump].get(letter) or Decimal("0")) + amt
