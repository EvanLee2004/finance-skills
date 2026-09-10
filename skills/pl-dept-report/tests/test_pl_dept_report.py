#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from __future__ import annotations

import json
import os
import subprocess
import sys
from decimal import Decimal
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
SCRIPTS = SKILL / "scripts"
CONFIG = SKILL / "config"
CONVERT = SCRIPTS / "convert.py"
POSTING_API = SKILL.parent / "kingdee-posting" / "scripts" / "kingdee_api.py"

for _name in ("inspect_inputs", "parse_export", "layout", "common", "formula_eval", "convert"):
    sys.modules.pop(_name, None)
sys.path.insert(0, str(SCRIPTS))
from formula_eval import eval_workbook  # noqa: E402
from layout import account_row_map, dept_col_letter, dept_columns, direct_children, load_layout  # noqa: E402
from convert import nature_amount, pick_dept_rows  # noqa: E402

for _name in ("inspect_inputs", "parse_export", "convert", "common"):
    sys.modules.pop(_name, None)


def _run(args: list[str], env: dict | None = None) -> subprocess.CompletedProcess:
    run_env = os.environ.copy()
    run_env["PL_DEPT_SKIP_API"] = "1"
    run_env["PL_DEPT_USE_CACHE"] = "0"
    run_env["KINGDEE_LOCAL_JSON"] = str(HERE / "_missing_kingdee.json")
    run_env["KINGDEE_PL_LOCAL_JSON"] = str(HERE / "_missing_pl.json")
    if env:
        run_env.update(env)
    return subprocess.run(
        [sys.executable, str(CONVERT), *args],
        capture_output=True,
        text=True,
        env=run_env,
    )


def _write_account(path: Path, company: str, rows: list[tuple[str, str, float | None, float | None]]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "科目余额表"
    ws["A1"] = "科目余额表"
    ws["A2"] = company
    ws["A3"] = "科目编码"
    ws["B3"] = "科目名称"
    ws["C3"] = "本期发生借方"
    ws["D3"] = "本期发生贷方"
    for i, (code, name, debit, credit) in enumerate(rows, start=4):
        ws.cell(i, 1).value = code
        ws.cell(i, 2).value = name
        ws.cell(i, 3).value = debit
        ws.cell(i, 4).value = credit
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _write_assist(path: Path, company: str, rows: list[tuple[str, str, str, float | None, float | None]]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "核算项目余额表"
    ws["A1"] = "核算项目余额表"
    ws["B1"] = "核算项目类别：部门"
    ws["A2"] = company
    ws["A3"] = "科目编码"
    ws["B3"] = "科目名称"
    ws["C3"] = "核算项目名称"
    ws["D3"] = "本期发生借方"
    ws["E3"] = "本期发生贷方"
    for i, (code, name, dept, debit, credit) in enumerate(rows, start=4):
        ws.cell(i, 1).value = code
        ws.cell(i, 2).value = name
        ws.cell(i, 3).value = dept
        ws.cell(i, 4).value = debit
        ws.cell(i, 5).value = credit
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _write_profit(path: Path, company: str, items: list[tuple[str, float]]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "利润表"
    ws["A1"] = "利润表"
    ws["A2"] = company
    ws["A3"] = "项目"
    ws["B3"] = "本月金额"
    for i, (name, amt) in enumerate(items, start=4):
        ws.cell(i, 1).value = name
        ws.cell(i, 2).value = amt
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def test_same_dept_across_books_is_summed(tmp_path: Path):
    _write_account(
        tmp_path / "wh.xlsx",
        "北京甲骨易文化传媒有限公司",
        [("550198", "活动团建费", 10.0, None)],
    )
    _write_assist(
        tmp_path / "wh_d.xlsx",
        "北京甲骨易文化传媒有限公司",
        [("550198", "活动团建费", "大客户", 10.0, None)],
    )
    _write_account(
        tmp_path / "sh.xlsx",
        "甲骨易智译（上海）科技有限公司",
        [("550198", "活动团建费", 15.0, None)],
    )
    _write_assist(
        tmp_path / "sh_d.xlsx",
        "甲骨易智译（上海）科技有限公司",
        [("550198", "活动团建费", "大客户", 15.0, None)],
    )
    out = tmp_path / "out.xlsx"
    _run(["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"])
    ws = openpyxl.load_workbook(out)["损益表"]
    layout = load_layout()
    row = account_row_map(layout)["550198"]
    ka = next(letter for name, letter in dept_columns(layout) if name == "KA")
    assert ws[f"{ka}{row}"].value == 25.0
    assert ws[f"E{row}"].value == 10
    assert ws[f"G{row}"].value == 15


def test_prior_period_profit_export_goes_to_last_month_not_current(tmp_path: Path):
    path = tmp_path / "wh_pl.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "利润表"
    ws["A1"] = "2026年7期利润表（月报）"
    ws["A2"] = "公司名称：北京甲骨易文化传媒有限公司"
    ws["A3"] = "项目"
    ws["B3"] = "本月金额"
    ws["A4"] = "一、营业收入"
    ws["B4"] = 40.0
    wb.save(path)
    _write_account(
        tmp_path / "wh.xlsx",
        "北京甲骨易文化传媒有限公司",
        [("5101", "主营业务收入", None, 10.0)],
    )
    out = tmp_path / "out.xlsx"
    _run(["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"])
    ws = openpyxl.load_workbook(out)["利润表"]
    layout = load_layout()
    labels = [row["label"] for row in layout["profit_rows"]]
    r = 2 + labels.index("收入")
    assert ws.cell(r, 3).value in (None, "")
    assert ws.cell(r, 11).value == 40.0


def test_filename_timestamp_is_not_report_period():
    from common import detect_period_text

    assert detect_period_text("期间：202608-202608") == "202608"
    assert detect_period_text("2026年7期利润表") == "202607"
    assert detect_period_text("核算项目余额表-20260901190115.xlsx") is None


def test_discover_input_dir_uses_explicit(tmp_path: Path):
    from common import discover_input_dir

    (tmp_path / "a.xlsx").write_bytes(b"not-xlsx")
    got = discover_input_dir(str(tmp_path))
    assert got == tmp_path


def test_leaf_code_folds_to_layout_parent(tmp_path: Path):
    _write_account(
        tmp_path / "wh.xlsx",
        "北京甲骨易文化传媒有限公司",
        [("510103", "主营业务收入_翻译语言服务", None, 30.0)],
    )
    out = tmp_path / "out.xlsx"
    _run(["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"])
    wb = openpyxl.load_workbook(out, data_only=False)
    layout = load_layout()
    r = account_row_map(layout)["5101"]
    assert wb["损益表"][f"F{r}"].value == 30.0
    report = out.with_name(out.stem + "_运行报告.txt").read_text(encoding="utf-8")
    assert "表外科目=510103" not in report


def test_xingchen_two_row_period_header(tmp_path: Path):
    path = tmp_path / "xingchen_assist.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "核算项目余额表"
    ws["A1"] = "核算项目余额表"
    ws["A2"] = "公司名称：北京甲骨易文化传媒有限公司"
    ws["A4"] = "期间"
    ws["B4"] = "部门编码"
    ws["C4"] = "部门名称"
    ws["D4"] = "科目编码"
    ws["E4"] = "科目名称"
    ws["H4"] = "本期发生额"
    ws["H5"] = "借方"
    ws["I5"] = "贷方"
    ws["A6"] = "202608"
    ws["C6"] = "本公司"
    ws["D6"] = "510103"
    ws["E6"] = "翻译语言服务"
    ws["H6"] = 12.5
    ws["I6"] = 30.0
    wb.save(path)
    sys.modules.pop("inspect_inputs", None)
    sys.modules.pop("parse_export", None)
    if str(SCRIPTS) in sys.path:
        sys.path.remove(str(SCRIPTS))
    sys.path.insert(0, str(SCRIPTS))
    from inspect_inputs import inspect_dir, header_map
    from layout import load_export_aliases
    from parse_export import parse_inspected

    aliases = load_export_aliases()
    headers = header_map(openpyxl.load_workbook(path).active, aliases)
    assert headers["period_debit"] == (5, 8)
    assert headers["period_credit"] == (5, 9)
    found = inspect_dir(tmp_path)
    assert found and found[0]["kind"] == "assist"
    assert found[0]["entity"] == "文化"
    parsed = parse_inspected(found)
    rows = [r for r in parsed["depts"] if r["code"] == "510103"]
    assert rows
    assert rows[0]["debit"] == Decimal("12.5")
    assert rows[0]["credit"] == Decimal("30.0")


def test_layout_json_has_structure_no_gold_amounts():
    layout = json.loads((CONFIG / "版式.json").read_text(encoding="utf-8"))
    assert layout["freeze"] == "C3"
    assert layout["entities"][:3] == ["甲骨易", "文化", "上海"]
    assert "山东分公司" in layout["no_xingchen"]
    codes = [a["code"] for a in layout["accounts"]]
    assert "5101" in codes and "5401" in codes and "5502" in codes
    text = (CONFIG / "版式.json").read_text(encoding="utf-8")
    assert "8166971" not in text
    kids = direct_children(layout)
    assert "540204" in kids["5402"]
    assert "540205" in kids["5402"]
    assert "550122" in kids["5501"]


def test_two_sheets_freeze_order_and_no_confirm(tmp_path: Path):
    _write_account(
        tmp_path / "foo.xlsx",
        "甲骨易（北京）语言科技股份有限公司",
        [("5101", "主营业务收入", None, 100.0), ("540103", "翻译语言服务", 40.0, None)],
    )
    _write_assist(
        tmp_path / "bar.xlsx",
        "甲骨易（北京）语言科技股份有限公司",
        [
            ("5101", "主营业务收入", "大客户", None, 100.0),
            ("540103", "翻译语言服务", "大客户", 40.0, None),
        ],
    )
    out = tmp_path / "out.xlsx"
    r = _run(["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"])
    assert out.is_file(), r.stdout + r.stderr
    wb = openpyxl.load_workbook(out)
    assert set(wb.sheetnames) == {"损益表", "利润表"}
    assert "确认情况" not in wb.sheetnames
    ws = wb["损益表"]
    assert ws.freeze_panes == "C3"
    assert ws["C1"].value == "甲骨易"
    assert ws["E1"].value == "文化"
    assert ws["I1"].value == "山东分公司"
    assert ws["Q1"].value == "济南子公司"
    depts = [ws.cell(2, c).value for c in range(21, 47)]
    assert depts[0] == "营销总监及助理"
    assert depts[1] == "KA"
    assert depts[3] == "本地化"
    assert depts[8] == "本地化"


def test_shandong_sichuan_jinan_empty_not_zero(tmp_path: Path):
    _write_account(
        tmp_path / "hq.xlsx",
        "甲骨易（北京）语言科技股份有限公司",
        [("540103", "翻译语言服务", 10.0, None)],
    )
    out = tmp_path / "out.xlsx"
    _run(["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"])
    ws = openpyxl.load_workbook(out)["损益表"]
    layout = load_layout()
    r = account_row_map(layout)["540103"]
    assert ws[f"C{r}"].value == 10
    assert ws[f"I{r}"].value is None
    assert ws[f"O{r}"].value is None
    assert ws[f"Q{r}"].value is None
    report = (tmp_path / "out_运行报告.txt").read_text(encoding="utf-8")
    assert "山东分公司" in report
    assert "四川分公司" in report
    assert "济南子公司" in report


def test_map_dakehu_to_ka_and_unmapped_stays_out(tmp_path: Path):
    _write_account(
        tmp_path / "a.xlsx",
        "甲骨易（北京）语言科技股份有限公司",
        [("550198", "活动团建费", 80.0, None)],
    )
    _write_assist(
        tmp_path / "b.xlsx",
        "甲骨易（北京）语言科技股份有限公司",
        [
            ("550198", "活动团建费", "大客户", 50.0, None),
            ("550198", "活动团建费", "神秘事业部", 30.0, None),
        ],
    )
    out = tmp_path / "out.xlsx"
    r = _run(["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"])
    ws = openpyxl.load_workbook(out)["损益表"]
    layout = load_layout()
    row = account_row_map(layout)["550198"]
    ka_col = None
    for name, letter in dept_columns(layout):
        if name == "KA":
            ka_col = letter
            break
    assert ka_col
    assert ws[f"{ka_col}{row}"].value == 50
    for name, letter in dept_columns(layout):
        val = ws[f"{letter}{row}"].value
        if name != "KA":
            assert val in (None, "") or (isinstance(val, str) and val.startswith("="))
            if val not in (None, "") and not str(val).startswith("="):
                raise AssertionError(name)
    report = (tmp_path / "out_运行报告.txt").read_text(encoding="utf-8")
    assert "神秘事业部" in report
    assert "80.0" not in r.stdout
    assert "30.00" not in r.stdout
    assert "50.0" not in r.stdout
    assert "核对非0未映射=550198" in r.stdout
    assert "核对非0其它=无" in r.stdout


def test_hq_chanpin_yingxiao_xiangmu_map_to_layout_cols(tmp_path: Path):
    _write_account(
        tmp_path / "hq.xlsx",
        "甲骨易（北京）语言科技股份有限公司",
        [
            ("550321", "软件服务费", 11.0, None),
            ("550198", "活动团建费", 13.0, None),
            ("540103", "翻译语言服务", 17.0, None),
        ],
    )
    _write_assist(
        tmp_path / "hq_d.xlsx",
        "甲骨易（北京）语言科技股份有限公司",
        [
            ("550321", "软件服务费", "产品部", 11.0, None),
            ("550198", "活动团建费", "营销二部", 13.0, None),
            ("540103", "翻译语言服务", "项目中心", 17.0, None),
        ],
    )
    out = tmp_path / "out.xlsx"
    r = _run(["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"])
    ws = openpyxl.load_workbook(out)["损益表"]
    layout = load_layout()
    row = account_row_map(layout)
    qudao = dept_col_letter(layout, "渠道开发中心", 1)
    shi = dept_col_letter(layout, "视听", 1)
    yizu = dept_col_letter(layout, "项目一组", 1)
    director = dept_col_letter(layout, "项目总监及助理", 1)
    assert qudao and shi and yizu and director
    assert ws[f"{qudao}{row['550321']}"].value == 11
    assert ws[f"{shi}{row['550198']}"].value == 13
    assert ws[f"{director}{row['540103']}"].value == 17
    assert ws[f"{yizu}{row['540103']}"].value in (None, "")
    report = (tmp_path / "out_运行报告.txt").read_text(encoding="utf-8")
    assert "产品部" not in report
    assert "营销二部" not in report
    assert "项目中心" not in report
    assert "11.0" not in r.stdout


def test_wenhua_shanghai_accounts_fill_right_when_assist_has_no_expense(tmp_path: Path):
    _write_account(
        tmp_path / "wh.xlsx",
        "北京甲骨易文化传媒有限公司",
        [
            ("550104", "工资", 21.0, None),
            ("550111", "差旅费", 8.0, None),
            ("540101", "工资", 14.0, None),
            ("550205", "工资", 16.0, None),
            ("5101", "主营业务收入", None, 9.0),
        ],
    )
    _write_assist(
        tmp_path / "wh_d.xlsx",
        "北京甲骨易文化传媒有限公司",
        [("510103", "主营业务收入_翻译语言服务", "本公司", None, 9.0)],
    )
    _write_account(
        tmp_path / "sh.xlsx",
        "甲骨易智译（上海）科技有限公司",
        [
            ("5504", "销售费用", 25.0, None),
            ("550405", "工资", 19.0, None),
            ("550408", "差旅费", 6.0, None),
        ],
    )
    _write_assist(
        tmp_path / "sh_d.xlsx",
        "甲骨易智译（上海）科技有限公司",
        [("510103", "主营业务收入_翻译语言服务", "本公司", None, 3.0)],
    )
    out = tmp_path / "out.xlsx"
    r = _run(["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"])
    ws = openpyxl.load_workbook(out)["损益表"]
    layout = load_layout()
    row = account_row_map(layout)
    local = dept_col_letter(layout, "本地化", 1)
    shi = dept_col_letter(layout, "视听", 1)
    yun = dept_col_letter(layout, "运营保障中心", 1)
    fan = dept_col_letter(layout, "翻译中心", 1)
    assert local and shi and yun and fan
    assert ws[f"{shi}{row['550101']}"].value == 21
    assert ws[f"{shi}{row['550111']}"].value == 8
    assert ws[f"{fan}{row['540109']}"].value == 14
    assert ws[f"{yun}{row['550201']}"].value == 16
    assert ws[f"{local}{row['550101']}"].value == 19
    assert ws[f"{local}{row['550111']}"].value == 6
    assert ws[f"{shi}{row['5101']}"].value in (None, "")
    assert ws[f"{local}{row['5101']}"].value in (None, "")
    director = dept_col_letter(layout, "营销总监及助理", 1)
    assert director
    assert ws[f"{director}{row['5101']}"].value == 12
    assert "21.0" not in r.stdout
    assert "19.0" not in r.stdout


def test_pick_dept_rows_prefers_hq_assist_file_over_vouchers():
    folded = [
        {"entity": "文化", "code": "550111", "dept": "本公司"},
        {"entity": "甲骨易", "code": "510103", "dept": "营销总监及助理"},
    ]
    api = [{"entity": "甲骨易", "code": "550321", "dept": "产品部"}]
    rows, note = pick_dept_rows(folded, api, hq_from_api=True)
    assert note
    assert any(r.get("code") == "510103" for r in rows)
    assert not any(r.get("code") == "550321" for r in rows)
    rows2, note2 = pick_dept_rows([folded[0]], api, hq_from_api=True)
    assert note2 is None
    assert any(r.get("code") == "550321" for r in rows2)


def test_hq_assist_income_leaf_folds_to_5101(tmp_path: Path):
    _write_account(
        tmp_path / "hq.xlsx",
        "甲骨易（北京）语言科技股份有限公司",
        [("5101", "主营业务收入", None, 20.0)],
    )
    _write_assist(
        tmp_path / "hq_d.xlsx",
        "甲骨易（北京）语言科技股份有限公司",
        [("510103", "主营业务收入_多语本地化服务", "营销总监及助理", None, 20.0)],
    )
    out = tmp_path / "out.xlsx"
    _run(["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"])
    ws = openpyxl.load_workbook(out)["损益表"]
    layout = load_layout()
    row = account_row_map(layout)["5101"]
    letter = dept_col_letter(layout, "营销总监及助理", 1)
    assert letter
    assert ws[f"{letter}{row}"].value == 20


def test_customer_assist_sheet_is_not_department(tmp_path: Path):
    from inspect_inputs import inspect_file

    path = tmp_path / "客户核算.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "核算项目余额表"
    ws["A1"] = "核算项目余额表"
    ws["A2"] = "公司名称：甲骨易（北京）语言科技股份有限公司"
    ws["A3"] = "期间：202608-202608"
    ws["A4"] = "客户编码"
    ws["B4"] = "客户名称"
    ws["C4"] = "科目编码"
    ws["D4"] = "科目名称"
    ws["E4"] = "本期发生额"
    ws["A5"] = "客户编码"
    ws["B5"] = "客户名称"
    ws["C5"] = "科目编码"
    ws["D5"] = "科目名称"
    ws["E5"] = "借方"
    ws["A6"] = "0101"
    ws["B6"] = "某客户"
    ws["C6"] = "113101"
    ws["D6"] = "应收账款"
    ws["E6"] = 9
    wb.save(path)
    found = inspect_file(path)
    assert found == []


def test_stdout_has_no_amount_or_secret(tmp_path: Path):
    secret = "SUPERSECRET_TOKEN_XYZ"
    fake = tmp_path / "kingdee.local.json"
    fake.write_text(json.dumps({"client_id": "1", "client_secret": secret, "app_key": "k", "app_secret": secret}), encoding="utf-8")
    _write_account(
        tmp_path / "a.xlsx",
        "甲骨易（北京）语言科技股份有限公司",
        [("540103", "翻译语言服务", 12345.67, None)],
    )
    out = tmp_path / "out.xlsx"
    r = _run(
        ["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"],
        env={"KINGDEE_LOCAL_JSON": str(fake)},
    )
    assert "12345.67" not in r.stdout
    assert secret not in r.stdout
    assert "期间=202608" in r.stdout
    assert "产物=" in r.stdout


def test_company_name_line_beats_related_party_in_rows(tmp_path: Path):
    path = tmp_path / "hunan_zgs.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "科目余额表"
    ws["A1"] = "科目余额表"
    ws["A2"] = "公司名称：甲骨易（湖南）科技有限公司"
    ws["A3"] = "科目编码"
    ws["B3"] = "科目名称"
    ws["C3"] = "本期发生借方"
    ws["D3"] = "本期发生贷方"
    ws["A4"] = "3001"
    ws["B4"] = "甲骨易（北京）语言科技股份有限公司"
    ws["C4"] = 1.0
    wb.save(path)
    sys.modules.pop("inspect_inputs", None)
    if str(SCRIPTS) in sys.path:
        sys.path.remove(str(SCRIPTS))
    sys.path.insert(0, str(SCRIPTS))
    from inspect_inputs import inspect_dir

    found = inspect_dir(tmp_path)
    assert found
    assert found[0]["entity"] == "湖南子公司"


def test_fetched_book_without_pl_codes_counts_as_source(tmp_path: Path):
    path = tmp_path / "hunan.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "科目余额表"
    ws["A1"] = "科目余额表"
    ws["A2"] = "公司名称：甲骨易（北京）语言科技股份有限公司湖南分公司"
    ws["A3"] = "科目编码"
    ws["B3"] = "科目名称"
    ws["C3"] = "本期发生借方"
    ws["D3"] = "本期发生贷方"
    ws["A4"] = "1002"
    ws["B4"] = "银行存款"
    ws["C4"] = 1.0
    wb.save(path)
    out = tmp_path / "out.xlsx"
    r = _run(["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"])
    assert "有源账套=" in r.stdout
    assert "湖南分公司" in r.stdout.split("有源账套=")[1].split("\n")[0]
    assert "已取但无损益科目=湖南分公司" in (tmp_path / "out_运行报告.txt").read_text(encoding="utf-8")


def test_no_key_with_export_still_builds(tmp_path: Path):
    _write_account(
        tmp_path / "wenhua.xlsx",
        "北京甲骨易文化传媒有限公司",
        [("5101", "主营业务收入", None, 20.0)],
    )
    out = tmp_path / "out.xlsx"
    r = _run(["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"])
    assert out.is_file()
    assert set(openpyxl.load_workbook(out).sheetnames) == {"损益表", "利润表"}
    assert "Client" not in r.stdout
    assert "secret" not in r.stdout.lower() or "[secret]" in r.stdout.lower()


def test_formulas_not_dead_numbers_balanced_and_unbalanced(tmp_path: Path):
    _write_account(
        tmp_path / "bal.xlsx",
        "甲骨易（北京）语言科技股份有限公司",
        [("540103", "翻译语言服务", 80.0, None), ("540109", "工资", 20.0, None)],
    )
    _write_assist(
        tmp_path / "bal_d.xlsx",
        "甲骨易（北京）语言科技股份有限公司",
        [
            ("540103", "翻译语言服务", "KA", 80.0, None),
            ("540109", "工资", "KA", 20.0, None),
        ],
    )
    out = tmp_path / "bal_out.xlsx"
    _run(["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"])
    wb = openpyxl.load_workbook(out, data_only=False)
    layout = load_layout()
    row = account_row_map(layout)["540103"]
    parent = account_row_map(layout)["5401"]
    assert str(wb["损益表"][f"S{row}"].value).startswith("=")
    assert str(wb["损益表"][f"AU{row}"].value).startswith("=")
    assert str(wb["损益表"][f"AV{row}"].value).startswith("=")
    assert str(wb["损益表"][f"U{parent}"].value).startswith("=")
    assert str(wb["利润表"]["R2"].value).startswith("=")
    book, _ = eval_workbook(wb)
    av = book.cell_value("损益表", row, 48)
    assert abs(Decimal(str(av))) <= Decimal("0.05")
    parent_u = book.cell_value("损益表", parent, 21)
    child_u = book.cell_value("损益表", row, 21)
    wage = account_row_map(layout)["540109"]
    wage_u = book.cell_value("损益表", wage, 21)
    # parent 5401 U = sum of direct children; KA is column V=22 not U
    ka_parent = book.cell_value("损益表", parent, 22)
    ka_child = book.cell_value("损益表", row, 22)
    ka_wage = book.cell_value("损益表", wage, 22)
    assert Decimal(str(ka_parent)) == Decimal(str(ka_child or 0)) + Decimal(str(ka_wage or 0))

    bad = tmp_path / "bad"
    bad.mkdir()
    _write_account(bad / "a.xlsx", "甲骨易（北京）语言科技股份有限公司", [("550198", "活动团建费", 80.0, None)])
    _write_assist(bad / "d.xlsx", "甲骨易（北京）语言科技股份有限公司", [("550198", "活动团建费", "KA", 50.0, None)])
    out2 = bad / "out.xlsx"
    _run(["--period", "202608", "--input-dir", str(bad), "--out", str(out2), "--no-api"])
    wb2 = openpyxl.load_workbook(out2, data_only=False)
    book2, _ = eval_workbook(wb2)
    av2 = book2.cell_value("损益表", account_row_map(layout)["550198"], 48)
    assert abs(Decimal(str(av2))) > Decimal("0.05")


def test_change_rate_formula_not_gold_wrong_ref(tmp_path: Path):
    _write_account(tmp_path / "a.xlsx", "甲骨易（北京）语言科技股份有限公司", [("5101", "主营业务收入", None, 10.0)])
    out = tmp_path / "out.xlsx"
    _run(["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"])
    ws = openpyxl.load_workbook(out)["利润表"]
    layout = load_layout()
    labels = [row["label"] for row in layout["profit_rows"]]
    r = 2 + labels.index("营业收入变动率")
    formula = str(ws.cell(r, 2).value)
    assert formula.startswith("=")
    assert "J22" not in formula
    assert "J2" in formula


def test_parent_includes_children_gold_missed():
    layout = load_layout()
    kids = direct_children(layout)
    assert "550122" in kids["5501"]
    src = (SCRIPTS / "convert.py").read_text(encoding="utf-8")
    assert "direct_children" in src


def test_does_not_import_posting_picker():
    text = (SCRIPTS / "kingdee_client.py").read_text(encoding="utf-8")
    assert "kingdee-posting" not in text
    assert "kingdee_api" not in text
    posting = POSTING_API.read_text(encoding="utf-8")
    assert "多账套同时授权时只认本机已绑的总部账套" in posting


def test_profit_label_fuzzy_match():
    from fetch_reports import _match_profit_label, map_profit_rows
    from layout import load_layout

    layout = load_layout()
    aliases = layout.get("profit_item_aliases") or {}
    assert _match_profit_label("一、营业收入", aliases) == "收入"
    assert _match_profit_label("投资收益（损失以“-”号填列）", aliases) == "+投资收益"
    assert _match_profit_label("减：营业成本", aliases) == "成本"
    rows = [
        {"item_name": "一、营业收入", "current_amount": 10},
        {"item_name": "公允价值变动收益（损失以“-”号填列）", "current_amount": 1},
    ]
    mapped = map_profit_rows(rows, layout)
    assert mapped["收入"] is not None
    assert mapped["+公允价值变动收益"] is not None


def test_default_period_august_when_september():
    from common import default_period
    from datetime import date

    assert default_period(date(2026, 9, 1)) == "202608"


def test_parent_formula_includes_gold_missed_children(tmp_path: Path):
    _write_account(tmp_path / "a.xlsx", "甲骨易（北京）语言科技股份有限公司", [("540204", "印花税", 1.0, None)])
    out = tmp_path / "out.xlsx"
    _run(["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"])
    ws = openpyxl.load_workbook(out)["损益表"]
    layout = load_layout()
    parent = account_row_map(layout)["5402"]
    child = account_row_map(layout)["540204"]
    formula = str(ws[f"U{parent}"].value)
    assert formula.startswith("=")
    assert f"U{child}" in formula
    extra = account_row_map(layout)["540205"]
    assert f"U{extra}" in formula


def test_pl_picker_skips_empty_book():
    from kingdee_client import pick_pl_authorize_rows

    rows = [
        {"status": 1, "accountId": "1783670301378479516", "serviceId": "7914379139221", "appSecret": "empty"},
        {"status": 1, "accountId": "1783803326505631821", "serviceId": "795589109148", "appSecret": "hq"},
    ]
    got = pick_pl_authorize_rows(rows, {"account_id": "1783803326505631821"})
    assert len(got) == 1
    assert got[0]["serviceId"] == "795589109148"


def test_fetch_balance_400_falls_back_to_voucher(monkeypatch):
    import fetch_reports
    import kingdee_client

    class FakeResp:
        def __init__(self, status, payload):
            self.status_code = status
            self._payload = payload

        def json(self):
            return self._payload

    def fake_request(method, path, creds, params=None, extra=None, timeout=30, retries=4):
        if path.endswith("profit_report"):
            return FakeResp(200, {"data": [{"item_name": "营业收入", "current_amount": 1}]})
        if path.endswith("account_balance_report"):
            return FakeResp(400, {"errcode": 400})
        if path.endswith("/fi/voucher"):
            return FakeResp(200, {"data": {"count": 0, "rows": []}})
        if path.endswith("department"):
            return FakeResp(200, {"data": []})
        return FakeResp(404, {})

    monkeypatch.setattr(kingdee_client, "request", fake_request)
    monkeypatch.setattr(fetch_reports, "request", fake_request)
    monkeypatch.setattr(fetch_reports, "get_app_token", lambda creds: ("tok", "https://example"))
    monkeypatch.setenv("PL_DEPT_USE_CACHE", "0")
    got = fetch_reports.fetch_hq("202608", {"client_id": "x"})
    assert any("400" in n for n in got["notes"])
    assert any("voucher_count" in n for n in got["notes"])
    assert "profit_rows=1" in got["notes"]


def test_parent_and_leaf_same_amount_does_not_double_count(tmp_path: Path):
    _write_account(
        tmp_path / "wh.xlsx",
        "北京甲骨易文化传媒有限公司",
        [
            ("5101", "主营业务收入", None, 249313.21),
            ("510103", "主营业务收入_翻译语言服务", None, 249313.21),
        ],
    )
    out = tmp_path / "out.xlsx"
    ran = _run(["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"])
    wb = openpyxl.load_workbook(out, data_only=False)
    layout = load_layout()
    r = account_row_map(layout)["5101"]
    assert wb["损益表"][f"F{r}"].value == 249313.21
    report = out.with_name(out.stem + "_运行报告.txt").read_text(encoding="utf-8")
    assert "表外科目=510103" not in report
    assert "249313.21" not in ran.stdout


def test_same_code_different_name_does_not_write_layout_row(tmp_path: Path):
    _write_account(
        tmp_path / "wh.xlsx",
        "北京甲骨易文化传媒有限公司",
        [
            ("540103", "住房公积金", 350.0, None),
            ("550104", "工资", 100.0, None),
            ("550204", "房租", 200.0, None),
            ("5503", "财务费用", 10.0, None),
        ],
    )
    _write_account(
        tmp_path / "sh.xlsx",
        "甲骨易智译（上海）科技有限公司",
        [("5504", "销售费用", 50.0, None)],
    )
    out = tmp_path / "out.xlsx"
    r = _run(["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"])
    assert out.is_file(), r.stdout + r.stderr
    ws = openpyxl.load_workbook(out)["损益表"]
    layout = load_layout()
    row = account_row_map(layout)
    assert ws[f"E{row['540103']}"].value in (None, "")
    assert ws[f"E{row['540112']}"].value == 350
    assert ws[f"E{row['550104']}"].value in (None, "")
    assert ws[f"E{row['550101']}"].value == 100
    assert ws[f"E{row['550204']}"].value in (None, "")
    assert ws[f"E{row['550212']}"].value == 200
    assert ws[f"E{row['5503']}"].value in (None, "")
    assert ws[f"E{row['5504']}"].value == 10
    assert ws[f"G{row['5504']}"].value in (None, "")
    assert ws[f"G{row['5501']}"].value == 50
    report = out.with_name(out.stem + "_运行报告.txt").read_text(encoding="utf-8")
    assert "同码不同名=540103->540112" in report
    assert "350.0" not in r.stdout
    assert "249313.21" not in r.stdout


def _yaml_description(skill_md: Path) -> str:
    text = skill_md.read_text(encoding="utf-8")
    parts = text.split("---", 2)
    assert len(parts) >= 3, skill_md
    return parts[1]


def test_two_trigger_phrases_do_not_steal_dept_expense_alloc():
    pl_md = (SKILL / "SKILL.md").read_text(encoding="utf-8")
    pl_yaml = _yaml_description(SKILL / "SKILL.md")
    other_md = (SKILL.parent / "dept-expense-alloc" / "SKILL.md").read_text(encoding="utf-8")
    other_yaml = _yaml_description(SKILL.parent / "dept-expense-alloc" / "SKILL.md")
    for phrase in ("月度损益表", "科目余额表"):
        assert phrase in pl_yaml
        assert phrase not in other_yaml
    assert "金蝶" in pl_yaml
    assert "用友" not in pl_yaml
    assert "pl-dept-report" in other_md
    assert "dept-expense-alloc" in pl_md
    assert "部门费用归集" in other_yaml


def test_skill_forbids_adhoc_openpyxl_and_default_skips_input_dir():
    text = (SKILL / "SKILL.md").read_text(encoding="utf-8")
    assert "python -c" in text
    assert "import openpyxl" in text
    assert "convert.py" in text
    assert "不要**加 `--input-dir`" in text or "不要加 `--input-dir`" in text
    assert "hq.xlsx" in text
    yaml = _yaml_description(SKILL / "SKILL.md")
    assert "月度损益表" in yaml and "科目余额表" in yaml


def test_shanghai_and_wenhua_bengongsi_special_mapping(tmp_path: Path):
    _write_account(
        tmp_path / "sh_acc.xlsx",
        "甲骨易智译（上海）科技有限公司",
        [("550101", "工资", 40.0, None), ("5101", "主营业务收入", None, 80.0), ("550212", "房租", 9.0, None)],
    )
    _write_assist(
        tmp_path / "sh_dept.xlsx",
        "甲骨易智译（上海）科技有限公司",
        [
            ("550101", "工资", "本公司", 40.0, None),
            ("5101", "主营业务收入", "本公司", None, 80.0),
            ("550212", "房租", "本公司", 9.0, None),
        ],
    )
    _write_account(
        tmp_path / "wh_acc.xlsx",
        "北京甲骨易文化传媒有限公司",
        [
            ("550111", "差旅费", 12.0, None),
            ("550201", "工资", 30.0, None),
            ("540109", "工资", 15.0, None),
            ("550212", "房租", 4.0, None),
            ("550242", "服务费", 6.0, None),
        ],
    )
    _write_assist(
        tmp_path / "wh_dept.xlsx",
        "北京甲骨易文化传媒有限公司",
        [
            ("550111", "差旅费", "本公司", 12.0, None),
            ("550201", "工资", "本公司", 30.0, None),
            ("540109", "工资", "本公司", 15.0, None),
            ("550212", "房租", "本公司", 4.0, None),
            ("550242", "服务费", "本公司", 6.0, None),
        ],
    )
    _write_assist(
        tmp_path / "hq_dept.xlsx",
        "甲骨易（北京）语言科技股份有限公司",
        [("550101", "工资", "大客户", 7.0, None)],
    )
    out = tmp_path / "out.xlsx"
    r = _run(["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"])
    assert out.is_file(), r.stdout + r.stderr
    ws = openpyxl.load_workbook(out)["损益表"]
    layout = load_layout()
    row = account_row_map(layout)
    local = dept_col_letter(layout, "本地化", 1)
    second_local = dept_col_letter(layout, "本地化", 2)
    shi = dept_col_letter(layout, "视听", 1)
    yun = dept_col_letter(layout, "运营保障中心", 1)
    fan = dept_col_letter(layout, "翻译中心", 1)
    ka = dept_col_letter(layout, "KA", 1)
    assert local and shi and yun and fan and ka and second_local
    assert ws[f"{local}{row['550101']}"].value == 40
    assert ws[f"{second_local}{row['550101']}"].value in (None, "")
    assert ws[f"{shi}{row['550111']}"].value == 12
    assert ws[f"{yun}{row['550201']}"].value == 30
    assert ws[f"{fan}{row['540109']}"].value == 15
    assert ws[f"{ka}{row['550101']}"].value == 7
    assert ws[f"{local}{row['5101']}"].value in (None, "")
    director = dept_col_letter(layout, "营销总监及助理", 1)
    assert director
    assert ws[f"{director}{row['5101']}"].value == 80
    assert ws[f"{local}{row['550212']}"].value in (None, "")
    hr = dept_col_letter(layout, "人力资源部", 1)
    assert hr
    assert ws[f"{yun}{row['550212']}"].value == 13
    assert ws[f"{hr}{row['550242']}"].value == 6
    report = out.with_name(out.stem + "_运行报告.txt").read_text(encoding="utf-8")
    assert "上海:本公司" in report or "本公司" in report


def _write_agency_profit(path: Path, company: str, period_text: str, month_mgmt: float) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "利润表"
    ws["A1"] = "利润表"
    ws["A2"] = f"编制单位：{company}"
    ws["B2"] = period_text
    ws["D2"] = "会小企02表"
    ws["A3"] = "项目"
    ws["B3"] = "行次"
    ws["C3"] = "本年累计金额"
    ws["D3"] = "本月金额"
    ws["A4"] = "一、营业收入"
    ws["A5"] = "减：营业成本"
    ws["A6"] = "管理费用"
    ws["D6"] = month_mgmt
    ws["A7"] = "其中：开办费"
    ws["A8"] = "二、营业利润（亏损以“-”号填列）"
    ws["D8"] = -month_mgmt
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()


def test_result_workbook_is_not_a_source(tmp_path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "损益表"
    ws["C1"] = "甲骨易"
    ws["I1"] = "山东分公司"
    ws["A2"] = "科目编码"
    ws["B2"] = "科目名称"
    ws["I2"] = "本期发生借方"
    ws["A3"] = "540109"
    ws["B3"] = "工资"
    ws["I3"] = 88.0
    wb.save(tmp_path / "2026年8月损益类部门科目余额表.xlsx")
    wb.close()
    _write_account(tmp_path / "hq.xlsx", "甲骨易（北京）语言科技股份有限公司", [("5101", "主营业务收入", None, 1.0)])
    out = tmp_path / "out.xlsx"
    r = _run(["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"])
    assert out.is_file(), r.stdout + r.stderr
    ws = openpyxl.load_workbook(out)["损益表"]
    layout = load_layout()
    row = account_row_map(layout)
    assert ws[f"I{row['540109']}"].value in (None, "")
    report = out.with_name(out.stem + "_运行报告.txt").read_text(encoding="utf-8")
    assert "忽略损益表成品" in report or "缺线下利润表" in report or "山东分公司" in (report.split("缺源账套=")[1].split("\n")[0] if "缺源账套=" in report else "")


def test_agency_profit_fills_shandong_sichuan_jinan(tmp_path: Path):
    _write_agency_profit(
        tmp_path / "sd.xls".replace(".xls", ".xlsx"),
        "甲骨易（北京）语言科技股份有限公司山东分公司",
        "2026年8期",
        88.0,
    )
    _write_agency_profit(
        tmp_path / "sc.xlsx",
        "甲骨易（北京）语言科技股份有限公司四川分公司",
        "期间：2026年08月",
        22.0,
    )
    _write_agency_profit(
        tmp_path / "jn.xlsx",
        "甲骨易(济南)科技有限公司",
        "2026-08",
        55.0,
    )
    _write_account(tmp_path / "hq.xlsx", "甲骨易（北京）语言科技股份有限公司", [("5101", "主营业务收入", None, 1.0)])
    out = tmp_path / "out.xlsx"
    r = _run(
        [
            "--period",
            "202608",
            "--input-dir",
            str(tmp_path),
            "--out",
            str(out),
            "--no-api",
        ]
    )
    assert out.is_file(), r.stdout + r.stderr
    assert "ask=" not in r.stdout
    wb = openpyxl.load_workbook(out)
    ws = wb["损益表"]
    pf = wb["利润表"]
    layout = load_layout()
    row = account_row_map(layout)
    assert ws[f"I{row['5401']}"].value == 88
    assert ws[f"I{row['540111']}"].value == 88
    assert ws[f"O{row['5502']}"].value == 22
    assert ws[f"O{row['550201']}"].value == 22
    assert ws[f"Q{row['5401']}"].value == 55
    assert ws[f"Q{row['540109']}"].value in (None, "")
    assert ws[f"Q{row['540123']}"].value == 55
    jinan_dept = dept_col_letter(layout, "济南分公司", 1)
    sichuan_dept = dept_col_letter(layout, "四川分公司", 1)
    jinan_zgs = dept_col_letter(layout, "济南子公司", 1)
    assert ws[f"{jinan_dept}{row['54011101']}"].value == 88
    finance = dept_col_letter(layout, "财务中心", 1)
    assert finance
    assert ws[f"{finance}{row['550201']}"].value == 22
    assert ws[f"{sichuan_dept}{row['550201']}"].value in (None, "")
    assert ws[f"{jinan_zgs}{row['540109']}"].value in (None, "")
    assert ws[f"{jinan_zgs}{row['540123']}"].value == 55
    assert pf["E6"].value == 88
    assert pf["H6"].value == 22
    assert pf["I6"].value == 55
    report = out.with_name(out.stem + "_运行报告.txt").read_text(encoding="utf-8")
    assert "线下利润表=山东分公司" in report
    assert "线下利润表=四川分公司" in report
    assert "线下利润表=济南子公司" in report


def test_missing_agency_profit_asks_and_leaves_empty(tmp_path: Path):
    _write_account(tmp_path / "hq.xlsx", "甲骨易（北京）语言科技股份有限公司", [("5101", "主营业务收入", None, 1.0)])
    out = tmp_path / "out.xlsx"
    r = _run(["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"])
    assert out.is_file(), r.stdout + r.stderr
    assert "ask=" in r.stdout
    assert "代账利润表" in r.stdout
    ws = openpyxl.load_workbook(out)["损益表"]
    layout = load_layout()
    row = account_row_map(layout)
    assert ws[f"I{row['5401']}"].value in (None, "")
    skip = _run(
        ["--period", "202608", "--input-dir", str(tmp_path), "--out", str(tmp_path / "skip.xlsx"), "--no-api", "--skip-offline"]
    )
    assert skip.returncode in (0, 2)
    assert "线下利润表=跳过" in (tmp_path / "skip_运行报告.txt").read_text(encoding="utf-8") or "跳过" in skip.stdout


def test_default_desktop_dir_helper(tmp_path: Path):
    from datetime import date
    from common import default_desktop_dir

    desktop = tmp_path / "Desktop"
    desktop.mkdir()
    got = default_desktop_dir("月度损益表", today=date(2026, 9, 2), home=tmp_path)
    assert got == desktop / "月度损益表_20260902"


def test_xingchen_creds_from_local_json_or_env(tmp_path: Path, monkeypatch):
    from xingchen_login import ASK_CREDS, _looks_logged_in, load_creds

    monkeypatch.delenv("XINGCHEN_USER", raising=False)
    monkeypatch.delenv("XINGCHEN_PASSWORD", raising=False)
    monkeypatch.setenv("XINGCHEN_LOCAL_JSON", str(tmp_path / "missing.json"))
    assert load_creds() is None
    path = tmp_path / "xingchen.local.json"
    path.write_text('{"username": "u", "password": "p"}\n', encoding="utf-8")
    monkeypatch.setenv("XINGCHEN_LOCAL_JSON", str(path))
    got = load_creds()
    assert got == {"username": "u", "password": "p"}
    monkeypatch.setenv("XINGCHEN_USER", "eu")
    monkeypatch.setenv("XINGCHEN_PASSWORD", "ep")
    got = load_creds()
    assert got == {"username": "eu", "password": "ep"}
    assert "xingchen.local.json" in ASK_CREDS
    assert _looks_logged_in("https://service.jdy.com/workbench/web/index.html", "进入使用")
    assert not _looks_logged_in("https://www.jdy.com/login/", "账号登录")


def _write_payroll(path: Path) -> None:
    wb = openpyxl.Workbook()
    org = wb.active
    org.title = "组织架构"
    org["A2"] = "姓名"
    org["B2"] = "组织架构-1"
    org["C2"] = "组织架构-2"
    org["A3"] = "甲"
    org["B3"] = "渠道开发中心"
    org["C3"] = "技术部"
    org["A4"] = "乙"
    org["B4"] = "营销中心"
    org["C4"] = "本地化事业部"
    org["A5"] = "丙"
    org["B5"] = "项目中心"
    org["C5"] = "本地化事业部"
    org["A6"] = "戊"
    org["B6"] = "运营保障中心"
    org["C6"] = "运营保障中心"
    wage = wb.create_sheet("202608甲骨易工资")
    for i, h in enumerate(["序号", "姓名", "组织架构1", "组织架构2", "部门", "地区", "身份证号码", "基本工资"], 1):
        wage.cell(1, i).value = h
    wage["A2"] = 1
    wage["B2"] = "甲"
    wage["C2"] = "渠道开发中心"
    wage["D2"] = "技术部"
    wage["H2"] = 10
    wage["A3"] = 2
    wage["B3"] = "乙"
    wage["C3"] = "营销中心"
    wage["D3"] = "本地化事业部"
    wage["H3"] = 20
    wage["A4"] = 3
    wage["B4"] = "丙"
    wage["C4"] = "项目中心"
    wage["D4"] = "本地化事业部"
    wage["H4"] = 30
    chengdu = wb.create_sheet("202608成都")
    chengdu["A1"] = "工资表"
    chengdu["A2"] = "编制单位：甲骨易（北京）语言科技股份有限公司成都分公司"
    chengdu["A3"] = "序号"
    chengdu["B3"] = "姓名"
    chengdu["C3"] = "部门"
    chengdu["F3"] = "基本工资"
    chengdu["A4"] = 1
    chengdu["B4"] = "丁"
    chengdu["C4"] = "财务中心"
    chengdu["F4"] = 22
    culture = wb.create_sheet("文化工资")
    for i, h in enumerate(["序号", "姓名", "组织架构1", "组织架构2", "部门", "地区", "身份证号码", "基本工资"], 1):
        culture.cell(1, i).value = h
    culture["A2"] = 1
    culture["B2"] = "戊"
    culture["C2"] = "运营保障中心"
    culture["D2"] = "运营保障中心"
    culture["H2"] = 5
    agency = wb.create_sheet("202608济南子")
    agency["A1"] = "甲骨易（北京）语言科技股份有限公司8月对帐单"
    agency["A2"] = "序号"
    agency["B2"] = "姓名"
    agency["H2"] = "养老保险"
    agency["J2"] = "失业保险"
    agency["L2"] = "工伤保险"
    agency["M2"] = "医疗保险"
    agency["O2"] = "公积金"
    agency["A3"] = None
    agency["H3"] = "雇主"
    agency["A5"] = "1"
    agency["B5"] = "己"
    agency["H5"] = 3
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()


def test_payroll_maps_rd_localization_sichuan_and_skips_double_count(tmp_path: Path):
    _write_account(
        tmp_path / "hq.xlsx",
        "甲骨易（北京）语言科技股份有限公司",
        [
            ("550301", "工资", 10.0, None),
            ("550101", "工资", 20.0, None),
            ("540109", "工资", 30.0, None),
        ],
    )
    _write_account(
        tmp_path / "wh.xlsx",
        "北京甲骨易文化传媒有限公司",
        [("550201", "工资", 99.0, None)],
    )
    _write_assist(
        tmp_path / "wh_dept.xlsx",
        "北京甲骨易文化传媒有限公司",
        [("550201", "工资", "本公司", 99.0, None)],
    )
    _write_agency_profit(
        tmp_path / "sc.xlsx",
        "甲骨易（北京）语言科技股份有限公司四川分公司",
        "期间：2026年08月",
        88.0,
    )
    _write_agency_profit(
        tmp_path / "jn.xlsx",
        "甲骨易(济南)科技有限公司",
        "2026-08",
        55.0,
    )
    _write_payroll(tmp_path / "202608_职工薪酬台账.xlsx")
    out = tmp_path / "out.xlsx"
    r = _run(["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"])
    assert out.is_file(), r.stdout + r.stderr
    ws = openpyxl.load_workbook(out)["损益表"]
    layout = load_layout()
    row = account_row_map(layout)
    tech = dept_col_letter(layout, "技术中心", 1)
    local1 = dept_col_letter(layout, "本地化", 1)
    local2 = dept_col_letter(layout, "本地化", 2)
    yun = dept_col_letter(layout, "运营保障中心", 1)
    finance = dept_col_letter(layout, "财务中心", 1)
    sichuan = dept_col_letter(layout, "四川分公司", 1)
    jinan_z = dept_col_letter(layout, "济南子公司", 1)
    assert ws[f"{tech}{row['550301']}"].value == 10
    assert ws[f"{local1}{row['550101']}"].value == 20
    assert ws[f"{local2}{row['540109']}"].value == 30
    assert ws[f"{yun}{row['550201']}"].value == 5
    assert ws[f"{finance}{row['550201']}"].value == 22
    assert ws[f"{sichuan}{row['550201']}"].value in (None, "")
    assert ws[f"O{row['550201']}"].value == 22
    assert ws[f"Q{row['540109']}"].value in (None, "")
    assert ws[f"{jinan_z}{row['540109']}"].value in (None, "")
    assert ws[f"{jinan_z}{row['54011101']}"].value == 3
    assert ws[f"{jinan_z}{row['540123']}"].value == 52
    assert ws[f"Q{row['540123']}"].value == 52
    assert ws[f"Q{row['5401']}"].value == 55
    pf = openpyxl.load_workbook(out)["利润表"]
    assert pf["E1"].value == "山东26年 8月"
    assert pf["H1"].value == "四川26年 8月"
    assert pf["I1"].value == "济南子公司26年8月"
    assert ws["A2"].font.name == "等线"
    assert ws["A2"].font.bold is not True
    assert ws["A2"].fill.fgColor.rgb.endswith("D6DCE4")
    assert ws[f"C{row['540109']}"].fill.fgColor.rgb.endswith("E2EFDA")


def test_resolve_dept_channel_and_localization():
    sys.path.insert(0, str(SCRIPTS))
    from payroll_ledger import resolve_dept, load_payroll_rules

    rules = load_payroll_rules()
    dept, occ, prefix = resolve_dept("渠道开发中心", "技术部", rules)
    assert dept == "技术中心"
    assert prefix is None
    assert (rules.get("center_to_prefix") or {}).get("渠道开发中心") == "5503"
    dept, occ, _ = resolve_dept("营销中心", "本地化事业部", rules)
    assert dept == "本地化" and occ == 1
    dept, occ, _ = resolve_dept("项目中心", "本地化事业部", rules)
    assert dept == "本地化" and occ == 2
    dept, _occ, prefix = resolve_dept("人力资源部", "", rules, entity="湖南分公司", extra_key="人力资源部")
    assert dept == "湖南分公司" and prefix == "5502"


def test_income_jinan_to_channel_bengongsi_to_director(tmp_path: Path):
    _write_account(
        tmp_path / "hq.xlsx",
        "甲骨易（北京）语言科技股份有限公司",
        [("5101", "主营业务收入", None, 100.0)],
    )
    _write_assist(
        tmp_path / "hq_d.xlsx",
        "甲骨易（北京）语言科技股份有限公司",
        [
            ("510107", "主营业务收入_其他", "济南分公司", None, 30.0),
            ("510102", "主营业务收入_游戏", "游戏", None, 70.0),
        ],
    )
    _write_account(
        tmp_path / "wh.xlsx",
        "北京甲骨易文化传媒有限公司",
        [("5101", "主营业务收入", None, 20.0)],
    )
    out = tmp_path / "out.xlsx"
    _run(["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"])
    ws = openpyxl.load_workbook(out, data_only=False)["损益表"]
    layout = load_layout()
    row = account_row_map(layout)["5101"]
    director = dept_col_letter(layout, "营销总监及助理", 1)
    channel = dept_col_letter(layout, "渠道开发中心", 1)
    game = dept_col_letter(layout, "游戏", 1)
    jinan = dept_col_letter(layout, "济南分公司", 1)
    assert ws[f"{channel}{row}"].value == 30
    assert ws[f"{game}{row}"].value == 70
    assert ws[f"{director}{row}"].value == 20
    assert ws[f"{jinan}{row}"].value in (None, "")
    book, _ = eval_workbook(openpyxl.load_workbook(out, data_only=False))
    av = book.cell_value("损益表", row, 48)
    assert abs(Decimal(str(av or 0))) <= Decimal("0.05")


def test_investment_income_goes_to_exec_office(tmp_path: Path):
    _write_account(
        tmp_path / "hq.xlsx",
        "甲骨易（北京）语言科技股份有限公司",
        [("520199", "其他", None, 5.0)],
    )
    out = tmp_path / "out.xlsx"
    _run(["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"])
    ws = openpyxl.load_workbook(out)["损益表"]
    layout = load_layout()
    row = account_row_map(layout)["520199"]
    exec_office = dept_col_letter(layout, "总经办", 1)
    assert ws[f"{exec_office}{row}"].value == 5
    book, _ = eval_workbook(openpyxl.load_workbook(out, data_only=False))
    av = book.cell_value("损益表", row, 48)
    assert abs(Decimal(str(av or 0))) <= Decimal("0.05")


def test_depreciation_amort_tax_finance_and_540103_forced_depts(tmp_path: Path):
    _write_account(
        tmp_path / "hq.xlsx",
        "甲骨易（北京）语言科技股份有限公司",
        [
            ("540103", "翻译语言服务", 40.0, None),
            ("550252", "折旧费", 12.0, None),
            ("550253", "无形资产摊销", 8.0, None),
            ("550319", "折旧费", 3.0, None),
            ("540201", "城市维护建设税", 5.0, None),
            ("550404", "手续费", 7.0, None),
        ],
    )
    _write_assist(
        tmp_path / "hq_d.xlsx",
        "甲骨易（北京）语言科技股份有限公司",
        [
            ("540103", "翻译语言服务", "项目一组", 25.0, None),
            ("540103", "翻译语言服务", "数据系统中心", 15.0, None),
            ("550404", "手续费", "KA", 7.0, None),
        ],
    )
    _write_account(
        tmp_path / "wh.xlsx",
        "北京甲骨易文化传媒有限公司",
        [("540201", "城市维护建设税", 2.0, None), ("550404", "手续费", 1.0, None)],
    )
    out = tmp_path / "out.xlsx"
    _run(["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"])
    ws = openpyxl.load_workbook(out, data_only=False)["损益表"]
    layout = load_layout()
    row = account_row_map(layout)
    director = dept_col_letter(layout, "项目总监及助理", 1)
    yizu = dept_col_letter(layout, "项目一组", 1)
    data = dept_col_letter(layout, "数据系统中心", 1)
    yun = dept_col_letter(layout, "运营保障中心", 1)
    finance = dept_col_letter(layout, "财务中心", 1)
    ka = dept_col_letter(layout, "KA", 1)
    assert director and yun and finance
    assert ws[f"{director}{row['540103']}"].value == 40
    assert ws[f"{yizu}{row['540103']}"].value in (None, "")
    assert ws[f"{data}{row['540103']}"].value in (None, "")
    assert ws[f"{yun}{row['550252']}"].value == 12
    assert ws[f"{yun}{row['550253']}"].value == 8
    assert ws[f"{yun}{row['550319']}"].value == 3
    assert ws[f"{finance}{row['540201']}"].value == 7
    assert ws[f"{finance}{row['550404']}"].value == 8
    assert ws[f"{ka}{row['550404']}"].value in (None, "")
    book, _ = eval_workbook(openpyxl.load_workbook(out, data_only=False))
    for code in ("540103", "550252", "550253", "550319", "540201", "550404"):
        av = book.cell_value("损益表", row[code], 48)
        assert abs(Decimal(str(av or 0))) <= Decimal("0.05"), code


def test_hunan_own_column_office_and_rent_without_assist(tmp_path: Path):
    _write_account(
        tmp_path / "hnz.xlsx",
        "甲骨易（湖南）科技有限公司",
        [("550210", "办公费", 4.0, None), ("550212", "房租", 5.0, None), ("550201", "工资", 8.0, None)],
    )
    _write_account(
        tmp_path / "hnf.xlsx",
        "甲骨易（北京）语言科技股份有限公司湖南分公司",
        [("550212", "房租", 2.0, None)],
    )
    _write_account(tmp_path / "hq.xlsx", "甲骨易（北京）语言科技股份有限公司", [("5101", "主营业务收入", None, 1.0)])
    out = tmp_path / "out.xlsx"
    r = _run(["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"])
    assert out.is_file(), r.stdout + r.stderr
    ws = openpyxl.load_workbook(out)["损益表"]
    layout = load_layout()
    row = account_row_map(layout)
    hnz = dept_col_letter(layout, "湖南子公司", 1)
    hnf = dept_col_letter(layout, "湖南分公司", 1)
    yun = dept_col_letter(layout, "运营保障中心", 1)
    assert hnz and hnf
    assert ws[f"{hnz}{row['550210']}"].value == 4
    assert ws[f"{hnz}{row['550212']}"].value == 5
    assert ws[f"{hnf}{row['550212']}"].value == 2
    assert ws[f"{hnz}{row['550201']}"].value in (None, "")
    assert ws[f"{yun}{row['550212']}"].value in (None, "")


def test_hunan_keeps_assist_and_does_not_double_office(tmp_path: Path):
    _write_account(
        tmp_path / "hnz.xlsx",
        "甲骨易（湖南）科技有限公司",
        [("550210", "办公费", 4.0, None)],
    )
    _write_assist(
        tmp_path / "hnz_d.xlsx",
        "甲骨易（湖南）科技有限公司",
        [("550210", "办公费", "运营保障中心", 4.0, None)],
    )
    out = tmp_path / "out.xlsx"
    _run(["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"])
    ws = openpyxl.load_workbook(out)["损益表"]
    layout = load_layout()
    row = account_row_map(layout)
    hnz = dept_col_letter(layout, "湖南子公司", 1)
    yun = dept_col_letter(layout, "运营保障中心", 1)
    assert ws[f"{yun}{row['550210']}"].value == 4
    assert ws[f"{hnz}{row['550210']}"].value in (None, "")


def test_jinan_service_adds_without_wiping_hq_hr():
    from convert import apply_jinan_service_leftover

    layout = load_layout()
    hr = dept_col_letter(layout, "人力资源部", 1)
    jn = dept_col_letter(layout, "济南子公司", 1)
    entity_amts = {
        "济南子公司": {
            "540109": {"debit": Decimal("6"), "credit": None},
            "540111": {"debit": Decimal("3"), "credit": None},
        }
    }
    dept_amts = {"540123": {hr: Decimal("9")}}
    notes: list[str] = []
    apply_jinan_service_leftover(
        entity_amts,
        dept_amts,
        [{"entity": "济南子公司", "values": {"管理费用": Decimal("10")}}],
        layout,
        notes,
    )
    assert dept_amts["540123"][jn] == Decimal("1")
    assert dept_amts["540123"][hr] == Decimal("9")
    assert entity_amts["济南子公司"]["540123"]["debit"] == Decimal("1")
    assert "济南服务费=管理费用-工资-社保" in notes


def test_rent_abstract_moves_from_yunbao_and_skips_without_source():
    from rent_abstract import apply_rent_abstract_split, lines_from_entries

    layout = load_layout()
    yun = dept_col_letter(layout, "运营保障中心", 1)
    gd = dept_col_letter(layout, "广东分公司", 1)
    hnf = dept_col_letter(layout, "湖南分公司", 1)
    lines = lines_from_entries(
        [
            {"account": "550212", "explanation": "深圳办公室房租", "debit": 3, "credit": None},
            {"account": "550212", "explanation": "长沙房租", "debit": 2, "credit": None},
            {"account": "550212", "explanation": "总部房租", "debit": 8, "credit": None},
        ]
    )
    dept_amts = {"550212": {yun: Decimal("20")}}
    notes: list[str] = []
    apply_rent_abstract_split(dept_amts, lines, layout, notes)
    assert dept_amts["550212"][gd] == Decimal("3")
    assert dept_amts["550212"][hnf] == Decimal("2")
    assert dept_amts["550212"][yun] == Decimal("15")
    assert "房租摘要=" in notes[0]
    empty_notes: list[str] = []
    left = {"550212": {yun: Decimal("20")}}
    apply_rent_abstract_split(left, [], layout, empty_notes)
    assert left["550212"][yun] == Decimal("20")
    assert left["550212"].get(gd) in (None, 0)
    assert "房租摘要=无源" in empty_notes


def test_rent_abstract_from_journal_xlsx(tmp_path: Path):
    _write_account(
        tmp_path / "hq.xlsx",
        "甲骨易（北京）语言科技股份有限公司",
        [("550212", "房租", 20.0, None)],
    )
    _write_assist(
        tmp_path / "hq_d.xlsx",
        "甲骨易（北京）语言科技股份有限公司",
        [("550212", "房租", "运营保障中心", 20.0, None)],
    )
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "序时账"
    ws["A1"] = "凭证号"
    ws["B1"] = "摘要"
    ws["C1"] = "科目代码"
    ws["D1"] = "科目名称"
    ws["E1"] = "借方"
    ws["F1"] = "贷方"
    ws["A2"] = "1"
    ws["B2"] = "深圳房租"
    ws["C2"] = "550212"
    ws["D2"] = "房租"
    ws["E2"] = 3
    ws["A3"] = "2"
    ws["B3"] = "长沙房租"
    ws["C3"] = "550212"
    ws["D3"] = "房租"
    ws["E3"] = 2
    wb.save(tmp_path / "序时账.xlsx")
    wb.close()
    out = tmp_path / "out.xlsx"
    r = _run(["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"])
    assert out.is_file(), r.stdout + r.stderr
    ws = openpyxl.load_workbook(out)["损益表"]
    layout = load_layout()
    row = account_row_map(layout)
    yun = dept_col_letter(layout, "运营保障中心", 1)
    gd = dept_col_letter(layout, "广东分公司", 1)
    hnf = dept_col_letter(layout, "湖南分公司", 1)
    assert ws[f"{gd}{row['550212']}"].value == 3
    assert ws[f"{hnf}{row['550212']}"].value == 2
    assert ws[f"{yun}{row['550212']}"].value == 15
    report = out.with_name(out.stem + "_运行报告.txt").read_text(encoding="utf-8")
    assert "房租摘要=" in report


def test_looks_like_id_or_phone_rejects_digits_not_money():
    sys.path.insert(0, str(SCRIPTS))
    from payroll_ledger import looks_like_id_or_phone

    assert looks_like_id_or_phone("111111111111111111")
    assert looks_like_id_or_phone("13800138000")
    assert looks_like_id_or_phone(111111111111111111)
    assert not looks_like_id_or_phone(3520.15)
    assert not looks_like_id_or_phone("880.50")
    assert not looks_like_id_or_phone(None)


def test_shanghai_payroll_id_in_amount_col_dropped_and_asks(tmp_path: Path):
    _write_account(
        tmp_path / "sh.xlsx",
        "甲骨易智译（上海）科技有限公司",
        [("5101", "主营业务收入", None, 1.0)],
    )
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "202608上海"
    ws["A1"] = "姓名"
    ws["B1"] = "养老(16%)"
    ws["C1"] = "医疗+生育(9%)"
    ws["A2"] = "甲"
    ws["B2"] = "111111111111111111"
    ws["C2"] = 12
    wb.save(tmp_path / "202608_职工薪酬台账.xlsx")
    wb.close()
    out = tmp_path / "out.xlsx"
    r = _run(["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"])
    assert out.is_file(), r.stdout + r.stderr
    assert "ask=" in r.stdout
    assert "202608上海" in r.stdout
    assert "薪酬台账拒读" in (out.with_name(out.stem + "_运行报告.txt").read_text(encoding="utf-8"))
    ws = openpyxl.load_workbook(out)["损益表"]
    layout = load_layout()
    row = account_row_map(layout)["55010301"]
    local = dept_col_letter(layout, "本地化", 1)
    assert ws[f"{local}{row}"].value in (None, "")


def test_skill_tells_agent_to_stop_on_payroll_ask():
    text = (SKILL / "SKILL.md").read_text(encoding="utf-8")
    assert "薪酬台账拒读" in text
    assert "禁止自己改列映射" in text or "不要自己改列映射" in text
    assert "原样问她" in text
    assert "核对非0其它" in text
    assert "假数" in text


def test_fake_payroll_filename_is_ignored(tmp_path: Path):
    _write_account(
        tmp_path / "wh.xlsx",
        "北京甲骨易文化传媒有限公司",
        [("550198", "活动团建费", 8.0, None)],
    )
    _write_assist(
        tmp_path / "wh_d.xlsx",
        "北京甲骨易文化传媒有限公司",
        [("550198", "活动团建费", "大客户", 8.0, None)],
    )
    _write_payroll(tmp_path / "202608_职工薪酬台账_假数.xlsx")
    out = tmp_path / "out.xlsx"
    r = _run(["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"])
    assert out.is_file(), r.stdout + r.stderr
    assert "核对非0其它=无" in r.stdout
    assert "薪酬台账=" not in r.stdout
    report = out.with_name(out.stem + "_运行报告.txt").read_text(encoding="utf-8")
    assert "薪酬台账=" not in report or "假数" not in report


def test_insert_row_then_id_in_amount_col_drops_sheet(tmp_path: Path):
    _write_account(
        tmp_path / "sh.xlsx",
        "甲骨易智译（上海）科技有限公司",
        [("5101", "主营业务收入", None, 1.0)],
    )
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "202608上海"
    ws["A1"] = "说明"
    ws["A2"] = "姓名"
    ws["B2"] = "养老(16%)"
    ws["C2"] = "医疗+生育(9%)"
    ws["A3"] = "甲"
    ws["B3"] = "111111111111111111"
    ws["C3"] = 12
    wb.save(tmp_path / "202608_职工薪酬台账.xlsx")
    wb.close()
    out = tmp_path / "out.xlsx"
    r = _run(["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"])
    assert out.is_file(), r.stdout + r.stderr
    assert "ask=" in r.stdout
    assert "202608上海" in r.stdout
    assert "核对非0其它=无" in r.stdout
    layout = load_layout()
    row = account_row_map(layout)["55010301"]
    local = dept_col_letter(layout, "本地化", 1)
    ws = openpyxl.load_workbook(out)["损益表"]
    assert ws[f"{local}{row}"].value in (None, "")


def test_mapped_office_expense_check_is_dash(tmp_path: Path):
    _write_account(
        tmp_path / "hq.xlsx",
        "甲骨易（北京）语言科技股份有限公司",
        [("550211", "办公费", 12.0, None)],
    )
    _write_assist(
        tmp_path / "hq_d.xlsx",
        "甲骨易（北京）语言科技股份有限公司",
        [("550211", "办公费", "运营保障中心", 12.0, None)],
    )
    out = tmp_path / "out.xlsx"
    r = _run(["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"])
    assert out.is_file(), r.stdout + r.stderr
    assert "核对非0其它=无" in r.stdout
    assert "核对非0绿行=无" in r.stdout
    from convert import list_nonzero_check_codes

    layout = load_layout()
    rows = list_nonzero_check_codes(openpyxl.load_workbook(out, data_only=False), layout)
    assert rows == []


def test_inspect_skips_product_xlsx(tmp_path: Path):
    sys.path.insert(0, str(SCRIPTS))
    from inspect_inputs import inspect_file

    product = tmp_path / "月度损益表_202608.xlsx"
    _write_account(product, "甲骨易（北京）语言科技股份有限公司", [("550211", "办公费", 1.0, None)])
    assert inspect_file(product) == []


def test_occurrence_uses_check_side_not_net():
    layout = load_layout()
    assert nature_amount("5101", Decimal("100"), Decimal("100"), layout) == Decimal("100")
    assert nature_amount("520199", Decimal("-5"), Decimal("-5"), layout) == Decimal("-5")
    assert nature_amount("5401", Decimal("80"), Decimal("80"), layout) == Decimal("80")
    assert nature_amount("5101", None, Decimal("12"), layout) == Decimal("12")
    assert nature_amount("540109", Decimal("9"), None, layout) == Decimal("9")
    assert nature_amount("5101", Decimal("3"), None, layout) is None
    assert nature_amount("5401", None, Decimal("3"), layout) is None


def test_closed_income_both_sides_still_checks_zero(tmp_path: Path):
    _write_account(
        tmp_path / "hq.xlsx",
        "甲骨易（北京）语言科技股份有限公司",
        [("5101", "主营业务收入", 100.0, 100.0)],
    )
    _write_assist(
        tmp_path / "hq_d.xlsx",
        "甲骨易（北京）语言科技股份有限公司",
        [
            ("510107", "主营业务收入_其他", "济南分公司", None, 30.0),
            ("510102", "主营业务收入_游戏", "游戏", None, 70.0),
        ],
    )
    _write_account(
        tmp_path / "wh.xlsx",
        "北京甲骨易文化传媒有限公司",
        [("5101", "主营业务收入", None, 20.0)],
    )
    out = tmp_path / "out.xlsx"
    r = _run(["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"])
    assert out.is_file(), r.stdout + r.stderr
    ws = openpyxl.load_workbook(out, data_only=False)["损益表"]
    layout = load_layout()
    row = account_row_map(layout)["5101"]
    director = dept_col_letter(layout, "营销总监及助理", 1)
    channel = dept_col_letter(layout, "渠道开发中心", 1)
    game = dept_col_letter(layout, "游戏", 1)
    assert ws[f"{channel}{row}"].value == 30
    assert ws[f"{game}{row}"].value == 70
    assert ws[f"{director}{row}"].value == 20
    book, _ = eval_workbook(openpyxl.load_workbook(out, data_only=False))
    av = book.cell_value("损益表", row, 48)
    assert abs(Decimal(str(av or 0))) <= Decimal("0.05")


def test_closed_investment_both_sides_goes_to_exec(tmp_path: Path):
    _write_account(
        tmp_path / "hq.xlsx",
        "甲骨易（北京）语言科技股份有限公司",
        [("520199", "其他", -5.0, -5.0)],
    )
    out = tmp_path / "out.xlsx"
    r = _run(["--period", "202608", "--input-dir", str(tmp_path), "--out", str(out), "--no-api"])
    assert out.is_file(), r.stdout + r.stderr
    ws = openpyxl.load_workbook(out)["损益表"]
    layout = load_layout()
    row = account_row_map(layout)["520199"]
    exec_office = dept_col_letter(layout, "总经办", 1)
    assert ws[f"{exec_office}{row}"].value == -5
    book, _ = eval_workbook(openpyxl.load_workbook(out, data_only=False))
    av = book.cell_value("损益表", row, 48)
    assert abs(Decimal(str(av or 0))) <= Decimal("0.05")
